from __future__ import annotations

from pathlib import Path
import json
import hashlib
import os
import subprocess
import shutil
import sys
import time
import threading
import traceback
import queue
import re

import customtkinter as ctk
import pandas as pd
from openpyxl import load_workbook
from PIL import Image as PILImage, ImageDraw, ImageFont, ImageFilter
from tkinter import filedialog, messagebox, simpledialog

from modules.shopify_parser import normalize_order_export, parse_shopify_orders
from modules.catalog_manager import (
    backup_file,
    clean_product_master,
    restore_product_master,
    restore_seed_product_master,
    save_product_candidate_list,
)
from modules.purchase_order_generator import safe_filename, split_embedded_size
from modules.purchase_order_files import is_final_purchase_order_pdf
from modules.paths import (
    data_dir,
    legacy_data_dir,
    legacy_product_master_path,
    product_master_path,
    product_master_update_marker_path,
    reports_dir,
    v12_data_dir,
)
from modules.master_sync import (
    live_product_master_path,
    product_master_signature,
)
from modules.packet_lock import review_routing_status
from modules.active_packet import (
    active_workbook_matches_state,
    bind_active_workbook,
    new_active_packet_state,
)
from modules.report_modes import GENERAL_SALES_PERIOD, UNIFORM_SIZING_EVENT
from modules.decoration_fulfillment import (
    STANDARD_ORCHID_WORKFLOW, ENTIRE_ORDER_OUTSOURCED, is_outsourced_decoration,
    normalize_decoration_fulfillment,
)
from modules.purchase_rules import apply_purchase_rule_defaults, normalize_bool, row_rules
from modules.routing_rules import enforce_permanent_vendor_overrides
from modules.product_intelligence import apply_product_intelligence
from modules.blank_garment_rules import BLANK_DECORATION_LABEL, is_blank_decoration
from modules.decoration_locations import (
    DECORATION_LOCATIONS,
    STANDARD_DECORATION_LOCATIONS,
    LEFT_CHEST,
    OTHER_CUSTOM,
    NOT_APPLICABLE_IN_HOUSE,
    NOT_APPLICABLE_NO_DECORATION,
    default_decoration_location,
    location_choice_and_custom,
    normalize_decoration_location,
    resolve_location_choice,
)
from modules.event_archive import validate_release_packet
from modules.employee_totals import save_employee_discounts
from modules.xlsx_reader import (
    load_decoration_fulfillment,
    load_event_name,
    load_report_mode,
    load_review_lines,
    load_system_info_value,
)
from modules.note_rules import decoration_note_requires_review, decoration_instruction_recommendation
from modules.internal_services import (
    HEMMING_ALTERATION_LABEL, SEW_ON_PATCH_LABEL,
    is_in_house_decoration, is_in_house_service_product,
)
from modules.product_candidate_sync import sync_review_product_candidates, sync_shopify_catalog_enrichment
from modules.process_identity import process_is_running, product_master_process_matches
from modules.product_master_launch import product_master_launch_command
from modules.scroll_support import install_native_scroll_support, register_scrollable
from modules.thread_ink_colors import load_thread_ink_colors
from modules.job_logo_image import prepare_artwork_for_preview, prepare_artwork_for_report
from modules.decoration_color_audit import (
    apply_product_master_decoration_colors,
    audit_group_key,
    audit_item_needs_attention,
    build_audit_view_snapshot,
    build_decoration_color_audit,
)
from modules.multi_csv_import import (
    merge_additional_orders,
    preview_additional_orders,
    safe_event_stem,
    write_combined_csv_atomic,
)
from modules.mission_control import (
    canonical_report_type,
    load_mission_control_snapshot,
    load_purchase_review_snapshot,
    line_block_reasons,
    load_po_overrides,
    save_po_overrides,
)


PROJECT = Path(__file__).resolve().parent


def resource_path(*parts: str) -> Path:
    base = Path(getattr(sys, "_MEIPASS", PROJECT))
    return base.joinpath(*parts)

DATA = data_dir()
MASTER = live_product_master_path()  # Display only; actions resolve the live path again.
REPORTS = reports_dir()
PRODUCT_MASTER_PID_FILE = DATA / "product_master.pid"
PRODUCT_MASTER_REQUEST_FILE = DATA / "product_master_open_request.json"
REVIEW_STATE_FILE = DATA / "current_review_state.json"
REGENERATE_HELPER = DATA / "Regenerate Purchase Review.command"
ACTIVE_PACKET_RESET_FILE = DATA / "active_packet_reset.json"
UI_PREFERENCES_FILE = DATA / "ui_preferences.json"

PURPLE = "#5A18C9"
PURPLE_DARK = "#17122F"
PURPLE_LIGHT = "#F2EAFD"
BORDER = "#E0D8EB"
SHADOW = "#E9E4F0"
TEXT = "#14111F"
MUTED = "#6A6375"
BG = "#F3F0F7"
WHITE = "#FFFFFF"
DASHBOARD_WATERMARK_OPACITY = 0.09
SUCCESS = "#16A34A"
WARNING = "#E56A00"
DANGER = "#9C0006"

DO_NOT_OUTSOURCE_DECISION = "Do Not Outsource — Ship to Orchid"
INSTRUCTION_DECISION_OPTIONS = [
    "Choose an instruction decision",
    "Follow Note as Written",
    "Keep Product Master Default",
    "No Decoration",
    "Embroidery",
    "Screen Print",
    "Sew On Patch",
    "Hemming / Alteration",
]
OUTSOURCE_INSTRUCTION_TERMS = (
    "do not outsource", "don't outsource", "ship to orchid", "send to orchid",
)


def _clean(value: object) -> str:
    return " ".join(str(value or "").split())


def _file_stamp(path: Path | None) -> tuple[str, int, int]:
    if not path:
        return ("", 0, 0)
    candidate = Path(path)
    try:
        stat = candidate.stat()
        return (str(candidate.resolve()), int(stat.st_mtime_ns), int(stat.st_size))
    except (FileNotFoundError, OSError):
        return (str(candidate), 0, 0)


def should_show_employee_totals(
    snapshot: dict, workbook_exists: bool, current_mode: str = "", imported_line_count: int = 0
) -> bool:
    """Employee Totals is an event-only sidebar feature."""
    report_mode = str(snapshot.get("report_mode", "") or current_mode or "")
    line_count = int(snapshot.get("line_count", 0) or imported_line_count or 0)
    return bool(workbook_exists and line_count > 0 and report_mode == UNIFORM_SIZING_EVENT)


_PRODUCT_MASTER_HEALTH_CACHE_KEY: tuple[str, int, int] | None = None
_PRODUCT_MASTER_HEALTH_CACHE_VALUE: dict[str, int] | None = None


def _product_master_health() -> dict[str, int]:
    global _PRODUCT_MASTER_HEALTH_CACHE_KEY, _PRODUCT_MASTER_HEALTH_CACHE_VALUE
    result = {"records": 0, "styles": 0, "complete": 0, "incomplete": 0, "percent": 0}
    master_path = live_product_master_path()
    cache_key = _file_stamp(master_path)
    if cache_key == _PRODUCT_MASTER_HEALTH_CACHE_KEY and _PRODUCT_MASTER_HEALTH_CACHE_VALUE is not None:
        return dict(_PRODUCT_MASTER_HEALTH_CACHE_VALUE)
    if not master_path.exists():
        return result
    try:
        frame = pd.read_csv(master_path, dtype=str).fillna("")
        frame = apply_product_intelligence(apply_purchase_rule_defaults(frame))
    except Exception:
        return result
    result["records"] = len(frame)
    if frame.empty:
        return result

    def style_key(row: pd.Series) -> str:
        style = _clean(row.get("Style Number", "")).casefold()
        name = _clean(row.get("Product Name", "")).casefold()
        return f"style:{style}" if style else f"product:{name}"

    complete = 0
    groups = list(frame.groupby(frame.apply(style_key, axis=1), sort=False))
    for _, rows in groups:
        first = rows.iloc[0]
        rules = row_rules(first)
        requires_color = normalize_bool(rules["Requires Color"], True)
        requires_decoration = normalize_bool(rules["Requires Decoration"], True)
        issues = []
        if "Setup Required" in rows.columns and rows["Setup Required"].astype(str).str.strip().str.casefold().isin({"yes", "y", "true", "1"}).any():
            issues.append("setup review")
        if not _clean(first.get("Product Name", "")):
            issues.append("name")
        if not _clean(first.get("Style Number", "")):
            issues.append("style")
        if not _clean(rules.get("Product Category", "")):
            issues.append("category")
        service_only = is_in_house_service_product(
            first.get("Product Name", ""),
            first.get("Original Line Item", ""),
            first.get("Decoration Type", ""),
            style_number=first.get("Style Number", ""),
        )
        if not service_only and rows.get("Vendor", pd.Series(dtype=str)).astype(str).str.strip().eq("").any():
            issues.append("vendor")
        if requires_color and rows.get("Garment Color", pd.Series(dtype=str)).astype(str).str.strip().eq("").any():
            issues.append("color")
        if rows.get("Decoration Type", pd.Series(dtype=str)).astype(str).str.strip().eq("").any():
            issues.append("decoration")
        if requires_decoration and "Decoration Location" in rows.columns:
            locations = rows["Decoration Location"].astype(str).str.strip()
            if locations.eq("").any() or locations.eq(OTHER_CUSTOM).any():
                issues.append("decoration location")
        if requires_decoration and "Decoration Color" in rows.columns:
            decorated = rows[~rows["Decoration Type"].map(lambda value: is_blank_decoration(value) or is_in_house_decoration(value))]
            if not decorated.empty and decorated["Decoration Color"].astype(str).str.strip().eq("").any():
                issues.append("decoration color")
        if not issues:
            complete += 1

    styles = len(groups)
    result.update({
        "styles": styles,
        "complete": complete,
        "incomplete": max(styles - complete, 0),
        "percent": round((complete / styles) * 100) if styles else 0,
    })
    _PRODUCT_MASTER_HEALTH_CACHE_KEY = cache_key
    _PRODUCT_MASTER_HEALTH_CACHE_VALUE = dict(result)
    return result


class OrchidPurchaseManager(ctk.CTk):
    PAGE_TITLES = {
        "dashboard": ("Dashboard", "Move the current purchase packet through the five purchasing stages."),
        "import": ("Current Event", "Start, add orders to, finish, and manage the active purchase packet."),
        "master": ("Product Master", "Search and maintain permanent vendor, category, color, and decoration rules."),
        "review": ("Purchase Review", "Complete only the order-specific decisions that remain after Product Master setup."),
        "audit": ("Decoration Color Audit", "Verify thread and ink colors for the products included in the current event."),
        "purchase": ("Purchase Orders", "Create final vendor purchase orders from the completed review."),
        "employees": ("Employee Totals", "View employee order totals, enter discounts, and save final event totals."),
        "settings": ("Settings", "Diagnostics, catalog maintenance, backups, and storage locations."),
    }

    def __init__(self):
        super().__init__()
        ctk.set_appearance_mode("light")
        self.title("Orchid Purchase Manager Professional 4.9.104")
        self.geometry("1400x900")
        self.minsize(1180, 760)
        self.configure(fg_color=BG)
        self.after(80, self._maximize_main_window)

        self.selected_csv: Path | None = None
        self.current_event_name = ""
        self.current_mode = ""
        self.current_decoration_fulfillment = STANDARD_ORCHID_WORKFLOW
        self.last_review_workbook: Path | None = None
        self.last_purchase_order_dir: Path | None = None
        self.imported_line_count = 0
        self.review_needed: int | None = None
        self.active_page = "import"
        self.nav_buttons: dict[str, ctk.CTkButton] = {}
        self.pages: dict[str, ctk.CTkFrame] = {}
        self.po_entry_vars: dict[str, tuple[str, str, ctk.StringVar]] = {}
        self.employee_discount_vars: dict[int, ctk.StringVar] = {}
        self.employee_total_labels: dict[int, ctk.CTkLabel] = {}
        self.employee_records_by_row: dict[int, dict] = {}
        self.last_snapshot: dict = {}
        self.product_master_process: subprocess.Popen | None = None
        self.product_master_launching = False
        self._product_master_was_open = False
        self._candidate_sync_workbook = ""
        self._regenerating_review = False
        self._auto_regenerate_after_id = None
        self.product_master_update_marker = product_master_update_marker_path()
        self._last_product_master_update = (
            self.product_master_update_marker.stat().st_mtime
            if self.product_master_update_marker.exists() else 0.0
        )
        self._launch_choice_shown = False
        self._suppress_empty_import_prompt = False
        self.success_sound_enabled = True
        self._success_animation_running = False
        self._success_animation_skip_requested = False
        self._reduced_motion_override = None
        ui_preferences = self._load_ui_preferences()
        self.dashboard_view = str(ui_preferences.get("dashboard_view", "guided") or "guided")
        if self.dashboard_view not in {"guided", "management"}:
            self.dashboard_view = "guided"
        self.interaction_animations_enabled = bool(ui_preferences.get("interaction_animations_enabled", True))
        self._logo_pulse_running = False
        self._logo_pulse_after_ids = []
        self._outsourced_logo_dialog_pending = False
        self._mission_snapshot_cache_key = None
        self._routing_status_cache_key = None
        self._routing_status_cache: dict = {"current": True, "reason": "", "changes": []}
        self._decoration_audit_status_cache_key = None
        self._decoration_audit_status_cache_value = None
        self._mission_snapshot_cache: dict = {}
        self._dashboard_refresh_key = None
        self._normalized_import_cache = None
        self._parsed_import_cache = None
        self._import_cache_key = None
        self._review_row_index_cache = {}
        self._review_value_cache = {}
        self._review_status_cache = None
        self._review_cache_key = None
        self._review_cache_data = {}
        self._review_fast_cache_key = None
        self._review_fast_snapshot_cache: dict = {}
        self._review_fast_values: dict = {}
        self._review_load_token = 0
        self._review_load_in_progress = False
        self._dashboard_warm_token = 0
        self._review_save_jobs: queue.Queue = queue.Queue()
        self._review_save_results: queue.Queue = queue.Queue()
        self._review_save_worker_started = False
        self._review_save_pending = 0
        self._review_save_sequence = 0
        self._review_save_polling = False
        self._review_exit_after_saves = False
        self._review_save_last_error = ""
        self.decoration_audit_records: list[dict] = []
        self.decoration_audit_row_vars: dict[str, dict] = {}
        self.decoration_audit_pending_edits: dict[str, dict] = {}
        self._decoration_audit_source_stamp = None
        self._decoration_audit_event_records_cache_key = None
        self._decoration_audit_event_records_cache: list[dict] = []
        self._decoration_audit_build_cache_key = None
        self._decoration_audit_build_cache: list[dict] = []

        try:
            enforce_permanent_vendor_overrides(live_product_master_path())
        except Exception:
            pass
        self.discover_existing_work()
        self.build_ui()
        self._install_global_click_pulse()
        install_native_scroll_support(
            self,
            [
                frame for frame in (
                    getattr(self, "mc_work_scroll", None),
                    getattr(self, "mc_employee_scroll", None),
                    getattr(self, "settings_scroll", None),
                    getattr(self, "review_decision_frame", None),
                ) if frame is not None
            ],
        )
        self.refresh_dashboard()
        self.show_page("import")
        self.after(650, self._prompt_resume_or_start_new)
        self.after(950, self._prompt_for_csv_if_empty)
        self.after(1100, self._poll_product_master_updates)
        self.after(1300, self._upgrade_active_review_for_decoration_notes)
        self.after(1550, self._resume_pending_review_saves)

    def _load_ui_preferences(self) -> dict:
        try:
            if UI_PREFERENCES_FILE.exists():
                payload = json.loads(UI_PREFERENCES_FILE.read_text(encoding="utf-8"))
                return payload if isinstance(payload, dict) else {}
        except Exception:
            pass
        return {}

    def _save_ui_preferences(self) -> None:
        try:
            UI_PREFERENCES_FILE.parent.mkdir(parents=True, exist_ok=True)
            UI_PREFERENCES_FILE.write_text(
                json.dumps({
                    "dashboard_view": getattr(self, "dashboard_view", "guided"),
                    "interaction_animations_enabled": bool(getattr(self, "interaction_animations_enabled", True)),
                }, indent=2, sort_keys=True),
                encoding="utf-8",
            )
        except Exception:
            pass

    def _build_sidebar_logo_frames(self, source: PILImage.Image) -> list[ctk.CTkImage]:
        """Create a fixed-size official logo plus a graceful purple radial pulse."""
        canvas_size = 108
        flower_size = 70
        flower = source.convert("RGBA").resize((flower_size, flower_size), PILImage.Resampling.LANCZOS)
        frames: list[ctk.CTkImage] = []
        purple_shades = [
            (93, 33, 201), (117, 55, 218), (142, 82, 229), (172, 121, 237),
        ]
        for frame_index in range(11):
            canvas = PILImage.new("RGBA", (canvas_size, canvas_size), (0, 0, 0, 0))
            if frame_index:
                draw = ImageDraw.Draw(canvas, "RGBA")
                progress = frame_index / 10.0
                for ring_index, color in enumerate(purple_shades):
                    phase = max(0.0, min(1.0, progress - ring_index * 0.10))
                    radius = 24 + phase * 28 + ring_index * 4
                    alpha = int(86 * (1.0 - progress * 0.72) * (1.0 - ring_index * 0.13))
                    box = (
                        canvas_size / 2 - radius,
                        canvas_size / 2 - radius,
                        canvas_size / 2 + radius,
                        canvas_size / 2 + radius,
                    )
                    draw.ellipse(box, outline=(*color, max(alpha, 0)), width=3)
            x = (canvas_size - flower_size) // 2
            y = (canvas_size - flower_size) // 2
            canvas.alpha_composite(flower, (x, y))
            frames.append(ctk.CTkImage(light_image=canvas, dark_image=canvas, size=(canvas_size, canvas_size)))
        return frames

    def _initialize_sidebar_logo(self, brand: ctk.CTkFrame) -> None:
        self.sidebar_brand_logo_label = None
        self.sidebar_brand_flower_frames = []
        white_flower_path = resource_path("assets", "orchid_flower_white.png")
        if not white_flower_path.exists():
            return
        source = PILImage.open(white_flower_path).convert("RGBA")
        self.sidebar_brand_flower_frames = self._build_sidebar_logo_frames(source)
        self.sidebar_brand_flower = self.sidebar_brand_flower_frames[0]
        self.sidebar_brand_logo_label = ctk.CTkLabel(
            brand, text="", image=self.sidebar_brand_flower,
            width=108, height=108, fg_color="transparent",
        )
        self.sidebar_brand_logo_label.grid(row=0, column=0, pady=(12, 0))

    def _initialize_header_wordmark(self, parent: ctk.CTkFrame) -> None:
        """Place the official Orchid wordmark directly on the purple header."""
        self.sidebar_brand_logo_label = None
        self.sidebar_brand_flower_frames = []
        logo_path = resource_path("assets", "orchid_logo.png")
        if not logo_path.exists():
            return
        source = PILImage.open(logo_path).convert("RGBA")
        # The lower portion of the source contains the retail tagline. The app
        # uses a dedicated PURCHASE MANAGER subtitle instead.
        wordmark = source.crop((0, 0, source.width, int(source.height * 0.81)))
        pixels = wordmark.load()
        for y in range(wordmark.height):
            for x in range(wordmark.width):
                red, green, blue, alpha = pixels[x, y]
                if not alpha:
                    continue
                is_orchid_purple = blue > red and red > green * 1.35
                pixels[x, y] = (red, green, blue, alpha) if is_orchid_purple else (248, 246, 252, alpha)
        bbox = wordmark.getchannel("A").getbbox()
        if bbox:
            wordmark = wordmark.crop(bbox)

        base_width = 560
        base_height = max(1, round(base_width * wordmark.height / wordmark.width))
        frames = []
        for index in range(11):
            # A very small scale bloom preserves the click animation without
            # disturbing the approved header proportions at rest.
            bloom = round(8 * (index / 10.0)) if index else 0
            size = (base_width + bloom, base_height + round(bloom * base_height / base_width))
            frames.append(ctk.CTkImage(light_image=wordmark, dark_image=wordmark, size=size))
        self.sidebar_brand_flower_frames = frames
        self.sidebar_brand_flower = frames[0]
        self.sidebar_brand_logo_label = ctk.CTkLabel(
            parent, text="", image=self.sidebar_brand_flower,
            width=590, height=142, fg_color="transparent",
        )
        self.sidebar_brand_logo_label.place(relx=0.5, rely=0.40, anchor="center")
        ctk.CTkLabel(
            parent, text="PURCHASE MANAGER", text_color="#F8F6FC",
            font=ctk.CTkFont(size=22, weight="bold"),
        ).place(relx=0.5, rely=0.80, anchor="center")

    def _finish_logo_pulse(self) -> None:
        self._logo_pulse_running = False
        for after_id in list(getattr(self, "_logo_pulse_after_ids", [])):
            try:
                self.after_cancel(after_id)
            except Exception:
                pass
        self._logo_pulse_after_ids = []
        label = getattr(self, "sidebar_brand_logo_label", None)
        frames = getattr(self, "sidebar_brand_flower_frames", [])
        if label is not None and frames:
            try:
                label.configure(image=frames[0])
            except Exception:
                pass

    def _play_logo_pulse(self) -> None:
        if not bool(getattr(self, "interaction_animations_enabled", True)):
            return
        if self._prefers_reduced_motion():
            return
        if bool(getattr(self, "_logo_pulse_running", False)):
            self._finish_logo_pulse()
        label = getattr(self, "sidebar_brand_logo_label", None)
        frames = getattr(self, "sidebar_brand_flower_frames", [])
        if label is None or len(frames) < 2:
            return
        self._logo_pulse_running = True
        sequence = list(range(1, len(frames))) + list(range(len(frames) - 2, 0, -1)) + [0]
        self._logo_pulse_after_ids = []
        for position, frame_index in enumerate(sequence):
            def show(index=frame_index, is_last=(position == len(sequence) - 1)):
                if not getattr(self, "_logo_pulse_running", False):
                    return
                try:
                    label.configure(image=frames[index])
                except Exception:
                    self._finish_logo_pulse()
                    return
                if is_last:
                    self._logo_pulse_running = False
                    self._logo_pulse_after_ids = []
            self._logo_pulse_after_ids.append(self.after(position * 36, show))

    def _install_global_click_pulse(self) -> None:
        """Pulse the official Orchid flower for every in-app mouse click."""
        try:
            self.bind_all("<Button-1>", self._on_global_click_pulse, add="+")
        except Exception:
            pass

    def _on_global_click_pulse(self, _event=None) -> None:
        if not bool(getattr(self, "interaction_animations_enabled", True)):
            return
        try:
            self.after_idle(self._play_logo_pulse)
        except Exception:
            self._play_logo_pulse()

    def _set_interaction_animations_enabled(self) -> None:
        enabled = bool(self.interaction_animation_var.get()) if hasattr(self, "interaction_animation_var") else True
        self.interaction_animations_enabled = enabled
        if not enabled:
            self._finish_logo_pulse()
        self._save_ui_preferences()

    def _maximize_main_window(self):
        """Open maximized while remaining a normal macOS window."""
        try:
            self.state("zoomed")
            return
        except Exception:
            pass
        try:
            self.update_idletasks()
            width = max(self.winfo_screenwidth(), self.winfo_width())
            height = max(self.winfo_screenheight() - 70, self.winfo_height())
            self.geometry(f"{width}x{height}+0+0")
        except Exception:
            pass

    def discover_existing_work(self):
        """Restore only the workbook explicitly bound to the active CSV packet.

        Previous versions selected the newest Review Workbook and PO folder by
        modification time.  That allowed an old event to reappear after Replace
        CSV. Saved PDFs remain on the computer for records, but they are never
        candidates for the active event.
        """
        state = self._load_review_state()
        workbook_text = _clean(state.get("active_review_workbook", ""))
        source_text = _clean(state.get("source_csv", ""))
        source_csv = Path(source_text).expanduser() if source_text else None

        if state and source_csv and source_csv.is_file() and not _clean(state.get("active_packet_id", "")):
            upgraded = new_active_packet_state(source_csv)
            active_packet_id = upgraded["active_packet_id"]
            upgraded.update(state)
            upgraded["active_packet_id"] = active_packet_id
            state = upgraded

        # One safe upgrade path for a packet created before active-workbook
        # identity existed: locate only a workbook whose protected lock records
        # the exact saved CSV hash.  No timestamp-only selection is permitted.
        if state and not workbook_text and source_csv and source_csv.is_file() and state.get("source_signature"):
            candidates = []
            for folder in (REPORTS / "Review Workbooks", REPORTS / "review_workbooks"):
                if folder.is_dir():
                    candidates.extend(folder.glob("*.xlsx"))
            for candidate in sorted(candidates, key=lambda item: item.stat().st_mtime_ns, reverse=True):
                try:
                    migrated = bind_active_workbook(state, candidate)
                except (OSError, ValueError):
                    continue
                state = migrated
                self._write_review_state_atomic(state)
                workbook_text = _clean(state.get("active_review_workbook", ""))
                break

        workbook = Path(workbook_text).expanduser() if workbook_text else None
        valid, reason = active_workbook_matches_state(state, workbook)
        if not valid:
            # v4.9.82 added a packet ID to the Review itself. Keep a valid
            # v4.9.81 state visible solely long enough to regenerate it; never
            # let that older workbook open reports in the meantime.
            if (
                workbook
                and workbook.is_file()
                and source_csv
                and source_csv.is_file()
                and "missing its protected packet identity" in reason.casefold()
            ):
                self.last_review_workbook = workbook
                self.selected_csv = source_csv
                self.current_event_name = _clean(state.get("event_name", ""))
                self.current_mode = _clean(state.get("report_mode", ""))
                self.current_decoration_fulfillment = normalize_decoration_fulfillment(
                    state.get("decoration_fulfillment", STANDARD_ORCHID_WORKFLOW)
                )
                state["packet_identity_upgrade_required"] = True
                self._write_review_state_atomic(state)
            return
        self.last_review_workbook = workbook
        self.selected_csv = source_csv
        self.current_event_name = _clean(state.get("event_name", ""))
        self.current_mode = _clean(state.get("report_mode", ""))
        self.current_decoration_fulfillment = normalize_decoration_fulfillment(
            state.get("decoration_fulfillment", STANDARD_ORCHID_WORKFLOW)
        )
        saved_output_text = _clean(state.get("active_purchase_order_dir", ""))
        saved_output = Path(saved_output_text).expanduser() if saved_output_text else None
        if saved_output and saved_output.is_dir() and (
            _clean(state.get("active_purchase_order_packet_lock_id", ""))
            == _clean(state.get("active_review_packet_lock_id", ""))
        ):
            self.last_purchase_order_dir = saved_output

    def _active_review_needs_decoration_note_upgrade(self) -> bool:
        if (
            self.current_mode != UNIFORM_SIZING_EVENT
            and normalize_decoration_fulfillment(self.current_decoration_fulfillment) != ENTIRE_ORDER_OUTSOURCED
        ):
            return False
        path = self.last_review_workbook
        if not path or not path.exists() or not REVIEW_STATE_FILE.exists():
            return False
        try:
            workbook = load_workbook(path, read_only=True, data_only=False)
            if "All PO Lines" not in workbook.sheetnames:
                workbook.close()
                return False
            sheet = workbook["All PO Lines"]
            header_row = None
            headers: dict[str, int] = {}
            for row_number in range(1, min(sheet.max_row, 20) + 1):
                values = [str(sheet.cell(row_number, col).value or "").strip() for col in range(1, sheet.max_column + 1)]
                if "Line ID" in values:
                    header_row = row_number
                    headers = {value: idx + 1 for idx, value in enumerate(values) if value}
                    break
            required_headers = {"Decoration Decision", "Mandatory Note Review", "Decoration Location", "Decoration Placement Instructions"}
            if not header_row or not required_headers.issubset(headers):
                workbook.close()
                return True

            # Professional 4.7.4 also upgrades active workbooks to the permanent/order-specific decoration-location schema.
            # Professional 4.4.6 broadened mandatory review beyond EMB wording to
            # placement, logo, screen-print, personalization, and alteration notes.
            # Rebuild only when an existing active row would now require review.
            instruction_col = headers.get("Purchase Instructions") or headers.get("Shopify Order Notes")
            mandatory_col = headers.get("Mandatory Note Review")
            if instruction_col and mandatory_col:
                for row_number in range(header_row + 1, sheet.max_row + 1):
                    note_text = _clean(sheet.cell(row_number, instruction_col).value)
                    mandatory = _clean(sheet.cell(row_number, mandatory_col).value).casefold() in {"yes", "y", "true", "1"}
                    if note_text and decoration_note_requires_review(note_text) and not mandatory:
                        workbook.close()
                        return True
            workbook.close()
            return False
        except Exception:
            return False

    def _upgrade_active_review_for_decoration_notes(self):
        """Rebuild an older active review so no customer instruction can be missed."""
        if self._regenerating_review or not self._active_review_needs_decoration_note_upgrade():
            return
        self.regenerate_current_review(silent=True)
        try:
            self.dashboard_status.configure(
                text="Purchase Review upgraded for expanded decoration locations, faster review navigation, and success sound controls.",
                text_color=SUCCESS,
            )
        except Exception:
            pass

    def _prompt_for_csv_if_empty(self):
        """Open the branded Current Event start page when no packet exists."""
        if self._suppress_empty_import_prompt:
            self._suppress_empty_import_prompt = False
            return
        if self.selected_csv or self.last_review_workbook or self.last_purchase_order_dir:
            return
        self.show_page("import")

    def _prompt_resume_or_start_new(self):
        """Explain unfinished saved work instead of silently landing mid-workflow."""
        if self._launch_choice_shown:
            return
        has_active_packet = bool(
            self.selected_csv
            or (self.last_review_workbook and self.last_review_workbook.exists())
        )
        if not has_active_packet:
            return
        self._launch_choice_shown = True

        # Create the window hidden, then present it after its final geometry is
        # known.  This prevents macOS from briefly drawing it at the default
        # top-left position before it is centered over Orchid.
        dialog = ctk.CTkToplevel(self)
        dialog.withdraw()
        dialog.title("Unfinished Purchase Packet")
        dialog.transient(self)
        dialog.resizable(False, False)
        dialog.configure(fg_color=BG)
        shell = ctk.CTkFrame(dialog, fg_color=WHITE, border_width=1, border_color="#D8CDE6", corner_radius=18)
        shell.pack(fill="both", expand=True, padx=18, pady=18)
        ctk.CTkLabel(
            shell, text="An unfinished purchase packet was found", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=24, weight="bold"), justify="center", wraplength=500,
        ).pack(padx=28, pady=(26, 8))
        packet_name = self.current_event_name or (
            self.last_review_workbook.stem if self.last_review_workbook else "the saved event"
        )
        ctk.CTkLabel(
            shell,
            text=f"Would you like to resume {packet_name}, or clear the active packet and begin a new event?\n\nProduct Master and already-created PDFs will not be deleted.",
            text_color=MUTED, font=ctk.CTkFont(size=14), justify="center", wraplength=520,
        ).pack(padx=30, pady=(0, 22))
        actions = ctk.CTkFrame(shell, fg_color="transparent")
        actions.pack(padx=24, pady=(0, 26))

        def resume():
            dialog.grab_release()
            dialog.destroy()
            # Resume the actual unfinished work. The review loader repairs older
            # saved rows and opens the first genuinely unresolved decision.
            if self.last_review_workbook and self.last_review_workbook.exists():
                self.show_page("review")
            else:
                self.show_page("import")

        def start_new():
            dialog.grab_release()
            dialog.destroy()
            self._suppress_empty_import_prompt = True
            self._clear_active_packet_state()
            self.refresh_current_event_page({})
            self.refresh_dashboard()
            self.show_page("import")
            self.after_idle(lambda: (self.refresh_current_event_page({}), self.refresh_dashboard()))

        ctk.CTkButton(
            actions, text="Resume Current Event", command=resume,
            width=210, height=48, fg_color=PURPLE, hover_color=PURPLE_DARK,
            font=ctk.CTkFont(size=14, weight="bold"),
        ).pack(side="left", padx=7)
        ctk.CTkButton(
            actions, text="Start New Event", command=start_new,
            width=190, height=48, fg_color=PURPLE_LIGHT, hover_color="#E6D9F5",
            border_width=1, border_color="#BFA9DD", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=14, weight="bold"),
        ).pack(side="left", padx=7)
        dialog.protocol("WM_DELETE_WINDOW", resume)
        dialog.update_idletasks()
        width, height = 610, 310
        x = max(self.winfo_rootx() + (self.winfo_width() - width) // 2, 0)
        y = max(self.winfo_rooty() + (self.winfo_height() - height) // 2, 0)
        dialog.geometry(f"{width}x{height}+{x}+{y}")

        def present_centered_dialog():
            if not dialog.winfo_exists():
                return
            dialog.deiconify()
            dialog.lift()
            dialog.focus_force()
            dialog.grab_set()

        dialog.after_idle(present_centered_dialog)

    @staticmethod
    def _friendly_modified(path: Path | None) -> str:
        if not path or not path.exists():
            return ""
        from datetime import datetime
        return datetime.fromtimestamp(path.stat().st_mtime).strftime("%B %-d, %Y at %-I:%M %p")

    @staticmethod
    def _source_file_signature(path: Path | None) -> str:
        if not path:
            return ""
        candidate = Path(path)
        if not candidate.is_file():
            return ""
        digest = hashlib.sha256()
        with candidate.open("rb") as handle:
            for chunk in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(chunk)
        return digest.hexdigest()

    def _remember_import_frames(self, path: Path, normalized, parsed) -> None:
        self._import_cache_key = (_file_stamp(path), _file_stamp(live_product_master_path()))
        self._normalized_import_cache = normalized.copy(deep=True)
        self._parsed_import_cache = parsed.copy(deep=True)

    def _cached_import_frames(self, path: Path):
        key = (_file_stamp(path), _file_stamp(live_product_master_path()))
        if key != self._import_cache_key:
            return None, None
        normalized = (
            self._normalized_import_cache.copy(deep=True)
            if self._normalized_import_cache is not None else None
        )
        parsed = (
            self._parsed_import_cache.copy(deep=True)
            if self._parsed_import_cache is not None else None
        )
        return normalized, parsed

    def _load_review_state(self) -> dict:
        try:
            if REVIEW_STATE_FILE.exists():
                payload = json.loads(REVIEW_STATE_FILE.read_text(encoding="utf-8"))
                return payload if isinstance(payload, dict) else {}
        except Exception:
            pass
        return {}

    def _update_review_state(self, **updates) -> dict:
        state = self._load_review_state()
        state.update(updates)
        try:
            REVIEW_STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
            REVIEW_STATE_FILE.write_text(json.dumps(state, indent=2), encoding="utf-8")
        except Exception:
            pass
        return state

    def _bind_active_review_workbook(self, workbook_path: Path) -> None:
        """Persist the one workbook allowed to represent the active CSV packet."""
        state = self._load_review_state()
        bound = bind_active_workbook(state, Path(workbook_path))
        self._write_review_state_atomic(bound)

    def _active_packet_is_valid(self, workbook_path: Path | None = None) -> tuple[bool, str]:
        """Reject a stale review before it can feed any active-event screen."""
        state = self._load_review_state()
        return active_workbook_matches_state(state, workbook_path or self.last_review_workbook)

    def _require_active_packet(self, action: str) -> bool:
        """Show one clear stop instead of silently falling back to an old event."""
        valid, reason = self._active_packet_is_valid()
        if valid:
            return True
        messagebox.showerror(
            "Active Event Verification Required",
            f"Orchid blocked {action} because the selected Purchase Review is not proven to belong to the active CSV.\n\n"
            f"{reason}\n\n"
            "No old event data was used. Return to Current Event and process the selected CSV again.",
        )
        return False

    @staticmethod
    def _review_issue_signature(issue: dict | None) -> str:
        """Stable signature used to distinguish an unchanged decision from a new one."""
        issue = issue or {}
        payload = "|".join(
            _clean(issue.get(field, "")).casefold()
            for field in ("order", "employee", "product", "description", "reason", "instructions")
        )
        return hashlib.sha256(payload.encode("utf-8")).hexdigest()[:24]

    @staticmethod
    def _is_journal_validation_shadow(issue: dict | None) -> bool:
        """Return True only for a legacy validation echo with no immutable identity.

        A source-identified issue is never treated as a harmless shadow. A new or
        changed customer instruction on the same order line must return to the
        Purchase Review queue unless its exact issue signature was completed.
        """
        issue = issue or {}
        if _clean(issue.get("source_id", "")) or _clean(issue.get("line_id", "")):
            return False
        source = _clean(issue.get("source", issue.get("fix_in", ""))).casefold()
        reason = _clean(issue.get("reason", "")).casefold()
        return bool(
            "validation" in source
            and (
                "customer decision required" in reason
                or "customer instruction decision required" in reason
            )
        )

    @staticmethod
    def _review_queue_identity(issue: dict | None) -> str:
        """Return the immutable identity for one required Purchase Review decision."""
        issue = issue or {}
        source_id = _clean(issue.get("source_id", ""))
        decision_key = _clean(issue.get("decision_key", ""))
        line_id = _clean(issue.get("line_id", ""))
        reason = _clean(issue.get("reason", "")).casefold()
        instructions = _clean(issue.get("instructions", "")).casefold()
        identity = source_id or decision_key or line_id or "|".join(
            _clean(issue.get(field, "")).casefold()
            for field in ("order", "employee", "product", "description")
        )
        return hashlib.sha256(
            f"{identity}|{reason}|{instructions}".encode("utf-8")
        ).hexdigest()[:32]

    def _build_purchase_review_queue(self, snapshot: dict | None = None) -> list[dict]:
        """Build the one authoritative, deduplicated Purchase Review queue.

        The review page, decision counter, Next navigation, and final PO preflight
        all consume this exact queue. No event-level completion marker can hide a
        source-identified or newly changed instruction decision.
        """
        snapshot = snapshot or {}
        queue_items: list[dict] = []
        seen: set[str] = set()
        for issue in snapshot.get("issues", []):
            if _clean(issue.get("fix_in", issue.get("source", ""))).casefold() == "product master":
                continue
            key = self._review_queue_identity(issue)
            if key in seen:
                continue
            seen.add(key)
            queue_items.append(dict(issue))
        return queue_items

    def _completed_review_records_from_state(self, workbook_path: Path | None = None) -> list[dict]:
        """Return completion markers for the active packet, even after regeneration.

        current_review_state.json is cleared whenever a different packet starts,
        so its registry already belongs to one active event. Older builds filtered
        records by the exact timestamped workbook path; regenerating or upgrading
        the event therefore hid valid completions. Keep all active-packet records
        and let stable decision/line identities perform the match.
        """
        state = self._load_review_state()
        records = state.get("completed_review_decisions", [])
        if not isinstance(records, list):
            return []
        return [dict(record) for record in records if isinstance(record, dict)]

    def _review_state_is_authoritatively_complete(self, workbook_path: Path | None = None) -> bool:
        """Return True when the active event reached zero decisions in Orchid.

        The active review state is cleared whenever a different purchase packet
        starts. Once the review screen records zero remaining decisions, that
        durable event-level marker must be the workflow source of truth even if
        the older Excel workbook still contains stale validation text.
        """
        state = self._load_review_state()
        if not bool(state.get("review_complete")):
            return False
        if not (state.get("review_completed_at") or state.get("completed_review_decisions")):
            return False
        saved_signature = _clean(state.get("source_signature", ""))
        if saved_signature and self.selected_csv and Path(self.selected_csv).exists():
            try:
                current_signature = self._source_file_signature(Path(self.selected_csv))
            except Exception:
                current_signature = ""
            if current_signature and current_signature != saved_signature:
                return False
        return True

    def _overlay_completed_review_state(self, snapshot: dict, workbook_path: Path) -> dict:
        """Overlay durable app decisions and honor an event-level completion marker.

        Individual journal records normally remove matching decisions. When the
        active event has already reached zero decisions, all remaining purchase-
        review-only issues are stale workbook shadows; Product Master issues are
        still preserved and must be corrected normally.
        """
        snapshot = dict(snapshot or {})
        issues = list(snapshot.get("issues", []))
        records = self._completed_review_records_from_state(workbook_path)
        authoritative_complete = self._review_state_is_authoritatively_complete(workbook_path)

        by_key: dict[str, dict] = {}
        by_line: dict[str, dict] = {}
        by_source: dict[str, dict] = {}
        for record in records:
            decision_key = _clean(record.get("decision_key", ""))
            line_id = _clean(record.get("line_id", ""))
            source_id = _clean(record.get("source_id", ""))
            if decision_key:
                by_key[decision_key] = record
            if line_id:
                by_line[line_id] = record
            if source_id:
                by_source[source_id] = record

        remaining = []
        removed = 0
        review_values = dict(snapshot.get("review_values", {}))
        for issue in issues:
            if _clean(issue.get("fix_in", issue.get("source", ""))).casefold() == "product master":
                remaining.append(issue)
                continue
            decision_key = _clean(issue.get("decision_key", ""))
            line_id = _clean(issue.get("line_id", ""))
            source_id = _clean(issue.get("source_id", ""))
            record = by_source.get(source_id) if source_id else (by_key.get(decision_key) or by_line.get(line_id))
            if record:
                stored_signature = _clean(record.get("issue_signature", ""))
                if stored_signature != self._review_issue_signature(issue):
                    # The same source line can acquire a new waterproof, logo,
                    # placement, or other customer instruction. Only the exact
                    # completed issue signature may remove it from the queue.
                    remaining.append(issue)
                    continue
                values = dict(record.get("values", {})) if isinstance(record.get("values"), dict) else {}
                values["Review Status"] = "Ready"
                if decision_key:
                    review_values[decision_key] = values
                if line_id:
                    review_values[line_id] = values
                if source_id:
                    review_values[source_id] = values
                removed += 1
                continue
            if authoritative_complete and not source_id and self._is_journal_validation_shadow(issue):
                # Legacy packets without Source IDs may retain a known validation
                # echo.  Never use the event-level marker to hide an unmatched
                # source-identified decision.
                removed += 1
                continue
            remaining.append(issue)

        if not removed and not authoritative_complete:
            return snapshot

        snapshot["issues"] = remaining
        purchase_issues = [
            issue for issue in remaining
            if _clean(issue.get("fix_in", issue.get("source", ""))).casefold() != "product master"
        ]
        product_issues = len(remaining) - len(purchase_issues)
        snapshot["review_values"] = review_values
        snapshot["review_count"] = len(purchase_issues)
        snapshot["purchase_review_count"] = len(purchase_issues)
        snapshot["product_master_review_count"] = product_issues
        snapshot["review_completed"] = max(
            int(snapshot.get("review_completed", 0)) + removed,
            len(records),
        )
        snapshot["review_remaining"] = len(purchase_issues)
        if authoritative_complete:
            snapshot["journal_authoritative_complete"] = True

        if not remaining:
            routes = []
            for route in snapshot.get("routes", []):
                updated_route = dict(route)
                updated_route["status"] = "Ready"
                routes.append(updated_route)
            snapshot["routes"] = routes
            snapshot["blocked_route_count"] = 0
            snapshot["ready_count"] = int(snapshot.get("line_count", 0) or 0)
        snapshot["workflow_blocked"] = bool(remaining or snapshot.get("blocked_route_count", 0))
        snapshot["resumed_from_journal"] = removed
        return snapshot

    def _clear_saved_review_draft(self, issue: dict | None = None) -> None:
        state = self._load_review_state()
        if not state:
            return
        if issue:
            saved_line = _clean(state.get("draft_line_id", ""))
            saved_key = _clean(state.get("draft_decision_key", ""))
            issue_line = _clean(issue.get("line_id", ""))
            issue_key = _clean(issue.get("decision_key", ""))
            if saved_line and issue_line and saved_line != issue_line and (not saved_key or saved_key != issue_key):
                return
        changed = False
        for key in ("draft_line_id", "draft_decision_key", "draft_saved_at", "draft_values"):
            if key in state:
                state.pop(key, None)
                changed = True
        if changed:
            try:
                REVIEW_STATE_FILE.write_text(json.dumps(state, indent=2), encoding="utf-8")
            except Exception:
                pass

    def _saved_review_draft_issue(self, issues: list[dict]) -> dict | None:
        state = self._load_review_state()
        saved_line = _clean(state.get("draft_line_id", ""))
        saved_key = _clean(state.get("draft_decision_key", ""))
        if not saved_line and not saved_key:
            return None
        for issue in issues:
            line_id = _clean(issue.get("line_id", ""))
            decision_key = _clean(issue.get("decision_key", ""))
            if (saved_line and line_id == saved_line) or (saved_key and decision_key == saved_key):
                return issue
        self._clear_saved_review_draft()
        return None

    # ---------- shell ----------
    def build_ui(self):
        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(0, weight=0)
        self.grid_rowconfigure(1, weight=1)
        self.build_sidebar()
        self.build_main_shell()

    @staticmethod
    def _purple_orchid_art(source: PILImage.Image, size: tuple[int, int]) -> PILImage.Image:
        """Return the Orchid flower in a lavender-to-purple treatment with no green."""
        flower = source.convert("RGBA")
        bbox = flower.getchannel("A").getbbox()
        if bbox:
            flower = flower.crop(bbox)
        flower.thumbnail(size, PILImage.Resampling.LANCZOS)
        alpha = flower.getchannel("A")
        detail = flower.convert("L")
        tinted = PILImage.new("RGBA", flower.size, (0, 0, 0, 0))
        height = max(1, flower.height - 1)
        top_color = (202, 126, 247)
        bottom_color = (132, 52, 213)
        pixels = []
        for y in range(flower.height):
            ratio = y / height
            base_color = tuple(
                round(top_color[channel] * (1.0 - ratio) + bottom_color[channel] * ratio)
                for channel in range(3)
            )
            for x in range(flower.width):
                pixel_alpha = alpha.getpixel((x, y))
                source_luma = detail.getpixel((x, y))
                # Keep the original petal and center linework as darker purple
                # while ensuring the entire mark stays in the purple family.
                shade = 0.66 + 0.34 * min(1.0, source_luma / 92.0)
                color = tuple(round(channel * shade) for channel in base_color)
                pixels.append((*color, pixel_alpha))
        tinted.putdata(pixels)
        return tinted

    @staticmethod
    def _sidebar_spotlight_background(width: int, height: int) -> PILImage.Image:
        """Continue the header color family behind Current Event, then fade to navigation."""
        width = max(1, int(width))
        height = max(1, int(height))
        sample_width = min(160, width)
        sample_height = min(132, height)

        base = PILImage.new("RGBA", (sample_width, sample_height), (0, 0, 0, 255))
        base_draw = ImageDraw.Draw(base)
        top = (17, 8, 40)
        bottom = (16, 12, 37)
        for y in range(sample_height):
            ratio = y / max(1, sample_height - 1)
            color = tuple(
                round(top[channel] * (1.0 - ratio) + bottom[channel] * ratio)
                for channel in range(3)
            )
            base_draw.line((0, y, sample_width, y), fill=(*color, 255))

        mask = PILImage.new("L", (sample_width, sample_height), 0)
        mask_pixels = []
        for y in range(sample_height):
            y_ratio = y / max(1, sample_height - 1)
            normalized_y = (y_ratio - 0.26) / 0.42
            vertical_fade = pow(max(0.0, 1.0 - y_ratio), 1.65)
            for x in range(sample_width):
                normalized_x = (x / max(1, sample_width - 1) - 0.50) / 0.48
                distance = normalized_x * normalized_x + normalized_y * normalized_y
                alpha = 132 * pow(2.718281828, -distance * 1.55) * vertical_fade
                mask_pixels.append(round(alpha))
        mask.putdata(mask_pixels)

        spotlight = PILImage.new("RGBA", base.size, (62, 18, 112, 255))
        background = PILImage.composite(spotlight, base, mask)
        return background.resize((width, height), PILImage.Resampling.LANCZOS)

    def build_sidebar(self):
        sidebar = ctk.CTkFrame(self, width=238, corner_radius=0, fg_color="#100C25")
        sidebar.grid(row=1, column=0, sticky="nsew")
        sidebar.grid_propagate(False)
        sidebar.grid_rowconfigure(2, weight=1)

        brand = ctk.CTkFrame(sidebar, fg_color="transparent", corner_radius=0, height=138)
        brand.grid(row=0, column=0, sticky="ew")
        brand.grid_propagate(False)
        brand.grid_columnconfigure(0, weight=1)
        self.sidebar_event_card = brand
        sidebar_spotlight = self._sidebar_spotlight_background(476, 276)
        self.sidebar_spotlight_image = ctk.CTkImage(
            light_image=sidebar_spotlight, dark_image=sidebar_spotlight, size=(238, 138)
        )
        self.sidebar_spotlight_label = ctk.CTkLabel(
            brand, text="", image=self.sidebar_spotlight_image,
            width=238, height=138, fg_color="transparent",
        )
        self.sidebar_spotlight_label.place(x=0, y=0)
        # Compatibility handle retained for state-refresh code; the sidebar is
        # intentionally text-only so the header flower remains the sole brand mark.
        self.sidebar_event_icon = ctk.CTkLabel(
            brand, text="", width=1, height=1, fg_color="transparent",
        )
        # Compatibility label retained for older refresh/state code, but the
        # duplicate packet title is intentionally not shown in the sidebar.
        self.sidebar_event_title = ctk.CTkLabel(
            brand, text="", width=1, height=1, fg_color="transparent",
        )
        self.sidebar_event_title.grid_remove()
        self.sidebar_event_meta = ctk.CTkLabel(
            brand, text=time.strftime("%b %d, %Y"), text_color="#E2D8F0",
            font=ctk.CTkFont(size=16, weight="bold"), justify="center",
        )
        self.sidebar_event_meta.grid(row=0, column=0, pady=(60, 0))

        nav = ctk.CTkFrame(sidebar, fg_color="transparent")
        # Current Event is the primary workspace. Lower the navigation slightly
        # so the sidebar feels balanced beneath the event summary.
        nav.grid(row=1, column=0, rowspan=2, sticky="nsew", padx=12, pady=(34, 0))
        nav.grid_columnconfigure(0, weight=1)
        self.sidebar_page_title = ctk.CTkLabel(nav, text="", width=1, height=1)
        self.sidebar_page_title.grid_remove()

        items = [
            ("import", "▣   Current Event"),
            ("master", "◇   Product Master"),
            ("review", "☑   Purchase Review"),
            ("audit", "◐   Decoration Colors"),
            ("purchase", "▤   Purchase Orders"),
            ("employees", "▦   Employee Totals"),
            ("job_logo", "◉   Logo"),
        ]
        for row, (page, label) in enumerate(items):
            button = ctk.CTkButton(
                nav, text=label, command=lambda p=page: self.show_page(p), anchor="w",
                height=42, corner_radius=11, fg_color="transparent", hover_color="#2B214A",
                text_color=WHITE, font=ctk.CTkFont(size=14),
            )
            button.grid(row=row, column=0, pady=3, sticky="ew")
            self.nav_buttons[page] = button
            if page == "employees":
                self.employee_nav_button = button
                button.grid_remove()
            if page == "job_logo":
                self.outsourced_logo_nav_button = button

        divider = ctk.CTkFrame(nav, height=1, fg_color="#40375A", corner_radius=0)
        divider.grid(row=7, column=0, sticky="ew", padx=10, pady=(16, 12))
        for row, page, label in [
            (8, "settings", "⚙   Settings & Tools"),
        ]:
            button = ctk.CTkButton(
                nav, text=label, command=lambda p=page: self.show_page(p), anchor="w",
                height=42, corner_radius=11, fg_color="transparent", hover_color="#2B214A",
                text_color=WHITE, font=ctk.CTkFont(size=14),
            )
            button.grid(row=row, column=0, pady=3, sticky="ew")
            self.nav_buttons[page] = button

        footer = ctk.CTkFrame(sidebar, fg_color="transparent")
        footer.grid(row=4, column=0, sticky="ew", padx=14, pady=(8, 14))
        ctk.CTkLabel(
            footer, text="v 4.9.104", text_color="#CFC4E0",
            font=ctk.CTkFont(size=12), justify="left", anchor="w",
        ).pack(anchor="w", padx=8, pady=(0, 10))
        ctk.CTkFrame(footer, height=1, fg_color="#40375A").pack(fill="x", padx=7, pady=(0, 11))
        ctk.CTkButton(
            footer, text="⇥  Exit", height=46, corner_radius=11,
            fg_color="transparent", hover_color="#2B214A", border_width=1,
            border_color="#5A4D78", text_color=WHITE, command=self.destroy,
            font=ctk.CTkFont(size=15, weight="bold"),
        ).pack(fill="x")

    def build_main_shell(self):
        main = ctk.CTkFrame(self, fg_color=BG, corner_radius=0)
        main.grid(row=1, column=1, sticky="nsew")
        main.grid_columnconfigure(0, weight=1)
        main.grid_rowconfigure(0, weight=1)

        header = ctk.CTkFrame(self, fg_color=BG, corner_radius=0, height=126)
        self.main_header = header
        header.grid(row=0, column=0, columnspan=2, sticky="ew")
        header.grid_propagate(False)
        header.grid_columnconfigure(0, weight=1)

        wave_path = resource_path("assets", "purple_wave_header.png")
        self._wave_render_job = None
        self._wave_render_size = (0, 0)
        self._wave_source = None
        if wave_path.exists():
            # Keep the full-resolution source and render a Retina-sized copy for
            # the exact window width. This prevents blur and avoids cropping the
            # complete wave artwork on larger displays.
            self._wave_source = PILImage.open(wave_path).convert("RGBA")
            self.wave_header_label = ctk.CTkLabel(
                header, text="", fg_color="transparent", cursor="hand2"
            )
            self.wave_header_label.place(x=0, y=0, relwidth=1, relheight=1)
            self.wave_header_label.bind("<Button-1>", lambda _event: self.show_page("import"), add="+")
            header.bind("<Configure>", self._schedule_wave_header_resize, add="+")
            self.after(40, self._render_current_wave_header)
        else:
            ctk.CTkFrame(header, fg_color="#5A18C9").place(relx=0, rely=0, relwidth=1, relheight=1)

        self.shell_event_card = self.sidebar_event_card
        self.shell_event_title = self.sidebar_event_title
        self.shell_event_meta = self.sidebar_event_meta
        icon = self.sidebar_event_icon
        event = self.sidebar_event_card

        # Existing refresh code uses these names; all now point to the single shell card.
        self.dashboard_summary = event
        self.current_event_summary = event
        self.dashboard_summary_icon = icon
        self.current_event_summary_icon = icon
        self.dashboard_summary_title = self.shell_event_title
        self.current_event_summary_title = self.shell_event_title
        self.dashboard_date_label = self.shell_event_meta
        self.current_event_date_label = self.shell_event_meta

        # Compatibility labels retained for the existing state-refresh code.
        self.page_title = ctk.CTkLabel(header, text="", width=1, height=1)
        self.page_title.grid_remove()
        self.page_subtitle = ctk.CTkLabel(header, text="", width=1, height=1)
        self.page_subtitle.grid_remove()
        self.header_chip = ctk.CTkLabel(header, text="", width=1, height=1)
        self.header_chip.grid_remove()
        self.header_date = ctk.CTkLabel(header, text="", width=1, height=1)
        self.header_date.grid_remove()

        self.page_container = ctk.CTkFrame(main, fg_color=BG, corner_radius=0)
        self.page_container.grid(row=0, column=0, sticky="nsew")
        self.page_container.grid_rowconfigure(0, weight=1)
        self.page_container.grid_columnconfigure(0, weight=1)

        self.build_dashboard_page()
        self._page_builders = {
            "import": self.build_import_page,
            "master": self.build_master_page,
            "review": self.build_review_page,
            "audit": self.build_decoration_audit_page,
            "purchase": self.build_purchase_page,
            "employees": self.build_employee_totals_page,
            "job_logo": self.build_outsourced_job_logo_page,
            "settings": self.build_settings_page,
        }

    def _ensure_page_built(self, page: str) -> None:
        if page in self.pages:
            return
        builder = getattr(self, "_page_builders", {}).get(page)
        if builder is None:
            return
        builder()
        if page in {"import", "master", "review", "audit", "purchase", "employees"}:
            self._dashboard_refresh_key = None
        scrollable_names = {
            "review": ("review_decision_frame",),
            "audit": ("decoration_audit_scroll",),
            "purchase": ("purchase_documents_frame",),
            "employees": ("mc_employee_scroll",),
            "settings": ("settings_scroll",),
        }
        for name in scrollable_names.get(page, ()):
            register_scrollable(self, getattr(self, name, None))

    def _schedule_wave_header_resize(self, event=None):
        """Debounce high-resolution header rendering while the window resizes."""
        if not getattr(self, "_wave_source", None):
            return
        width = int(getattr(event, "width", 0) or self.main_header.winfo_width())
        height = int(getattr(event, "height", 0) or self.main_header.winfo_height())
        if width < 80 or height < 40:
            return
        if self._wave_render_job is not None:
            try:
                self.after_cancel(self._wave_render_job)
            except Exception:
                pass
        self._wave_render_job = self.after(
            70, lambda w=width, h=height: self._render_wave_header(w, h)
        )

    def _render_current_wave_header(self):
        if not getattr(self, "_wave_source", None):
            return
        width = max(80, int(self.main_header.winfo_width() or 1600))
        height = max(40, int(self.main_header.winfo_height() or 126))
        self._render_wave_header(width, height)

    @staticmethod
    def _spotlight_header_background(width: int, height: int) -> PILImage.Image:
        """Create the approved dark aubergine header with a soft central spotlight."""
        width = max(1, int(width))
        height = max(1, int(height))

        # Build the effect at a small resolution and upscale it smoothly. This
        # keeps window resizing fast while avoiding visible gradient banding.
        sample_width = min(320, width)
        sample_height = min(48, height)
        base = PILImage.new("RGBA", (sample_width, sample_height), (0, 0, 0, 255))
        base_draw = ImageDraw.Draw(base)
        top = (13, 6, 31)
        bottom = (19, 8, 43)
        for y in range(sample_height):
            ratio = y / max(1, sample_height - 1)
            color = tuple(
                round(top[channel] * (1.0 - ratio) + bottom[channel] * ratio)
                for channel in range(3)
            )
            base_draw.line((0, y, sample_width, y), fill=(*color, 255))

        mask = PILImage.new("L", (sample_width, sample_height), 0)
        mask_pixels = []
        for y in range(sample_height):
            normalized_y = (y / max(1, sample_height - 1) - 0.43) / 0.78
            for x in range(sample_width):
                normalized_x = (x / max(1, sample_width - 1) - 0.50) / 0.36
                distance = normalized_x * normalized_x + normalized_y * normalized_y
                mask_pixels.append(round(150 * pow(2.718281828, -distance * 1.65)))
        mask.putdata(mask_pixels)

        spotlight = PILImage.new("RGBA", base.size, (66, 20, 143, 255))
        background = PILImage.composite(spotlight, base, mask)
        return background.resize((width, height), PILImage.Resampling.LANCZOS)

    def _render_wave_header(self, width: int, height: int):
        """Render the approved spotlight header at 2x display resolution."""
        self._wave_render_job = None
        width = max(80, int(width))
        height = max(40, int(height))
        if self._wave_render_size == (width, height):
            return
        try:
            retina_size = (width * 2, height * 2)
            rendered = self._spotlight_header_background(*retina_size)
            pulse_scales = (0.95, 0.958, 0.966, 0.974)
            self.wave_header_pulse_frames = []
            for scale in pulse_scales:
                branded = self._composite_header_branding(rendered.copy(), scale=scale)
                self.wave_header_pulse_frames.append(ctk.CTkImage(
                    light_image=branded, dark_image=branded, size=(width, height)
                ))
            self.wave_header_image = self.wave_header_pulse_frames[0]
            self.wave_header_label.configure(image=self.wave_header_image)
            # Reuse the established interaction-animation controller with the
            # fully composited header frames. No transparent overlay is needed.
            self.sidebar_brand_logo_label = self.wave_header_label
            self.sidebar_brand_flower_frames = self.wave_header_pulse_frames
            self._wave_render_size = (width, height)
        except Exception:
            # The solid purple header remains visible if image rendering fails.
            pass

    @staticmethod
    def _tracked_text_width(draw: ImageDraw.ImageDraw, text: str, font, tracking: int) -> int:
        widths = [draw.textlength(character, font=font) for character in text]
        return round(sum(widths) + tracking * max(len(text) - 1, 0))

    def _composite_header_branding(self, canvas: PILImage.Image, scale: float = 1.0) -> PILImage.Image:
        """Render the approved flower / ORCHID | PURCHASE MANAGER header lockup."""
        logo_path = resource_path("assets", "orchid_logo.png")
        flower_path = resource_path("assets", "orchid_flower_purple.png")
        if not logo_path.exists() or not flower_path.exists():
            return canvas
        source = PILImage.open(logo_path).convert("RGBA")

        # Isolate the thin ORCHID lettering without the retail tagline or flower.
        letters = source.crop((0, 90, source.width, 325))
        bbox = letters.getchannel("A").getbbox()
        if bbox:
            letters = letters.crop(bbox)
        pixels = letters.load()
        for y in range(letters.height):
            for x in range(letters.width):
                red, green, blue, alpha = pixels[x, y]
                if alpha:
                    pixels[x, y] = (248, 246, 252, alpha)

        # Use a compact, left-aligned brand lockup so the header supports the
        # workspace instead of visually dominating it.
        branding_height = min(canvas.height, 256)
        target_width = max(1, round(canvas.width * 0.13 * scale))
        target_height = max(1, round(branding_height * 0.23 * scale))
        letters = letters.resize((target_width, target_height), PILImage.Resampling.LANCZOS)
        draw = ImageDraw.Draw(canvas)
        font_size = max(12, round(branding_height * 0.11 * scale))
        font_candidates = (
            "/System/Library/Fonts/Supplemental/Arial Bold.ttf",
            "/System/Library/Fonts/Helvetica.ttc",
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        )
        font = None
        for candidate in font_candidates:
            try:
                font = ImageFont.truetype(candidate, font_size)
                break
            except Exception:
                continue
        if font is None:
            font = ImageFont.load_default()
        subtitle = "PURCHASE MANAGER"
        tracking = max(1, round(font_size * 0.24))
        subtitle_width = self._tracked_text_width(draw, subtitle, font, tracking)

        flower_source = PILImage.open(flower_path).convert("RGBA")
        flower_limit = max(1, round(branding_height * 0.33 * scale))
        flower = self._purple_orchid_art(flower_source, (flower_limit, flower_limit))

        flower_gap = round(canvas.width * 0.009 * scale)
        divider_gap_left = round(canvas.width * 0.012 * scale)
        divider_gap_right = round(canvas.width * 0.014 * scale)
        divider_width = max(2, round(canvas.width * 0.0008))
        group_width = (
            flower.width + flower_gap + target_width + divider_gap_left
            + divider_width + divider_gap_right + subtitle_width
        )
        # Anchor the compact lockup to the left of the workspace. Measurements
        # are expressed in display pixels and scaled for the Retina render.
        display_scale = canvas.height / 148
        left = round(18 * display_scale)
        center_y = round(canvas.height * 0.50)

        # Subtle spotlight: bright enough to make the flower/ORCHID lockup feel
        # special, but restrained so it does not wash out the header artwork.
        glow_layer = PILImage.new("RGBA", canvas.size, (0, 0, 0, 0))
        glow_draw = ImageDraw.Draw(glow_layer)
        glow_left = left - round(18 * display_scale)
        glow_right = left + flower.width + flower_gap + target_width + round(18 * display_scale)
        glow_top = center_y - round(40 * display_scale)
        glow_bottom = center_y + round(40 * display_scale)
        glow_draw.ellipse(
            (glow_left, glow_top, glow_right, glow_bottom),
            fill=(158, 78, 250, 104),
        )
        glow_layer = glow_layer.filter(ImageFilter.GaussianBlur(max(8, round(18 * display_scale))))
        canvas.alpha_composite(glow_layer)

        flower_y = center_y - flower.height // 2
        canvas.alpha_composite(flower, (left, flower_y))
        letters_left = left + flower.width + flower_gap
        letters_top = center_y - target_height // 2
        canvas.alpha_composite(letters, (letters_left, letters_top))

        divider_x = letters_left + target_width + divider_gap_left
        divider_height = max(1, round(branding_height * 0.46 * scale))
        divider_top = center_y - divider_height // 2
        draw.rounded_rectangle(
            (divider_x, divider_top, divider_x + divider_width, divider_top + divider_height),
            radius=max(1, divider_width // 2), fill=(220, 201, 244, 218),
        )

        x = divider_x + divider_width + divider_gap_right
        text_box = draw.textbbox((0, 0), subtitle, font=font)
        text_height = text_box[3] - text_box[1]
        y = center_y - text_height // 2 - text_box[1]
        for character in subtitle:
            draw.text((x, y), character, font=font, fill=(248, 246, 252, 255))
            x += round(draw.textlength(character, font=font)) + tracking
        return canvas

    def show_page(self, page: str):
        self._ensure_page_built(page)
        if page not in self.pages:
            page = "import"
        self.active_page = page
        for frame in self.pages.values():
            frame.grid_remove()
        self.pages[page].grid(row=0, column=0, sticky="nsew")
        for name, button in self.nav_buttons.items():
            if name == "import":
                button.grid()
            active = name == page
            button.configure(
                fg_color="#6A24D4" if active else "transparent",
                font=ctk.CTkFont(size=15, weight="bold" if active else "normal"),
            )
        active_event = self.current_event_name or "No Active Purchase Packet"
        self.header_chip.configure(text=active_event if len(active_event) <= 34 else active_event[:31] + "...")
        self.main_header.grid()

        # Purchase Review is intentionally decoupled from the full dashboard
        # refresh. Show the page first, then load its small review-only snapshot
        # on a worker thread. This prevents large events from beachballing while
        # Mission Control scans All PO Lines, totals, and vendor routes.
        if page == "review":
            self._load_review_page_async()
            return

        self.refresh_dashboard()
        if page == "job_logo":
            self.refresh_outsourced_job_logo_page()
        elif page == "audit":
            self.refresh_decoration_audit_page()

    # ---------- reusable UI ----------

    # ---------- reusable UI ----------
    def new_page(self, name: str) -> ctk.CTkFrame:
        page = ctk.CTkFrame(self.page_container, fg_color=BG, corner_radius=0)
        page.grid_columnconfigure(0, weight=1)
        page.grid_rowconfigure(9, weight=1)
        self.pages[name] = page
        return page

    def stat_card(self, parent, row, column, title, value, subtitle, columns=4):
        parent.grid_columnconfigure(tuple(range(columns)), weight=1)
        shadow = ctk.CTkFrame(parent, fg_color=SHADOW, corner_radius=15, width=1, height=1)
        shadow.grid(row=row, column=column, sticky="nsew", padx=6, pady=(9, 4))
        frame = ctk.CTkFrame(parent, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=15)
        frame.grid(row=row, column=column, sticky="nsew", padx=6, pady=(5, 8))
        value_label = ctk.CTkLabel(frame, text=value, text_color=PURPLE,
                                   font=ctk.CTkFont(size=24, weight="bold"), anchor="w")
        value_label.pack(anchor="w", padx=18, pady=(16, 2))
        title_label = ctk.CTkLabel(frame, text=title, text_color=TEXT,
                                   font=ctk.CTkFont(size=13, weight="bold"), anchor="w")
        title_label.pack(anchor="w", padx=18)
        subtitle_label = ctk.CTkLabel(frame, text=subtitle, text_color=MUTED,
                                      font=ctk.CTkFont(size=11), anchor="w", wraplength=230, justify="left")
        subtitle_label.pack(anchor="w", padx=18, pady=(3, 15))
        return value_label, title_label, subtitle_label

    def section_card(self, parent, row, title, body):
        shadow = ctk.CTkFrame(parent, fg_color=SHADOW, corner_radius=16, width=1, height=1)
        shadow.grid(row=row, column=0, sticky="ew", padx=30, pady=(12, 4))
        frame = ctk.CTkFrame(parent, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=16)
        frame.grid(row=row, column=0, sticky="ew", padx=30, pady=(8, 8))
        frame.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(frame, text=title, text_color=TEXT,
                     font=ctk.CTkFont(size=20, weight="bold"), anchor="w").grid(
            row=0, column=0, padx=24, pady=(19, 5), sticky="w"
        )
        ctk.CTkLabel(frame, text=body, text_color=MUTED,
                     font=ctk.CTkFont(size=13), wraplength=830,
                     justify="left", anchor="w").grid(
            row=1, column=0, padx=24, pady=(0, 18), sticky="w"
        )
        return frame

    def primary_button(self, parent, text, command, row=2, column=0, padx=(22, 8), sticky="w"):
        button = ctk.CTkButton(parent, text=text, command=command, height=42,
                               fg_color=PURPLE, hover_color=PURPLE_DARK,
                               font=ctk.CTkFont(size=13, weight="bold"))
        button.grid(row=row, column=column, padx=padx, pady=(0, 20), sticky=sticky)
        return button

    def secondary_button(self, parent, text, command, row=2, column=1, padx=(8, 22), sticky="w"):
        button = ctk.CTkButton(parent, text=text, command=command, height=42,
                               fg_color=PURPLE_LIGHT, hover_color="#E5D9F4",
                               border_width=1, border_color=BORDER, text_color=PURPLE_DARK,
                               font=ctk.CTkFont(size=13, weight="bold"))
        button.grid(row=row, column=column, padx=padx, pady=(0, 20), sticky=sticky)
        return button


    def _add_orchid_watermark(
        self, parent, size=(520, 520), relx=0.5, rely=0.56,
        anchor="center", opacity=0.055,
    ):
        """Place a large, centered, high-resolution orchid watermark behind page content."""
        try:
            path = resource_path("assets", "orchid_flower_purple.png")
            if not path.exists():
                return None
            source = PILImage.open(path).convert("RGBA")
            alpha = source.getchannel("A").point(
                lambda value: max(0, min(255, int(value * opacity)))
            )
            source.putalpha(alpha)
            image = ctk.CTkImage(light_image=source, dark_image=source, size=size)
            if not hasattr(self, "_workspace_watermark_images"):
                self._workspace_watermark_images = []
            self._workspace_watermark_images.append(image)
            label = ctk.CTkLabel(parent, text="", image=image, fg_color="transparent")
            label.place(relx=relx, rely=rely, anchor=anchor)
            return label
        except Exception:
            return None

    def _build_event_summary(self, page, row: int):
        """Create the compact current-event card without oversized empty space."""
        shadow = ctk.CTkFrame(page, fg_color=SHADOW, corner_radius=16, width=465, height=82)
        shadow.grid(row=row, column=0, sticky="e", padx=32, pady=(15, 5))
        shadow.grid_propagate(False)
        summary = ctk.CTkFrame(
            page, fg_color=WHITE, border_width=1, border_color=BORDER,
            corner_radius=16, width=465, height=76,
        )
        summary.grid(row=row, column=0, sticky="e", padx=32, pady=(11, 9))
        summary.pack_propagate(False)

        icon = ctk.CTkLabel(
            summary, text="▣", width=46, height=46, corner_radius=11,
            fg_color=PURPLE_LIGHT, text_color=PURPLE,
            font=ctk.CTkFont(size=20, weight="bold"),
        )
        icon.pack(side="left", padx=(13, 10), pady=14)

        text_block = ctk.CTkFrame(summary, fg_color="transparent")
        text_block.pack(side="left", fill="y")
        ctk.CTkLabel(
            text_block, text="CURRENT EVENT", text_color=PURPLE,
            font=ctk.CTkFont(size=9, weight="bold"), anchor="w",
        ).pack(anchor="w", pady=(15, 0))
        title = ctk.CTkLabel(
            text_block, text="No Active Purchase Packet", text_color=TEXT,
            font=ctk.CTkFont(size=13, weight="bold"), anchor="w", width=275,
        )
        title.pack(anchor="w", pady=(1, 0))

        date = ctk.CTkLabel(
            summary, text=time.strftime("%b %d, %Y"), text_color=MUTED,
            font=ctk.CTkFont(size=10, weight="bold"), anchor="e",
        )
        date.pack(side="right", padx=(10, 14))
        ctk.CTkFrame(summary, width=1, height=46, fg_color=BORDER).pack(
            side="right", fill="y", padx=(10, 0), pady=14
        )
        return summary, icon, title, date

    def _build_workflow_tracker(self, page, row: int, storage_attr: str, padx: int = 52):
        """Build the five-stage purchasing workflow tracker."""
        panel = ctk.CTkFrame(
            page, fg_color="transparent", corner_radius=0, height=126,
        )
        panel.grid(row=row, column=0, sticky="ew", padx=max(padx, 24), pady=(18, 12))
        panel.grid_propagate(False)
        stages = ctk.CTkFrame(panel, fg_color="transparent")
        stages.pack(fill="both", expand=True)
        for column in (0, 2, 4, 6, 8):
            stages.grid_columnconfigure(column, weight=1, uniform="workflow_step", minsize=168)
        for column in (1, 3, 5, 7):
            stages.grid_columnconfigure(column, weight=0, minsize=18)
        stages.grid_rowconfigure(0, weight=1)

        if not hasattr(self, "_workflow_icon_images"):
            self._workflow_icon_images = {}
            for key, filename in {
                "import": "workflow_import",
                "master": "workflow_shirt",
                "review": "workflow_review",
                "audit": "workflow_audit",
                "purchase": "workflow_po",
            }.items():
                purple_path = resource_path("assets", f"{filename}_purple.png")
                white_path = resource_path("assets", f"{filename}_white.png")
                purple_source = PILImage.open(purple_path).convert("RGBA")
                white_source = PILImage.open(white_path).convert("RGBA")
                self._workflow_icon_images[(key, "purple")] = ctk.CTkImage(
                    light_image=purple_source, dark_image=purple_source, size=(40, 40)
                )
                self._workflow_icon_images[(key, "white")] = ctk.CTkImage(
                    light_image=white_source, dark_image=white_source, size=(40, 40)
                )

        stage_widgets = []
        stage_names = ("Import Orders", "Product Master", "Purchase Review", "Decoration Colors", "Purchase Orders")
        stage_keys = ("import", "master", "review", "audit", "purchase")
        target_pages = ("import", "master", "review", "audit", "purchase")
        for index, (label, stage_key, target_page) in enumerate(zip(stage_names, stage_keys, target_pages)):
            step = ctk.CTkFrame(
                stages, fg_color=WHITE, border_width=1, border_color=BORDER,
                corner_radius=16, height=98,
            )
            step.grid(row=0, column=index * 2, sticky="nsew", pady=4)
            step.grid_propagate(False)
            step.grid_columnconfigure(1, weight=1)
            step.grid_rowconfigure((0, 1), weight=1)
            icon = ctk.CTkButton(
                step, text="", image=self._workflow_icon_images[(stage_key, "purple")],
                width=58, height=58, corner_radius=13,
                fg_color="transparent", hover_color=PURPLE_LIGHT,
                border_width=0,
                command=lambda p=target_page: self.show_page(p),
            )
            icon.grid(row=0, column=0, rowspan=2, padx=(12, 7), pady=18)
            title = ctk.CTkLabel(
                step, text=label, text_color=TEXT, anchor="w",
                font=ctk.CTkFont(size=11, weight="bold"), wraplength=105,
            )
            title.grid(row=0, column=1, sticky="sw", padx=(0, 8), pady=(14, 1))
            body_label = ctk.CTkLabel(step, text="", width=1, height=1)
            body_label.grid_remove()
            status = ctk.CTkLabel(
                step, text="Waiting", text_color=MUTED, fg_color="transparent", anchor="w",
                font=ctk.CTkFont(size=10, weight="bold"),
            )
            status.grid(row=1, column=1, sticky="nw", padx=(0, 8), pady=(1, 14))
            connector = None
            if index < 4:
                connector = ctk.CTkFrame(stages, height=2, width=18, fg_color="#D7D0E0", corner_radius=2)
                connector.grid(row=0, column=index * 2 + 1, sticky="ew", padx=1)
            stage_widgets.append({
                "card": step, "circle": icon, "title": title,
                "body": body_label, "status": status, "connector": connector,
                "stage_key": stage_key,
            })
        setattr(self, storage_attr, stage_widgets)
        return panel

    def build_dashboard_page(self):
        page = self.new_page("dashboard")
        page.grid_rowconfigure(2, weight=1)

        self.dashboard_view_toggle = ctk.CTkSegmentedButton(
            page, values=["Guided View", "Management View"], command=self._set_dashboard_view,
            width=300, height=36, corner_radius=11, fg_color="#EEE9F5",
            selected_color="#C7A7FF", selected_hover_color="#B58AF5",
            unselected_color="#EEE9F5", unselected_hover_color="#E3D9EF",
            text_color=TEXT, font=ctk.CTkFont(size=12, weight="bold"),
        )
        self.dashboard_view_toggle.place(relx=0.975, y=10, anchor="ne")
        self.mc_guided_tracker = self._build_workflow_tracker(page, 0, "mc_stage_labels", padx=52)
        self.mc_guided_tracker.grid_configure(pady=(46, 12))

        hero_shadow = ctk.CTkFrame(page, fg_color=SHADOW, corner_radius=22, height=430)
        self.mc_guided_shadow = hero_shadow
        hero_shadow.grid(row=1, column=0, sticky="new", padx=58, pady=(4, 18))
        hero = ctk.CTkFrame(
            page, fg_color=WHITE, border_width=1, border_color=BORDER,
            corner_radius=22, height=424,
        )
        hero.grid(row=1, column=0, sticky="new", padx=58, pady=(0, 22))
        hero.grid_propagate(False)
        hero.grid_columnconfigure(0, weight=1)
        self.mc_hero = hero
        # The dashboard watermark was removed to keep the expanded workspace clean.

        self.mc_animation_frame = ctk.CTkFrame(hero, fg_color="transparent", height=42)
        self.mc_animation_frame.grid(row=0, column=0, sticky="ew", padx=28, pady=(8, 0))
        self.mc_animation_frame.grid_propagate(False)
        self.success_flower_labels = []
        purple_flower = resource_path("assets", "orchid_flower_purple.png")
        if purple_flower.exists():
            flower_source = PILImage.open(purple_flower).convert("RGBA")
            self.success_flower_image = ctk.CTkImage(light_image=flower_source, dark_image=flower_source, size=(42, 42))
            self.success_bloom_image = ctk.CTkImage(light_image=flower_source, dark_image=flower_source, size=(64, 64))
            for _ in range(7):
                label = ctk.CTkLabel(self.mc_animation_frame, text="", image=self.success_flower_image,
                                     fg_color="transparent", width=44, height=44)
                self.success_flower_labels.append(label)
            self.success_bloom_label = ctk.CTkLabel(self.mc_animation_frame, text="", image=self.success_bloom_image,
                                                    fg_color="transparent", width=66, height=66)
        self.mc_animation_controls = ctk.CTkFrame(self.mc_animation_frame, fg_color="transparent")
        self.mc_sound_toggle = ctk.CTkButton(
            self.mc_animation_controls, text="Sound: On", command=self._toggle_success_sound,
            width=98, height=28, corner_radius=10, fg_color=PURPLE_LIGHT, hover_color="#E6D9F5",
            border_width=1, border_color="#CDBBE5", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=12, weight="bold"),
        )
        self.mc_sound_toggle.pack(side="left", padx=(0, 6))
        self.mc_skip_animation_button = ctk.CTkButton(
            self.mc_animation_controls, text="Skip Animation", command=self._skip_orchid_success_animation,
            width=122, height=28, corner_radius=10, fg_color="transparent", hover_color=PURPLE_LIGHT,
            border_width=1, border_color="#CDBBE5", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=12, weight="bold"), state="disabled",
        )
        self.mc_skip_animation_button.pack(side="left")
        self.mc_animation_controls.place(relx=1.0, rely=0.05, anchor="ne")
        self.mc_animation_controls.place_forget()
        self._update_success_control_buttons()

        self.mc_step_label = ctk.CTkLabel(hero, text="", width=1, height=1)
        self.mc_step_label.grid_remove()
        self.mc_banner_icon = ctk.CTkLabel(
            hero, text="⧉", width=64, height=64, corner_radius=32,
            fg_color=PURPLE_LIGHT, text_color=PURPLE,
            font=ctk.CTkFont(size=29, weight="bold"),
        )
        self.mc_banner_icon.grid(row=1, column=0, pady=(12, 12))
        self.mc_banner_title = ctk.CTkLabel(
            hero, text="Import Orders", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=27, weight="bold"), justify="center", wraplength=790,
        )
        self.mc_banner_title.grid(row=2, column=0, padx=50, pady=(0, 8))
        self.mc_banner_subtitle = ctk.CTkLabel(
            hero, text="Choose the Shopify or Report Toaster export to start a new purchase packet.",
            text_color=MUTED, font=ctk.CTkFont(size=15), justify="center", wraplength=820,
        )
        self.mc_banner_subtitle.grid(row=3, column=0, padx=56, pady=(0, 14))
        self.dashboard_status = self.mc_banner_subtitle
        self.mc_event_label = ctk.CTkLabel(
            hero, text="No order CSV selected", text_color=TEXT,
            fg_color="#F3F1F6", corner_radius=14, height=31,
            font=ctk.CTkFont(size=12, weight="bold"), justify="center",
        )
        self.mc_event_label.grid(row=4, column=0, padx=40, pady=(0, 16))
        metrics = ctk.CTkFrame(hero, fg_color="#FAF8FC", corner_radius=14)
        metrics.grid(row=5, column=0, sticky="ew", padx=90, pady=(0, 16))
        metrics.grid_columnconfigure((0, 1), weight=1, uniform="metric")
        self.mc_metrics_frame = metrics
        self.mc_metric_widgets = {}
        for column, (key, caption) in enumerate((("employees", "Employees"), ("total", "Final Total"))):
            metric = ctk.CTkFrame(metrics, fg_color="transparent")
            metric.grid(row=0, column=column, padx=10, pady=12, sticky="nsew")
            value = ctk.CTkLabel(metric, text="—", text_color=PURPLE_DARK, font=ctk.CTkFont(size=23, weight="bold"))
            value.pack(pady=(0, 2))
            ctk.CTkLabel(metric, text=caption, text_color=MUTED, font=ctk.CTkFont(size=10, weight="bold")).pack()
            self.mc_metric_widgets[key] = {"frame": metric, "value": value}
        metrics.grid_remove()
        button_row = ctk.CTkFrame(hero, fg_color="transparent")
        button_row.grid(row=6, column=0, pady=(0, 22))
        self.mc_banner_button = ctk.CTkButton(
            button_row, text="Import Orders", command=lambda: self.show_page("import"),
            width=330, height=56, corner_radius=11, fg_color=PURPLE, hover_color="#46109F",
            font=ctk.CTkFont(size=15, weight="bold"),
        )
        self.mc_banner_button.pack(side="left", padx=7)
        self.mc_banner_secondary_button = ctk.CTkButton(
            button_row, text="Start New Event", command=self.start_new_purchase_packet,
            width=205, height=54, corner_radius=9, fg_color=PURPLE_LIGHT, hover_color="#E6D9F5",
            border_width=1, border_color="#BFA9DD", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=14, weight="bold"),
        )
        self.mc_banner_secondary_button.pack(side="left", padx=7)
        self.mc_banner_secondary_button.pack_forget()

        # Compatibility controls retained for the existing workflow engine.
        hidden = ctk.CTkFrame(page, fg_color="transparent")
        self.mc_master_card = self._mission_card(hidden, 0, "Product Master", "Permanent product catalog")
        self.mc_master_status = self.mc_master_card["status"]
        self.mc_master_detail = self.mc_master_card["detail"]
        self.mc_master_button = ctk.CTkButton(self.mc_master_card["frame"], text="Open Product Master", command=self.open_product_master)
        self.mc_master_button.grid(row=5, column=0)
        self.mc_review_card = self._mission_card(hidden, 1, "Purchase Review", "Only decisions that require attention")
        self.mc_review_status = self.mc_review_card["status"]
        self.mc_review_detail = self.mc_review_card["detail"]
        self.mc_review_button = ctk.CTkButton(self.mc_review_card["frame"], text="Open Purchase Review", command=lambda: self.show_page("review"))
        self.mc_review_button.grid(row=5, column=0)
        self.mc_po_card = self._mission_card(hidden, 2, "Purchase Orders", "PO numbers and final vendor PDFs")
        self.mc_po_status = self.mc_po_card["status"]
        self.mc_po_detail = self.mc_po_card["detail"]
        self.mc_po_card_button = ctk.CTkButton(self.mc_po_card["frame"], text="Edit PO Numbers", command=self.open_po_number_editor)
        self.mc_po_card_button.grid(row=5, column=0)
        self.mc_work_summary = ctk.CTkLabel(hidden, text="")
        self.mc_work_open_button = ctk.CTkButton(hidden, text="Open Review", command=self.open_latest_review)
        self.mc_work_scroll = ctk.CTkScrollableFrame(hidden, width=2, height=2)
        self.mc_work_scroll.grid_columnconfigure(1, weight=1)
        self.mc_action_title = ctk.CTkLabel(hidden, text="")
        self.mc_action_detail = ctk.CTkLabel(hidden, text="")
        self.mc_action_button = ctk.CTkButton(hidden, text="", command=lambda: None)

        self._build_management_dashboard(page)
        self.after_idle(self._apply_dashboard_view)

    def _set_dashboard_view(self, value: str) -> None:
        normalized = "management" if str(value).strip().casefold().startswith("management") else "guided"
        if normalized != getattr(self, "dashboard_view", "guided"):
            self.dashboard_view = normalized
            self._save_ui_preferences()
        self._apply_dashboard_view()
        self.refresh_dashboard()

    def _apply_dashboard_view(self) -> None:
        management = getattr(self, "dashboard_view", "guided") == "management"
        if hasattr(self, "dashboard_view_toggle"):
            self.dashboard_view_toggle.set("Management View" if management else "Guided View")
        guided_widgets = [
            getattr(self, "mc_guided_tracker", None),
            getattr(self, "mc_guided_shadow", None),
            getattr(self, "mc_hero", None),
        ]
        if management:
            for widget in guided_widgets:
                if widget is not None:
                    widget.grid_remove()
            if hasattr(self, "mc_management_frame"):
                self.mc_management_frame.place(relx=0, rely=0, relwidth=1, relheight=1)
        else:
            if hasattr(self, "mc_management_frame"):
                self.mc_management_frame.place_forget()
            for widget in guided_widgets:
                if widget is not None:
                    widget.grid()
        if hasattr(self, "dashboard_view_toggle"):
            self.dashboard_view_toggle.lift()

    def _management_metric_card(self, parent, column: int, title: str, key: str) -> None:
        card = ctk.CTkFrame(parent, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=16)
        card.grid(row=0, column=column, sticky="nsew", padx=6, pady=5)
        value = ctk.CTkLabel(card, text="—", text_color=PURPLE, font=ctk.CTkFont(size=24, weight="bold"))
        value.pack(pady=(14, 2))
        ctk.CTkLabel(card, text=title, text_color=MUTED, font=ctk.CTkFont(size=11, weight="bold")).pack(pady=(0, 14))
        self.mc_management_metric_labels[key] = value

    def _management_detail_card(self, parent, row: int, column: int, title: str) -> tuple[ctk.CTkFrame, ctk.CTkLabel]:
        card = ctk.CTkFrame(parent, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=17)
        card.grid(row=row, column=column, sticky="nsew", padx=7, pady=7)
        card.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(card, text=title, text_color=PURPLE_DARK, font=ctk.CTkFont(size=16, weight="bold"), anchor="w").grid(
            row=0, column=0, padx=18, pady=(16, 8), sticky="ew"
        )
        detail = ctk.CTkLabel(card, text="", text_color=TEXT, font=ctk.CTkFont(size=12), justify="left", anchor="nw", wraplength=330)
        detail.grid(row=1, column=0, padx=18, pady=(0, 16), sticky="nsew")
        card.grid_rowconfigure(1, weight=1)
        return card, detail

    def _build_management_dashboard(self, page: ctk.CTkFrame) -> None:
        frame = ctk.CTkScrollableFrame(page, fg_color=BG, corner_radius=0)
        self.mc_management_frame = frame
        frame.grid_columnconfigure((0, 1, 2), weight=1, uniform="management")

        heading = ctk.CTkFrame(frame, fg_color="transparent")
        heading.grid(row=0, column=0, columnspan=3, sticky="ew", padx=14, pady=(52, 6))
        heading.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(heading, text="Management Dashboard", text_color=PURPLE_DARK, font=ctk.CTkFont(size=29, weight="bold"), anchor="w").grid(row=0, column=0, sticky="w")
        self.mc_management_heading = ctk.CTkLabel(heading, text="Experienced view • the same live event and workflow", text_color=MUTED, font=ctk.CTkFont(size=12), anchor="w")
        self.mc_management_heading.grid(row=1, column=0, sticky="w", pady=(2, 0))

        metrics = ctk.CTkFrame(frame, fg_color="transparent")
        metrics.grid(row=1, column=0, columnspan=3, sticky="ew", padx=8, pady=(4, 4))
        for column in range(7):
            metrics.grid_columnconfigure(column, weight=1, uniform="management_metric")
        self.mc_management_metric_labels = {}
        for column, (title, key) in enumerate((
            ("Employees", "employees"), ("Line Items", "lines"), ("Ready", "ready"),
            ("Decisions", "decisions"), ("Blocked", "blocked"), ("Vendors", "vendors"),
            ("Ship to Orchid", "manual"),
        )):
            self._management_metric_card(metrics, column, title, key)

        event_card, self.mc_management_event_detail = self._management_detail_card(frame, 2, 0, "Event Summary")
        review_card, self.mc_management_review_detail = self._management_detail_card(frame, 2, 1, "Purchase Review Overview")

        actions = ctk.CTkFrame(frame, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=17)
        actions.grid(row=2, column=2, sticky="nsew", padx=7, pady=7)
        actions.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(actions, text="Quick Actions", text_color=PURPLE_DARK, font=ctk.CTkFont(size=16, weight="bold"), anchor="w").grid(row=0, column=0, padx=18, pady=(16, 8), sticky="ew")
        quick_actions = [
            ("Open Purchase Review", lambda: self.show_page("review")),
            ("Regenerate Purchase Review", self.regenerate_current_review),
            ("Employee Totals", lambda: self.show_page("employees")),
            ("Receiving & Decoration Report", lambda: self._open_generated_report("in-house")),
            ("Outsourced Embroidery Report", lambda: self._open_generated_report("outsourced-embroidery")),
            ("Outsourced Screen-Printing Report", lambda: self._open_generated_report("outsourced-screen-printing")),
            ("Ship to Orchid Purchase Order", lambda: self._open_generated_report("non-included")),
            ("Purchase Documents", self.open_latest_purchase_orders),
        ]
        self.mc_management_action_buttons = {}
        for row_index, (label, command) in enumerate(quick_actions, start=1):
            button = ctk.CTkButton(
                actions, text=label, command=command, anchor="w", height=36, corner_radius=9,
                fg_color="#F7F3FC", hover_color=PURPLE_LIGHT, border_width=1, border_color="#D8CBE8",
                text_color=PURPLE_DARK, font=ctk.CTkFont(size=12, weight="bold"),
            )
            button.grid(row=row_index, column=0, padx=16, pady=3, sticky="ew")
            self.mc_management_action_buttons[label] = button
        ctk.CTkFrame(actions, fg_color="transparent", height=9).grid(row=len(quick_actions) + 1, column=0)

        recent = ctk.CTkFrame(frame, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=17)
        recent.grid(row=3, column=0, columnspan=3, sticky="ew", padx=7, pady=(7, 22))
        recent.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(recent, text="Recent Activity", text_color=PURPLE_DARK, font=ctk.CTkFont(size=16, weight="bold"), anchor="w").grid(row=0, column=0, padx=18, pady=(15, 6), sticky="ew")
        self.mc_management_recent_detail = ctk.CTkLabel(recent, text="No event activity yet.", text_color=TEXT, font=ctk.CTkFont(size=12), justify="left", anchor="w")
        self.mc_management_recent_detail.grid(row=1, column=0, padx=18, pady=(0, 16), sticky="ew")
        frame.place_forget()

    def _open_generated_report(self, report_kind: str) -> None:
        root, _using_archive_fallback = self._purchase_order_release_location()
        if not root or not root.exists():
            messagebox.showinfo("Report Not Generated", "Generate purchase orders and reports first.")
            return
        tokens = {
            "in-house": ("receiving", "decoration", "report"),
            "outsourced": ("outsourced", "decoration"),
            "outsourced-embroidery": ("outsourced", "embroidery", "job"),
            "outsourced-screen-printing": ("outsourced", "screen", "printing", "job"),
            "non-included": ("ship_to_orchid", "ship to orchid", "non-included", "non included"),
        }.get(report_kind, (report_kind,))
        candidates = sorted(root.rglob("*.pdf"), key=lambda path: path.stat().st_mtime, reverse=True)
        # A department report must match every part of its filename; matching
        # only "outsourced" would otherwise open whichever department happened
        # to sort first.
        if report_kind == "non-included":
            match = next((path for path in candidates if any(token in path.name.casefold() for token in tokens)), None)
        else:
            match = next((path for path in candidates if all(token in path.name.casefold() for token in tokens)), None)
        if match is None:
            messagebox.showinfo("Report Not Generated", "That report has not been generated for the current event yet.")
            return
        subprocess.run(["open", str(match)], check=False)

    def _management_manual_piece_count(self) -> int:
        workbook = self.last_review_workbook
        if not workbook or not workbook.is_file():
            return 0
        try:
            total = 0
            for row in load_review_lines(workbook):
                include = _clean(row.get("Include", "Yes")).casefold()
                manual = _clean(row.get("Do Not Outsource", "No")).casefold()
                if include in {"yes", "y", "true", "1", "include"} and manual in {"yes", "y", "true", "1"}:
                    try:
                        total += max(int(float(row.get("Quantity", 0) or 0)), 0)
                    except (TypeError, ValueError):
                        pass
            return total
        except Exception:
            return 0

    def _management_recent_activity(self) -> str:
        entries = []
        paths = [
            (self.selected_csv, "Orders imported"),
            (live_product_master_path(), "Product Master updated"),
            (self.last_review_workbook, "Purchase Review updated"),
            (self.last_purchase_order_dir, "Purchase documents generated"),
        ]
        for path, label in paths:
            try:
                if path and Path(path).exists():
                    modified = Path(path).stat().st_mtime
                    entries.append((modified, f"{time.strftime('%b %d, %I:%M %p', time.localtime(modified))}  •  {label}"))
            except Exception:
                pass
        entries.sort(key=lambda item: item[0], reverse=True)
        return "\n".join(text for _, text in entries[:5]) or "No event activity yet."

    def _refresh_management_dashboard(self, snapshot: dict, purchase_generated: bool = False) -> None:
        if not hasattr(self, "mc_management_metric_labels"):
            return
        line_count = int(snapshot.get("line_count", self.imported_line_count or 0) or 0)
        ready = int(snapshot.get("ready_count", 0) or 0)
        decisions = int(snapshot.get("purchase_review_count", snapshot.get("review_count", 0)) or 0)
        blocked = int(snapshot.get("blocked_route_count", 0) or 0)
        product_setup = int(snapshot.get("product_master_count", 0) or 0)
        employees = int(snapshot.get("employee_count", 0) or 0)
        routes = list(snapshot.get("routes", []))
        vendors = len({_clean(route.get("vendor", "")).casefold() for route in routes if _clean(route.get("vendor", ""))})
        manual = self._management_manual_piece_count()
        values = {
            "employees": employees, "lines": line_count, "ready": ready, "decisions": decisions,
            "blocked": blocked, "vendors": vendors, "manual": manual,
        }
        for key, value in values.items():
            self.mc_management_metric_labels[key].configure(text=f"{value:,}")

        event_name = self.current_event_name or snapshot.get("event_name") or "No Active Purchase Packet"
        mode = self.current_mode or snapshot.get("report_mode") or "—"
        fulfillment = normalize_decoration_fulfillment(self.current_decoration_fulfillment)
        final_total = float(snapshot.get("employee_grand_total", 0.0) or 0.0)
        po_state = "Generated" if purchase_generated else ("Ready to generate" if line_count and not decisions and not blocked and not product_setup else "In progress")
        self.mc_management_event_detail.configure(text=(
            f"Event:  {event_name}\n"
            f"Order mode:  {mode}\n"
            f"Decoration workflow:  {fulfillment}\n"
            f"Employee total:  ${final_total:,.2f}\n"
            f"Purchase documents:  {po_state}"
        ))
        self.mc_management_review_detail.configure(text=(
            f"Product Master setup:  {product_setup:,}\n"
            f"Purchase Review decisions:  {decisions:,}\n"
            f"Blocked routes:  {blocked:,}\n"
            f"Ready purchase lines:  {ready:,} of {line_count:,}\n"
            f"Ship-to-Orchid exception pieces:  {manual:,}"
        ))
        self.mc_management_recent_detail.configure(text=self._management_recent_activity())
        employee_button = self.mc_management_action_buttons.get("Employee Totals")
        if employee_button is not None:
            employee_button.configure(state="normal" if should_show_employee_totals(snapshot, bool(self.last_review_workbook), self.current_mode, self.imported_line_count) else "disabled")
        report_root, _using_archive_fallback = self._purchase_order_release_location()
        report_state = "normal" if report_root else "disabled"
        for label in (
            "Receiving & Decoration Report",
            "Outsourced Embroidery Report",
            "Outsourced Screen-Printing Report",
            "Ship to Orchid Purchase Order",
        ):
            button = self.mc_management_action_buttons.get(label)
            if button is not None:
                button.configure(state=report_state)

    def _mission_card(self, parent, column: int, title: str, subtitle: str) -> dict:
        frame = ctk.CTkFrame(parent, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=15)
        frame.grid(row=0, column=column, sticky="nsew", padx=7, pady=5)
        frame.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(frame, text=title, text_color=TEXT,
                     font=ctk.CTkFont(size=19, weight="bold")).grid(row=0, column=0, padx=16, pady=(16, 3))
        ctk.CTkLabel(frame, text=subtitle, text_color=MUTED,
                     font=ctk.CTkFont(size=12)).grid(row=1, column=0, padx=16)
        status = ctk.CTkLabel(frame, text="Not started", height=28, corner_radius=14,
                              fg_color="#ECEAF0", text_color=MUTED,
                              font=ctk.CTkFont(size=13, weight="bold"))
        status.grid(row=2, column=0, padx=16, pady=(10, 6))
        detail = ctk.CTkLabel(frame, text="", text_color=TEXT,
                              font=ctk.CTkFont(size=12), justify="center", wraplength=290)
        detail.grid(row=3, column=0, padx=16, pady=(0, 6))
        return {"frame": frame, "status": status, "detail": detail}

    def _highlight_mission_card(self, active: str | None):
        cards = {
            "master": self.mc_master_card["frame"],
            "review": self.mc_review_card["frame"],
            "purchase": self.mc_po_card["frame"],
        }
        for name, frame in cards.items():
            frame.configure(
                border_width=3 if name == active else 1,
                border_color=PURPLE if name == active else BORDER,
            )

    def _sync_current_product_candidates(self) -> dict:
        """Add deduplicated review candidates to the live Product Master once.

        This makes genuinely new Report Toaster/Shopify styles immediately
        available in the Product Master Needs Setup queue without overwriting
        any saved manual assignments.
        """
        workbook = self.last_review_workbook
        if not workbook or not Path(workbook).exists():
            return {"changed": False}
        workbook_key = str(Path(workbook).resolve())
        if self._candidate_sync_workbook == workbook_key:
            return {"changed": False, "cached": True}
        try:
            result = sync_review_product_candidates(Path(workbook), live_product_master_path())
            self._candidate_sync_workbook = workbook_key
            if result.get("changed"):
                try:
                    self.product_master_update_marker.parent.mkdir(parents=True, exist_ok=True)
                    self.product_master_update_marker.write_text(str(time.time()), encoding="utf-8")
                    self._last_product_master_update = self.product_master_update_marker.stat().st_mtime
                except Exception:
                    pass
            return result
        except Exception as error:
            print(f"Product candidate sync warning: {error}", file=sys.stderr)
            return {"changed": False, "error": str(error)}

    def _purchase_review_routing_status(self, force: bool = False) -> dict:
        """Return whether Product Master can still change this event's PO routes.

        This intentionally ignores a thread/ink color or catalog-only update.
        It becomes stale only when an active line's vendor, decoration
        type/location, or Never Outsource result would change.
        """
        workbook = self.last_review_workbook
        if not workbook or not Path(workbook).is_file():
            return {"current": True, "reason": "", "changes": []}
        master = live_product_master_path()
        override = Path(master).parent / "never_outsource_overrides.json"
        cache_key = (_file_stamp(Path(workbook)), _file_stamp(master), _file_stamp(override))
        if not force and cache_key == self._routing_status_cache_key:
            return dict(self._routing_status_cache)
        try:
            status = review_routing_status(Path(workbook))
            if not isinstance(status, dict):
                raise RuntimeError("Routing check returned an invalid result.")
        except Exception as error:
            status = {
                "current": False,
                "reason": f"Orchid could not verify current Product Master routing: {error}",
                "changes": [],
            }
        self._routing_status_cache_key = cache_key
        self._routing_status_cache = dict(status)
        return dict(status)

    def _review_is_stale(self) -> bool:
        if not self._is_archived_report_reopen() and self.last_review_workbook:
            valid, _reason = self._active_packet_is_valid(self.last_review_workbook)
            if not valid:
                return True
        return not bool(self._purchase_review_routing_status().get("current", False))

    def _poll_product_master_updates(self):
        try:
            existing_pid = self._existing_product_master_pid()
            if existing_pid:
                self._product_master_was_open = True
            elif self._product_master_was_open:
                # Product Master has just closed. Rebuild the review once so all
                # saved setup decisions are applied without asking the employee
                # to hunt for a regenerate action.
                self._product_master_was_open = False
                if self._review_is_stale() and self.last_review_workbook:
                    self.regenerate_current_review(silent=True)

            if self.product_master_update_marker.exists():
                modified = self.product_master_update_marker.stat().st_mtime
                if modified > self._last_product_master_update:
                    self._last_product_master_update = modified
                    self.refresh_dashboard()
                    self._schedule_auto_review_regeneration()
        except Exception:
            pass
        try:
            self.after(900, self._poll_product_master_updates)
        except Exception:
            pass

    def _schedule_auto_review_regeneration(self):
        """Debounce Product Master saves into one silent Purchase Review refresh."""
        if not self.last_review_workbook or not REVIEW_STATE_FILE.exists() or not self._review_is_stale():
            return
        if self._auto_regenerate_after_id:
            try:
                self.after_cancel(self._auto_regenerate_after_id)
            except Exception:
                pass
        self._auto_regenerate_after_id = self.after(700, self._run_auto_review_regeneration)

    def _run_auto_review_regeneration(self):
        self._auto_regenerate_after_id = None
        if self._regenerating_review or not self._review_is_stale():
            return
        self.regenerate_current_review(silent=True)

    def reopen_unresolved_purchase_review(self):
        """Fallback: clear app-only completion markers and show workbook issues again."""
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            messagebox.showinfo("No Purchase Review", "Create a Purchase Review first.")
            return
        if not messagebox.askyesno(
            "Reopen Unresolved Purchase Review?",
            "Use this only if you prefer to review the unresolved workbook rows again.\n\n"
            "Orchid will remove the event-level completion marker and the app recovery journal. "
            "Rows already saved successfully in Excel remain completed; only unresolved workbook rows return.",
            default=messagebox.NO,
        ):
            return
        state = self._load_review_state()
        for key in (
            "review_complete", "review_completed_at", "review_workbook",
            "completed_review_decisions", "pending_review_saves",
            "draft_line_id", "draft_decision_key", "draft_saved_at", "draft_values",
        ):
            state.pop(key, None)
        try:
            REVIEW_STATE_FILE.write_text(json.dumps(state, indent=2), encoding="utf-8")
        except Exception as error:
            messagebox.showerror("Unable to Reopen Purchase Review", str(error))
            return
        self._review_save_pending = 0
        self._invalidate_review_cache()
        self._invalidate_fast_review_snapshot()
        self._mission_snapshot_cache_key = None
        self._mission_snapshot_cache = {}
        self.show_page("review")
        self._load_review_page_async(force=True)

    def regenerate_current_review(self, silent: bool = False, return_to_purchase: bool = False):
        from modules.review_workbook import generate_review_workbook

        if self._regenerating_review:
            return
        self._regenerating_review = True
        try:
            if not REVIEW_STATE_FILE.exists():
                messagebox.showinfo("No Saved Review", "Create a Purchase Review first.")
                return
            state = json.loads(REVIEW_STATE_FILE.read_text(encoding="utf-8"))
            source_csv = Path(state.get("source_csv", "")).expanduser()
            if not source_csv.exists():
                messagebox.showerror("Orders File Missing", f"The saved order CSV can no longer be found:\n{source_csv}")
                return
            if not _clean(state.get("active_packet_id", "")):
                upgraded = new_active_packet_state(source_csv)
                active_packet_id = upgraded["active_packet_id"]
                upgraded.update(state)
                upgraded["active_packet_id"] = active_packet_id
                state = upgraded
                self._write_review_state_atomic(state)
            result = generate_review_workbook(
                source_csv, live_product_master_path(), REPORTS,
                report_mode=state.get("report_mode") or GENERAL_SALES_PERIOD,
                event_name=state.get("event_name") or "",
                active_packet_id=_clean(state.get("active_packet_id", "")),
                decoration_fulfillment=state.get("decoration_fulfillment") or STANDARD_ORCHID_WORKFLOW,
                regenerate_helper=REGENERATE_HELPER,
                previous_review_path=self.last_review_workbook,
            )
            self.selected_csv = source_csv
            self.last_review_workbook = Path(result["output_path"])
            self._bind_active_review_workbook(self.last_review_workbook)
            self._reset_review_navigation()
            self.current_event_name = result.get("event_name") or self.current_event_name or "General Sales Period"
            self.current_mode = result.get("report_mode") or self.current_mode
            self.current_decoration_fulfillment = result.get("decoration_fulfillment") or self.current_decoration_fulfillment
            self.imported_line_count = int(result.get("lines", 0))
            self.review_needed = int(result.get("review_decisions", result.get("review_lines", 0)))
            refreshed_state = self._load_review_state()
            refreshed_state.pop("decoration_color_audit_requires_regeneration", None)
            refreshed_state.pop("additional_import_pending", None)
            self._write_review_state_atomic(refreshed_state)
            self._routing_status_cache_key = None
            self._sync_current_product_candidates()
            self.refresh_dashboard()
            if return_to_purchase:
                self.show_page("purchase")
                self.refresh_purchase_page()
                snapshot = self._fresh_snapshot()
                review_count = int(snapshot.get("review_count", 0) or 0)
                blocked = int(snapshot.get("blocked_route_count", 0) or 0)
                if review_count or blocked:
                    messagebox.showwarning(
                        "Purchase Review Needs Attention",
                        "Product Master routing was refreshed, but the updated event has decisions that need attention before Purchase Orders can be created."
                    )
                    self.show_page("review")
                else:
                    messagebox.showinfo(
                        "Purchase Review Refreshed",
                        "Current routing was applied. You are back on Purchase Orders and can continue PO generation."
                    )
            elif silent:
                self.dashboard_status.configure(
                    text="Purchase Review refreshed automatically after Product Master was saved.",
                    text_color=SUCCESS,
                )
            else:
                messagebox.showinfo(
                    "Purchase Review Refreshed",
                    "The current Product Master changes were applied. Continue the review inside Orchid Purchase Manager.",
                )
        except Exception as error:
            if not silent:
                messagebox.showerror("Unable to Regenerate Purchase Review", str(error))
            else:
                self.dashboard_status.configure(text=f"Unable to refresh Purchase Review: {error}", text_color=DANGER)
        finally:
            self._regenerating_review = False

    def _set_pill(self, widget, text: str, state: str):
        styles = {
            "complete": (SUCCESS, "#EAF8EF"),
            "attention": (WARNING, "#FFF1DC"),
            "pending": (MUTED, "#ECEAF0"),
        }
        color, fill = styles[state]
        widget.configure(text=text, text_color=color, fg_color=fill)

    def _set_stage(self, index: int, state: str) -> None:
        if state == "complete":
            circle_fill, circle_text, border, symbol = SUCCESS, WHITE, SUCCESS, "✓"
            title_color, status_color, status_text = "#166534", "#166534", "Complete"
            enabled = "normal"
        elif state == "active":
            circle_fill, circle_text, border, symbol = PURPLE, WHITE, PURPLE, str(index + 1)
            title_color, status_color, status_text = PURPLE, PURPLE, "Start Here" if index == 0 else "Current"
            enabled = "normal"
        elif state == "attention":
            circle_fill, circle_text, border, symbol = "#FFF1DC", "#9A4D00", WARNING, "!"
            title_color, status_color, status_text = "#9A4D00", "#9A4D00", "Needs Attention"
            enabled = "normal"
        else:
            circle_fill, circle_text, border, symbol = "#F4EEFC", PURPLE, "#C8B9DA", str(index + 1)
            title_color, status_color, status_text = "#342B45", "#756D82", "Waiting"
            enabled = "normal"
        for attr in (
            "mc_stage_labels", "current_event_stage_labels", "master_stage_labels",
            "review_stage_labels", "audit_stage_labels", "purchase_stage_labels",
        ):
            tracker = getattr(self, attr, None)
            if not tracker or index >= len(tracker):
                continue
            stage = tracker[index]
            stage["circle"].configure(
                text=symbol, fg_color=circle_fill, hover_color=PURPLE_LIGHT,
                text_color=circle_text, border_color=border, state=enabled,
            )
            stage["title"].configure(text_color=title_color)
            stage["status"].configure(text=status_text, text_color=status_color, fg_color="transparent")
            connector = stage.get("connector")
            if connector is not None:
                connector.configure(fg_color=PURPLE if state == "complete" else "#C8BCD8")

    def _update_dashboard_metrics(self, snapshot: dict, show: bool) -> None:
        if not hasattr(self, "mc_metrics_frame"):
            return
        if not show:
            self.mc_metrics_frame.grid_remove()
            self.mc_hero.configure(height=376)
            self.mc_guided_shadow.configure(height=382)
            return
        self.mc_hero.configure(height=450)
        self.mc_guided_shadow.configure(height=456)
        employee_count = int(snapshot.get("employee_count", 0) or 0)
        final_total = float(snapshot.get("employee_grand_total", 0.0) or 0.0)
        values = {
            "employees": f"{employee_count:,}" if employee_count else "—",
            "total": f"${final_total:,.2f}" if final_total else "—",
        }
        for key, value in values.items():
            self.mc_metric_widgets[key]["value"].configure(text=value)
        self.mc_metrics_frame.grid()

    def _prefers_reduced_motion(self) -> bool:
        env = os.environ.get("ORCHID_REDUCED_MOTION", "").strip().casefold()
        if env in {"1", "true", "yes", "on"}:
            return True
        if env in {"0", "false", "no", "off"}:
            return False
        if getattr(self, "_reduced_motion_override", None) is not None:
            return bool(self._reduced_motion_override)
        if sys.platform == "darwin":
            try:
                value = subprocess.check_output(
                    ["defaults", "read", "com.apple.universalaccess", "reduceMotion"],
                    stderr=subprocess.DEVNULL,
                    text=True,
                    timeout=1.2,
                ).strip()
                return value == "1"
            except Exception:
                return False
        return False

    def _toggle_success_sound(self) -> None:
        self.success_sound_enabled = not bool(getattr(self, "success_sound_enabled", True))
        self._update_success_control_buttons()

    def _update_success_control_buttons(self) -> None:
        if hasattr(self, "mc_sound_toggle"):
            self.mc_sound_toggle.configure(
                text="Sound: On" if bool(getattr(self, "success_sound_enabled", True)) else "Sound: Off"
            )
        if hasattr(self, "mc_skip_animation_button"):
            self.mc_skip_animation_button.configure(
                state="normal" if bool(getattr(self, "_success_animation_running", False)) else "disabled"
            )

    def _play_success_sound(self) -> None:
        if not bool(getattr(self, "success_sound_enabled", True)):
            return
        sound_path = resource_path("assets", "success_chime.wav")
        try:
            if sys.platform == "darwin":
                if sound_path.exists():
                    subprocess.Popen(
                        ["afplay", str(sound_path)],
                        stdout=subprocess.DEVNULL,
                        stderr=subprocess.DEVNULL,
                    )
                else:
                    self.bell()
            elif sys.platform.startswith("win"):
                try:
                    import winsound
                    if sound_path.exists():
                        winsound.PlaySound(str(sound_path), winsound.SND_FILENAME | winsound.SND_ASYNC)
                    else:
                        winsound.MessageBeep()
                except Exception:
                    self.bell()
            else:
                self.bell()
        except Exception:
            try:
                self.bell()
            except Exception:
                pass

    def _finish_success_animation(self) -> None:
        for label in getattr(self, "success_flower_labels", []):
            try:
                label.place_forget()
            except Exception:
                pass
        bloom = getattr(self, "success_bloom_label", None)
        if bloom is not None:
            try:
                bloom.place_forget()
            except Exception:
                pass
        self._success_animation_running = False
        self._update_success_control_buttons()

    def _skip_orchid_success_animation(self) -> None:
        if not bool(getattr(self, "_success_animation_running", False)):
            return
        self._success_animation_skip_requested = True
        self._finish_success_animation()

    def _play_orchid_success_animation(self) -> None:
        labels = getattr(self, "success_flower_labels", [])
        if getattr(self, "_success_animation_running", False):
            return
        self._success_animation_running = True
        self._success_animation_skip_requested = False
        self._update_success_control_buttons()
        for label in labels:
            label.place_forget()
        bloom = getattr(self, "success_bloom_label", None)
        if bloom is not None:
            bloom.place_forget()

        reduced = self._prefers_reduced_motion()
        positions = (0.10, 0.23, 0.36, 0.49, 0.62, 0.75, 0.88)
        walk_delay = 180 if reduced else 320
        visible_ms = 640 if reduced else 1200

        for index, (label, relx) in enumerate(zip(labels, positions)):
            def show_flower(widget=label, x=relx):
                if not getattr(self, "_success_animation_running", False) or getattr(self, "_success_animation_skip_requested", False):
                    return
                widget.place(relx=x, rely=0.5, anchor="center")
                widget.lift()
            def hide_flower(widget=label):
                if getattr(self, "_success_animation_skip_requested", False):
                    return
                widget.place_forget()
            self.after(index * walk_delay, show_flower)
            self.after(index * walk_delay + visible_ms, hide_flower)

        bloom_start = 90 if reduced else len(labels) * walk_delay + 140

        def show_bloom():
            if not getattr(self, "_success_animation_running", False) or getattr(self, "_success_animation_skip_requested", False):
                return
            bloom_widget = getattr(self, "success_bloom_label", None)
            if bloom_widget is not None:
                bloom_widget.place(relx=0.5, rely=0.5, anchor="center")
                bloom_widget.lift()
            self._play_success_sound()

        def hide_bloom():
            if getattr(self, "_success_animation_skip_requested", False):
                return
            bloom_widget = getattr(self, "success_bloom_label", None)
            if bloom_widget is not None:
                bloom_widget.place_forget()
            self._finish_success_animation()

        self.after(bloom_start, show_bloom)
        self.after(bloom_start + (650 if reduced else 1050), hide_bloom)

    def _mission_snapshot(self, force: bool = False) -> dict:
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            self._mission_snapshot_cache_key = ("none",)
            self._mission_snapshot_cache = {
                "issues": [], "routes": [], "review_count": 0, "blocked_route_count": 0,
                "workflow_blocked": False, "line_count": 0, "ready_count": 0,
                "event_name": "", "report_mode": "", "employee_totals": [],
                "employee_count": 0, "employee_grand_total": 0.0,
            }
            return self._mission_snapshot_cache
        if REVIEW_STATE_FILE.exists() and not self._is_archived_report_reopen():
            active_valid, active_reason = self._active_packet_is_valid(self.last_review_workbook)
            if not active_valid:
                self._mission_snapshot_cache_key = ("unverified-active-packet", _file_stamp(self.last_review_workbook))
                self._mission_snapshot_cache = {
                    "issues": [], "routes": [], "review_count": 0, "blocked_route_count": 0,
                    "workflow_blocked": True, "line_count": 0, "ready_count": 0,
                    "event_name": "", "report_mode": "", "employee_totals": [],
                    "employee_count": 0, "employee_grand_total": 0.0,
                    "active_packet_error": active_reason,
                }
                return self._mission_snapshot_cache
        cache_key = (
            _file_stamp(self.last_review_workbook),
            _file_stamp(DATA / "mission_control_state.json"),
            _file_stamp(REVIEW_STATE_FILE),
            # Product Master is part of the live review decision.  A workbook
            # warning can be resolved after the workbook was created, so its
            # saved file stamp must invalidate the snapshot as well.  Without
            # this, the dashboard can continue to show an old "Needs Setup"
            # count even while Product Master correctly says every style is
            # complete.
            _file_stamp(live_product_master_path()),
        )
        if not force and cache_key == self._mission_snapshot_cache_key:
            snapshot = self._overlay_completed_review_state(
                self._mission_snapshot_cache, self.last_review_workbook
            )
            self._mission_snapshot_cache = snapshot
            return snapshot
        try:
            snapshot = load_mission_control_snapshot(
                self.last_review_workbook, DATA, live_product_master_path()
            )
            snapshot = self._overlay_completed_review_state(
                snapshot, self.last_review_workbook
            )
            self._mission_snapshot_cache_key = cache_key
            self._mission_snapshot_cache = snapshot
            return snapshot
        except Exception as error:
            return {"issues": [], "routes": [], "review_count": 0, "blocked_route_count": 0,
                    "workflow_blocked": False, "line_count": 0, "ready_count": 0,
                    "event_name": "", "report_mode": "", "employee_totals": [],
                    "employee_count": 0, "employee_grand_total": 0.0, "error": str(error)}

    def _clear_children(self, frame):
        for child in frame.winfo_children():
            child.destroy()

    def _configure_dashboard_workspace(self, snapshot: dict) -> bool:
        """Expose Employee Totals as a sidebar page only for sizing events."""
        show_employee_totals = should_show_employee_totals(
            snapshot,
            bool(self.last_review_workbook and self.last_review_workbook.exists()),
            self.current_mode,
            self.imported_line_count,
        )
        if hasattr(self, "employee_nav_button"):
            if show_employee_totals:
                self.employee_nav_button.grid()
            else:
                self.employee_nav_button.grid_remove()
        if hasattr(self, "outsourced_logo_nav_button"):
            # The job-report logo is used by standard outsourced screen-print
            # reports as well as Entire Order Outsourced packets, so keep this
            # setup page permanently available in the sidebar.
            self.outsourced_logo_nav_button.grid()
        return show_employee_totals

    def _render_work_queue(self, snapshot: dict):
        self._clear_children(self.mc_work_scroll)
        issues = list(snapshot.get("issues", []))
        routes = list(snapshot.get("routes", []))
        line_count = int(snapshot.get("line_count", 0) or 0)
        blocked_routes = [
            route for route in routes
            if str(route.get("status", "")).casefold() != "ready"
        ]
        if not self.last_review_workbook or line_count == 0:
            self.mc_work_summary.configure(text="No active purchase packet.")
            ctk.CTkLabel(
                self.mc_work_scroll,
                text="Import Shopify orders and create a Purchase Review to begin.",
                text_color=MUTED, font=ctk.CTkFont(size=12),
            ).grid(row=0, column=0, padx=12, pady=22, sticky="w")
            return

        product_master_issues = [
            issue for issue in issues
            if str(issue.get("fix_in", "")).casefold() == "product master"
        ]
        purchase_review_issues = [
            issue for issue in issues
            if str(issue.get("fix_in", "")).casefold() != "product master"
        ]
        if not product_master_issues and not purchase_review_issues and not blocked_routes:
            self.mc_work_summary.configure(text="No blockers remain. This purchase packet is ready.")
            ctk.CTkLabel(
                self.mc_work_scroll, text="✓  Purchase Review complete", text_color=SUCCESS,
                font=ctk.CTkFont(size=14, weight="bold"),
            ).grid(row=0, column=0, padx=12, pady=22, sticky="w")
            return

        self.mc_work_summary.configure(
            text=(
                f"{len(product_master_issues)} new product setup item(s)  •  "
                f"{len(purchase_review_issues)} order review item(s)  •  "
                f"{len(blocked_routes)} blocked route(s)"
            )
        )
        self.mc_work_scroll.grid_columnconfigure(0, weight=1)
        row_cursor = 0

        def add_section(title: str, subtitle: str, rows: list[dict], kind: str, limit: int = 6):
            nonlocal row_cursor
            if not rows:
                return
            section = ctk.CTkFrame(
                self.mc_work_scroll, fg_color=WHITE, border_width=1,
                border_color="#E5DDEE", corner_radius=10,
            )
            section.grid(row=row_cursor, column=0, sticky="ew", padx=4, pady=(4, 8))
            section.grid_columnconfigure(0, weight=1)
            header = ctk.CTkFrame(section, fg_color=PURPLE_LIGHT, corner_radius=8)
            header.grid(row=0, column=0, sticky="ew", padx=7, pady=7)
            header.grid_columnconfigure(0, weight=1)
            ctk.CTkLabel(
                header, text=f"{title}  ({len(rows)})", text_color=PURPLE_DARK,
                font=ctk.CTkFont(size=13, weight="bold"), anchor="w",
            ).grid(row=0, column=0, padx=10, pady=(7, 1), sticky="w")
            ctk.CTkLabel(
                header, text=subtitle, text_color=MUTED,
                font=ctk.CTkFont(size=10), anchor="w",
            ).grid(row=1, column=0, padx=10, pady=(0, 7), sticky="w")

            for index, item in enumerate(rows[:limit], start=1):
                row_frame = ctk.CTkFrame(section, fg_color="#FCFBFE", corner_radius=8)
                row_frame.grid(row=index, column=0, sticky="ew", padx=7, pady=(0, 6))
                row_frame.grid_columnconfigure(1, weight=1)
                if kind == "master":
                    badge_text, badge_color, badge_fill = "NEW STYLE", WARNING, "#FFF1DC"
                    product = " — ".join(part for part in [item.get("product"), item.get("description")] if part)
                    action_text = item.get("instructions") or item.get("reason") or "Complete Product Master setup"
                    product_number = item.get("product") or ""
                    command = lambda value=product_number: self.open_product_master(search_product=value)
                    button_text = "Open Product"
                elif kind == "route":
                    badge_text, badge_color, badge_fill = "BLOCKED ROUTE", DANGER, "#FDECEC"
                    product = f"{item.get('vendor', 'Unknown vendor')} — {item.get('report_type', 'Needs routing')}"
                    action_text = "Resolve the remaining product or order decision that is blocking this purchase order."
                    command = self.open_latest_review
                    button_text = "Open Review"
                else:
                    badge_text, badge_color, badge_fill = "ORDER REVIEW", WARNING, "#FFF1DC"
                    product = " — ".join(part for part in [item.get("product"), item.get("description")] if part)
                    person = " • ".join(part for part in [item.get("employee"), item.get("order")] if part)
                    if person:
                        product = product + ("\n" if product else "") + person
                    action_text = item.get("instructions") or item.get("reason") or "Complete the order-specific correction"
                    command = self.open_latest_review
                    button_text = "Open Review"
                ctk.CTkLabel(
                    row_frame, text=badge_text, width=96, height=25, corner_radius=12,
                    fg_color=badge_fill, text_color=badge_color,
                    font=ctk.CTkFont(size=9, weight="bold"),
                ).grid(row=0, column=0, rowspan=2, padx=(9, 8), pady=9, sticky="w")
                ctk.CTkLabel(
                    row_frame, text=product or "Unknown item", text_color=TEXT,
                    font=ctk.CTkFont(size=11, weight="bold"), anchor="w", justify="left",
                ).grid(row=0, column=1, padx=4, pady=(7, 1), sticky="ew")
                ctk.CTkLabel(
                    row_frame, text=action_text, text_color=MUTED,
                    font=ctk.CTkFont(size=10), anchor="w", justify="left", wraplength=570,
                ).grid(row=1, column=1, padx=4, pady=(0, 7), sticky="ew")
                ctk.CTkButton(
                    row_frame, text=button_text, width=105, height=30,
                    fg_color=PURPLE_LIGHT, hover_color="#E4D6F5", text_color=PURPLE_DARK,
                    command=command,
                ).grid(row=0, column=2, rowspan=2, padx=9, pady=9, sticky="e")
            if len(rows) > limit:
                ctk.CTkLabel(
                    section, text=f"+ {len(rows) - limit} more", text_color=MUTED,
                    font=ctk.CTkFont(size=10, weight="bold"),
                ).grid(row=limit + 1, column=0, padx=12, pady=(0, 8), sticky="w")
            row_cursor += 1

        add_section(
            "Product Master", "Permanent setup saved for future orders.",
            product_master_issues, "master",
        )
        add_section(
            "Purchase Review", "Corrections that apply only to the current orders.",
            purchase_review_issues, "review",
        )
        add_section(
            "Blocked Purchase Orders", "Routes waiting on a remaining decision.",
            blocked_routes, "route", limit=5,
        )

    @staticmethod
    def _money_text(value: object) -> str:
        try:
            return f"${float(value):,.2f}"
        except (TypeError, ValueError):
            return "$0.00"

    @staticmethod
    def _discount_value(value: object) -> float:
        text = str(value or "").replace("$", "").replace(",", "").strip()
        if not text:
            return 0.0
        amount = float(text)
        if amount < 0:
            raise ValueError("Discounts cannot be negative.")
        return amount

    def _render_employee_totals(self, snapshot: dict):
        records = list(snapshot.get("employee_totals", []))
        render_key = (
            _file_stamp(self.last_review_workbook),
            str(snapshot.get("report_mode", "") or ""),
            tuple(
                (
                    record.get("row"), record.get("employee"), record.get("company"),
                    record.get("order_numbers"), record.get("order_total"), record.get("discount"), record.get("total"),
                )
                for record in records
            ),
        )
        # Avoid rebuilding the editable table when nothing in the saved source
        # changed.  This keeps the page responsive and preserves a discount the
        # user is actively typing until they choose Save Totals.
        if render_key == getattr(self, "_employee_totals_render_key", None):
            return
        self._clear_children(self.mc_employee_scroll)
        self.employee_discount_vars = {}
        self.employee_total_labels = {}
        self._employee_totals_render_key = render_key
        self.employee_records_by_row = {
            int(record.get("row", index + 1)): dict(record)
            for index, record in enumerate(records)
        }

        report_mode = str(snapshot.get("report_mode", "") or "")
        if not self.last_review_workbook:
            self.mc_employee_summary.configure(
                text="Create a Uniform Sizing Event Purchase Review to see employee totals."
            )
            self.mc_employee_grand_total.configure(text="Grand Total  $0.00")
            self.mc_employee_save_button.configure(state="disabled")
            ctk.CTkLabel(
                self.mc_employee_scroll,
                text="Employee totals will appear here after the Purchase Review is created.",
                text_color=MUTED, font=ctk.CTkFont(size=12), wraplength=360, justify="left",
            ).grid(row=0, column=0, padx=12, pady=22, sticky="w")
            return
        if report_mode != UNIFORM_SIZING_EVENT:
            self.mc_employee_summary.configure(text="Employee totals are available for Uniform Sizing Events.")
            self.mc_employee_grand_total.configure(text="Grand Total  $0.00")
            self.mc_employee_save_button.configure(state="disabled")
            ctk.CTkLabel(
                self.mc_employee_scroll,
                text="This purchase packet is using General Sales Period mode.",
                text_color=MUTED, font=ctk.CTkFont(size=12), wraplength=360, justify="left",
            ).grid(row=0, column=0, padx=12, pady=22, sticky="w")
            return
        if not records:
            self.mc_employee_summary.configure(text="No employee totals were found in this Purchase Review.")
            self.mc_employee_grand_total.configure(text="Grand Total  $0.00")
            self.mc_employee_save_button.configure(state="disabled")
            return

        self.mc_employee_summary.configure(
            text=f"{len(records)} employee{'s' if len(records) != 1 else ''} • Enter discounts below, then save."
        )
        self.mc_employee_save_button.configure(state="normal")
        headers = ["Employee", "Order Total", "Discount", "Total"]
        for column, header in enumerate(headers):
            ctk.CTkLabel(
                self.mc_employee_scroll, text=header, text_color=PURPLE_DARK,
                font=ctk.CTkFont(size=10, weight="bold"), anchor="w",
            ).grid(row=0, column=column, padx=5, pady=(4, 7), sticky="ew")
        self.mc_employee_scroll.grid_columnconfigure(0, weight=1)

        for index, record in enumerate(records, start=1):
            row_number = int(record.get("row", index))
            employee = str(record.get("employee", "") or "").strip()
            company = str(record.get("company", "") or "").strip()
            orders = str(record.get("order_numbers", "") or "").strip()
            primary = employee or company or "Unnamed employee"
            secondary_parts = []
            if company and company.casefold() != primary.casefold():
                secondary_parts.append(company)
            if orders:
                secondary_parts.append(orders)
            display = primary + ("\n" + " • ".join(secondary_parts) if secondary_parts else "")
            ctk.CTkLabel(
                self.mc_employee_scroll, text=display, text_color=TEXT,
                font=ctk.CTkFont(size=10, weight="bold"), anchor="w", justify="left", wraplength=170,
            ).grid(row=index, column=0, padx=5, pady=5, sticky="w")
            ctk.CTkLabel(
                self.mc_employee_scroll, text=self._money_text(record.get("order_total", 0)),
                text_color=TEXT, font=ctk.CTkFont(size=10), anchor="e",
            ).grid(row=index, column=1, padx=5, pady=5, sticky="e")
            discount = float(record.get("discount", 0) or 0)
            variable = ctk.StringVar(value=f"{discount:.2f}")
            self.employee_discount_vars[row_number] = variable
            entry = ctk.CTkEntry(
                self.mc_employee_scroll, textvariable=variable, width=82, height=30,
                border_color="#D9D0E6", fg_color=WHITE, justify="right",
            )
            entry.grid(row=index, column=2, padx=5, pady=5, sticky="e")
            total_label = ctk.CTkLabel(
                self.mc_employee_scroll, text=self._money_text(record.get("total", 0)),
                text_color=PURPLE_DARK, font=ctk.CTkFont(size=10, weight="bold"), anchor="e",
            )
            total_label.grid(row=index, column=3, padx=5, pady=5, sticky="e")
            self.employee_total_labels[row_number] = total_label
            variable.trace_add("write", lambda *_args, row=row_number: self._update_employee_total_row(row))

        self._refresh_employee_grand_total()

    def _update_employee_total_row(self, row_number: int):
        record = self.employee_records_by_row.get(row_number)
        variable = self.employee_discount_vars.get(row_number)
        label = self.employee_total_labels.get(row_number)
        if not record or variable is None or label is None:
            return
        try:
            discount = self._discount_value(variable.get())
            total = float(record.get("order_total", 0) or 0) - discount
            label.configure(text=self._money_text(total), text_color=PURPLE_DARK)
        except (TypeError, ValueError):
            label.configure(text="Invalid", text_color=DANGER)
        self._refresh_employee_grand_total()

    def _refresh_employee_grand_total(self):
        grand_total = 0.0
        valid = True
        for row_number, record in self.employee_records_by_row.items():
            variable = self.employee_discount_vars.get(row_number)
            try:
                discount = self._discount_value(variable.get() if variable else 0)
                grand_total += float(record.get("order_total", 0) or 0) - discount
            except (TypeError, ValueError):
                valid = False
        if valid:
            self.mc_employee_grand_total.configure(
                text=f"Grand Total  {self._money_text(grand_total)}", text_color=PURPLE_DARK,
            )
        else:
            self.mc_employee_grand_total.configure(text="Grand Total  Check discounts", text_color=DANGER)

    def save_dashboard_employee_totals(self):
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            messagebox.showinfo("No Purchase Review", "Create a Purchase Review first.")
            return
        discounts: dict[int, float] = {}
        try:
            for row_number, variable in self.employee_discount_vars.items():
                discounts[row_number] = self._discount_value(variable.get())
        except ValueError as error:
            messagebox.showwarning("Invalid Discount", str(error) or "Enter a valid discount amount.")
            return
        try:
            backup = save_employee_discounts(self.last_review_workbook, discounts)
        except PermissionError:
            messagebox.showwarning(
                "Close the Excel Workbook",
                "Excel is currently using this Purchase Review. Close the workbook, then click Save Totals again.",
            )
            return
        except Exception as error:
            messagebox.showerror("Employee Totals Not Saved", str(error))
            return
        self.refresh_dashboard()
        messagebox.showinfo(
            "Employee Totals Saved",
            f"Saved {len(discounts)} employee discount adjustment(s).\n\nA backup was created at:\n{backup}",
        )

    def _employee_totals_pdf_path(self):
        from modules.final_po_generator import generate_employee_totals_pdf

        if not self.last_review_workbook or not self.last_review_workbook.exists():
            raise FileNotFoundError("Create a Uniform Sizing Event purchase packet first.")
        if not self._require_active_packet("Employee Totals"):
            raise RuntimeError("Employee Totals was blocked because the active packet could not be verified.")
        # Save onscreen discounts first so the shared PDF is current.
        if self.employee_discount_vars:
            discounts = {row: self._discount_value(var.get()) for row, var in self.employee_discount_vars.items()}
            save_employee_discounts(self.last_review_workbook, discounts)
        return generate_employee_totals_pdf(self.last_review_workbook)

    def view_employee_totals_pdf(self):
        try:
            path = self._employee_totals_pdf_path()
            subprocess.run(["open", str(path)], check=False)
        except Exception as error:
            messagebox.showerror("Unable to View Employee Totals", str(error))

    def share_employee_totals_pdf(self):
        try:
            path = self._employee_totals_pdf_path()
            subject = f"Employee Totals — {self.current_event_name or 'Uniform Sizing Event'}"
            script = '''
            on run argv
                set attachmentFile to POSIX file (item 1 of argv)
                set messageSubject to item 2 of argv
                tell application "Mail"
                    activate
                    set newMessage to make new outgoing message with properties {subject:messageSubject, visible:true}
                    tell content of newMessage to make new attachment with properties {file name:attachmentFile} at after the last paragraph
                end tell
            end run
            '''
            subprocess.run(["osascript", "-e", script, str(path), subject], check=True)
        except Exception as error:
            messagebox.showerror("Unable to Share Employee Totals", str(error))

    def _outsourced_job_logo_path(self):
        if not self.last_review_workbook:
            return None
        stem = self.last_review_workbook.stem
        event_key = safe_filename(self.current_event_name or "Current Event")
        for ext in (".png", ".jpg", ".jpeg", ".webp"):
            stable = self.last_review_workbook.parent / f"{event_key}__Outsourced_Job_Logo{ext}"
            if stable.exists():
                return stable
            candidate = self.last_review_workbook.parent / f"{stem}__Outsourced_Job_Logo{ext}"
            if candidate.exists():
                return candidate
        return None

    def _outsourced_job_logo_metadata_path(self, workbook_path: Path | None = None, event_name: str = ""):
        """Return the sidecar file that keeps outsourced-cover setup with the event."""
        workbook = Path(workbook_path or self.last_review_workbook) if (workbook_path or self.last_review_workbook) else None
        if not workbook:
            return None
        event_key = safe_filename(event_name or self.current_event_name or "Current Event")
        return workbook.parent / f"{event_key}__Outsourced_Job_Logo_Metadata.json"

    def _outsourced_job_metadata(self) -> dict:
        """Load the saved job-cover setup, including per-color artwork files.

        Older events have a metadata file containing only ``outsourced_job_name``.
        Treat that as a valid first version so archived packets continue to open.
        """
        if not self.last_review_workbook:
            return {}
        workbook = Path(self.last_review_workbook)
        event_key = safe_filename(self.current_event_name or "Current Event")
        candidates = [
            workbook.parent / f"{event_key}__Outsourced_Job_Logo_Metadata.json",
            workbook.parent / f"{workbook.stem}__Outsourced_Job_Logo_Metadata.json",
        ]
        for candidate in candidates:
            try:
                payload = json.loads(candidate.read_text(encoding="utf-8"))
            except (OSError, ValueError, TypeError, AttributeError):
                continue
            if isinstance(payload, dict):
                payload.setdefault("artwork", {})
                if not isinstance(payload["artwork"], dict):
                    payload["artwork"] = {}
                return payload
        return {"artwork": {}}

    def _save_outsourced_job_metadata(self, payload: dict) -> bool:
        metadata_path = self._outsourced_job_logo_metadata_path()
        if not metadata_path:
            return False
        payload = dict(payload or {})
        payload["updated_at"] = time.time()
        payload.setdefault("artwork", {})
        try:
            metadata_path.parent.mkdir(parents=True, exist_ok=True)
            temporary = metadata_path.with_name(
                f".{metadata_path.name}.{os.getpid()}.{int(time.time() * 1000)}.tmp"
            )
            try:
                temporary.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
                os.replace(temporary, metadata_path)
            finally:
                if temporary.exists():
                    temporary.unlink(missing_ok=True)
            return True
        except OSError:
            return False

    def _outsourced_job_name(self) -> str:
        """Load the optional production name recognized by the outside decorator."""
        payload = self._outsourced_job_metadata()
        return _clean(payload.get("outsourced_job_name", payload.get("job_name", "")))

    def _outsourced_cover_notes(self) -> str:
        return _clean(self._outsourced_job_metadata().get("cover_notes", ""))

    def _outsourced_to_for_department(self, department: str = "", payload: dict | None = None) -> str:
        """Return the event's destination for one outsourced report department.

        The destination is intentionally an event setting rather than a Product
        Master field: the same product can go to a different outside decorator
        from one job to the next.  Older packets did not save a destination, so
        they keep Orchid's established Stitch N Print default.
        """
        department = self._outsourced_report_department(department) or "embroidery"
        metadata = payload if isinstance(payload, dict) else self._outsourced_job_metadata()
        saved = metadata.get("outsourced_to", {}) if isinstance(metadata, dict) else {}
        if isinstance(saved, dict):
            value = _clean(saved.get(department, ""))
        else:
            # Be tolerant of an early single-value draft if one is ever saved.
            value = _clean(saved)
        return value or "Stitch N Print"

    def save_outsourced_job_name(self, show_confirmation: bool = True) -> str:
        """Persist the printable job/logo name beside the active event logo."""
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            if show_confirmation:
                messagebox.showinfo("Outsourced Job Name", "Create the Purchase Review before saving the outsourced job name.")
            return ""
        if hasattr(self, "outsourced_logo_name_var"):
            name = _clean(self.outsourced_logo_name_var.get())[:160]
        else:
            name = self._outsourced_job_name()
        payload = self._outsourced_job_metadata()
        payload["outsourced_job_name"] = name
        if not self._save_outsourced_job_metadata(payload):
            if show_confirmation:
                messagebox.showerror("Unable to Save Job Name", "Orchid could not save the outsourced job-cover details.")
            return ""
        if hasattr(self, "dashboard_status"):
            summary = name or "Use the event name"
            self.dashboard_status.configure(text=f"Outsourced report name saved: {summary}", text_color=SUCCESS)
        if show_confirmation:
            messagebox.showinfo(
                "Outsourced Job Name Saved",
                (f"The outsourced report will use “{name}” as its job name."
                 if name else "The outsourced report will use the event name until a job name is entered."),
            )
        self.refresh_outsourced_job_logo_page()
        return name

    @staticmethod
    def _artwork_label(decoration_type: str, color: str) -> str:
        suffix = "Thread" if "embroider" in decoration_type.casefold() else "Ink" if "screen" in decoration_type.casefold() else "Artwork"
        return f"{_clean(color)} {suffix}".strip()

    @staticmethod
    def _outsourced_report_department(value: str) -> str:
        text = _clean(value).casefold()
        if "embroider" in text:
            return "embroidery"
        if "screen" in text:
            return "screen-printing"
        return ""

    def _active_outsourced_report_department(self) -> str:
        selection = (
            self.outsourced_report_type_var.get()
            if hasattr(self, "outsourced_report_type_var")
            else "Embroidery"
        )
        return self._outsourced_report_department(selection) or "embroidery"

    def _select_outsourced_report_department(self, selection: str) -> None:
        """Switch the Logo-page department with an unmistakable visual state."""
        if selection not in {"Embroidery", "Screen Printing"}:
            return
        self.outsourced_report_type_var.set(selection)
        self.refresh_outsourced_job_logo_page()

    def _configure_outsourced_report_department_tabs(self, state: str) -> None:
        """Give embroidery and screen printing independent, high-contrast tab states."""
        if not hasattr(self, "outsourced_embroidery_tab") or not hasattr(self, "outsourced_screen_printing_tab"):
            return
        selected = self.outsourced_report_type_var.get()
        embroidery_selected = selected == "Embroidery"
        screen_printing_selected = selected == "Screen Printing"
        disabled = state == "disabled"
        self.outsourced_embroidery_tab.configure(
            state=state,
            fg_color="#5B21B6" if embroidery_selected and not disabled else ("#E8E1F3" if disabled else "#F2ECFA"),
            hover_color="#4C1D95" if embroidery_selected else "#E4D8F3",
            border_color="#5B21B6" if embroidery_selected else "#C8B4E4",
            text_color=WHITE if embroidery_selected and not disabled else ("#8A8490" if disabled else "#43206F"),
        )
        self.outsourced_screen_printing_tab.configure(
            state=state,
            fg_color="#1E4F7A" if screen_printing_selected and not disabled else ("#E1E7EC" if disabled else "#EAF3FA"),
            hover_color="#173E62" if screen_printing_selected else "#D7E8F5",
            border_color="#1E4F7A" if screen_printing_selected else "#AFC7D9",
            text_color=WHITE if screen_printing_selected and not disabled else ("#8A8490" if disabled else "#173E62"),
        )

    @staticmethod
    def _outsourced_report_department_label(department: str) -> str:
        return "Screen Printing" if department == "screen-printing" else "Embroidery"

    def _outsourced_cover_artwork_reference_spec(self, department: str) -> dict:
        """Return the always-available artwork row for one report department.

        A job can need its artwork loaded before its finalized thread/ink colors
        are present in the Purchase Review.  This department-specific reference
        keeps both tabs useful at every stage while preserving the optional,
        color-specific artwork rows below it.
        """
        selected = self._outsourced_report_department(department)
        label = self._outsourced_report_department_label(selected)
        return {
            "key": f"{selected}::artwork-reference",
            "decoration_type": label,
            "color": "",
            "label": f"{label} Artwork Reference",
            "locations": set(),
            "department": selected,
            "reference_only": True,
        }

    def _outsourced_cover_artwork_specs(self, department: str = "") -> list[dict]:
        """Return the actual color/type artwork rows needed by the active event."""
        if not self.last_review_workbook or not Path(self.last_review_workbook).is_file():
            return []
        try:
            fulfillment = load_decoration_fulfillment(self.last_review_workbook)
            records = load_review_lines(self.last_review_workbook)
        except Exception:
            return []
        selected_department = self._outsourced_report_department(department)
        grouped: dict[str, dict] = {}
        for row in records:
            decoration_type = _clean(row.get("Operational Decoration Type") or row.get("Decoration Type"))
            color = _clean(row.get("Operational Decoration Color") or row.get("Decoration Color"))
            location = _clean(row.get("Operational Decoration Location") or row.get("Decoration Location"))
            if not decoration_type or not color:
                continue
            report_department = self._outsourced_report_department(decoration_type)
            if not report_department or (selected_department and report_department != selected_department):
                continue
            if not is_outsourced_decoration(
                decoration_type, fulfillment,
                do_not_outsource=_clean(row.get("Do Not Outsource", "")).casefold() in {"yes", "true", "1"},
            ):
                continue
            key = f"{decoration_type.casefold()}::{color.casefold()}"
            item = grouped.setdefault(key, {
                "key": key,
                "decoration_type": decoration_type,
                "color": color,
                "label": self._artwork_label(decoration_type, color),
                "locations": set(),
                "department": report_department,
            })
            if location:
                item["locations"].add(location)
        return sorted(
            grouped.values(),
            key=lambda item: (item["decoration_type"].casefold(), item["color"].casefold()),
        )

    def _outsourced_cover_artwork_entries(self, department: str = "") -> list[dict]:
        """Resolve the artwork that belongs on an outsourced report cover."""
        if not self.last_review_workbook:
            return []
        metadata = self._outsourced_job_metadata()
        artwork = metadata.get("artwork", {})
        parent = Path(self.last_review_workbook).parent
        entries = []
        departments = (
            [self._outsourced_report_department(department)]
            if department else ["embroidery", "screen-printing"]
        )
        for selected_department in departments:
            department_entries = []
            for spec in self._outsourced_cover_artwork_specs(selected_department):
                saved = artwork.get(spec["key"], {}) if isinstance(artwork, dict) else {}
                filename = _clean(saved.get("filename", "")) if isinstance(saved, dict) else ""
                candidate = parent / filename if filename and Path(filename).name == filename else None
                if candidate and candidate.is_file():
                    department_entries.append({**spec, "path": candidate})
            # A department-wide reference is the practical fallback when a job
            # has only one logo or when its final thread/ink colors are not
            # available yet. It deliberately stays off that department's cover
            # once color-specific artwork has been supplied, avoiding duplicate
            # artwork on the same report.
            if not department_entries:
                reference = self._outsourced_cover_artwork_reference_spec(selected_department)
                saved = artwork.get(reference["key"], {}) if isinstance(artwork, dict) else {}
                filename = _clean(saved.get("filename", "")) if isinstance(saved, dict) else ""
                candidate = parent / filename if filename and Path(filename).name == filename else None
                if candidate and candidate.is_file():
                    department_entries.append({**reference, "path": candidate})
            entries.extend(department_entries)
        # Legacy events with one uploaded job logo continue to render usable
        # artwork until their next report regeneration.
        if not entries:
            legacy = self._outsourced_job_logo_path()
            if legacy and legacy.is_file():
                entries.append({
                    "key": "legacy-job-logo", "decoration_type": "", "color": "",
                    "label": "Artwork Reference", "locations": set(), "path": legacy,
                })
        return entries

    def save_outsourced_cover_details(self, show_confirmation: bool = True) -> bool:
        if not self.last_review_workbook or not Path(self.last_review_workbook).is_file():
            if show_confirmation:
                messagebox.showinfo("Outsourced Report Cover", "Create the Purchase Review before saving cover details.")
            return False
        payload = self._outsourced_job_metadata()
        payload["outsourced_job_name"] = _clean(self.outsourced_logo_name_var.get())[:160]
        payload["cover_notes"] = _clean(self.outsourced_cover_notes_text.get("1.0", "end"))[:1200]
        department = self._active_outsourced_report_department()
        destinations = payload.get("outsourced_to", {})
        if not isinstance(destinations, dict):
            destinations = {}
        destinations[department] = _clean(self.outsourced_to_var.get())[:160] or "Stitch N Print"
        payload["outsourced_to"] = destinations
        if not self._save_outsourced_job_metadata(payload):
            if show_confirmation:
                messagebox.showerror("Unable to Save Cover Details", "Orchid could not save the outsourced report cover details.")
            return False
        if show_confirmation:
            destination = self._outsourced_to_for_department(department, payload)
            messagebox.showinfo(
                "Outsourced Report Cover Saved",
                f"Job name, {self._outsourced_report_department_label(department)} destination ({destination}), and special notes are ready for the next report generation.",
            )
        self.refresh_outsourced_job_logo_page()
        return True

    def build_outsourced_job_logo_page(self):
        page = self.new_page("job_logo")
        card = ctk.CTkFrame(page, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=18)
        card.grid(row=0, column=0, sticky="nsew", padx=34, pady=30)
        card.grid_columnconfigure(0, weight=1)
        page.grid_rowconfigure(0, weight=1)
        ctk.CTkLabel(
            card, text="Outsourced Report Covers", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=28, weight="bold"),
        ).grid(row=0, column=0, padx=30, pady=(34, 6))
        ctk.CTkLabel(
            card,
            text="Choose a report department, set its outsourced destination, then add only the artwork that belongs on that job report.",
            text_color=MUTED, font=ctk.CTkFont(size=14), wraplength=820, justify="center",
        ).grid(row=1, column=0, padx=30, pady=(0, 24))
        self.outsourced_report_type_var = ctk.StringVar(value="Embroidery")
        self.outsourced_report_type_tabs = ctk.CTkFrame(
            card, fg_color="#ECE7F1", border_width=1, border_color="#D4C9E3", corner_radius=13,
        )
        self.outsourced_report_type_tabs.grid(row=2, column=0, padx=30, pady=(0, 18))
        self.outsourced_embroidery_tab = ctk.CTkButton(
            self.outsourced_report_type_tabs, text="Embroidery",
            command=lambda: self._select_outsourced_report_department("Embroidery"),
            width=180, height=42, corner_radius=10, border_width=1,
            font=ctk.CTkFont(size=14, weight="bold"),
        )
        self.outsourced_embroidery_tab.grid(row=0, column=0, padx=(4, 2), pady=4)
        self.outsourced_screen_printing_tab = ctk.CTkButton(
            self.outsourced_report_type_tabs, text="Screen Printing",
            command=lambda: self._select_outsourced_report_department("Screen Printing"),
            width=180, height=42, corner_radius=10, border_width=1,
            font=ctk.CTkFont(size=14, weight="bold"),
        )
        self.outsourced_screen_printing_tab.grid(row=0, column=1, padx=(2, 4), pady=4)
        self._configure_outsourced_report_department_tabs("normal")
        cover_identity = ctk.CTkFrame(card, fg_color="#FAF7FD", border_width=1, border_color="#E1D5F0", corner_radius=14)
        cover_identity.grid(row=3, column=0, sticky="ew", padx=54, pady=(0, 12))
        cover_identity.grid_columnconfigure(0, weight=1)
        cover_identity.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(
            cover_identity, text="Job Name", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=13, weight="bold"), anchor="w",
        ).grid(row=0, column=0, sticky="ew", padx=(16, 8), pady=(12, 4))
        self.outsourced_to_heading = ctk.CTkLabel(
            cover_identity, text="Outsourced To — Embroidery", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=13, weight="bold"), anchor="w",
        )
        self.outsourced_to_heading.grid(row=0, column=1, sticky="ew", padx=(8, 16), pady=(12, 4))
        self.outsourced_logo_name_var = ctk.StringVar()
        self.outsourced_logo_name_entry = ctk.CTkEntry(
            cover_identity, textvariable=self.outsourced_logo_name_var, height=42,
            placeholder_text="Example: Utilities Department", border_color="#BFA9DD",
        )
        self.outsourced_logo_name_entry.grid(row=1, column=0, sticky="ew", padx=(16, 8), pady=(0, 12))
        self.outsourced_logo_name_entry.bind("<Return>", lambda _event: self.save_outsourced_cover_details())
        self.outsourced_to_var = ctk.StringVar()
        self.outsourced_to_entry = ctk.CTkEntry(
            cover_identity, textvariable=self.outsourced_to_var, height=42,
            placeholder_text="Stitch N Print", border_color=PURPLE, border_width=2,
        )
        self.outsourced_to_entry.grid(row=1, column=1, sticky="ew", padx=(8, 16), pady=(0, 12))
        self.outsourced_to_entry.bind("<Return>", lambda _event: self.save_outsourced_cover_details())
        self.outsourced_logo_name_save_button = ctk.CTkButton(
            card, text="Save Cover Details", command=self.save_outsourced_cover_details,
            width=190, height=42, fg_color=PURPLE_LIGHT, hover_color="#E6D9F5",
            border_width=1, border_color="#BFA9DD", text_color=PURPLE_DARK,
        )
        self.outsourced_logo_name_save_button.grid(row=4, column=0, padx=30, pady=(0, 8))
        ctk.CTkLabel(
            card, text="Job Name is the large production name. Outsourced To is saved separately for Embroidery and Screen Printing and prints prominently on that cover.",
            text_color=MUTED, font=ctk.CTkFont(size=12),
        ).grid(row=5, column=0, padx=30, pady=(0, 4))
        details = ctk.CTkFrame(card, fg_color="transparent")
        details.grid(row=6, column=0, sticky="ew", padx=54, pady=(14, 8))
        details.grid_columnconfigure(0, weight=1)
        details.grid_columnconfigure(1, weight=1)
        self.outsourced_logo_location_heading = ctk.CTkLabel(details, text="Embroidery Locations", text_color=PURPLE_DARK, font=ctk.CTkFont(size=13, weight="bold"), anchor="w")
        self.outsourced_logo_location_heading.grid(row=0, column=0, sticky="ew", padx=(0, 12), pady=(0, 5))
        ctk.CTkLabel(details, text="Special Notes", text_color=PURPLE_DARK, font=ctk.CTkFont(size=13, weight="bold"), anchor="w").grid(row=0, column=1, sticky="ew", padx=(12, 0), pady=(0, 5))
        self.outsourced_logo_location_summary = ctk.CTkLabel(details, text="Create a Purchase Review to read locations.", fg_color="#F7F3FC", corner_radius=12, text_color=TEXT, font=ctk.CTkFont(size=13, weight="bold"), justify="left", anchor="w", wraplength=355, height=92)
        self.outsourced_logo_location_summary.grid(row=1, column=0, sticky="nsew", padx=(0, 12))
        self.outsourced_cover_notes_text = ctk.CTkTextbox(details, height=92, fg_color="#F7F3FC", border_width=1, border_color="#E1D5F0", text_color=TEXT, font=ctk.CTkFont(size=13), wrap="word")
        self.outsourced_cover_notes_text.grid(row=1, column=1, sticky="nsew", padx=(12, 0))
        self.outsourced_artwork_heading = ctk.CTkLabel(card, text="Embroidery Artwork by Thread Color", text_color=PURPLE_DARK, font=ctk.CTkFont(size=17, weight="bold"))
        self.outsourced_artwork_heading.grid(row=7, column=0, padx=30, pady=(14, 2))
        self.outsourced_logo_status = ctk.CTkLabel(
            card, text="", text_color=TEXT, font=ctk.CTkFont(size=13, weight="bold"),
        )
        self.outsourced_logo_status.grid(row=8, column=0, padx=30, pady=(0, 8))
        self.outsourced_artwork_rows = ctk.CTkScrollableFrame(card, fg_color="#FBF9FE", border_width=1, border_color="#E1D5F0", corner_radius=14, height=262)
        self.outsourced_artwork_rows.grid(row=9, column=0, sticky="ew", padx=54, pady=(0, 16))
        self.outsourced_artwork_rows.grid_columnconfigure(0, weight=1)
        report_buttons = ctk.CTkFrame(card, fg_color="transparent")
        report_buttons.grid(row=10, column=0, pady=(0, 28))
        self.outsourced_logo_regenerate_button = ctk.CTkButton(
            report_buttons, text="Regenerate Purchase Review & Reports",
            command=self.regenerate_review_and_reports,
            width=300, height=46, fg_color=PURPLE, hover_color=PURPLE_DARK,
            font=ctk.CTkFont(size=13, weight="bold"),
        )
        self.outsourced_logo_regenerate_button.pack(side="left", padx=6)
        self.outsourced_logo_open_report_button = ctk.CTkButton(
            report_buttons, text="Open Embroidery Job Report",
            command=lambda: self._open_generated_report("outsourced-embroidery"),
            width=235, height=46, fg_color=PURPLE_LIGHT, hover_color="#E6D9F5",
            border_width=1, border_color="#BFA9DD", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=13, weight="bold"),
        )
        self.outsourced_logo_open_report_button.pack(side="left", padx=6)

    def refresh_outsourced_job_logo_page(self):
        if not hasattr(self, "outsourced_logo_status"):
            return
        applicable = bool(self.last_review_workbook and self.last_review_workbook.exists())
        department = self._active_outsourced_report_department()
        department_label = self._outsourced_report_department_label(department)
        artwork_kind = "Thread" if department == "embroidery" else "Ink"
        payload = self._outsourced_job_metadata() if applicable else {}
        job_name = _clean(payload.get("outsourced_job_name", payload.get("job_name", "")))
        notes = _clean(payload.get("cover_notes", ""))
        outsourced_to = self._outsourced_to_for_department(department, payload)
        state = "normal" if applicable else "disabled"
        self.outsourced_logo_name_entry.configure(state=state)
        self.outsourced_to_entry.configure(state=state)
        self.outsourced_logo_name_save_button.configure(state=state)
        self.outsourced_logo_name_var.set(job_name)
        self.outsourced_to_var.set(outsourced_to)
        self.outsourced_to_heading.configure(text=f"Outsourced To — {department_label}")
        self._configure_outsourced_report_department_tabs(state)
        self.outsourced_logo_location_heading.configure(text=f"{department_label} Locations")
        self.outsourced_artwork_heading.configure(text=f"{department_label} Artwork by {artwork_kind} Color")
        self.outsourced_cover_notes_text.configure(state="normal")
        self.outsourced_cover_notes_text.delete("1.0", "end")
        self.outsourced_cover_notes_text.insert("1.0", notes)
        self.outsourced_cover_notes_text.configure(state=state)
        self.outsourced_logo_regenerate_button.configure(
            state=state,
            text=(
                "Regenerate Archived Reports"
                if self._is_archived_report_reopen()
                else "Regenerate Purchase Review & Reports"
            ),
        )
        self.outsourced_logo_open_report_button.configure(
            text=f"Open {department_label} Job Report",
            command=lambda kind=f"outsourced-{department}": self._open_generated_report(kind),
            state=state,
        )
        self._clear_children(self.outsourced_artwork_rows)
        if not applicable:
            self.outsourced_logo_location_summary.configure(text="Create a Purchase Review to read this department's locations.")
            self.outsourced_logo_status.configure(text="Import orders and create the Purchase Review before setting up the report cover.", text_color=WARNING)
            ctk.CTkLabel(self.outsourced_artwork_rows, text=f"{department_label} artwork choices will appear once Orchid can see the job's outsourced {artwork_kind.casefold()} colors.", text_color=MUTED, font=ctk.CTkFont(size=13), wraplength=760, justify="center").grid(row=0, column=0, padx=20, pady=42)
            return
        specs = self._outsourced_cover_artwork_specs(department)
        # Keep a full Replace / Open / Remove row on *both* tabs, even before
        # Orchid can derive a department's color-specific artwork list.
        reference_spec = self._outsourced_cover_artwork_reference_spec(department)
        display_specs = [reference_spec, *specs]
        locations = sorted({location for spec in specs for location in spec["locations"]}, key=str.casefold)
        self.outsourced_logo_location_summary.configure(text=" • ".join(locations) if locations else f"No outsourced {department_label.casefold()} locations are currently in this job.")
        artwork = payload.get("artwork", {}) if isinstance(payload.get("artwork", {}), dict) else {}
        uploaded_count = 0
        for row_index, spec in enumerate(display_specs):
            saved = artwork.get(spec["key"], {}) if isinstance(artwork, dict) else {}
            filename = _clean(saved.get("filename", "")) if isinstance(saved, dict) else ""
            path = Path(self.last_review_workbook).parent / filename if filename and Path(filename).name == filename else None
            exists = bool(path and path.is_file())
            uploaded_count += int(exists)
            row = ctk.CTkFrame(self.outsourced_artwork_rows, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=11)
            row.grid(row=row_index, column=0, sticky="ew", padx=8, pady=5)
            row.grid_columnconfigure(0, weight=1)
            ctk.CTkLabel(row, text=spec["label"], text_color=PURPLE_DARK, font=ctk.CTkFont(size=14, weight="bold"), anchor="w").grid(row=0, column=0, sticky="ew", padx=(14, 8), pady=(9, 0))
            context = (
                f"Use one file for this entire {department_label.casefold()} job."
                if spec.get("reference_only")
                else " • ".join(filter(None, [spec["decoration_type"], ", ".join(sorted(spec["locations"], key=str.casefold))]))
            )
            ctk.CTkLabel(row, text=context or "Outsourced artwork", text_color=MUTED, font=ctk.CTkFont(size=11), anchor="w").grid(row=1, column=0, sticky="ew", padx=(14, 8), pady=(0, 9))
            ctk.CTkLabel(row, text="Artwork uploaded" if exists else "Artwork needed", text_color=SUCCESS if exists else WARNING, font=ctk.CTkFont(size=11, weight="bold")).grid(row=0, column=1, rowspan=2, padx=8, pady=8)
            ctk.CTkButton(row, text="Replace" if exists else "Upload", command=lambda item=spec: self.manage_outsourced_cover_artwork(item, "upload"), width=92, height=33, fg_color=PURPLE, hover_color=PURPLE_DARK).grid(row=0, column=2, rowspan=2, padx=4, pady=8)
            ctk.CTkButton(row, text="Open", command=lambda item=spec: self.manage_outsourced_cover_artwork(item, "preview"), width=74, height=33, state="normal" if exists else "disabled", fg_color=PURPLE_LIGHT, hover_color="#E6D9F5", border_width=1, border_color="#BFA9DD", text_color=PURPLE_DARK).grid(row=0, column=3, rowspan=2, padx=4, pady=8)
            ctk.CTkButton(row, text="Remove", command=lambda item=spec: self.manage_outsourced_cover_artwork(item, "remove"), width=82, height=33, state="normal" if exists else "disabled", fg_color="transparent", hover_color=PURPLE_LIGHT, border_width=1, border_color="#BFA9DD", text_color=PURPLE_DARK).grid(row=0, column=4, rowspan=2, padx=(4, 10), pady=8)
        if not specs:
            self.outsourced_logo_status.configure(
                text=(
                    f"Add the {department_label.casefold()} artwork now. "
                    "Color-specific rows will appear automatically once this job has them."
                ),
                text_color=MUTED,
            )
        else:
            self.outsourced_logo_status.configure(
                text=(
                    f"{uploaded_count} of {len(display_specs)} {department_label.casefold()} artwork file(s) uploaded. "
                    "Use the artwork reference when one file covers the whole job, or add color-specific artwork below it."
                ),
                text_color=SUCCESS if uploaded_count else MUTED,
            )

    def manage_outsourced_cover_artwork(self, spec: dict, action: str) -> None:
        if not self.last_review_workbook:
            return
        payload = self._outsourced_job_metadata()
        artwork = payload.setdefault("artwork", {})
        saved = artwork.get(spec["key"], {}) if isinstance(artwork, dict) else {}
        filename = _clean(saved.get("filename", "")) if isinstance(saved, dict) else ""
        current = Path(self.last_review_workbook).parent / filename if filename and Path(filename).name == filename else None
        if action == "preview":
            if current and current.is_file():
                subprocess.run(["open", str(current)], check=False)
            return
        if action == "remove":
            if not current or not current.is_file() or not messagebox.askyesno("Remove Cover Artwork?", f"Remove the {spec['label']} artwork from this event?"):
                return
            current.unlink(missing_ok=True)
            artwork.pop(spec["key"], None)
            self._save_outsourced_job_metadata(payload)
            self.refresh_outsourced_job_logo_page()
            return
        self.save_outsourced_cover_details(show_confirmation=False)
        self._outsourced_cover_artwork_dialog_spec = spec
        self.after(1, self._open_outsourced_cover_artwork_dialog)

    def _open_outsourced_cover_artwork_dialog(self) -> None:
        spec = getattr(self, "_outsourced_cover_artwork_dialog_spec", None)
        if not isinstance(spec, dict) or not self.last_review_workbook:
            return
        try:
            selected = self._choose_outsourced_job_logo_file()
            if not selected:
                return
            source = Path(selected)
            if source.suffix.casefold() not in {".png", ".jpg", ".jpeg", ".webp"}:
                messagebox.showwarning("Unsupported Artwork File", "Choose a PNG, JPG, JPEG, or WEBP image for the report cover.")
                return
            payload = self._outsourced_job_metadata()
            artwork = payload.setdefault("artwork", {})
            old = artwork.get(spec["key"], {}) if isinstance(artwork, dict) else {}
            old_name = _clean(old.get("filename", "")) if isinstance(old, dict) else ""
            event_key = safe_filename(self.current_event_name or "Current Event")
            # Save a normalized transparent PNG for the outsourced cover.
            # The original customer upload remains untouched; only Orchid's
            # event copy is cleaned so rectangular image backgrounds do not
            # print behind the logo.
            filename = f"{event_key}__Outsourced_Artwork__{safe_filename(spec['key'])}.png"
            target = Path(self.last_review_workbook).parent / filename
            prepared = target.with_name(f".{target.name}.{os.getpid()}.tmp")
            try:
                background_removed = prepare_artwork_for_report(source, prepared)
                os.replace(prepared, target)
            finally:
                if prepared.exists():
                    prepared.unlink(missing_ok=True)
            artwork[spec["key"]] = {"label": spec["label"], "filename": filename}
            if not self._save_outsourced_job_metadata(payload):
                raise RuntimeError("The artwork file copied, but its event record could not be saved.")
            if old_name and old_name != filename and Path(old_name).name == old_name:
                (Path(self.last_review_workbook).parent / old_name).unlink(missing_ok=True)
            self.dashboard_status.configure(
                text=(
                    f"Saved clean cover artwork: {spec['label']}"
                    if background_removed else f"Saved cover artwork: {spec['label']}"
                ),
                text_color=SUCCESS,
            )
        except Exception as error:
            messagebox.showerror("Unable to Save Cover Artwork", f"Orchid could not save that artwork.\n\n{error}")
        finally:
            self._outsourced_cover_artwork_dialog_spec = None
            self.after(120, self.refresh_outsourced_job_logo_page)

    def manage_outsourced_job_logo(self, action: str = "upload"):
        if not self.last_review_workbook or not self.last_review_workbook.is_file():
            messagebox.showinfo("Job Report Logo", "Create the Purchase Review before uploading the event logo.")
            return
        current = self._outsourced_job_logo_path()
        if action == "preview":
            if current:
                subprocess.run(["open", str(current)], check=False)
            return
        if action == "remove":
            if not current:
                return
            if not messagebox.askyesno("Remove Outsourced Job Logo?", f"Remove {current.name} from this event?"):
                return
            current.unlink(missing_ok=True)
            self.refresh_outsourced_job_logo_page()
            return

        # A user can type the production name and then upload the image without
        # needing a separate save step. The sidecar travels with the logo.
        self.save_outsourced_job_name(show_confirmation=False)

        # Let the button release first, then open exactly one chooser. On macOS
        # we use Finder's native chooser rather than Tk's picker; the Tk picker
        # can fail silently behind a CustomTkinter window on some Macs.
        if bool(getattr(self, "_outsourced_logo_dialog_pending", False)):
            return
        self._outsourced_logo_dialog_pending = True
        try:
            self.outsourced_logo_upload_button.configure(state="disabled")
            self.outsourced_logo_status.configure(
                text="Opening your Mac file picker…", text_color=MUTED
            )
            self.update_idletasks()
        except Exception:
            pass
        self.after(1, self._open_outsourced_job_logo_dialog)

    @staticmethod
    def _choose_outsourced_job_logo_file() -> str:
        """Return an image chosen in Finder, or an empty string when cancelled."""
        if sys.platform == "darwin":
            chooser_script = """
try
    POSIX path of (choose file with prompt \"Choose Job Report Logo\")
on error number -128
    return \"\"
end try
"""
            result = subprocess.run(
                ["osascript", "-e", chooser_script],
                text=True,
                capture_output=True,
                check=False,
            )
            if result.returncode == 0:
                return result.stdout.strip()
            message = (result.stderr or result.stdout or "Unable to open Finder.").strip()
            raise RuntimeError(message)
        return filedialog.askopenfilename(
            title="Choose Job Report Logo",
            filetypes=[
                ("Logo images", "*.png *.jpg *.jpeg *.webp"),
                ("PNG images", "*.png"), ("JPEG images", "*.jpg *.jpeg"),
            ],
        )

    def _open_outsourced_job_logo_dialog(self):
        try:
            selected = self._choose_outsourced_job_logo_file()
            if not selected:
                return
            source = Path(selected)
            if source.suffix.casefold() not in {".png", ".jpg", ".jpeg", ".webp"}:
                messagebox.showwarning(
                    "Unsupported Logo File",
                    "Choose a PNG, JPG, JPEG, or WEBP image. Embroidery machine files such as EXP cannot be placed on the PDF report.",
                )
                return
            event_key = safe_filename(self.current_event_name or "Current Event")
            old_paths = []
            for ext in (".png", ".jpg", ".jpeg", ".webp"):
                old = self.last_review_workbook.parent / f"{self.last_review_workbook.stem}__Outsourced_Job_Logo{ext}"
                stable_old = self.last_review_workbook.parent / f"{event_key}__Outsourced_Job_Logo{ext}"
                old_paths.extend((old, stable_old))
            # Store a clean transparent event copy. This is the same
            # non-destructive cleanup used for uploaded outsourced-cover art,
            # so Current Event and the PDF never display a rectangular canvas.
            target = self.last_review_workbook.parent / f"{event_key}__Outsourced_Job_Logo.png"
            prepared = target.with_name(f".{target.name}.{os.getpid()}.tmp")
            try:
                background_removed = prepare_artwork_for_report(source, prepared)
                os.replace(prepared, target)
            finally:
                if prepared.exists():
                    prepared.unlink(missing_ok=True)
            for old_path in old_paths:
                if old_path != target:
                    old_path.unlink(missing_ok=True)
            self.dashboard_status.configure(
                text=(
                    "Saved clean transparent job logo"
                    if background_removed else "Saved transparent job logo"
                ),
                text_color=SUCCESS,
            )
        except Exception as error:
            try:
                self.outsourced_logo_status.configure(
                    text="Unable to open the logo file picker. Please try Upload / Replace Logo again.",
                    text_color=WARNING,
                )
            except Exception:
                pass
            messagebox.showerror(
                "Unable to Open Logo Picker",
                f"Orchid could not open the file picker.\n\n{error}",
            )
        finally:
            self._outsourced_logo_dialog_pending = False
            # Refresh the Logo page and Current Event preview after Finder fully
            # releases its modal mouse grab.
            try:
                self.after(120, self.refresh_outsourced_job_logo_page)
                self.after(160, self.refresh_current_event_page)
            except Exception:
                self.refresh_outsourced_job_logo_page()
                self.refresh_current_event_page()

    def _render_po_manager(self, snapshot: dict):
        self._clear_children(self.mc_po_scroll)
        self.po_entry_vars = {}
        routes = snapshot.get("routes", [])
        if not routes:
            self.mc_po_summary.configure(text="Create a Purchase Review to manage PO numbers here.")
            ctk.CTkLabel(self.mc_po_scroll, text="No purchase-order routes available yet.", text_color=MUTED,
                         font=ctk.CTkFont(size=12)).grid(row=0, column=0, padx=12, pady=22)
            self.mc_generate_button.configure(state="disabled")
            return
        ready = sum(1 for row in routes if str(row.get("status", "")).casefold() == "ready")
        self.mc_po_summary.configure(text=f"{ready} ready  •  {len(routes) - ready} need review  •  Edit PO numbers directly below.")
        headers = ["Vendor", "Report", "PO Number", "Status"]
        for col, header in enumerate(headers):
            ctk.CTkLabel(self.mc_po_scroll, text=header, text_color=PURPLE_DARK,
                         font=ctk.CTkFont(size=11, weight="bold"), anchor="w").grid(
                row=0, column=col, padx=8, pady=(5, 7), sticky="ew")
        for row_index, route in enumerate(routes, start=1):
            vendor = route.get("vendor", "")
            report = route.get("report_type", "")
            key = route.get("key", "")
            var = ctk.StringVar(value=route.get("po_number", ""))
            self.po_entry_vars[key] = (vendor, report, var)
            ctk.CTkLabel(self.mc_po_scroll, text=vendor, text_color=TEXT,
                         font=ctk.CTkFont(size=11, weight="bold"), anchor="w").grid(
                row=row_index, column=0, padx=8, pady=4, sticky="ew")
            ctk.CTkLabel(self.mc_po_scroll, text=report, text_color=MUTED,
                         font=ctk.CTkFont(size=10), anchor="w").grid(
                row=row_index, column=1, padx=8, pady=4, sticky="ew")
            ctk.CTkEntry(self.mc_po_scroll, textvariable=var, height=32,
                         border_color="#D9D0E6", fg_color=WHITE).grid(
                row=row_index, column=2, padx=8, pady=4, sticky="ew")
            status = route.get("status", "Ready")
            is_ready = str(status).casefold() == "ready"
            ctk.CTkLabel(
                self.mc_po_scroll, text=("✓  " if is_ready else "⚠  ") + status,
                text_color="#166534" if is_ready else "#8A3F00",
                fg_color="#EAF8EF" if is_ready else "#FFF1DC",
                height=28, corner_radius=14,
                font=ctk.CTkFont(size=11, weight="bold"), anchor="center",
            ).grid(row=row_index, column=3, padx=8, pady=4, sticky="ew")
        blockers = int(snapshot.get("review_count", 0))
        blocked_routes = sum(1 for row in routes if str(row.get("status", "")).casefold() != "ready")
        missing_po = any(not str(row.get("po_number", "")).strip() for row in routes)
        self.mc_generate_button.configure(
            state="normal" if blockers == 0 and blocked_routes == 0 and not missing_po else "disabled"
        )

    def refresh_dashboard(self):
        if self.last_review_workbook and self.last_review_workbook.exists():
            self._sync_current_product_candidates()
        health = _product_master_health()
        snapshot = self._mission_snapshot()
        self.last_snapshot = snapshot
        if self.last_review_workbook and self.last_review_workbook.exists():
            self.review_needed = int(snapshot.get("review_count", 0))
            self.imported_line_count = int(snapshot.get("line_count", self.imported_line_count or 0))
            self.current_event_name = snapshot.get("event_name") or self.current_event_name or "General Sales Period"
            self.current_mode = snapshot.get("report_mode") or self.current_mode

        refresh_key = (
            self._mission_snapshot_cache_key,
            _file_stamp(live_product_master_path()),
            _file_stamp(self.selected_csv),
            _file_stamp(self.last_purchase_order_dir),
            _file_stamp(REVIEW_STATE_FILE),
            self.current_event_name,
            self.current_mode,
            self.current_decoration_fulfillment,
        )
        if refresh_key == self._dashboard_refresh_key:
            self._apply_dashboard_view()
            return

        review_count = int(snapshot.get("review_count", 0))
        product_master_count = int(snapshot.get("product_master_count", 0))
        purchase_review_count = int(snapshot.get("purchase_review_count", review_count))
        routes = list(snapshot.get("routes", []))
        blocked_routes = [
            route for route in routes
            if str(route.get("status", "")).casefold() != "ready"
        ]
        blocked_route_count = len(blocked_routes)
        workflow_blocked = bool(review_count or blocked_route_count)
        missing_routes = [r for r in routes if not str(r.get("po_number", "")).strip()]
        entered = len(routes) - len(missing_routes)
        all_po_ready = bool(routes) and not workflow_blocked
        line_count = int(snapshot.get("line_count", 0) or 0)
        has_active_packet = bool(line_count > 0 or routes or review_count or product_master_count or purchase_review_count)
        has_active_review = bool(self.last_review_workbook and has_active_packet)
        # While permanent setup decisions remain, Step 2 owns the workflow.
        # A review is considered stale only after those setup items are complete.
        review_stale = bool(has_active_review and product_master_count == 0 and self._review_is_stale())
        review_state = self._load_review_state()
        audit_requires_regeneration = bool(
            review_state.get("decoration_color_audit_requires_regeneration", False)
            and (review_stale or review_state.get("additional_import_pending", False))
        )
        purchase_generated_for_current = bool(
            self.last_purchase_order_dir
            and self.last_purchase_order_dir.exists()
            and self.last_review_workbook
            and self.last_review_workbook.exists()
            and not audit_requires_regeneration
            and self.last_purchase_order_dir.stat().st_mtime >= self.last_review_workbook.stat().st_mtime
        )

        import_complete = bool(has_active_review)
        master_complete = bool(has_active_review and product_master_count == 0)
        review_complete = bool(has_active_review and product_master_count == 0 and purchase_review_count == 0 and blocked_route_count == 0)
        audit_status = self._decoration_audit_status() if review_complete else {}
        audit_complete = bool(review_complete and audit_status.get("complete"))
        audit_attention = int(audit_status.get("needs_attention", 0) or 0)
        po_complete = bool(audit_complete and purchase_generated_for_current)
        self._set_stage(0, "complete" if import_complete else "active")
        if not import_complete:
            self._set_stage(1, "pending")
            self._set_stage(2, "pending")
            self._set_stage(3, "pending")
            self._set_stage(4, "pending")
        else:
            self._set_stage(1, "attention" if product_master_count else ("complete" if master_complete else "active"))
            self._set_stage(2, "attention" if purchase_review_count or blocked_route_count else ("complete" if review_complete else "pending"))
            self._set_stage(3, "attention" if review_complete and audit_attention else ("complete" if audit_complete else "pending"))
            self._set_stage(4, "complete" if po_complete else ("active" if audit_complete else "pending"))

        # Product Master owns reusable vendor and decoration setup.
        if product_master_count:
            self._set_pill(
                self.mc_master_status,
                f"{product_master_count} NEW STYLE{'S' if product_master_count != 1 else ''} NEED SETUP",
                "attention",
            )
            self.mc_master_detail.configure(
                text="Complete these new style details once. Orchid will refresh the review when Product Master closes."
            )
        elif review_stale:
            self._set_pill(self.mc_master_status, "UPDATES NOT APPLIED", "attention")
            self.mc_master_detail.configure(text="Product Master changed after this Purchase Review was created")
        else:
            self._set_pill(self.mc_master_status, f"{health['styles']} PRODUCTS CONFIGURED", "complete")
            self.mc_master_detail.configure(text="Permanent vendor and decoration catalog")

        # Purchase Review card is reserved for order-specific decisions.
        if product_master_count:
            self._set_pill(self.mc_review_status, "WAITING FOR PRODUCT MASTER", "pending")
            self.mc_review_detail.configure(text="Step 3 unlocks after permanent product setup is complete")
        elif review_stale:
            self._set_pill(self.mc_review_status, "REFRESH REQUIRED", "attention")
            self.mc_review_detail.configure(text="Regenerate Purchase Review to apply the latest Product Master changes")
        elif not has_active_review:
            self._set_pill(self.mc_review_status, "NOT CREATED", "pending")
            self.mc_review_detail.configure(text="Process the imported orders first")
        elif purchase_review_count:
            self._set_pill(self.mc_review_status, f"{purchase_review_count} ORDER DECISION{'S' if purchase_review_count != 1 else ''} REQUIRED", "attention")
            self.mc_review_detail.configure(
                text=f"{purchase_review_count} order-specific correction{'s' if purchase_review_count != 1 else ''} remain"
            )
        elif blocked_route_count:
            self._set_pill(
                self.mc_review_status,
                f"WAITING ON {blocked_route_count} DECISION{'S' if blocked_route_count != 1 else ''}",
                "attention",
            )
            self.mc_review_detail.configure(text="Open the Work Queue to see where each route must be corrected")
        else:
            self._set_pill(self.mc_review_status, "COMPLETE", "complete")
            self.mc_review_detail.configure(text="No purchasing decisions or blocked routes remain")

        # Purchase Orders card: a simple status and one edit/generate button.
        if product_master_count:
            self._set_pill(self.mc_po_status, "WAITING FOR PRODUCT MASTER", "pending")
            self.mc_po_detail.configure(text="Complete Step 2 before purchase orders can continue")
            self.mc_po_card_button.configure(text="Set Up Products", command=lambda: self.show_page("master"), state="normal")
        elif review_stale:
            self._set_pill(self.mc_po_status, "WAITING FOR REFRESH", "pending")
            self.mc_po_detail.configure(text="Apply Product Master updates before continuing")
            self.mc_po_card_button.configure(text="Regenerate Review", command=self.regenerate_current_review, state="normal")
        elif not routes:
            self._set_pill(self.mc_po_status, "NOT STARTED", "pending")
            self.mc_po_detail.configure(text="Purchase-order routes appear after review is created")
            self.mc_po_card_button.configure(text="Edit PO Numbers", command=self.open_po_number_editor, state="disabled")
        elif blocked_route_count:
            self._set_pill(
                self.mc_po_status,
                f"WAITING ON {blocked_route_count} DECISION{'S' if blocked_route_count != 1 else ''}",
                "attention",
            )
            self.mc_po_detail.configure(text="Resolve the remaining Product Master or order-specific decisions")
            self.mc_po_card_button.configure(text="View Blocked Routes", command=lambda: self.show_page("review"), state="normal")
        else:
            self._set_pill(self.mc_po_status, "READY", "complete")
            self.mc_po_detail.configure(text=f"{len(routes)} vendor purchase order{'s' if len(routes) != 1 else ''} ready")
            self.mc_po_card_button.configure(text="Generate Purchase Orders", command=self.generate_latest_purchase_orders, state="normal")

        # Centered hero and one primary workflow action.
        active_card = None
        completed = False
        if not has_active_packet and not self.selected_csv:
            step_number = 1
            title = "Import Orders"
            detail = "Choose the Shopify or Report Toaster export to start a new purchase packet."
            button_text = "Import Orders"
            command = lambda: self.show_page("import")
        elif not has_active_review:
            step_number = 1
            title = "Process Imported Orders"
            detail = "The order CSV is selected. Process the import so Orchid can check Product Master before Purchase Review."
            button_text = "Process Imported Orders"
            command = self.create_review_workbook
        elif product_master_count:
            step_number = 2
            title = "Complete Product Setup"
            detail = f"{product_master_count} new style{'s' if product_master_count != 1 else ''} need permanent setup before continuing."
            button_text = "Set Up Next Product"
            command = self.open_next_product_setup
            active_card = "master"
        elif review_stale:
            step_number = 2
            title = "Apply Product Master Updates"
            detail = "Regenerate Purchase Review to apply the latest routing and decoration changes."
            button_text = "Regenerate Purchase Review"
            command = self.regenerate_current_review
            active_card = "master"
        elif purchase_review_count or blocked_route_count:
            step_number = 3
            title = "Complete Purchase Review"
            if purchase_review_count:
                detail = f"{purchase_review_count} order-specific decision{'s' if purchase_review_count != 1 else ''} still require attention."
            else:
                detail = f"{blocked_route_count} purchase-order route{'s are' if blocked_route_count != 1 else ' is'} still blocked."
            button_text = "Open Purchase Review"
            command = lambda: self.show_page("review")
            active_card = "review"
        elif not purchase_generated_for_current:
            step_number = 4
            title = "Generate Purchase Orders"
            detail = "Purchase Review is complete. Enter one PO number for each vendor, then create the final PDFs."
            button_text = "Generate Purchase Orders"
            command = self.generate_latest_purchase_orders
            active_card = "purchase"
        else:
            step_number = 4
            completed = True
            title = "Purchase Orders Generated Successfully"
            detail = "The final vendor PDFs have been created and the event has been archived."
            button_text = "Open Purchase Orders"
            command = self.open_latest_purchase_orders
            active_card = "purchase"

        self._highlight_mission_card(active_card)
        self.mc_step_label.configure(text="")
        hero_icons = {1: "⧉", 2: "◇", 3: "☑", 4: "▤"}
        processing_import = bool(step_number == 1 and self.selected_csv and not has_active_review)
        hero_icon = "✦" if processing_import else hero_icons.get(step_number, "◆")
        self.mc_banner_icon.configure(text="✓" if completed else hero_icon)
        hero_colors = {
            1: ("#2563EB", "#EAF2FF"),
            2: (PURPLE, PURPLE_LIGHT),
            3: ("#B85C00", "#FFF1DC"),
            4: (SUCCESS, "#EAF8EF"),
        }
        hero_color, hero_fill = hero_colors[step_number]
        if completed:
            hero_color, hero_fill = "#166534", "#EAF8EF"
        self.mc_step_label.configure(text_color=hero_color)
        self.mc_banner_icon.configure(fg_color=hero_fill, text_color=hero_color)
        self.mc_banner_title.configure(text=title)
        self.mc_banner_subtitle.configure(text=detail, text_color=MUTED)
        event_name = self.current_event_name or snapshot.get("event_name") or ""
        line_count_text = int(snapshot.get("line_count", self.imported_line_count or 0) or 0)
        if event_name:
            packet_line = event_name
            if line_count_text and not completed:
                packet_line += f"  •  {line_count_text:,} purchase lines"
            self.mc_event_label.configure(text=packet_line)
            self.mc_event_label.grid()
        elif self.selected_csv:
            self.mc_event_label.configure(text=self.selected_csv.name)
            self.mc_event_label.grid()
        else:
            self.mc_event_label.grid_remove()
        self.mc_banner_button.configure(text=button_text, command=command, state="normal")
        if completed:
            self.mc_banner_secondary_button.pack(side="left", padx=7)
            if hasattr(self, "mc_animation_controls"):
                self.mc_animation_controls.place(relx=1.0, rely=0.05, anchor="ne")
        else:
            self.mc_banner_secondary_button.pack_forget()
            if hasattr(self, "mc_animation_controls"):
                self.mc_animation_controls.place_forget()
            self._success_animation_skip_requested = True
            self._finish_success_animation()
        self._update_success_control_buttons()
        self._update_dashboard_metrics(snapshot, show=completed)
        # Avoid repeating the same primary action in multiple dashboard locations.
        if product_master_count:
            self.mc_master_button.grid_remove()
            self.mc_work_open_button.configure(text="Set Up Next Product", command=self.open_next_product_setup)
        elif review_stale:
            self.mc_master_button.grid()
            self.mc_master_button.configure(text="Open Product Master", command=self.open_product_master, state="normal")
            self.mc_work_open_button.configure(text="Regenerate Review", command=self.regenerate_current_review)
        else:
            self.mc_master_button.grid()
            if purchase_review_count or blocked_route_count:
                self.mc_work_open_button.configure(text="Open Purchase Review", command=lambda: self.show_page("review"))
            else:
                self.mc_work_open_button.configure(text="Open Review", command=lambda: self.show_page("review"))
        if purchase_review_count or (blocked_route_count and not product_master_count):
            self.mc_review_button.grid_remove()
        else:
            self.mc_review_button.grid()
        self.mc_action_title.configure(text=title.title())
        self.mc_action_detail.configure(text=detail)
        self.mc_action_button.configure(text=button_text, command=command, state="normal")

        show_employee_totals = self._configure_dashboard_workspace(snapshot)
        self._render_work_queue(snapshot)
        if show_employee_totals and hasattr(self, "mc_employee_scroll"):
            self._render_employee_totals(snapshot)
        self._refresh_management_dashboard(snapshot, purchase_generated_for_current)
        self._apply_dashboard_view()

        selected_text = self.selected_csv.name if self.selected_csv else "No order CSV selected"
        self.refresh_master_page()
        if hasattr(self, "import_file_label"):
            self.import_file_label.configure(text=selected_text)
        self.refresh_purchase_page()
        self.refresh_current_event_page(snapshot)
        self._dashboard_refresh_key = refresh_key

    def open_product_master_search(self):
        value = " ".join(self.mc_product_search_var.get().split()) if hasattr(self, "mc_product_search_var") else ""
        self.open_product_master(search_product=value)

    def open_po_number_editor(
        self,
        generate_after_save: bool = False,
        workbook_path: Path | None = None,
        state_workbook_path: Path | None = None,
    ) -> bool:
        workbook = Path(workbook_path) if workbook_path else self.last_review_workbook
        state_workbook = Path(state_workbook_path) if state_workbook_path else workbook
        if not workbook or not workbook.exists():
            messagebox.showinfo("No Purchase Review", "Create a Purchase Review first.")
            return False
        snapshot = load_mission_control_snapshot(workbook, DATA, live_product_master_path())
        snapshot = self._overlay_completed_review_state(snapshot, state_workbook)
        routes = list(snapshot.get("routes", []))
        if not routes:
            messagebox.showinfo("No Purchase Orders", "No vendor purchase orders are available yet.")
            return False

        routes.sort(key=lambda route: str(route.get("vendor", "")).casefold())
        window = ctk.CTkToplevel(self)
        window.title("Assign Purchase Order Numbers" if generate_after_save else "Edit Purchase Order Numbers")
        window.geometry("700x600")
        window.minsize(620, 480)
        window.transient(self)
        window.grab_set()
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(2, weight=1)

        heading = "Assign Purchase Order Numbers" if generate_after_save else "Edit Purchase Order Numbers"
        ctk.CTkLabel(
            window, text=heading, text_color=TEXT,
            font=ctk.CTkFont(size=23, weight="bold"), anchor="w",
        ).grid(row=0, column=0, padx=24, pady=(22, 4), sticky="ew")
        subtitle = (
            "Enter one PO number for each vendor in this receiving packet. These numbers are saved with "
            "this Purchase Review and appear on the vendor PDFs and the Receiving & Decoration Report."
        )
        ctk.CTkLabel(
            window, text=subtitle, text_color=MUTED,
            font=ctk.CTkFont(size=13), anchor="w", justify="left", wraplength=620,
        ).grid(row=1, column=0, padx=24, pady=(0, 12), sticky="ew")

        scroll = ctk.CTkScrollableFrame(window, fg_color="#FCFBFE", corner_radius=12)
        scroll.grid(row=2, column=0, padx=24, pady=(0, 14), sticky="nsew")
        register_scrollable(self, scroll)
        scroll.grid_columnconfigure(1, weight=1)
        for col, label in enumerate(("Vendor", "PO Number")):
            ctk.CTkLabel(
                scroll, text=label, text_color=PURPLE_DARK,
                font=ctk.CTkFont(size=12, weight="bold"), anchor="w",
            ).grid(row=0, column=col, padx=12, pady=(10, 8), sticky="ew")

        editor_vars = []
        first_missing_entry = None
        for idx, route in enumerate(routes, start=1):
            vendor = str(route.get("vendor", "")).strip()
            original = str(route.get("po_number", "")).strip()
            variable = ctk.StringVar(value=original)
            editor_vars.append((vendor, variable))
            missing = not original
            row_fill = "#FFF8EC" if missing else "transparent"
            ctk.CTkLabel(
                scroll, text=vendor, text_color=TEXT, fg_color=row_fill,
                font=ctk.CTkFont(size=12, weight="bold"), anchor="w",
            ).grid(row=idx, column=0, padx=12, pady=6, sticky="ew")
            entry = ctk.CTkEntry(
                scroll, textvariable=variable, height=38,
                placeholder_text="Enter PO number",
                border_color=WARNING if missing else "#D9D0E6", fg_color=WHITE,
            )
            entry.grid(row=idx, column=1, padx=12, pady=6, sticky="ew")
            if missing and first_missing_entry is None:
                first_missing_entry = entry

        button_bar = ctk.CTkFrame(window, fg_color="transparent")
        button_bar.grid(row=3, column=0, padx=24, pady=(0, 22), sticky="e")
        ctk.CTkButton(
            button_bar, text="Cancel", width=110, height=42,
            fg_color=PURPLE_LIGHT, hover_color="#E4D6F5", text_color=PURPLE_DARK,
            command=window.destroy,
        ).pack(side="left", padx=4)

        def save_numbers():
            values: dict[tuple[str, str], str] = {}
            missing_vendors = []
            for vendor, variable in editor_vars:
                number = " ".join(variable.get().split())
                if not number:
                    missing_vendors.append(vendor)
                else:
                    values[(vendor.casefold(), "combined vendor order")] = number

            if generate_after_save and missing_vendors:
                messagebox.showwarning(
                    "PO Numbers Required",
                    "Enter a PO number for every vendor before generating purchase orders:\n\n"
                    + "\n".join(missing_vendors),
                    parent=window,
                )
                return

            save_po_overrides(DATA, state_workbook, values)
            window.destroy()
            if generate_after_save:
                self._run_final_purchase_order_generation(
                    workbook, values, active_workbook_path=state_workbook
                )
            elif missing_vendors:
                self.refresh_dashboard()
                messagebox.showwarning(
                    "PO Numbers Saved",
                    "The entered PO numbers were saved. Still missing:\n\n"
                    + "\n".join(missing_vendors),
                )
            else:
                self.refresh_dashboard()
                messagebox.showinfo(
                    "PO Numbers Saved",
                    "The vendor PO numbers were saved with this Purchase Review.",
                )

        button_text = "Save & Generate Purchase Orders" if generate_after_save else "Save PO Numbers"
        ctk.CTkButton(
            button_bar, text=button_text, width=245 if generate_after_save else 170, height=42,
            fg_color=PURPLE, hover_color=PURPLE_DARK,
            font=ctk.CTkFont(size=13, weight="bold"),
            command=save_numbers,
        ).pack(side="left", padx=4)
        def bring_editor_forward():
            try:
                if not window.winfo_exists():
                    return
                window.deiconify()
                window.lift()
                window.attributes("-topmost", True)
                window.focus_force()
                window.after(180, lambda: window.attributes("-topmost", False) if window.winfo_exists() else None)
                if first_missing_entry is not None:
                    first_missing_entry.focus_set()
            except Exception:
                pass

        window.after(180, bring_editor_forward)
        return True

    def save_dashboard_po_numbers(self, quiet: bool = False) -> dict[tuple[str, str], str]:
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            if not quiet:
                messagebox.showinfo("No Purchase Review", "Create a Purchase Review first.")
            return {}
        values: dict[tuple[str, str], str] = {}
        for _key, (vendor, report, variable) in self.po_entry_vars.items():
            number = " ".join(variable.get().split())
            if number:
                values[(vendor.casefold(), canonical_report_type(report))] = number
        save_po_overrides(DATA, self.last_review_workbook, values)
        if not quiet:
            self.dashboard_status.configure(text=f"Saved {len(values)} PO number(s) in Orchid Purchase Manager.", text_color=SUCCESS)
            messagebox.showinfo("PO Numbers Saved", f"Saved {len(values)} PO number(s) for the latest Purchase Review.")
        return values

    def _journal_completion_jobs_for_workbook(self, workbook: Path) -> list[dict]:
        """Build exact replay jobs for every completed Purchase Review decision.

        Source ID is the immutable packet identity and must remain attached when
        a decision is replayed into a regenerated or journal-reconciled workbook.
        Earlier builds dropped Source ID here and fell back to timestamp-specific
        Line IDs / Decision Keys.  A valid saved decision could therefore become
        a silent no-op, leaving customer-instruction rows unresolved at final PO
        preflight even though Purchase Review had reached zero.
        """
        jobs: list[dict] = []
        def completion_time(item: dict) -> float:
            try:
                return float(item.get("completed_at", 0) or 0)
            except (TypeError, ValueError):
                return 0.0

        records = sorted(
            self._completed_review_records_from_state(workbook),
            key=completion_time,
        )
        for index, record in enumerate(records, start=1):
            values = record.get("values", {})
            if not isinstance(values, dict) or not values:
                continue
            source_id = _clean(record.get("source_id", ""))
            issue = dict(record.get("issue", {})) if isinstance(record.get("issue"), dict) else {}
            source_only = bool(record.get("source_only", False) or issue.get("source_only", False))
            jobs.append({
                "job_id": _clean(record.get("job_id", "")) or f"reconcile-{index}",
                "workbook": str(workbook),
                "source_id": source_id,
                "line_id": _clean(record.get("line_id", "")),
                "decision_key": _clean(record.get("decision_key", "")),
                "issue_signature": _clean(record.get("issue_signature", "")),
                "values": dict(values),
                "issue": issue,
                "is_instruction_decision": bool(_clean(values.get("Decoration Decision", ""))),
                "editing_search": not bool(source_id) or source_only,
                # Source-only decisions legitimately have no Review & Edit row,
                # but must still update exactly one immutable All PO Lines row.
                "allow_missing": not bool(source_id) or source_only,
                "require_event_match": bool(source_id),
                "source_only": source_only,
            })
        return jobs

    def _audit_override_jobs_for_workbook(self, workbook: Path) -> list[dict]:
        """Return durable decoration-color corrections for final reconciliation.

        Purchase Review decisions can contain an older decoration color. Audit
        overrides are intentionally appended after those decisions so the last
        verified style/color correction remains authoritative.
        """
        state = self._load_review_state()
        overrides = state.get("decoration_color_audit_overrides", [])
        if not isinstance(overrides, list):
            return []
        jobs = []
        for index, override in enumerate(overrides, start=1):
            if not isinstance(override, dict):
                continue
            color = _clean(override.get("decoration_color", ""))
            if not color:
                continue
            jobs.append({
                "job_id": _clean(override.get("job_id", "")) or f"audit-color-{index}",
                "workbook": str(workbook),
                "line_id": _clean(override.get("line_id", "")),
                "decision_key": _clean(override.get("decision_key", "")),
                "values": {
                    "Decoration Color": color,
                    "Product #": _clean(override.get("style", "")),
                    "Description": _clean(override.get("product_name", "")),
                    "Garment Color": _clean(override.get("garment_color", "")),
                    "Decoration Type": _clean(override.get("decoration_type", "")),
                    "Decoration Location": _clean(override.get("location", "")),
                    "Purchase Vendor": _clean(override.get("purchase_vendor", "")),
                    "Size": _clean(override.get("size", "")),
                    "Quantity": _clean(override.get("quantity", "")),
                    "Order Number": _clean(override.get("order_number", "")),
                    "Employee Name": _clean(override.get("employee_name", "")),
                },
                "issue": {
                    "order": _clean(override.get("order_number", "")),
                    "employee": _clean(override.get("employee_name", "")),
                },
                "is_instruction_decision": False,
                # A line ID can change after Purchase Review is regenerated.
                # Keep the verified style / garment / decoration-type audit
                # decision applicable to the matching current-event rows.
                "audit_override": True,
                "editing_search": True,
                "allow_missing": True,
                # A verification is valid only when it can be reapplied to the
                # current event. Do not let final reconciliation silently drop
                # a saved audit approval if neither its line ID nor its exact
                # audit identity is present in the workbook.
                "require_event_match": True,
            })
        return jobs

    def _store_decoration_color_audit_changes(self, changes: list[dict]) -> None:
        """Persist audit overrides and update existing recovery-journal values."""
        state = self._load_review_state()
        overrides = state.get("decoration_color_audit_overrides", [])
        overrides = [dict(item) for item in overrides if isinstance(item, dict)] if isinstance(overrides, list) else []
        by_line = {_clean(item.get("line_id", "")): item for item in overrides if _clean(item.get("line_id", ""))}
        for change in changes:
            color = _clean(change.get("decoration_color", ""))
            # The audit location is the current user-approved value. Using the
            # pre-save event record here would reintroduce the old location when
            # reconciliation later reapplies the durable audit override.
            location = _clean(change.get("location", ""))
            for record in change.get("line_records", []):
                line_id = _clean(record.get("Line ID", ""))
                entry = {
                    "job_id": f"audit-{line_id or hashlib.sha256(str(record).encode()).hexdigest()[:12]}",
                    "line_id": line_id,
                    "decision_key": _clean(record.get("Decision Key", "")),
                    "style": _clean(record.get("Product #", change.get("style", ""))),
                    "product_name": _clean(record.get("Description", change.get("product_name", ""))),
                    "garment_color": _clean(record.get("Garment Color", change.get("garment_color", ""))),
                    "decoration_type": _clean(record.get("Decoration Type", change.get("decoration_type", ""))),
                    "location": location or _clean(record.get("Decoration Location", "")),
                    "decoration_color": color,
                    "purchase_vendor": _clean(record.get("Purchase Vendor", "")),
                    "size": _clean(record.get("Size", "")),
                    "quantity": _clean(record.get("Quantity", "")),
                    "order_number": _clean(record.get("Order Number", "")),
                    "employee_name": _clean(record.get("Employee Name", "")),
                    "updated_at": time.time(),
                }
                if line_id:
                    by_line[line_id] = entry
                else:
                    overrides.append(entry)
        line_entries = list(by_line.values())
        no_line_entries = [item for item in overrides if not _clean(item.get("line_id", ""))]
        state["decoration_color_audit_overrides"] = (line_entries + no_line_entries)[-10000:]

        completed = state.get("completed_review_decisions", [])
        completed = [dict(item) for item in completed if isinstance(item, dict)] if isinstance(completed, list) else []
        audit_values_by_line = {
            _clean(item.get("line_id", "")): {
                "color": _clean(item.get("decoration_color", "")),
                "location": _clean(item.get("location", "")),
            }
            for item in state["decoration_color_audit_overrides"]
            if _clean(item.get("line_id", ""))
        }
        for item in completed:
            line_id = _clean(item.get("line_id", ""))
            if line_id and line_id in audit_values_by_line:
                values = item.get("values", {})
                values = dict(values) if isinstance(values, dict) else {}
                values["Decoration Color"] = audit_values_by_line[line_id]["color"]
                values["Decoration Location"] = audit_values_by_line[line_id]["location"]
                item["values"] = values
        state["completed_review_decisions"] = completed

        pending = state.get("pending_review_saves", [])
        pending = [dict(item) for item in pending if isinstance(item, dict)] if isinstance(pending, list) else []
        for item in pending:
            line_id = _clean(item.get("line_id", ""))
            if line_id and line_id in audit_values_by_line:
                values = item.get("values", {})
                values = dict(values) if isinstance(values, dict) else {}
                values["Decoration Color"] = audit_values_by_line[line_id]["color"]
                values["Decoration Location"] = audit_values_by_line[line_id]["location"]
                item["values"] = values
        if pending:
            state["pending_review_saves"] = pending
        try:
            REVIEW_STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
            REVIEW_STATE_FILE.write_text(json.dumps(state, indent=2), encoding="utf-8")
        except Exception:
            pass

    @staticmethod
    def _clear_completed_instruction_shadows_in_workbook(workbook_path: Path) -> int:
        """Clear stale customer-decision flags after the event reached zero.

        This intentionally clears only customer-instruction validation shadows.
        Missing product, vendor, size, color, quantity, and Product Master issues
        remain untouched and can still block final purchase orders.
        """
        workbook_path = Path(workbook_path)
        workbook = load_workbook(workbook_path)
        changed = 0

        def headers_for(sheet):
            for row_number in range(1, min(sheet.max_row, 20) + 1):
                values = [
                    str(sheet.cell(row_number, column).value or "").strip()
                    for column in range(1, sheet.max_column + 1)
                ]
                if "Line ID" in values:
                    return row_number, {value: index + 1 for index, value in enumerate(values) if value}
            return None, {}

        for sheet_name in ("Review & Edit", "All PO Lines"):
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            header_row, headers = headers_for(sheet)
            if not header_row:
                continue
            for row_number in range(header_row + 1, sheet.max_row + 1):
                reason_parts = []
                for field in ("Event Review Reason", "Review Reason", "Action Required"):
                    column = headers.get(field)
                    if column:
                        reason_parts.append(_clean(sheet.cell(row_number, column).value))
                reason_text = " | ".join(reason_parts).casefold()
                if not (
                    "customer decision required" in reason_text
                    or "customer instruction decision required" in reason_text
                ):
                    continue
                permanent_reason = ""
                if headers.get("Permanent Review Reason"):
                    permanent_reason = _clean(
                        sheet.cell(row_number, headers["Permanent Review Reason"]).value
                    )
                if headers.get("Event Review Reason"):
                    sheet.cell(row_number, headers["Event Review Reason"]).value = ""
                if headers.get("Review Reason"):
                    sheet.cell(row_number, headers["Review Reason"]).value = permanent_reason
                if headers.get("Action Required") and not permanent_reason:
                    sheet.cell(row_number, headers["Action Required"]).value = ""
                if headers.get("Review Status"):
                    sheet.cell(row_number, headers["Review Status"]).value = "Ready"
                if headers.get("Resolution"):
                    sheet.cell(row_number, headers["Resolution"]).value = "Recovered from completed Purchase Review"
                if headers.get("Fix In"):
                    sheet.cell(row_number, headers["Fix In"]).value = (
                        "Product Master" if permanent_reason else "Purchase Review"
                    )
                changed += 1

        temp_path = workbook_path.with_name(
            f".{workbook_path.stem}.orchid-recovery{workbook_path.suffix}"
        )
        try:
            workbook.save(temp_path)
            workbook.close()
            os.replace(temp_path, workbook_path)
        finally:
            try:
                workbook.close()
            except Exception:
                pass
            try:
                if temp_path.exists():
                    temp_path.unlink()
            except Exception:
                pass
        return changed

    def _build_journal_reconciled_workbook(self, source_workbook: Path) -> Path:
        """Create a final-generation copy with durable decisions applied."""
        source_workbook = Path(source_workbook)
        jobs = self._journal_completion_jobs_for_workbook(source_workbook)
        jobs.extend(self._audit_override_jobs_for_workbook(source_workbook))
        # Final reconciliation may apply only exact saved decisions and verified
        # decoration-color overrides. An event-level "complete" marker is never
        # allowed to erase a newly discovered or changed customer-instruction
        # decision. The same reconciled validator therefore drives both Purchase
        # Review completion and final Purchase Order generation.
        if not jobs:
            return source_workbook
        target_dir = DATA / "reconciled_review_workbooks"
        target_dir.mkdir(parents=True, exist_ok=True)
        target = target_dir / f"{source_workbook.stem}_journal_reconciled.xlsx"
        building = target_dir / f".{source_workbook.stem}_journal_reconciled.building.xlsx"
        if building.exists():
            building.unlink()
        shutil.copy2(source_workbook, building)
        if jobs:
            reconciled_jobs = []
            for job in jobs:
                updated = dict(job)
                updated["workbook"] = str(building)
                reconciled_jobs.append(updated)
            # A regenerated packet can legitimately omit a line that was
            # completed in an earlier copy of the same event.  Reconciliation
            # must retain every possibly-current decision, but a decision with
            # no source, line, key, or unique semantic trace in either sheet is
            # an orphaned historical marker and cannot block audit/report work.
            self._apply_review_save_jobs(
                reconciled_jobs, skip_orphaned_recovery_jobs=True
            )
        os.replace(building, target)
        return target

    def _prepare_purchase_orders_from_journal(self) -> None:
        if getattr(self, "_po_reconcile_in_progress", False):
            return
        if not self.last_review_workbook or not self.last_review_workbook.is_file():
            messagebox.showinfo("No Purchase Review", "Create a Purchase Review first.")
            return
        # A reopened completed archive is intentionally report-only.  It has no
        # journal workbook to reconcile, so it must never enter the normal PO
        # preparation route (where an empty legacy path could resolve to ".").
        if self._is_archived_report_reopen():
            self.open_archived_report_regeneration()
            messagebox.showinfo(
                "Archived Reports",
                "This completed event is open only for job-name/logo updates and report regeneration. "
                "The original archived event remains unchanged.",
            )
            return
        routing_status = self._purchase_review_routing_status()
        if not routing_status.get("current", False):
            messagebox.showwarning(
                "Purchase Review Refresh Needed",
                "Product Master changed a route used by this event. No PDFs were created.\n\n"
                + _clean(routing_status.get("reason", ""))
                + "\n\nUse Regenerate Purchase Review on this Purchase Orders page. Orchid will return you here afterward.",
            )
            self.show_page("purchase")
            self.refresh_purchase_page()
            return
        active_workbook = Path(self.last_review_workbook)
        snapshot = self._fresh_snapshot()
        review_count = int(snapshot.get("review_count", 0))
        blocked_route_count = int(snapshot.get("blocked_route_count", 0))
        journal_jobs = self._journal_completion_jobs_for_workbook(active_workbook)
        authoritative_complete = self._review_state_is_authoritatively_complete(active_workbook)
        if (review_count or blocked_route_count) and not journal_jobs and not authoritative_complete:
            summary = self._blocker_summary(snapshot)
            messagebox.showwarning(
                "Purchase Review Incomplete",
                "Purchase Review did not pass the final readiness check.\n\n"
                + (summary or f"{review_count} decision(s) and {blocked_route_count} blocked route(s) remain."),
            )
            self.show_page("review")
            return

        if not self._require_completed_decoration_audit():
            return

        self._po_reconcile_in_progress = True
        dialog = ctk.CTkToplevel(self)
        dialog.title("Preparing Purchase Orders")
        dialog.geometry("500x220")
        dialog.resizable(False, False)
        dialog.transient(self)
        dialog.grab_set()
        dialog.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            dialog, text="Preparing Purchase Orders…", text_color=TEXT,
            font=ctk.CTkFont(size=22, weight="bold"),
        ).grid(row=0, column=0, padx=28, pady=(30, 8), sticky="ew")
        status = ctk.CTkLabel(
            dialog,
            text="Reconciling your saved Purchase Review decisions.",
            text_color=MUTED, font=ctk.CTkFont(size=13), wraplength=430,
        )
        status.grid(row=1, column=0, padx=28, pady=(0, 16), sticky="ew")
        progress = ctk.CTkProgressBar(dialog, mode="indeterminate", width=400)
        progress.grid(row=2, column=0, padx=45, pady=(0, 24), sticky="ew")
        progress.start()
        result_queue: queue.Queue = queue.Queue()

        def worker():
            try:
                effective = self._build_journal_reconciled_workbook(active_workbook)
                effective_snapshot = load_mission_control_snapshot(
                    effective, DATA, live_product_master_path()
                )
                effective_snapshot = self._overlay_completed_review_state(
                    effective_snapshot, active_workbook
                )
                result_queue.put(("ok", effective, effective_snapshot))
            except Exception as error:
                result_queue.put(("error", error, traceback.format_exc()))

        def close_dialog():
            try:
                progress.stop()
                dialog.grab_release()
                dialog.destroy()
            except Exception:
                pass
            self._po_reconcile_in_progress = False

        def poll():
            try:
                kind, value, detail = result_queue.get_nowait()
            except queue.Empty:
                if dialog.winfo_exists():
                    self.after(80, poll)
                return
            close_dialog()
            if kind == "error":
                print(detail, file=sys.stderr)
                messagebox.showerror(
                    "Unable to Prepare Purchase Orders",
                    str(value) + "\n\nYour completed decisions remain safe in Orchid's recovery journal.",
                )
                return
            effective = Path(value)
            effective_snapshot = detail
            remaining = len(self._build_purchase_review_queue(effective_snapshot))
            blocked = int(effective_snapshot.get("blocked_route_count", 0))
            if remaining or blocked:
                messagebox.showwarning(
                    "Purchase Review Still Requires Attention",
                    "The reconciled workbook still contains a genuine unresolved decision.\n\n"
                    + (self._blocker_summary(effective_snapshot) or f"{remaining} decision(s) and {blocked} blocked route(s) remain."),
                )
                self.show_page("review")
                return
            def open_assignment_window():
                try:
                    opened = self.open_po_number_editor(
                        generate_after_save=True,
                        workbook_path=effective,
                        state_workbook_path=active_workbook,
                    )
                    if not opened:
                        messagebox.showwarning(
                            "PO Number Assignment Did Not Open",
                            "Purchase Review reconciliation completed, but Orchid could not open the PO-number assignment window. "
                            "No reports were generated. Select Prepare & Assign PO Numbers again.",
                        )
                except Exception as error:
                    traceback.print_exc()
                    messagebox.showerror(
                        "Unable to Open PO Number Assignment",
                        str(error) + "\n\nNo purchase-order reports were generated.",
                    )

            # On macOS, opening a new modal in the same callback that destroys
            # the progress modal can leave the PO-number editor hidden behind
            # the main window. Hand the next modal back to Tk's event loop so
            # it is reliably visible and focused.
            self.after(220, open_assignment_window)

        threading.Thread(target=worker, name="orchid-po-reconcile", daemon=True).start()
        self.after(80, poll)

    def generate_latest_purchase_orders(self):
        self._prepare_purchase_orders_from_journal()

    def generate_test_purchase_orders(self) -> None:
        """Generate an audited test packet while holding unfinished review lines.

        This is deliberately separate from normal PO generation.  It never
        marks Purchase Review complete and cannot place an unanswered review
        line into a vendor order.
        """
        if self._review_save_pending:
            messagebox.showinfo(
                "Finishing Purchase Review Saves",
                "Wait for the current Purchase Review save to finish before generating a test packet.",
            )
            return
        if not self.last_review_workbook or not self.last_review_workbook.is_file():
            messagebox.showinfo("No Purchase Review", "Process the imported orders before generating a test packet.")
            return
        if not self._require_active_packet("Test purchase-order generation"):
            return
        routing_status = self._purchase_review_routing_status()
        if not routing_status.get("current", False):
            messagebox.showwarning(
                "Purchase Review Refresh Needed",
                "Product Master changed a route used by this event. Regenerate Purchase Review first so the test uses the current saved vendor and decoration setup.\n\n"
                + _clean(routing_status.get("reason", "")),
            )
            return
        snapshot = self._fresh_snapshot()
        product_master_count = int(snapshot.get("product_master_count", 0) or 0)
        if product_master_count:
            messagebox.showwarning(
                "Product Setup Still Required",
                "Test purchase orders cannot guess missing permanent purchasing information. Complete the remaining Product Master setup first, then use Test Generate POs.\n\n"
                + (self._blocker_summary(snapshot) or f"{product_master_count} product setup item(s) remain."),
            )
            return
        pending = int(snapshot.get("review_count", 0) or 0)
        blocked = int(snapshot.get("blocked_route_count", 0) or 0)
        held = pending + blocked
        if not messagebox.askyesno(
            "Generate Audited Test Purchase Orders?",
            "This creates TEST-numbered purchase-order PDFs for the fully configured lines only.\n\n"
            f"{held} unfinished Purchase Review line{'s' if held != 1 else ''} will be held out of vendor orders and listed in a separate Held for Review report.\n\n"
            "Before release, Orchid will verify that every included garment is accounted for exactly once. This does not mark Purchase Review complete. Continue?",
            parent=self,
        ):
            return
        self._run_final_purchase_order_generation(
            self.last_review_workbook,
            po_overrides={},
            active_workbook_path=self.last_review_workbook,
            hold_unresolved_review=True,
            test_mode=True,
        )

    def open_archived_report_regeneration(self) -> None:
        """Open the safe logo/name screen for a report-only archived event."""
        if not self.last_review_workbook or not self.last_review_workbook.is_file():
            messagebox.showinfo("No Archived Report", "Choose an archived event first.")
            return
        if not self._is_archived_report_reopen():
            messagebox.showinfo(
                "Current Event",
                "Use the normal Purchase Orders actions for the current event.",
            )
            return
        self.show_page("job_logo")
        self.after(80, self.refresh_outsourced_job_logo_page)

    def regenerate_review_and_reports(self):
        """Refresh the active review, then rebuild reports with the saved event logo."""
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            messagebox.showinfo("No Purchase Review", "Create a Purchase Review first.")
            return
        # Save a typed job/logo name before the report worker reads the sidecar.
        # This makes Regenerate a safe one-click action after entering a name.
        if hasattr(self, "outsourced_logo_name_var"):
            self.save_outsourced_job_name(show_confirmation=False)
        # An archived event is reopened as a report-only working copy. Its final
        # workbook is already complete, so do not rebuild Purchase Review from
        # an older CSV just to add a logo and regenerate the PDFs.
        if self._is_archived_report_reopen():
            self.regenerate_current_purchase_orders()
            return
        self.regenerate_current_review(silent=True)
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            return
        snapshot = self._fresh_snapshot()
        review_count = int(snapshot.get("review_count", 0))
        blocked_route_count = int(snapshot.get("blocked_route_count", 0))
        if review_count or blocked_route_count:
            messagebox.showwarning(
                "Purchase Review Requires Attention",
                "The Purchase Review was regenerated, but reports cannot be rebuilt until the remaining decisions are completed.\n\n"
                + (self._blocker_summary(snapshot) or f"{review_count} decision(s) and {blocked_route_count} blocked route(s) remain."),
            )
            self.show_page("review")
            return
        if not self._require_completed_decoration_audit():
            return
        overrides = load_po_overrides(DATA, self.last_review_workbook)
        self._run_final_purchase_order_generation(self.last_review_workbook, overrides)

    def regenerate_current_purchase_orders(self):
        """Rebuild final purchase orders from the latest completed Purchase Review."""
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            messagebox.showinfo("No Purchase Review", "Create a Purchase Review first.")
            return
        if not self._require_active_packet("Purchase Order regeneration"):
            return
        if hasattr(self, "outsourced_logo_name_var"):
            self.save_outsourced_job_name(show_confirmation=False)
        report_only_archive = self._is_archived_report_reopen()
        if report_only_archive:
            # Report-only archives deliberately have no recovery journal to
            # reconcile. Use their protected working workbook directly, which
            # avoids treating a blank legacy source path as the current folder.
            overrides = load_po_overrides(DATA, self.last_review_workbook)
            self._run_final_purchase_order_generation(
                self.last_review_workbook,
                overrides,
                active_workbook_path=self.last_review_workbook,
            )
            return
        if not report_only_archive:
            # Refresh only when Product Master changed an active PO route.
            # Thread/ink and catalog-only saves do not interrupt regeneration.
            if self._review_is_stale():
                self.regenerate_current_review(silent=True)
                if not self.last_review_workbook or not self.last_review_workbook.exists():
                    return
            snapshot = self._fresh_snapshot()
            review_count = int(snapshot.get("review_count", 0))
            blocked_route_count = int(snapshot.get("blocked_route_count", 0))
            if review_count or blocked_route_count:
                messagebox.showwarning(
                    "Purchase Review Incomplete",
                    "Purchase orders cannot be regenerated until Purchase Review is complete.\n\n"
                    + (self._blocker_summary(snapshot) or f"{review_count} decision(s) and {blocked_route_count} blocked route(s) remain."),
                )
                self.show_page("review")
                return
            if not self._require_completed_decoration_audit():
                return
        overrides = load_po_overrides(DATA, self.last_review_workbook)
        self._run_final_purchase_order_generation(self.last_review_workbook, overrides)

    def guided_stage_card(self, page, step_label: str, title: str, body: str):
        """Create the large white action card used throughout the workflow."""
        page.grid_rowconfigure(2, weight=1)
        shadow = ctk.CTkFrame(page, fg_color=SHADOW, corner_radius=24, height=450)
        shadow.grid(row=1, column=0, sticky="ew", padx=48, pady=(8, 18))
        shadow.grid_propagate(False)
        frame = ctk.CTkFrame(
            page, fg_color=WHITE, border_width=1, border_color=BORDER,
            corner_radius=24, height=444,
        )
        frame.grid(row=1, column=0, sticky="ew", padx=48, pady=(3, 23))
        frame.grid_propagate(False)
        frame.grid_columnconfigure(0, weight=1)
        self._add_orchid_watermark(frame, size=(520, 520), relx=0.5, rely=0.57, opacity=0.14)

        badge = ctk.CTkLabel(frame, text="", width=1, height=1)
        badge.grid_remove()
        icon_map = {
            "Complete Product Setup": "◇",
            "Complete Purchase Review": "☑",
            "Generate Purchase Orders": "▤",
        }
        icon = ctk.CTkLabel(
            frame, text=icon_map.get(title, "◆"), width=72, height=72, corner_radius=36,
            fg_color=PURPLE_LIGHT, text_color=PURPLE,
            font=ctk.CTkFont(size=28, weight="bold"),
        )
        icon.grid(row=1, column=0, pady=(29, 13))
        title_label = ctk.CTkLabel(
            frame, text=title, text_color=TEXT,
            font=ctk.CTkFont(size=29, weight="bold"), justify="center", wraplength=900,
        )
        title_label.grid(row=2, column=0, padx=50, pady=(0, 8))
        body_label = ctk.CTkLabel(
            frame, text=body, text_color=MUTED,
            font=ctk.CTkFont(size=15), justify="center", wraplength=840,
        )
        body_label.grid(row=3, column=0, padx=56, pady=(0, 15))
        return frame, badge, title_label, body_label

    def build_import_page(self):
        page = self.new_page("import")
        page.grid_rowconfigure(1, weight=1)
        page.grid_columnconfigure(0, weight=1)

        self._build_workflow_tracker(page, 0, "current_event_stage_labels", padx=48)

        split = ctk.CTkFrame(page, fg_color="transparent")
        split.grid(row=1, column=0, sticky="nsew", padx=48, pady=(2, 8))
        split.grid_columnconfigure(0, weight=11, uniform="current_event_split")
        split.grid_columnconfigure(1, weight=9, uniform="current_event_split")
        split.grid_rowconfigure(0, weight=1)
        self.current_event_split = split

        glow_source = PILImage.open(resource_path("assets", "current_event_glow.png")).convert("RGBA")
        self.current_event_glow_image = ctk.CTkImage(
            light_image=glow_source, dark_image=glow_source, size=(920, 520)
        )
        self.current_event_glow = ctk.CTkLabel(
            split, text="", image=self.current_event_glow_image, fg_color="transparent"
        )
        self.current_event_glow.place(relx=0.52, rely=0.47, anchor="center")

        illustration_card = ctk.CTkFrame(
            split, fg_color="transparent", border_width=0,
            corner_radius=0,
        )
        illustration_card.grid(row=0, column=0, sticky="nsew", padx=(0, 2))
        illustration_card.grid_rowconfigure(0, weight=1)
        illustration_card.grid_columnconfigure(0, weight=1)
        self.current_event_illustration_card = illustration_card
        illustration_source = PILImage.open(resource_path("assets", "current_event_cats.png")).convert("RGBA")
        self.current_event_cats_image = ctk.CTkImage(
            light_image=illustration_source, dark_image=illustration_source, size=(472, 489)
        )
        ctk.CTkLabel(
            illustration_card, text="", image=self.current_event_cats_image,
            fg_color="transparent",
        ).grid(row=0, column=0, sticky="e", padx=0, pady=(0, 4))

        # The event action area intentionally blends into the shared Current
        # Event canvas instead of looking like a separate floating window.
        card_shadow = ctk.CTkFrame(split, fg_color="transparent", corner_radius=0)
        card_shadow.grid_remove()
        card = ctk.CTkFrame(
            split, fg_color="transparent", border_width=0,
            corner_radius=0,
        )
        card.grid(row=0, column=1, sticky="nsew", padx=(0, 64), pady=(0, 6))
        card.grid_columnconfigure(0, weight=1)
        card.grid_rowconfigure(0, weight=1)
        self.current_event_card = card
        self.current_event_card_shadow = card_shadow

        # The Current Event actions already make importing and adding orders
        # explicit.  Keep this hero focused on the active job and its logo;
        # a second decorative import icon above it is redundant and distracting.
        self.current_event_title = ctk.CTkLabel(
            card, text="No Active Purchase Packet", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=33, weight="bold"), justify="center", wraplength=440,
        )
        self.current_event_title.grid(row=1, column=0, padx=30, pady=(22, 8))
        self.current_event_job_logo = ctk.CTkLabel(
            card, text="", width=1, height=1, fg_color="transparent",
        )
        self.current_event_job_logo.grid(row=2, column=0, padx=30, pady=(0, 7))
        self.current_event_job_logo.grid_remove()
        self.current_event_body = ctk.CTkLabel(
            card, text="Import the Shopify or Report Toaster CSV for the orders you want to process. Product Master and saved settings are preserved.",
            text_color="#51485F", font=ctk.CTkFont(size=16), justify="center", wraplength=430,
        )
        self.current_event_body.grid(row=3, column=0, padx=38, pady=(0, 12))
        self.import_file_label = ctk.CTkLabel(
            card, text="No order CSV selected", text_color=TEXT,
            fg_color="#EEEAF4", corner_radius=15, height=32,
            font=ctk.CTkFont(size=13, weight="bold"), justify="center",
        )
        self.import_file_label.grid(row=4, column=0, padx=34, pady=(0, 10))
        self.current_event_details = ctk.CTkLabel(
            card, text="", text_color=MUTED, font=ctk.CTkFont(size=12),
            justify="center", wraplength=390,
        )
        self.current_event_details.grid(row=5, column=0, padx=30, pady=(0, 6))
        buttons = ctk.CTkFrame(card, fg_color="transparent")
        buttons.grid(row=6, column=0, sticky="ew", padx=46, pady=(0, 4))
        buttons.grid_columnconfigure((0, 1), weight=1, uniform="event_actions")
        self.current_event_buttons = buttons
        self.current_event_primary = ctk.CTkButton(
            buttons, text="Import Shopify CSV", command=self.choose_csv,
            width=270, height=60, corner_radius=10, fg_color=PURPLE, hover_color=PURPLE_DARK,
            font=ctk.CTkFont(size=15, weight="bold"),
        )
        self.current_event_primary.grid(row=0, column=0, columnspan=2, padx=5, pady=5)
        self.current_event_add = ctk.CTkButton(
            buttons, text="Add Orders CSV", command=self.add_orders_to_current_event,
            width=190, height=40, corner_radius=10, fg_color=PURPLE, hover_color=PURPLE_DARK,
            border_width=0, text_color=WHITE,
            font=ctk.CTkFont(size=13, weight="bold"),
        )
        self.current_event_add.grid(row=0, column=1, sticky="ew", padx=5, pady=5)
        self.current_event_replace = ctk.CTkButton(
            buttons, text="Start New Event", command=self.start_new_purchase_packet,
            width=190, height=40, corner_radius=10, fg_color=PURPLE, hover_color=PURPLE_DARK,
            border_width=0, text_color=WHITE,
            font=ctk.CTkFont(size=12, weight="bold"),
        )
        # Replacing an import must never compete with the longer Review button
        # for horizontal space.  On the Current Event card the two buttons used
        # to share a narrow row, which clipped the "CSV" part of the label on
        # compact laptop windows.  Give each maintenance action its own stable,
        # full-width row instead.
        self.current_event_replace.grid(row=1, column=0, columnspan=2, sticky="ew", padx=5, pady=5)
        self.current_event_regenerate = ctk.CTkButton(
            buttons, text="Regenerate Purchase Review", command=self.regenerate_current_review,
            width=220, height=40, corner_radius=10, fg_color=PURPLE, hover_color=PURPLE_DARK,
            border_width=0, text_color=WHITE,
        )
        self.current_event_regenerate.grid(row=2, column=0, columnspan=2, sticky="ew", padx=5, pady=5)
        self.current_event_regenerate_po = ctk.CTkButton(
            buttons, text="Regenerate Purchase Orders", command=self.regenerate_current_purchase_orders,
            width=220, height=40, corner_radius=10, fg_color=PURPLE, hover_color=PURPLE_DARK,
            border_width=0, text_color=WHITE,
            font=ctk.CTkFont(size=12, weight="bold"),
        )
        self.current_event_regenerate_po.grid(row=3, column=0, sticky="ew", padx=5, pady=5)
        self.current_event_edit_po = ctk.CTkButton(
            buttons, text="Edit PO Numbers", command=self.open_po_number_editor,
            width=180, height=40, corner_radius=10, fg_color=PURPLE, hover_color=PURPLE_DARK,
            border_width=0, text_color=WHITE,
            font=ctk.CTkFont(size=12, weight="bold"),
        )
        self.current_event_edit_po.grid(row=3, column=1, sticky="ew", padx=5, pady=5)
        self.current_event_discard = ctk.CTkButton(
            buttons, text="Finish & Clear Event", command=self.finish_and_clear_current_event,
            width=210, height=38, corner_radius=9, fg_color="#FFF1F1", hover_color="#FADDDD",
            border_width=1, border_color="#D8A5A5", text_color=DANGER,
            font=ctk.CTkFont(size=13, weight="bold"),
        )
        self.current_event_discard.grid(row=4, column=0, columnspan=2, sticky="ew", padx=5, pady=(6, 0))
        self.current_event_master = ctk.CTkButton(
            card, text="Open Product Master", command=lambda: self.show_page("master"),
            width=250, height=40, corner_radius=10, fg_color=PURPLE,
            hover_color=PURPLE_DARK, text_color=WHITE,
            font=ctk.CTkFont(size=14, weight="bold"),
        )
        self.current_event_master.grid(row=7, column=0, pady=(6, 8))

        # The empty state is a single, centered composition.  Keep the normal
        # two-column layout available for active purchase packets, where the
        # customer/job logo belongs above the event details.
        self._current_event_empty_layout_visible = False

    def _show_current_event_empty_layout(self) -> None:
        """Show the reliable centered empty-state layout without the glow seam."""
        if not hasattr(self, "current_event_split"):
            return
        self._current_event_empty_layout_visible = True
        self.current_event_glow.place_forget()
        # Keep the illustration and actions under one geometry manager.  The
        # 4.9.102 placement conversion can leave both widgets unmapped on
        # macOS, producing an entirely blank Current Event page.
        self.current_event_illustration_card.place_forget()
        self.current_event_card.place_forget()
        self.current_event_split.grid_columnconfigure(0, weight=0, minsize=474, uniform="")
        self.current_event_split.grid_columnconfigure(1, weight=0, minsize=536, uniform="")
        self.current_event_split.grid_anchor("center")
        self.current_event_illustration_card.grid()
        self.current_event_card.grid()

    def _show_current_event_active_layout(self) -> None:
        """Restore the logo-ready layout used by active purchase packets."""
        if not hasattr(self, "current_event_split"):
            return
        self._current_event_empty_layout_visible = False
        self.current_event_illustration_card.place_forget()
        self.current_event_card.place_forget()
        self.current_event_split.grid_columnconfigure(0, weight=11, minsize=0, uniform="current_event_split")
        self.current_event_split.grid_columnconfigure(1, weight=9, minsize=0, uniform="current_event_split")
        self.current_event_illustration_card.grid()
        self.current_event_card.grid()
        self.current_event_glow.place(relx=0.52, rely=0.47, anchor="center")

    def _refresh_current_event_job_logo(self) -> None:
        """Show the saved customer/job logo below the active event name."""
        if not hasattr(self, "current_event_job_logo"):
            return
        logo_path = self._outsourced_job_logo_path()
        if not logo_path or not logo_path.is_file():
            self.current_event_job_logo.configure(image=None, text="")
            self.current_event_job_logo.grid_remove()
            return
        try:
            image, _background_removed = prepare_artwork_for_preview(logo_path)
            image.thumbnail((220, 86), PILImage.Resampling.LANCZOS)
            self.current_event_job_logo_image = ctk.CTkImage(
                light_image=image, dark_image=image, size=image.size,
            )
            self.current_event_job_logo.configure(image=self.current_event_job_logo_image, text="")
        except Exception:
            self.current_event_job_logo.configure(
                image=None, text="Customer / Job Logo", text_color=MUTED,
                font=ctk.CTkFont(size=12, weight="bold"),
            )
        self.current_event_job_logo.grid()

    def _set_current_event_stage(self, index: int, state: str, status_text: str) -> None:
        if state == "complete":
            card_fill, border, icon_fill, icon_variant = "#EAF7EE", "#A9D9B6", SUCCESS, "white"
            title_color, status_color = "#166534", "#166534"
            enabled = "normal"
        elif state == "attention":
            card_fill, border, icon_fill, icon_variant = "#FFF4E5", WARNING, WARNING, "white"
            title_color, status_color = "#9A4D00", "#9A4D00"
            enabled = "normal"
        elif state == "active":
            card_fill, border, icon_fill, icon_variant = PURPLE, PURPLE, "transparent", "white"
            title_color, status_color = WHITE, WHITE
            enabled = "normal"
        else:
            card_fill, border, icon_fill, icon_variant = "#FBF9FD", "#C6B7D8", "#EEE5F8", "purple"
            title_color, status_color = "#2D243D", "#665C74"
            enabled = "normal"
        for attr in (
            "mc_stage_labels", "current_event_stage_labels", "master_stage_labels",
            "review_stage_labels", "audit_stage_labels", "purchase_stage_labels",
        ):
            tracker = getattr(self, attr, None)
            if not tracker or index >= len(tracker):
                continue
            stage = tracker[index]
            stage_key = stage.get("stage_key", ("import", "master", "review", "audit", "purchase")[index])
            stage["card"].configure(fg_color=card_fill, border_color=border)
            stage["circle"].configure(
                image=self._workflow_icon_images[(stage_key, icon_variant)],
                fg_color=icon_fill, hover_color=PURPLE_LIGHT,
                state=enabled,
            )
            stage["title"].configure(text_color=title_color)
            stage["status"].configure(text=status_text, text_color=status_color, fg_color="transparent")
            connector = stage.get("connector")
            if connector is not None:
                connector.configure(fg_color=PURPLE if state == "complete" else "#C8BCD8")

    def _refresh_current_event_progress(self, snapshot: dict | None = None) -> None:
        snapshot = snapshot or {}
        has_review = bool(self.last_review_workbook and self.last_review_workbook.exists())
        has_csv = bool(self.selected_csv) or has_review
        review_state = self._load_review_state()
        routing_stale = bool(has_review and self._review_is_stale())
        audit_requires_regeneration = bool(
            review_state.get("decoration_color_audit_requires_regeneration", False)
            and (routing_stale or review_state.get("additional_import_pending", False))
        )
        has_purchase_orders = bool(
            self.last_purchase_order_dir and self.last_purchase_order_dir.exists()
            and not audit_requires_regeneration
        )
        product_issues = sum(
            1 for issue in snapshot.get("issues", [])
            if _clean(issue.get("fix_in", issue.get("source", ""))).casefold() == "product master"
        )
        review_issues = int(snapshot.get("review_count", 0) or 0)
        blocked_routes = int(snapshot.get("blocked_route_count", 0) or 0)
        review_complete = bool(has_review and not product_issues and not review_issues and not blocked_routes)
        audit_status = self._decoration_audit_status() if review_complete else {}
        audit_complete = bool(review_complete and audit_status.get("complete"))
        audit_remaining = int(audit_status.get("needs_attention", 0) or 0)

        if not has_csv:
            states = (
                ("active", "Start Here"), ("pending", "Waiting"),
                ("pending", "Waiting"), ("pending", "Waiting"), ("pending", "Waiting"),
            )
        elif not has_review:
            states = (
                ("active", "CSV Selected"), ("pending", "Waiting"),
                ("pending", "Waiting"), ("pending", "Waiting"), ("pending", "Waiting"),
            )
        elif product_issues:
            states = (
                ("complete", "Complete"), ("attention", "Needs Setup"),
                ("pending", "Waiting"), ("pending", "Waiting"), ("pending", "Waiting"),
            )
        elif review_issues or blocked_routes:
            states = (
                ("complete", "Complete"), ("complete", "Complete"),
                ("attention", "Needs Attention"), ("pending", "Waiting"), ("pending", "Waiting"),
            )
        elif not audit_complete:
            states = (
                ("complete", "Complete"), ("complete", "Complete"),
                ("complete", "Complete"), ("attention", f"{audit_remaining} to Review"),
                ("pending", "Waiting"),
            )
        elif has_purchase_orders:
            states = tuple(("complete", "Complete") for _ in range(5))
        else:
            states = (
                ("complete", "Complete"), ("complete", "Complete"),
                ("complete", "Complete"), ("complete", "Complete"),
                ("active", "Ready"),
            )
        for index, (state, status) in enumerate(states):
            self._set_current_event_stage(index, state, status)

    def build_master_page(self):
        page = self.new_page("master")
        page.grid_rowconfigure(9, weight=0, minsize=0)
        page.grid_rowconfigure(1, weight=1, minsize=520)
        self._build_workflow_tracker(page, 0, "master_stage_labels", padx=34)

        workspace = ctk.CTkFrame(page, fg_color="transparent")
        workspace.grid(row=1, column=0, sticky="nsew", padx=24, pady=(4, 12))
        workspace.grid_columnconfigure(0, weight=1)
        workspace.grid_rowconfigure(2, weight=1)
        self.master_workspace = workspace

        header = ctk.CTkFrame(workspace, fg_color="transparent")
        header.grid(row=0, column=0, sticky="ew", padx=8, pady=(2, 10))
        header.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            header, text="Product Master", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=30, weight="bold"), anchor="w",
        ).grid(row=0, column=0, sticky="w")
        ctk.CTkLabel(
            header, text="Manage permanent purchasing and decoration defaults used for every future purchase packet.",
            text_color=MUTED, font=ctk.CTkFont(size=14), anchor="w", justify="left",
        ).grid(row=1, column=0, sticky="w", pady=(2, 0))
        self.master_step_badge = ctk.CTkLabel(
            header, text="Checking Product Master…", fg_color=PURPLE_LIGHT,
            text_color=PURPLE_DARK, corner_radius=14, height=34,
            font=ctk.CTkFont(size=13, weight="bold"),
        )
        self.master_step_badge.grid(row=0, column=1, rowspan=2, padx=(18, 0), sticky="e")

        cards = ctk.CTkFrame(workspace, fg_color="transparent")
        cards.grid(row=1, column=0, sticky="ew")
        cards.grid_columnconfigure((0, 1), weight=1, uniform="master_overview")

        status_shadow = ctk.CTkFrame(cards, fg_color=SHADOW, corner_radius=20, width=1, height=1)
        status_shadow.grid(row=0, column=0, sticky="nsew", padx=(8, 7), pady=(5, 9))
        status = ctk.CTkFrame(cards, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=20)
        status.grid(row=0, column=0, sticky="nsew", padx=(8, 7), pady=(1, 13))
        status.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            status, text="Product Master Status", text_color=TEXT,
            font=ctk.CTkFont(size=21, weight="bold"), anchor="w",
        ).grid(row=0, column=0, padx=24, pady=(22, 2), sticky="w")
        self.master_details = ctk.CTkLabel(
            status, text="Permanent catalog readiness", text_color=MUTED,
            font=ctk.CTkFont(size=12), anchor="w",
        )
        self.master_details.grid(row=1, column=0, padx=24, pady=(0, 14), sticky="w")
        self.master_health_large = ctk.CTkLabel(
            status, text="—", text_color=PURPLE,
            font=ctk.CTkFont(size=34, weight="bold"), anchor="w",
        )
        self.master_health_large.grid(row=2, column=0, padx=24, pady=(0, 8), sticky="w")
        self.master_progress = ctk.CTkProgressBar(
            status, height=13, corner_radius=7, progress_color=PURPLE,
            fg_color="#EEE8F5", border_width=0,
        )
        self.master_progress.grid(row=3, column=0, sticky="ew", padx=24, pady=(0, 18))
        self.master_progress.set(0)

        metrics = ctk.CTkFrame(status, fg_color="#FAF8FC", corner_radius=14)
        metrics.grid(row=4, column=0, sticky="ew", padx=20, pady=(0, 20))
        metrics.grid_columnconfigure((0, 1, 2), weight=1, uniform="master_metrics")
        metric_specs = (
            ("master_products_value", "Products"),
            ("master_need_setup_value", "Need Setup"),
            ("master_ready_value", "Ready"),
        )
        for column, (attribute, label) in enumerate(metric_specs):
            value = ctk.CTkLabel(
                metrics, text="—", text_color=PURPLE_DARK,
                font=ctk.CTkFont(size=23, weight="bold"),
            )
            value.grid(row=0, column=column, pady=(14, 1))
            setattr(self, attribute, value)
            ctk.CTkLabel(
                metrics, text=label, text_color=MUTED,
                font=ctk.CTkFont(size=11, weight="bold"),
            ).grid(row=1, column=column, pady=(0, 13))

        actions_shadow = ctk.CTkFrame(cards, fg_color=SHADOW, corner_radius=20, width=1, height=1)
        actions_shadow.grid(row=0, column=1, sticky="nsew", padx=(7, 8), pady=(5, 9))
        actions = ctk.CTkFrame(cards, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=20)
        actions.grid(row=0, column=1, sticky="nsew", padx=(7, 8), pady=(1, 13))
        actions.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            actions, text="Product Actions", text_color=TEXT,
            font=ctk.CTkFont(size=21, weight="bold"), anchor="w",
        ).grid(row=0, column=0, padx=24, pady=(22, 3), sticky="w")
        ctk.CTkLabel(
            actions, text="Open the catalog, add a style, or handle products that need setup.",
            text_color=MUTED, font=ctk.CTkFont(size=12), anchor="w", justify="left", wraplength=430,
        ).grid(row=1, column=0, padx=24, pady=(0, 14), sticky="w")
        self.master_page_open_button = ctk.CTkButton(
            actions, text="Open Product Catalog", command=self.open_product_master,
            height=50, corner_radius=10, fg_color=PURPLE, hover_color=PURPLE_DARK,
            font=ctk.CTkFont(size=15, weight="bold"),
        )
        self.master_page_open_button.grid(row=2, column=0, sticky="ew", padx=24, pady=(0, 9))
        self.master_page_secondary_button = ctk.CTkButton(
            actions, text="+ Add New Product", command=self.add_new_product,
            height=46, corner_radius=10, fg_color=PURPLE_LIGHT, hover_color="#E6D9F5",
            border_width=1, border_color="#BFA9DD", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=14, weight="bold"),
        )
        self.master_page_secondary_button.grid(row=3, column=0, sticky="ew", padx=24, pady=(0, 9))
        self.master_page_review_button = ctk.CTkButton(
            actions, text="No Products Need Setup", command=self.open_next_product_setup,
            height=44, corner_radius=10, fg_color="transparent", hover_color=PURPLE_LIGHT,
            border_width=1, border_color=BORDER, text_color=MUTED,
            font=ctk.CTkFont(size=13, weight="bold"), state="disabled",
        )
        self.master_page_review_button.grid(row=4, column=0, sticky="ew", padx=24, pady=(0, 20))

        readiness_shadow = ctk.CTkFrame(workspace, fg_color=SHADOW, corner_radius=20, width=1, height=1)
        readiness_shadow.grid(row=2, column=0, sticky="nsew", padx=8, pady=(5, 4))
        readiness = ctk.CTkFrame(
            workspace, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=20,
        )
        readiness.grid(row=2, column=0, sticky="nsew", padx=8, pady=(1, 8))
        readiness.grid_columnconfigure(1, weight=1)
        readiness.grid_rowconfigure(0, weight=1)
        self._add_orchid_watermark(
            readiness, size=(260, 260), relx=0.93, rely=0.60,
            anchor="center", opacity=0.105,
        )
        self.master_ready_icon = ctk.CTkLabel(
            readiness, text="✓", width=68, height=68, corner_radius=34,
            fg_color="#EAF7EE", text_color=SUCCESS,
            font=ctk.CTkFont(size=30, weight="bold"),
        )
        self.master_ready_icon.grid(row=0, column=0, rowspan=4, padx=(28, 20), pady=24, sticky="w")
        self.master_state_title = ctk.CTkLabel(
            readiness, text="Your Product Master is ready", text_color=TEXT,
            font=ctk.CTkFont(size=23, weight="bold"), anchor="w", justify="left",
        )
        self.master_state_title.grid(row=0, column=1, padx=(0, 250), pady=(25, 2), sticky="sw")
        self.master_state_body = ctk.CTkLabel(
            readiness, text="All permanent purchasing and decoration settings are complete.",
            text_color=MUTED, font=ctk.CTkFont(size=14), anchor="w", justify="left", wraplength=680,
        )
        self.master_state_body.grid(row=1, column=1, padx=(0, 250), pady=(0, 5), sticky="nw")
        self.master_issue_identity = ctk.CTkLabel(
            readiness, text="", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=16, weight="bold"), anchor="w", justify="left", wraplength=680,
        )
        self.master_issue_identity.grid(row=2, column=1, padx=(0, 250), pady=(0, 2), sticky="nw")
        self.master_issue_reason = ctk.CTkLabel(
            readiness, text="", text_color=WARNING,
            font=ctk.CTkFont(size=13, weight="bold"), anchor="w", justify="left", wraplength=680,
        )
        self.master_issue_reason.grid(row=3, column=1, padx=(0, 250), pady=(0, 24), sticky="nw")
        self._active_product_setup_issue = None

    def _product_master_issues(self) -> list[dict]:
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            return []
        snapshot = self.last_snapshot or self._mission_snapshot()
        return [
            issue for issue in snapshot.get("issues", [])
            if _clean(issue.get("source", "")) == "Product Master"
        ]

    def open_next_product_setup(self):
        issues = self._product_master_issues()
        if not issues:
            self.open_product_master()
            return
        issue = issues[0]
        product = _clean(issue.get("product", "")) or _clean(issue.get("description", ""))
        self.open_product_master(setup_product=product)

    def refresh_master_page(self):
        if not hasattr(self, "master_state_title"):
            return
        health = _product_master_health()
        issues = self._product_master_issues()
        percent = max(0, min(100, int(health.get("percent", 0) or 0)))
        styles = int(health.get("styles", 0) or 0)
        complete = int(health.get("complete", 0) or 0)
        incomplete = int(health.get("incomplete", 0) or 0)
        self.master_health_large.configure(text=f"{percent}% Ready", text_color=PURPLE if incomplete == 0 else WARNING)
        self.master_progress.set(percent / 100 if styles else 0)
        self.master_progress.configure(progress_color=SUCCESS if incomplete == 0 and styles else (WARNING if incomplete else PURPLE))
        self.master_products_value.configure(text=f"{styles:,}")
        self.master_need_setup_value.configure(text=f"{incomplete:,}", text_color=WARNING if incomplete else SUCCESS)
        self.master_ready_value.configure(text=f"{complete:,}", text_color=SUCCESS)
        self.master_details.configure(
            text=f"{styles:,} permanent product style{'s' if styles != 1 else ''} in the catalog"
        )

        if issues:
            issue = issues[0]
            self._active_product_setup_issue = issue
            count = len(issues)
            self.master_step_badge.configure(
                text=f"{count} Need Setup", fg_color="#FFF1DD", text_color="#9A4D00",
            )
            self.master_ready_icon.configure(text="!", fg_color="#FFF1DD", text_color=WARNING)
            self.master_state_title.configure(text="Product setup needs attention")
            self.master_state_body.configure(
                text="Complete the next product once and Orchid will reuse those settings on every future purchase packet."
            )
            colors = _clean(issue.get("garment_colors", ""))
            identity_parts = [_clean(issue.get("product", "")), _clean(issue.get("description", ""))]
            identity = " — ".join(dict.fromkeys(part for part in identity_parts if part))
            detail_parts = [_clean(issue.get("reason", ""))]
            if colors:
                detail_parts.append(f"Colors in this order: {colors}")
            self.master_issue_identity.configure(text=identity or "Next product setup item")
            self.master_issue_reason.configure(text="  •  ".join(part for part in detail_parts if part))
            self.master_page_open_button.configure(
                text=f"Review {count} Product{'s' if count != 1 else ''} Needing Setup",
                command=self.open_next_product_setup, state="normal",
            )
            self.master_page_secondary_button.configure(
                text="Open Product Catalog", command=self.open_product_master, state="normal",
            )
            self.master_page_review_button.configure(
                text="+ Add New Product", command=self.add_new_product, state="normal",
                text_color=PURPLE_DARK,
            )
        elif self.last_review_workbook and self._review_is_stale():
            self._active_product_setup_issue = None
            self.master_step_badge.configure(
                text="Saved Changes Ready", fg_color=PURPLE_LIGHT, text_color=PURPLE_DARK,
            )
            self.master_ready_icon.configure(text="↻", fg_color=PURPLE_LIGHT, text_color=PURPLE)
            self.master_state_title.configure(text="Product Master changes are ready")
            self.master_state_body.configure(
                text="Apply the saved catalog changes to the active purchase packet before continuing Purchase Review."
            )
            self.master_issue_identity.configure(text="")
            self.master_issue_reason.configure(text="")
            self.master_page_open_button.configure(
                text="Apply Changes to Purchase Review", command=self.regenerate_current_review, state="normal",
            )
            self.master_page_secondary_button.configure(
                text="Open Product Catalog", command=self.open_product_master, state="normal",
            )
            self.master_page_review_button.configure(
                text="+ Add New Product", command=self.add_new_product, state="normal",
                text_color=PURPLE_DARK,
            )
        else:
            self._active_product_setup_issue = None
            ready = bool(styles and incomplete == 0)
            self.master_step_badge.configure(
                text="All Products Ready" if ready else "Product Master",
                fg_color="#EAF7EE" if ready else PURPLE_LIGHT,
                text_color="#166534" if ready else PURPLE_DARK,
            )
            self.master_ready_icon.configure(
                text="✓" if ready else "◇",
                fg_color="#EAF7EE" if ready else PURPLE_LIGHT,
                text_color=SUCCESS if ready else PURPLE,
            )
            self.master_state_title.configure(
                text="Your Product Master is ready" if ready else "Build your permanent product catalog"
            )
            self.master_state_body.configure(
                text=(f"All {styles:,} products have complete purchasing and decoration settings."
                      if ready else "Add products and complete permanent purchasing and decoration settings once."),
            )
            self.master_issue_identity.configure(text="")
            self.master_issue_reason.configure(text="")
            self.master_page_open_button.configure(
                text="Open Product Catalog", command=self.open_product_master, state="normal",
            )
            self.master_page_secondary_button.configure(
                text="+ Add New Product", command=self.add_new_product, state="normal",
            )
            self.master_page_review_button.configure(
                text="No Products Need Setup" if ready else "Review Product Setup",
                command=self.open_next_product_setup,
                state="disabled" if ready else "normal",
                text_color=MUTED if ready else PURPLE_DARK,
            )

    def build_review_page(self):
        page = self.new_page("review")
        # new_page reserves row 9 for some dashboard layouts. Purchase Review
        # uses row 1 as its expanding workspace, so remove that empty weighted
        # row and give every available vertical pixel to the editor.
        page.grid_rowconfigure(9, weight=0, minsize=0)
        page.grid_rowconfigure(1, weight=1, minsize=430)
        self._build_workflow_tracker(page, 0, "review_stage_labels", padx=34)
        review_shadow = ctk.CTkFrame(page, fg_color=SHADOW, corner_radius=24, width=1, height=1)
        review_shadow.grid(row=1, column=0, sticky="nsew", padx=20, pady=(1, 3))
        shell = ctk.CTkFrame(page, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=24)
        shell.grid(row=1, column=0, sticky="nsew", padx=20, pady=(0, 4))
        shell.grid_columnconfigure(0, weight=1)
        # Keep the editable decision form roomy on smaller Mac displays while
        # leaving the action bar fixed below it.
        shell.grid_rowconfigure(4, weight=1, minsize=410)
        self.review_shell = shell
        self._add_orchid_watermark(shell, size=(520, 520), relx=0.5, rely=0.58)

        self.review_step_label = ctk.CTkLabel(shell, text="", text_color=PURPLE,
                                              font=ctk.CTkFont(size=1))
        self.review_state_title = ctk.CTkLabel(shell, text="Complete Purchase Review", text_color=PURPLE_DARK,
                                               font=ctk.CTkFont(size=29, weight="bold"), justify="center")
        self.review_state_title.grid(row=0, column=0, padx=24, pady=(7, 1))
        self.review_state_body = ctk.CTkLabel(shell, text="Resolve only the order decisions that need attention.",
                                              text_color=MUTED, font=ctk.CTkFont(size=13), justify="center")
        self.review_state_body.grid(row=1, column=0, padx=24, pady=(0, 2))
        self.review_progress = ctk.CTkLabel(shell, text="", text_color=TEXT,
                                            font=ctk.CTkFont(size=13, weight="bold"))
        self.review_progress.grid(row=2, column=0, pady=(0, 1))

        self.review_decision_frame = ctk.CTkScrollableFrame(
            shell, fg_color="#FCFBFE", border_width=1, border_color=BORDER,
            corner_radius=16, scrollbar_button_color="#CDBBE5", scrollbar_button_hover_color=PURPLE,
        )
        self.review_decision_frame.grid(row=4, column=0, sticky="nsew", padx=8, pady=(1, 0))
        self.review_decision_frame.grid_columnconfigure(1, weight=1)

        self.review_identity = ctk.CTkLabel(self.review_decision_frame, text="", text_color=TEXT,
                                            font=ctk.CTkFont(size=18, weight="bold"), anchor="w", justify="left")
        self.review_identity.grid(row=0, column=0, columnspan=2, padx=14, pady=(8, 3), sticky="ew")
        self.review_note_callout = ctk.CTkFrame(
            self.review_decision_frame, fg_color="#FFF6D9", border_width=2,
            border_color="#E2A21A", corner_radius=14,
        )
        self.review_note_callout.grid(row=1, column=0, columnspan=2, padx=14, pady=(3, 6), sticky="ew")
        self.review_note_callout.grid_columnconfigure(0, weight=1)
        self.review_note_callout_title = ctk.CTkLabel(
            self.review_note_callout, text="⚠  CUSTOMER INSTRUCTION — REVIEW REQUIRED",
            text_color="#8A5200", font=ctk.CTkFont(size=15, weight="bold"), anchor="w", justify="left",
        )
        self.review_note_callout_title.grid(row=0, column=0, padx=14, pady=(9, 3), sticky="ew")
        self.review_note_callout_text = ctk.CTkLabel(
            self.review_note_callout, text="", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=18, weight="bold"), anchor="w", justify="left", wraplength=900,
        )
        self.review_note_callout_text.grid(row=1, column=0, padx=14, pady=(0, 10), sticky="ew")
        self.review_note_callout.grid_remove()

        self.review_reason = ctk.CTkLabel(self.review_decision_frame, text="", text_color=WARNING,
                                          font=ctk.CTkFont(size=14, weight="bold"), anchor="w", justify="left", wraplength=820)
        self.review_reason.grid(row=1, column=0, columnspan=2, padx=14, pady=(0, 3), sticky="ew")
        self.review_instructions = ctk.CTkLabel(self.review_decision_frame, text="", text_color=MUTED,
                                                font=ctk.CTkFont(size=13), anchor="w", justify="left", wraplength=900)
        self.review_instructions.grid(row=2, column=0, columnspan=2, padx=14, pady=(0, 5), sticky="ew")
        self.review_recommendation = ctk.CTkLabel(
            self.review_decision_frame, text="", fg_color="#EFE7FB", text_color=PURPLE_DARK,
            corner_radius=10, font=ctk.CTkFont(size=13, weight="bold"),
            anchor="w", justify="left", wraplength=900,
        )
        self.review_recommendation.grid(row=3, column=0, columnspan=2, padx=14, pady=(0, 6), sticky="ew")
        self.review_recommendation.grid_remove()

        self.review_vars = {
            "Product #": ctk.StringVar(), "Garment Color": ctk.StringVar(), "Size": ctk.StringVar(),
            "Quantity": ctk.StringVar(), "Purchase Vendor": ctk.StringVar(),
            "Decoration Type": ctk.StringVar(), "Decoration Location": ctk.StringVar(value=LEFT_CHEST),
            "Decoration Placement Instructions": ctk.StringVar(), "Decoration Color": ctk.StringVar(),
            "Include": ctk.StringVar(value="Yes"),
            "Do Not Outsource": ctk.StringVar(value="No"),
            "Decoration Decision": ctk.StringVar(value="Choose an instruction decision"),
        }
        self.review_custom_location_var = ctk.StringVar()
        self.review_vars["Decoration Type"].trace_add("write", self._on_review_decoration_type_changed)
        self.review_decoration_decision_label = ctk.CTkLabel(
            self.review_decision_frame, text="Instruction Decision", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=13, weight="bold"), anchor="w"
        )
        self.review_decoration_decision_label.grid(row=4, column=0, padx=(22, 12), pady=6, sticky="w")
        self.review_decoration_decision_menu = ctk.CTkOptionMenu(
            self.review_decision_frame, variable=self.review_vars["Decoration Decision"],
            values=INSTRUCTION_DECISION_OPTIONS,
            command=self._apply_review_decoration_decision, height=44,
            fg_color=PURPLE, button_color=PURPLE_DARK,
        )
        self.review_decoration_decision_menu.grid(row=4, column=1, padx=(0, 22), pady=6, sticky="ew")

        self.review_entries = {}
        labels = (
            ("Product #", 5), ("Garment Color", 6), ("Size", 7), ("Quantity", 8),
            ("Purchase Vendor", 9), ("Decoration Type", 10), ("Decoration Location", 11),
            ("Decoration Placement Instructions", 13), ("Decoration Color", 14),
        )
        for label, row in labels:
            display_label = "Placement Instructions" if label == "Decoration Placement Instructions" else label
            ctk.CTkLabel(self.review_decision_frame, text=display_label, text_color=TEXT,
                         font=ctk.CTkFont(size=13, weight="bold"), anchor="w").grid(
                row=row, column=0, padx=(14, 10), pady=4, sticky="w")
            if label == "Decoration Location":
                self.review_location_menu = ctk.CTkOptionMenu(
                    self.review_decision_frame, variable=self.review_vars[label],
                    values=DECORATION_LOCATIONS, command=self._apply_review_location_choice,
                    height=42, fg_color=PURPLE, button_color=PURPLE_DARK,
                )
                self.review_location_menu.grid(row=row, column=1, padx=(0, 14), pady=4, sticky="ew")
            elif label == "Decoration Type":
                self.review_decoration_type_menu = ctk.CTkOptionMenu(
                    self.review_decision_frame, variable=self.review_vars[label],
                    values=["Embroidery", "Screen Print", SEW_ON_PATCH_LABEL, HEMMING_ALTERATION_LABEL, BLANK_DECORATION_LABEL],
                    command=lambda _choice: self._on_review_decoration_type_changed(),
                    height=42, fg_color=PURPLE, button_color=PURPLE_DARK,
                )
                self.review_decoration_type_menu.grid(row=row, column=1, padx=(0, 14), pady=4, sticky="ew")
            else:
                entry = ctk.CTkEntry(
                    self.review_decision_frame, textvariable=self.review_vars[label], height=42,
                    fg_color=WHITE, border_color=BORDER, border_width=1,
                )
                entry.grid(row=row, column=1, padx=(0, 14), pady=4, sticky="ew")
                self.review_entries[label] = entry

        # The reserved panel gives the expanded location menu a clean area instead
        # of letting it visually cover Include Item. It also holds custom input.
        self.review_location_detail = ctk.CTkFrame(self.review_decision_frame, fg_color="#F7F3FC", corner_radius=10, height=64)
        self.review_location_detail.grid(row=12, column=0, columnspan=2, padx=14, pady=(1, 5), sticky="ew")
        self.review_location_detail.grid_columnconfigure(1, weight=1)
        self.review_location_detail.grid_propagate(False)
        self.review_location_helper = ctk.CTkLabel(
            self.review_location_detail,
            text="Location routes are separated on the vendor PO. Choose Other / Custom for a rare placement.",
            text_color=MUTED, font=ctk.CTkFont(size=12), anchor="w", justify="left", wraplength=850,
        )
        self.review_location_helper.grid(row=0, column=0, columnspan=2, padx=12, pady=(10, 5), sticky="ew")
        self.review_custom_location_label = ctk.CTkLabel(
            self.review_location_detail, text="Custom Location *", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=13, weight="bold"), anchor="w",
        )
        self.review_custom_location_entry = ctk.CTkEntry(
            self.review_location_detail, textvariable=self.review_custom_location_var, height=42,
            placeholder_text="Example: Right Chest or Center Back",
        )
        self.review_custom_location_label.grid(row=1, column=0, padx=(12, 10), pady=(3, 10), sticky="w")
        self.review_custom_location_entry.grid(row=1, column=1, padx=(0, 12), pady=(3, 10), sticky="ew")
        self.review_custom_location_label.grid_remove()
        self.review_custom_location_entry.grid_remove()

        ctk.CTkLabel(self.review_decision_frame, text="Include Item", text_color=TEXT,
                     font=ctk.CTkFont(size=13, weight="bold"), anchor="w").grid(
            row=15, column=0, padx=(14, 10), pady=4, sticky="w")
        self.review_include_menu = ctk.CTkOptionMenu(
            self.review_decision_frame, variable=self.review_vars["Include"], values=["Yes", "No"], height=42
        )
        self.review_include_menu.grid(row=15, column=1, padx=(0, 14), pady=4, sticky="w")

        # Keep review actions pinned at the bottom of the full-height page.
        # The form itself scrolls independently, so the buttons stay available
        # and no longer consume space inside the editable field list.
        action_bar = ctk.CTkFrame(shell, fg_color="transparent")
        action_bar.grid(row=5, column=0, padx=12, pady=(4, 10), sticky="e")
        self.review_product_master_button = ctk.CTkButton(
            action_bar, text="Open Product Master", command=self.open_product_master, width=180, height=44,
            fg_color=PURPLE_LIGHT, hover_color="#E6D9F5", border_width=1,
            border_color="#BFA9DD", text_color=PURPLE_DARK,
        )
        self.review_product_master_button.pack(side="left", padx=4)
        self.review_exit_button = ctk.CTkButton(
            action_bar, text="Save Progress & Exit", command=self.save_review_progress_and_exit,
            width=175, height=42, fg_color="transparent", hover_color=PURPLE_LIGHT,
            border_width=1, border_color="#BFA9DD", text_color=PURPLE_DARK,
        )
        self.review_exit_button.pack(side="left", padx=4)
        self.review_previous_button = ctk.CTkButton(
            action_bar, text="← Previous Decision", command=self.open_previous_review_decision,
            width=180, height=44, fg_color="transparent", hover_color=PURPLE_LIGHT,
            border_width=1, border_color="#BFA9DD", text_color=PURPLE_DARK, state="disabled",
        )
        self.review_previous_button.pack(side="left", padx=4)
        self.review_search_button = ctk.CTkButton(
            action_bar, text="Search Review", command=self.open_purchase_review_search,
            width=145, height=44, fg_color=PURPLE_LIGHT, hover_color="#E6D9F5",
            border_width=1, border_color="#BFA9DD", text_color=PURPLE_DARK,
        )
        self.review_search_button.pack(side="left", padx=4)
        self.review_export_button = ctk.CTkButton(
            action_bar, text="Export/Open Excel Copy", command=self.open_latest_review, width=190, height=44,
            fg_color="transparent", hover_color=PURPLE_LIGHT, border_width=1, border_color=BORDER, text_color=MUTED,
        )
        self.review_export_button.pack(side="left", padx=4)
        self.review_test_generate_button = ctk.CTkButton(
            action_bar, text="Test Generate POs", command=self.generate_test_purchase_orders,
            width=165, height=44, fg_color="#FFF1DC", hover_color="#F5DEB6",
            border_width=1, border_color="#E2A21A", text_color="#8A5200",
        )
        self.review_test_generate_button.pack(side="left", padx=4)
        self.review_save_button = ctk.CTkButton(
            action_bar, text="Save & Next", command=self.save_current_review_decision,
            width=160, height=46, fg_color=PURPLE, hover_color=PURPLE_DARK,
            font=ctk.CTkFont(size=14, weight="bold"),
        )
        self.review_save_button.pack(side="left", padx=4)
        self._review_enter_last_at = 0.0
        self.bind("<Return>", self._handle_review_enter, add="+")
        self.bind("<KP_Enter>", self._handle_review_enter, add="+")
        self._active_review_issue = None
        self._review_decision_history = []
        self._review_history_index = None
        self._review_search_issue = None
    def _set_review_location_value(self, location: str):
        decoration_type = self.review_vars["Decoration Type"].get()
        normalized = normalize_decoration_location(
            location, decoration_type,
            requires_decoration=not is_blank_decoration(decoration_type) and not is_in_house_decoration(decoration_type),
            product_name=(self._active_review_issue or {}).get("description", ""),
        )
        choice, custom = location_choice_and_custom(normalized, decoration_type)
        self.review_vars["Decoration Location"].set(choice)
        self.review_custom_location_var.set(custom)
        self._apply_review_location_choice(choice)

    def _apply_review_location_choice(self, choice: str):
        choice = _clean(choice)
        is_custom = choice == OTHER_CUSTOM
        if is_custom:
            self.review_location_detail.configure(height=112)
            self.review_location_helper.configure(
                text="Enter the exact custom route below. It will receive its own labeled PO section/page."
            )
            self.review_custom_location_label.grid()
            self.review_custom_location_entry.grid()
            self.review_custom_location_entry.configure(state="normal")
        else:
            self.review_location_detail.configure(height=64)
            self.review_custom_location_var.set("")
            self.review_custom_location_label.grid_remove()
            self.review_custom_location_entry.grid_remove()
            if choice in {NOT_APPLICABLE_NO_DECORATION, NOT_APPLICABLE_IN_HOUSE}:
                self.review_location_helper.configure(text=choice)
            else:
                self.review_location_helper.configure(
                    text="Location routes are separated on the vendor PO. Choose Other / Custom for a rare placement."
                )

    def _resolved_review_location(self, decoration_type: str) -> str:
        if is_blank_decoration(decoration_type) or is_in_house_decoration(decoration_type):
            return ""
        choice = self.review_vars["Decoration Location"].get()
        location = resolve_location_choice(choice, self.review_custom_location_var.get())
        return normalize_decoration_location(
            location, decoration_type, requires_decoration=True,
            product_name=(self._active_review_issue or {}).get("description", ""),
        )

    def _on_review_decoration_type_changed(self, *_):
        if not hasattr(self, "review_location_menu"):
            return
        decoration_type = self.review_vars["Decoration Type"].get()
        in_house = is_in_house_decoration(decoration_type)
        no_decoration = is_blank_decoration(decoration_type)
        disabled = in_house or no_decoration
        if disabled:
            display = NOT_APPLICABLE_IN_HOUSE if in_house else NOT_APPLICABLE_NO_DECORATION
            self.review_vars["Decoration Location"].set(display)
            self.review_custom_location_var.set("")
            self.review_vars["Decoration Placement Instructions"].set("")
            self.review_vars["Decoration Color"].set("")
            self.review_location_menu.configure(state="disabled")
            self.review_custom_location_entry.configure(state="disabled")
            placement = self.review_entries.get("Decoration Placement Instructions")
            color = self.review_entries.get("Decoration Color")
            if placement is not None:
                placement.configure(state="disabled")
            if color is not None:
                color.configure(state="disabled")
            self._apply_review_location_choice(display)
            self._update_instruction_decision_options()
        else:
            current = self.review_vars["Decoration Location"].get().strip()
            if not current or current in {NOT_APPLICABLE_NO_DECORATION, NOT_APPLICABLE_IN_HOUSE}:
                self._set_review_location_value(default_decoration_location(
                    (self._active_review_issue or {}).get("description", ""), "", decoration_type
                ))
            self.review_location_menu.configure(state="normal")
            placement = self.review_entries.get("Decoration Placement Instructions")
            color = self.review_entries.get("Decoration Color")
            if placement is not None:
                placement.configure(state="normal")
            if color is not None:
                color.configure(state="normal")
            self._update_instruction_decision_options()

    def _review_outsourcing_exception_applicable(self, decoration_type: str | None = None) -> bool:
        """Return whether this line can be redirected from an outside decorator.

        In an Entire Order Outsourced packet, every included physical garment is
        normally routed to the outside decorator. The exception must therefore
        remain available even when the user changes the decoration type to
        Blank Garment (No Decoration), because that garment still needs to be
        purchased separately and shipped to Orchid. Internal service-only lines
        do not use this routing exception.
        """
        decoration_value = _clean(
            decoration_type if decoration_type is not None else self.review_vars["Decoration Type"].get()
        )
        decoration_key = decoration_value.casefold()
        entire_order_outsourced = (
            normalize_decoration_fulfillment(self.current_decoration_fulfillment)
            == ENTIRE_ORDER_OUTSOURCED
        )
        if entire_order_outsourced:
            return not is_in_house_decoration(decoration_value)
        return "screen" in decoration_key

    def _review_instruction_decision_options(self, decoration_type: str | None = None) -> list[str]:
        instruction_kind = getattr(self, "_review_instruction_kind", "decoration")
        if instruction_kind == "general":
            return [
                "Choose an instruction decision",
                "Follow Note as Written",
                "Keep Product Master Default",
            ]
        options = list(INSTRUCTION_DECISION_OPTIONS)
        if self._review_outsourcing_exception_applicable(decoration_type):
            options.append(DO_NOT_OUTSOURCE_DECISION)
        return options

    def _update_instruction_decision_options(self):
        if not hasattr(self, "review_decoration_decision_menu"):
            return
        options = self._review_instruction_decision_options()
        self.review_decoration_decision_menu.configure(values=options)
        current = _clean(self.review_vars["Decoration Decision"].get())
        if current not in options:
            if current == DO_NOT_OUTSOURCE_DECISION:
                self.review_vars["Do Not Outsource"].set("No")
            self.review_vars["Decoration Decision"].set("Choose an instruction decision")

    def _save_current_review_draft(self) -> bool:
        """Save the current in-progress form to the lightweight recovery file.

        Draft saving no longer rewrites the complete Excel workbook. The values
        are restored when Purchase Review is reopened and are synchronized to the
        workbook only after the user completes the decision with Save & Next.
        """
        issue = getattr(self, "_active_review_issue", None)
        if not issue or not self.last_review_workbook or not self.last_review_workbook.exists():
            return True
        if _clean(issue.get("fix_in", issue.get("source", ""))).casefold() == "product master":
            return True
        values = {field: _clean(variable.get()) for field, variable in self.review_vars.items()}
        if values.get("Decoration Location") == OTHER_CUSTOM and _clean(self.review_custom_location_var.get()):
            values["Decoration Location"] = _clean(self.review_custom_location_var.get())
        self._update_review_state(
            draft_line_id=_clean(issue.get("line_id", "")),
            draft_decision_key=_clean(issue.get("decision_key", "")),
            draft_values=values,
            draft_saved_at=time.time(),
        )
        return True

    def save_review_progress_and_exit(self):
        if not self._save_current_review_draft():
            return
        self.dashboard_status.configure(
            text="Purchase Review draft saved. Reopen Purchase Review to continue from this line.",
            text_color=SUCCESS,
        )
        if self._review_save_pending:
            self._review_exit_after_saves = True
            self.review_state_title.configure(text="Finishing Purchase Review Saves…")
            self.review_state_body.configure(text="Your draft is protected. Orchid is finishing the completed decisions before returning to Current Event.")
            self.review_save_button.configure(text="Finishing Saves…", state="disabled")
            self.review_exit_button.configure(state="disabled")
            return
        self.refresh_dashboard()
        self.show_page("import")

    def _apply_review_decoration_decision(self, choice: str):
        """Apply the user's explicit instruction decision to the editable order fields."""
        choice = _clean(choice)
        if choice == DO_NOT_OUTSOURCE_DECISION:
            if not self._review_outsourcing_exception_applicable():
                messagebox.showwarning(
                    "Do Not Outsource Is Not Applicable",
                    "This exception is only available when the current decoration route would normally be outsourced.",
                )
                self.review_vars["Do Not Outsource"].set("No")
                self.review_vars["Decoration Decision"].set("Choose an instruction decision")
                return
            self.review_vars["Include"].set("Yes")
            self.review_vars["Do Not Outsource"].set("Yes")
            return

        # Preserve Product Master's automatic route unless the user explicitly
        # chooses the one-item outsourcing exception.
        self.review_vars["Do Not Outsource"].set(
            getattr(self, "_review_default_do_not_outsource", "No") or "No"
        )
        if choice == "Keep Product Master Default":
            self.review_vars["Decoration Type"].set(getattr(self, "_review_default_decoration_type", ""))
            self._set_review_location_value(getattr(self, "_review_default_decoration_location", LEFT_CHEST))
            self.review_vars["Decoration Placement Instructions"].set(
                getattr(self, "_review_default_placement_instructions", "")
            )
            self.review_vars["Decoration Color"].set(getattr(self, "_review_default_decoration_color", ""))
        elif choice == "No Decoration":
            self.review_vars["Decoration Type"].set(BLANK_DECORATION_LABEL)
            self.review_vars["Decoration Location"].set("")
            self.review_custom_location_var.set("")
            self.review_vars["Decoration Placement Instructions"].set("")
            self.review_vars["Decoration Color"].set("")
            self._apply_review_location_choice("")
        elif choice in {"Embroidery", "Screen Print"}:
            self.review_vars["Decoration Type"].set(choice)
            if not self.review_vars["Decoration Location"].get().strip():
                self._set_review_location_value(default_decoration_location(
                    (self._active_review_issue or {}).get("description", ""), "", choice
                ))
        elif choice == "Sew On Patch":
            self.review_vars["Decoration Type"].set(SEW_ON_PATCH_LABEL)
        elif choice == "Hemming / Alteration":
            self.review_vars["Decoration Type"].set(HEMMING_ALTERATION_LABEL)
        elif choice == "Follow Note as Written":
            # Following an instruction note must never infer or rewrite the
            # product's decoration location. The location already loaded into
            # Purchase Review comes from Product Master (or an explicit manual
            # order override) and remains unchanged unless the user edits the
            # Decoration Location field directly. This prevents unrelated note
            # wording such as "black hats" from changing outerwear to Hat.
            pass

    def _update_previous_button(self):
        if not hasattr(self, "review_previous_button"):
            return
        if self._review_history_index is not None:
            enabled = self._review_history_index > 0
        else:
            enabled = bool(self._review_decision_history)
        self.review_previous_button.configure(state="normal" if enabled else "disabled")

    def _search_purchase_review_records(self, query: str) -> list[dict]:
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            return []
        needle = _clean(query).casefold()
        if not needle:
            return []
        try:
            records = load_review_lines(self.last_review_workbook)
        except Exception:
            return []

        matches = []
        seen = set()
        for record in records:
            line_id = _clean(record.get("Line ID", ""))
            identity = line_id or "|".join(_clean(record.get(field, "")) for field in ("Order Number", "Product #", "Garment Color", "Size"))
            if identity in seen:
                continue
            searchable = " ".join(_clean(record.get(field, "")) for field in (
                "Product #", "Description", "Employee Name", "Company", "Order Number",
                "Garment Color", "Purchase Vendor",
            )).casefold()
            if needle not in searchable:
                continue
            seen.add(identity)
            style = _clean(record.get("Product #", ""))
            exact_style = style.casefold() == needle
            record_copy = dict(record)
            record_copy["_exact_style"] = exact_style
            matches.append(record_copy)
        matches.sort(key=lambda record: (
            0 if record.get("_exact_style") else 1,
            _clean(record.get("Product #", "")).casefold(),
            _clean(record.get("Employee Name", "")).casefold(),
            _clean(record.get("Order Number", "")).casefold(),
        ))
        return matches[:150]

    def open_purchase_review_search(self):
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            messagebox.showinfo("No Purchase Review", "Process imported orders before searching Purchase Review.")
            return

        window = ctk.CTkToplevel(self)
        window.title("Search Purchase Review")
        window.geometry("980x720")
        window.minsize(760, 560)
        window.transient(self)
        window.grab_set()
        window.configure(fg_color=BG)
        window.grid_columnconfigure(0, weight=1)
        window.grid_rowconfigure(3, weight=1)

        ctk.CTkLabel(
            window, text="Search Purchase Review", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=25, weight="bold"), anchor="w",
        ).grid(row=0, column=0, padx=24, pady=(20, 4), sticky="ew")
        ctk.CTkLabel(
            window,
            text="Search by product/style number, employee, company/department, or order number. Select a result to edit that exact line.",
            text_color=MUTED, font=ctk.CTkFont(size=13), anchor="w", justify="left", wraplength=900,
        ).grid(row=1, column=0, padx=24, pady=(0, 10), sticky="ew")

        search_row = ctk.CTkFrame(window, fg_color="transparent")
        search_row.grid(row=2, column=0, padx=24, pady=(0, 12), sticky="ew")
        search_row.grid_columnconfigure(0, weight=1)
        query_var = ctk.StringVar()
        entry = ctk.CTkEntry(
            search_row, textvariable=query_var, height=42,
            placeholder_text="Example: ST350, J799S, Brock Bennett, or 1513",
        )
        entry.grid(row=0, column=0, padx=(0, 8), sticky="ew")
        results_frame = ctk.CTkScrollableFrame(
            window, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=14,
            scrollbar_button_color="#CDBBE5", scrollbar_button_hover_color=PURPLE,
        )
        results_frame.grid(row=3, column=0, padx=24, pady=(0, 20), sticky="nsew")
        results_frame.grid_columnconfigure(0, weight=1)
        register_scrollable(self, results_frame)

        def open_result(record: dict):
            self._review_history_index = None
            self._review_search_issue = {
                "line_id": _clean(record.get("Line ID", "")),
                "decision_key": _clean(record.get("Decision Key", "")),
                "order": f"#{_clean(record.get('Order Number', ''))}" if _clean(record.get("Order Number", "")) else "",
                "employee": _clean(record.get("Employee Name", "")),
                "product": _clean(record.get("Product #", "")),
                "description": _clean(record.get("Description", "")),
                "reason": "Search result selected for correction",
                "instructions": "Update the incorrect order information, then save the line.",
                "fix_in": "Purchase Review",
                "source": "Purchase Review Search",
            }
            window.grab_release()
            window.destroy()
            self.show_page("review")
            self.refresh_review_page()

        def render_results(*_args):
            for child in results_frame.winfo_children():
                child.destroy()
            query = _clean(query_var.get())
            if not query:
                ctk.CTkLabel(
                    results_frame, text="Enter a style, employee, or order number to search.",
                    text_color=MUTED, font=ctk.CTkFont(size=14),
                ).grid(row=0, column=0, padx=20, pady=30)
                return
            matches = self._search_purchase_review_records(query)
            if not matches:
                ctk.CTkLabel(
                    results_frame, text=f"No Purchase Review lines matched ‘{query}’.",
                    text_color=WARNING, font=ctk.CTkFont(size=14, weight="bold"),
                ).grid(row=0, column=0, padx=20, pady=30)
                return
            ctk.CTkLabel(
                results_frame, text=f"{len(matches)} matching line{'s' if len(matches) != 1 else ''}",
                text_color=PURPLE_DARK, font=ctk.CTkFont(size=13, weight="bold"), anchor="w",
            ).grid(row=0, column=0, padx=12, pady=(10, 4), sticky="ew")
            for index, record in enumerate(matches, start=1):
                card = ctk.CTkFrame(
                    results_frame, fg_color="#FCFBFE", border_width=1, border_color=BORDER, corner_radius=12,
                )
                card.grid(row=index, column=0, padx=10, pady=5, sticky="ew")
                card.grid_columnconfigure(0, weight=1)
                style = _clean(record.get("Product #", "")) or "No style"
                description = _clean(record.get("Description", ""))
                employee = _clean(record.get("Employee Name", "")) or "Employee not listed"
                company = _clean(record.get("Company", ""))
                order = _clean(record.get("Order Number", ""))
                color = _clean(record.get("Garment Color", "")) or "No color"
                size = _clean(record.get("Size", "")) or "No size"
                qty = _clean(record.get("Quantity", "")) or "0"
                vendor = _clean(record.get("Purchase Vendor", "")) or "Vendor not assigned"
                decoration = " — ".join(filter(None, [
                    _clean(record.get("Decoration Type", "")),
                    _clean(record.get("Decoration Location", "")),
                ])) or "Decoration not assigned"
                status = _clean(record.get("Review Status", "")) or "Current line"
                ctk.CTkLabel(
                    card, text=f"{style}  •  {description}", text_color=TEXT,
                    font=ctk.CTkFont(size=15, weight="bold"), anchor="w", justify="left",
                ).grid(row=0, column=0, padx=14, pady=(10, 2), sticky="ew")
                ctk.CTkLabel(
                    card,
                    text=f"{employee}{' — ' + company if company else ''}  •  Order #{order or '—'}  •  {color}  •  {size}  •  Qty {qty}",
                    text_color=PURPLE_DARK, font=ctk.CTkFont(size=12), anchor="w", justify="left",
                ).grid(row=1, column=0, padx=14, pady=1, sticky="ew")
                ctk.CTkLabel(
                    card, text=f"{vendor}  •  {decoration}  •  Status: {status}",
                    text_color=MUTED, font=ctk.CTkFont(size=12), anchor="w", justify="left",
                ).grid(row=2, column=0, padx=14, pady=(1, 10), sticky="ew")
                ctk.CTkButton(
                    card, text="Open Line", command=lambda item=dict(record): open_result(item),
                    width=115, height=34, fg_color=PURPLE, hover_color=PURPLE_DARK,
                ).grid(row=0, column=1, rowspan=3, padx=12, pady=10)

        ctk.CTkButton(
            search_row, text="Search", command=render_results, width=120, height=42,
            fg_color=PURPLE, hover_color=PURPLE_DARK,
        ).grid(row=0, column=1)
        entry.bind("<Return>", render_results)
        render_results()
        entry.focus_set()

    def open_previous_review_decision(self):
        """Edit the immediately preceding saved decision without reopening the queue."""
        if not self._review_decision_history:
            return
        if self._review_history_index is None:
            self._review_history_index = len(self._review_decision_history) - 1
        elif self._review_history_index > 0:
            self._review_history_index -= 1
        self.refresh_review_page()

    def _reset_review_navigation(self):
        self._review_decision_history = []
        self._review_history_index = None
        self._review_search_issue = None
        self._update_previous_button()

    def _set_decoration_decision_visibility(self, visible: bool):
        if visible:
            self._update_instruction_decision_options()
            self.review_decoration_decision_label.grid()
            self.review_decoration_decision_menu.grid()
            self.review_decoration_decision_menu.configure(state="normal")
        else:
            self.review_decoration_decision_label.grid_remove()
            self.review_decoration_decision_menu.grid_remove()
            if hasattr(self, "review_note_callout"):
                self._set_review_instruction_callout("", False)

    def _set_review_instruction_callout(self, note_text: str = "", visible: bool = False):
        """Show the customer instruction as the most prominent part of the review."""
        if visible and _clean(note_text):
            self._review_instruction_text = _clean(note_text)
            self.review_reason.grid_remove()
            self.review_note_callout_text.configure(text=_clean(note_text))
            self.review_note_callout.grid()
            self.review_decoration_decision_label.configure(text="Instruction Decision")
        else:
            self._review_instruction_text = ""
            self.review_note_callout.grid_remove()
            self.review_note_callout_text.configure(text="")
            self.review_reason.grid()
            self.review_decoration_decision_label.configure(text="Instruction Decision")


    def _handle_review_enter(self, event=None):
        """Use Return as Save & Next only on the active Purchase Review page."""
        if getattr(self, "active_page", "") != "review":
            return None
        if not getattr(self, "_active_review_issue", None):
            return None
        try:
            if str(self.review_save_button.cget("state")) != "normal":
                return "break"
        except Exception:
            return None
        now = time.monotonic()
        if now - float(getattr(self, "_review_enter_last_at", 0.0)) < 0.25:
            return "break"
        self._review_enter_last_at = now
        self.save_current_review_decision()
        return "break"

    def _reset_review_field_highlights(self) -> None:
        for entry in getattr(self, "review_entries", {}).values():
            try:
                entry.configure(fg_color=WHITE, border_color=BORDER, border_width=1)
            except Exception:
                pass

    def _highlight_review_required_fields(self, reason_text: str, values: dict) -> None:
        """Purple-highlight every missing field responsible for this decision."""
        self._reset_review_field_highlights()
        reason = _clean(reason_text).casefold()
        fields: list[str] = []
        checks = (
            ("missing product number", "Product #"),
            ("missing description", "Product #"),
            ("missing garment color", "Garment Color"),
            ("missing size", "Size"),
            ("missing purchase vendor", "Purchase Vendor"),
            ("missing decoration color", "Decoration Color"),
            ("quantity must be greater than zero", "Quantity"),
        )
        for phrase, field in checks:
            if phrase in reason and field not in fields:
                fields.append(field)
        for field in fields:
            entry = self.review_entries.get(field)
            if entry is not None:
                entry.configure(fg_color="#F4ECFF", border_color=PURPLE, border_width=3)
        if fields:
            first = self.review_entries.get(fields[0])
            if first is not None:
                def focus_first():
                    try:
                        first.focus_set()
                        first.icursor("end")
                    except Exception:
                        pass
                self.after_idle(focus_first)

    def _set_review_fields_state(self, state: str):
        for entry in self.review_entries.values():
            entry.configure(state=state)
        if hasattr(self, "review_decoration_type_menu"):
            self.review_decoration_type_menu.configure(state=state)
        if hasattr(self, "review_location_menu"):
            self.review_location_menu.configure(state=state)
        if hasattr(self, "review_custom_location_entry") and self.review_vars["Decoration Location"].get() == OTHER_CUSTOM:
            self.review_custom_location_entry.configure(state=state)
        if state == "normal":
            self._on_review_decoration_type_changed()

    def _invalidate_review_cache(self) -> None:
        self._review_cache_key = None
        self._review_cache_data = {}
        self._review_value_cache = {}
        self._review_status_cache = None

    def _load_review_cache(self) -> dict:
        """Read review rows once and reuse them across one UI refresh/save cycle."""
        path = self.last_review_workbook
        if not path or not path.exists():
            self._invalidate_review_cache()
            return {}
        key = _file_stamp(path)
        if key == self._review_cache_key and self._review_cache_data:
            return self._review_cache_data
        workbook = load_workbook(path, read_only=True, data_only=False)
        data = {"sheets": {}, "values": {}, "completed": 0, "remaining": 0}
        decision_statuses = {}
        context_fields = list(getattr(self, "review_vars", {})) + [
            "Purchase Instructions", "Shopify Order Notes", "Shopify Line Notes",
            "Review Reason", "Action Required", "Resolution", "Review Status",
        ]
        for sheet_name in ("Review & Edit", "All PO Lines"):
            if sheet_name not in workbook.sheetnames:
                continue
            sheet = workbook[sheet_name]
            header_row, headers = None, {}
            for row_number in range(1, min(sheet.max_row, 20) + 1):
                row_values = [str(sheet.cell(row_number, col).value or "").strip() for col in range(1, sheet.max_column + 1)]
                if "Line ID" in row_values:
                    header_row = row_number
                    headers = {value: idx + 1 for idx, value in enumerate(row_values) if value}
                    break
            if not header_row:
                continue
            line_rows, decision_rows = {}, {}
            for row_number in range(header_row + 1, sheet.max_row + 1):
                line_id = _clean(sheet.cell(row_number, headers.get("Line ID", 0)).value) if headers.get("Line ID") else ""
                decision_key = _clean(sheet.cell(row_number, headers.get("Decision Key", 0)).value) if headers.get("Decision Key") else ""
                if line_id:
                    line_rows.setdefault(line_id, []).append(row_number)
                if decision_key:
                    decision_rows.setdefault(decision_key, []).append(row_number)
                cache_key = decision_key or line_id
                if cache_key and cache_key not in data["values"]:
                    data["values"][cache_key] = {
                        field: (_clean(sheet.cell(row_number, headers[field]).value) if field in headers else "")
                        for field in context_fields
                    }
                if sheet_name == "Review & Edit":
                    include = _clean(sheet.cell(row_number, headers.get("Include", 0)).value) if headers.get("Include") else "Yes"
                    if include.casefold() in {"yes", "y", "true", "1", "include"}:
                        status = _clean(sheet.cell(row_number, headers.get("Review Status", 0)).value).casefold() if headers.get("Review Status") else ""
                        decision_statuses.setdefault(cache_key or f"row:{row_number}", []).append(status)
            data["sheets"][sheet_name] = {
                "header_row": header_row, "headers": headers,
                "line_rows": line_rows, "decision_rows": decision_rows,
            }
        workbook.close()
        total = len(decision_statuses)
        completed = sum(1 for statuses in decision_statuses.values() if statuses and all(status == "ready" for status in statuses))
        data["completed"] = completed
        data["remaining"] = max(total - completed, 0)
        self._review_cache_key = key
        self._review_cache_data = data
        return data

    def _review_sheet_context(self):
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            return None, None, None
        workbook = load_workbook(self.last_review_workbook)
        if "Review & Edit" not in workbook.sheetnames:
            workbook.close()
            return None, None, None
        sheet = workbook["Review & Edit"]
        header_row = None
        headers = {}
        for row in range(1, min(sheet.max_row, 20) + 1):
            values = [str(sheet.cell(row, col).value or "").strip() for col in range(1, sheet.max_column + 1)]
            if "Line ID" in values and "Review Status" in values:
                header_row = row
                headers = {value: idx + 1 for idx, value in enumerate(values) if value}
                break
        if not header_row:
            workbook.close()
            return None, None, None
        return workbook, sheet, (header_row, headers)

    def _read_review_line_values(self, issue: dict) -> dict:
        target_key = _clean(issue.get("decision_key", ""))
        target_line = _clean(issue.get("line_id", ""))
        target_source = _clean(issue.get("source_id", ""))
        values = (
            self._review_fast_values.get(target_source)
            or self._review_fast_values.get(target_key)
            or self._review_fast_values.get(target_line)
        )
        if values is None:
            cache = self._load_review_cache()
            values = cache.get("values", {}).get(target_key or target_line)
            if values is None:
                # A decision key can change after a correction; fall back to its line ID.
                values = cache.get("values", {}).get(target_line, {})
        normalized = dict(values or {})
        color, size = split_embedded_size(
            normalized.get("Garment Color", ""), normalized.get("Size", "")
        )
        normalized["Garment Color"] = color
        normalized["Size"] = size
        return normalized

    def _scroll_review_to_top(self) -> None:
        frame = getattr(self, "review_decision_frame", None)
        if frame is None:
            return
        canvas = getattr(frame, "_parent_canvas", None)
        if canvas is None:
            return
        try:
            canvas.yview_moveto(0)
        except Exception:
            pass

    def _invalidate_fast_review_snapshot(self) -> None:
        self._review_fast_cache_key = None
        self._review_fast_snapshot_cache = {}
        self._review_fast_values = {}

    def _show_review_loading_state(self) -> None:
        if not hasattr(self, "review_state_title"):
            return
        self.review_state_title.configure(text="Loading Purchase Review…")
        self.review_state_body.configure(
            text="Opening the first decision now. The rest of the review queue is loading in the background."
        )
        self.review_progress.configure(text="Reading the saved review queue…")
        self.review_identity.configure(text="")
        self.review_reason.configure(text="")
        self.review_instructions.configure(text="")
        self.review_recommendation.grid_remove()
        self._set_review_instruction_callout("", False)
        self._set_decoration_decision_visibility(False)
        self._set_review_fields_state("disabled")
        self.review_include_menu.configure(state="disabled")
        self.review_product_master_button.pack_forget()
        self.review_previous_button.configure(state="disabled")
        self.review_search_button.configure(state="disabled")
        self.review_save_button.configure(text="Loading…", state="disabled")

    def _cache_fast_review_snapshot(self, snapshot: dict, cache_key=None) -> None:
        if cache_key is None and self.last_review_workbook and self.last_review_workbook.exists():
            cache_key = _file_stamp(self.last_review_workbook)
        self._review_fast_cache_key = cache_key
        self._review_fast_snapshot_cache = snapshot or {}
        self._review_fast_values = dict((snapshot or {}).get("review_values", {}))

    def _apply_loaded_review_snapshot(self, token: int, path: Path, cache_key, snapshot: dict | None, error: Exception | None = None) -> None:
        if token != self._review_load_token:
            return
        self._review_load_in_progress = False
        if path != self.last_review_workbook or not path.exists():
            return
        if error is not None:
            self.review_state_title.configure(text="Unable to Open Purchase Review")
            self.review_state_body.configure(text="The saved review workbook could not be read.")
            self.review_progress.configure(text=str(error))
            self.review_save_button.configure(text="Try Again", state="normal", command=lambda: self._load_review_page_async(force=True))
            return
        snapshot = self._overlay_completed_review_state(snapshot or {}, path)
        self._cache_fast_review_snapshot(snapshot, cache_key)
        # Keep lightweight data available to the review editor without replacing
        # a richer mission-control snapshot that may already be cached.
        self.review_search_button.configure(state="normal")
        self.review_save_button.configure(command=self.save_current_review_decision)
        self.refresh_review_page(snapshot=snapshot)
        self._warm_dashboard_snapshot_async(path)

    def _load_review_page_async(self, force: bool = False) -> None:
        if not hasattr(self, "review_state_title"):
            return
        path = self.last_review_workbook
        if not path or not path.exists():
            self.refresh_review_page(snapshot={})
            return
        if not self._is_archived_report_reopen() and not self._active_packet_is_valid(path)[0]:
            self.review_state_title.configure(text="Active Event Verification Required")
            self.review_state_body.configure(
                text="This Purchase Review is not proven to match the selected CSV, so Orchid will not display its decisions."
            )
            self.review_progress.configure(text="Return to Current Event and process the selected CSV again.")
            self.review_save_button.configure(text="Return to Current Event", state="normal", command=lambda: self.show_page("import"))
            return
        cache_key = _file_stamp(path)
        if not force and cache_key == self._review_fast_cache_key and self._review_fast_snapshot_cache:
            snapshot = self._overlay_completed_review_state(self._review_fast_snapshot_cache, path)
            self._cache_fast_review_snapshot(snapshot, cache_key)
            self.refresh_review_page(snapshot=snapshot)
            self._warm_dashboard_snapshot_async(path)
            return
        self._review_load_token += 1
        token = self._review_load_token
        self._review_load_in_progress = True
        self._show_review_loading_state()
        result_queue = queue.Queue()

        def worker():
            try:
                result_queue.put(("success", load_purchase_review_snapshot(path, live_product_master_path())))
            except Exception as exc:
                result_queue.put(("error", exc))

        def poll_worker():
            if token != self._review_load_token:
                return
            try:
                kind, payload = result_queue.get_nowait()
            except queue.Empty:
                if self.active_page == "review" and self.winfo_exists():
                    self.after(60, poll_worker)
                return
            if kind == "success":
                self._apply_loaded_review_snapshot(token, path, cache_key, payload, None)
            else:
                self._apply_loaded_review_snapshot(token, path, cache_key, None, payload)

        threading.Thread(target=worker, name="orchid-review-loader", daemon=True).start()
        self.after(20, poll_worker)

    def _warm_dashboard_snapshot_async(self, path: Path | None = None) -> None:
        candidate = path or self.last_review_workbook
        if not candidate:
            return
        path = Path(candidate)
        if not path.exists():
            return
        cache_key = (
            _file_stamp(path),
            _file_stamp(DATA / "mission_control_state.json"),
            _file_stamp(REVIEW_STATE_FILE),
            _file_stamp(live_product_master_path()),
        )
        if cache_key == self._mission_snapshot_cache_key:
            return
        self._dashboard_warm_token += 1
        token = self._dashboard_warm_token
        result_queue = queue.Queue()

        def worker():
            try:
                result_queue.put(load_mission_control_snapshot(path, DATA, live_product_master_path()))
            except Exception as exc:
                result_queue.put(exc)

        def poll_worker():
            if token != self._dashboard_warm_token:
                return
            try:
                result = result_queue.get_nowait()
            except queue.Empty:
                if self.winfo_exists():
                    self.after(120, poll_worker)
                return
            if isinstance(result, Exception):
                return
            if path != self.last_review_workbook:
                return
            result = self._overlay_completed_review_state(result, path)
            self._mission_snapshot_cache_key = cache_key
            self._mission_snapshot_cache = result
            self.last_snapshot = result

        threading.Thread(target=worker, name="orchid-dashboard-warmup", daemon=True).start()
        self.after(120, poll_worker)

    def _fresh_snapshot(self) -> dict:
        """Reload the workbook so every workflow screen uses the same live state."""
        snapshot = self._mission_snapshot()
        self.last_snapshot = snapshot
        return snapshot

    def _current_purchase_review_issues(self, snapshot: dict | None = None) -> list[dict]:
        snapshot = snapshot or self._review_fast_snapshot_cache or self._fresh_snapshot()
        return self._build_purchase_review_queue(snapshot)

    def _blocker_summary(self, snapshot: dict, limit: int = 5) -> str:
        lines: list[str] = []
        for issue in snapshot.get("issues", [])[:limit]:
            identity = " • ".join(filter(None, [
                _clean(issue.get("order", "")),
                _clean(issue.get("employee", "")),
                _clean(issue.get("product", "")) or _clean(issue.get("description", "")),
            ]))
            reason = _clean(issue.get("reason", "")) or "Needs review"
            lines.append(f"• {identity or 'Purchase line'} — {reason}")
        remaining = max(int(snapshot.get("review_count", 0)) - len(lines), 0)
        if remaining:
            lines.append(f"• Plus {remaining} additional decision(s)")
        if not lines and int(snapshot.get("blocked_route_count", 0)):
            for route in snapshot.get("routes", [])[:limit]:
                if _clean(route.get("status", "")).casefold() == "ready":
                    continue
                lines.append(f"• {_clean(route.get('vendor', 'Vendor route'))} — route still blocked")
        return "\n".join(lines)

    def _review_status_counts(self) -> tuple[int, int]:
        if self._review_fast_snapshot_cache:
            return (
                int(self._review_fast_snapshot_cache.get("review_completed", 0)),
                int(self._review_fast_snapshot_cache.get("review_remaining", 0)),
            )
        cache = self._load_review_cache()
        return int(cache.get("completed", 0)), int(cache.get("remaining", 0))

    def refresh_review_page(self, snapshot: dict | None = None):
        if not hasattr(self, "review_state_title"):
            return
        self._reset_review_field_highlights()
        if hasattr(self, "review_exit_button") and not self._review_exit_after_saves:
            self.review_exit_button.configure(state="normal")
        self.after_idle(self._scroll_review_to_top)
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            self.review_state_title.configure(text="Process Imported Orders")
            self.review_state_body.configure(text="Import orders first, then process the imported CSV. Orchid will check Product Master before Purchase Review.")
            self.review_progress.configure(text="No active review")
            self.review_identity.configure(text="The imported orders have not been processed yet.")
            self.review_reason.configure(text="")
            self.review_instructions.configure(text="")
            self.review_save_button.configure(text="Process Orders", command=self.create_review_workbook, state="normal")
            self.review_product_master_button.pack_forget()
            self.review_recommendation.grid_remove()
            self._set_decoration_decision_visibility(False)
            self._update_previous_button()
            return
        snapshot = snapshot or self._review_fast_snapshot_cache or self._fresh_snapshot()
        issues = self._current_purchase_review_issues(snapshot)
        product_issues = [
            issue for issue in snapshot.get("issues", [])
            if _clean(issue.get("fix_in", issue.get("source", ""))).casefold() == "product master"
        ]
        search_issue = getattr(self, "_review_search_issue", None)
        if product_issues and not search_issue:
            issue = product_issues[0]
            self._active_review_issue = issue
            self.review_state_title.configure(text="Product Setup Required")
            self.review_state_body.configure(text="Complete the permanent product setup before order-specific review can continue.")
            self.review_progress.configure(text=f"{len(product_issues)} product setup decision(s)")
            self.review_identity.configure(text=f"{issue.get('product','')}  {issue.get('description','')}")
            self.review_reason.configure(text=issue.get("reason", "Product Master detail needed"))
            self.review_instructions.configure(text=issue.get("instructions", "Open Product Master to complete this item."))
            self._set_review_fields_state("disabled")
            self.review_include_menu.configure(state="disabled")
            self.review_save_button.configure(state="disabled", text="Save & Next")
            self.review_product_master_button.pack(side="left", padx=4)
            self.review_recommendation.grid_remove()
            self._set_decoration_decision_visibility(False)
            self.review_previous_button.configure(state="disabled")
            return
        self.review_product_master_button.pack_forget()
        if not issues and self._review_history_index is None and not search_issue:
            self._active_review_issue = None
            if not product_issues:
                self._update_review_state(
                    review_complete=True,
                    review_completed_at=time.time(),
                    review_workbook=str(self.last_review_workbook or ""),
                )
            if self._review_save_pending:
                self.review_state_title.configure(text="Purchase Review Complete")
                self.review_state_body.configure(text="All decisions are safely recorded. Excel synchronization is continuing quietly in the background.")
                self.review_progress.configure(text=f"0 decisions remaining  •  {self._review_save_pending} background sync pending")
                self.review_identity.configure(text="✓ Continue to Decoration Colors")
                self.review_reason.configure(text="")
                self.review_instructions.configure(text="Orchid will reconcile the recovery journal automatically before final purchase orders are generated.")
                self._set_review_fields_state("disabled")
                self.review_include_menu.configure(state="disabled")
                self.review_save_button.configure(text="Go to Decoration Colors", command=lambda: self.show_page("audit"), state="normal")
                self.review_recommendation.grid_remove()
                self._set_decoration_decision_visibility(False)
                self._update_previous_button()
                return
            if int(snapshot.get("blocked_route_count", 0)):
                self.review_state_title.configure(text="Purchase Review Still Required")
                self.review_state_body.configure(text="A purchase-order route is still blocked. The same final readiness check used before the Decoration Color Audit is shown below.")
                self.review_progress.configure(text=f"{int(snapshot.get('blocked_route_count', 0))} blocked route(s)")
                self.review_identity.configure(text="Purchase-order preflight did not pass")
                self.review_reason.configure(text=self._blocker_summary(snapshot) or "A vendor route still needs correction.")
                self.review_instructions.configure(text="Return to the dashboard Work Queue to open the exact blocked route.")
                self._set_review_fields_state("disabled")
                self.review_include_menu.configure(state="disabled")
                self.review_save_button.configure(text="Return to Current Event", command=lambda: self.show_page("import"), state="normal")
                self.review_recommendation.grid_remove()
                self._set_decoration_decision_visibility(False)
                self._update_previous_button()
                return
            self.review_state_title.configure(text="Purchase Review Complete")
            self.review_state_body.configure(text="No order decisions or blocked purchase-order rows remain. Continue to the Decoration Color Audit before Purchase Orders.")
            self.review_progress.configure(text="0 decisions remaining")
            self.review_identity.configure(text="✓ Final purchase-order preflight passed")
            self.review_reason.configure(text="")
            self.review_instructions.configure(text="The dashboard will advance automatically to Step 4: Decoration Colors.")
            self._set_review_fields_state("disabled")
            self.review_include_menu.configure(state="disabled")
            self.review_save_button.configure(text="Go to Decoration Colors", command=lambda: self.show_page("audit"), state="normal")
            self.review_recommendation.grid_remove()
            self._set_decoration_decision_visibility(False)
            self._update_previous_button()
            return
        editing_previous = self._review_history_index is not None
        editing_search = bool(search_issue)
        draft_issue = self._saved_review_draft_issue(issues) if not editing_previous and not editing_search else None
        issue = search_issue if editing_search else (
            self._review_decision_history[self._review_history_index] if editing_previous else (draft_issue or issues[0])
        )
        self._active_review_issue = issue
        if editing_search:
            title_text = "Edit Search Result"
            body_text = "Correct this exact order line, then save to return to the normal review queue."
        elif editing_previous:
            title_text = "Edit Previous Decision"
            body_text = "Correct the saved choice, then return to your current place in the review queue."
        else:
            title_text = "Complete Purchase Review"
            body_text = "Correct the highlighted information, then save to continue to the next decision."
        self.review_state_title.configure(text=title_text)
        self.review_state_body.configure(text=body_text)
        completed_count, remaining_count = self._review_status_counts()
        self.review_progress.configure(
            text="Search result selected" if editing_search else (
                f"Previous decision {self._review_history_index + 1} of {len(self._review_decision_history)}"
                if editing_previous else f"{completed_count} completed  •  {remaining_count} remaining"
            )
        )
        identity = "  •  ".join(filter(None, [issue.get("order", ""), issue.get("employee", ""), issue.get("product", ""), issue.get("description", "")]))
        self.review_identity.configure(text=identity or "Current order decision")
        reason_text = issue.get("reason", "Needs review")
        self.review_reason.configure(text=reason_text)
        values = self._read_review_line_values(issue)
        if draft_issue is not None and not editing_previous and not editing_search:
            draft_state = self._load_review_state()
            draft_values = draft_state.get("draft_values", {})
            if isinstance(draft_values, dict):
                values.update({field: _clean(value) for field, value in draft_values.items()})
        instruction_text = issue.get("instructions", "Complete the missing order information.")
        note_text = values.get("Purchase Instructions") or values.get("Shopify Order Notes")
        is_customer_decision = "customer decision required" in _clean(reason_text).casefold()
        note_key = _clean(note_text).casefold()
        is_outsource_instruction = any(term in note_key for term in OUTSOURCE_INSTRUCTION_TERMS)
        is_waterproof_verification = (
            "waterproof/rainwear item:" in note_key
            or "waterproof style " in note_key
        )
        is_decoration_instruction = bool(
            is_customer_decision
            and (decoration_note_requires_review(note_text) or is_waterproof_verification)
        )
        is_instruction_decision = bool(is_customer_decision)
        self._review_instruction_kind = (
            "decoration" if is_decoration_instruction else
            "outsource" if is_outsource_instruction else
            "general"
        )
        if is_instruction_decision:
            self._set_review_instruction_callout(note_text, True)
            instruction_text = (
                "Choose an explicit instruction decision below, update any affected order fields, then save."
            )
            if is_waterproof_verification:
                recommendation = "Keep Product Master Default"
                recommendation_reason = (
                    "Confirm the saved decoration, or explicitly choose No Decoration, Embroidery, or Screen Print."
                )
            elif is_decoration_instruction:
                recommendation, recommendation_reason = decoration_instruction_recommendation(
                    note_text, values.get("Decoration Type", "")
                )
            elif is_outsource_instruction and self._review_outsourcing_exception_applicable(values.get("Decoration Type", "")):
                recommendation = DO_NOT_OUTSOURCE_DECISION
                recommendation_reason = "The note directs this one item to Orchid instead of the outside decorator."
            else:
                recommendation = "Follow Note as Written"
                recommendation_reason = "Update the product, size, color, vendor, or other order-specific fields to match the instruction."
            if recommendation:
                self.review_recommendation.configure(
                    text=f"Recommended: {recommendation} — {recommendation_reason}", padx=14, pady=10
                )
                self.review_recommendation.grid()
            else:
                self.review_recommendation.grid_remove()
        else:
            self._review_instruction_kind = ""
            self.review_recommendation.grid_remove()
        self.review_instructions.configure(text=instruction_text)
        self._review_default_decoration_type = values.get("Decoration Type", "")
        self._review_default_decoration_location = normalize_decoration_location(
            values.get("Decoration Location", ""), values.get("Decoration Type", ""),
            requires_decoration=not is_blank_decoration(values.get("Decoration Type", "")) and not is_in_house_decoration(values.get("Decoration Type", "")),
            product_name=issue.get("description", ""),
        )
        self._review_default_placement_instructions = values.get("Decoration Placement Instructions", "")
        self._review_default_decoration_color = values.get("Decoration Color", "")
        actual_location = normalize_decoration_location(
            values.get("Decoration Location", ""), values.get("Decoration Type", ""),
            requires_decoration=not is_blank_decoration(values.get("Decoration Type", "")) and not is_in_house_decoration(values.get("Decoration Type", "")),
            product_name=issue.get("description", ""),
        )
        location_choice, custom_location = location_choice_and_custom(actual_location, values.get("Decoration Type", ""))
        values["Decoration Location"] = location_choice
        for field, variable in self.review_vars.items():
            default = "Yes" if field == "Include" else ("No" if field == "Do Not Outsource" else "")
            if field == "Decoration Decision" and is_instruction_decision:
                default = "Choose an instruction decision"
            variable.set(values.get(field, "") or default)
        self._review_default_do_not_outsource = values.get("Do Not Outsource", "No") or "No"
        if (
            is_instruction_decision
            and self._review_instruction_kind == "outsource"
            and _clean(values.get("Do Not Outsource", "No")).casefold() in {"yes", "y", "true", "1"}
        ):
            self.review_vars["Decoration Decision"].set(DO_NOT_OUTSOURCE_DECISION)
        self.review_custom_location_var.set(custom_location)
        self._apply_review_location_choice(location_choice)
        self._set_decoration_decision_visibility(is_instruction_decision)
        self._set_review_fields_state("normal")
        self._highlight_review_required_fields(reason_text, values)
        self.review_include_menu.configure(state="normal")
        self.review_save_button.configure(
            text="Save Changes & Return" if (editing_previous or editing_search) else "Save & Next",
            command=self.save_current_review_decision, state="normal"
        )
        self._update_previous_button()

    def _pending_review_jobs_from_state(self) -> list[dict]:
        state = self._load_review_state()
        jobs = state.get("pending_review_saves", [])
        return [dict(job) for job in jobs if isinstance(job, dict)] if isinstance(jobs, list) else []

    def _write_pending_review_jobs(self, jobs: list[dict]) -> None:
        state = self._load_review_state()
        if jobs:
            state["pending_review_saves"] = jobs
        else:
            state.pop("pending_review_saves", None)
        try:
            REVIEW_STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
            REVIEW_STATE_FILE.write_text(json.dumps(state, indent=2), encoding="utf-8")
        except Exception:
            pass

    def _journal_review_save_job(self, job: dict) -> None:
        """Persist both the pending Excel job and its completed-review marker."""
        state = self._load_review_state()
        jobs = state.get("pending_review_saves", [])
        jobs = [dict(item) for item in jobs if isinstance(item, dict)] if isinstance(jobs, list) else []
        job_id = _clean(job.get("job_id", ""))
        jobs = [item for item in jobs if _clean(item.get("job_id", "")) != job_id]
        jobs.append(job)
        state["pending_review_saves"] = jobs

        records = state.get("completed_review_decisions", [])
        records = [dict(item) for item in records if isinstance(item, dict)] if isinstance(records, list) else []
        target_key = _clean(job.get("decision_key", ""))
        target_line = _clean(job.get("line_id", ""))
        target_source = _clean(job.get("source_id", ""))
        target_workbook = _clean(job.get("workbook", ""))
        def same_record(item: dict) -> bool:
            if _clean(item.get("workbook", "")) != target_workbook:
                return False
            item_source = _clean(item.get("source_id", ""))
            if target_source:
                return item_source == target_source
            return bool(
                (target_key and _clean(item.get("decision_key", "")) == target_key)
                or (target_line and _clean(item.get("line_id", "")) == target_line)
            )
        records = [item for item in records if not same_record(item)]
        records.append({
            "workbook": target_workbook,
            "source_signature": _clean(state.get("source_signature", "")),
            "event_name": _clean(state.get("event_name", self.current_event_name)),
            "line_id": target_line,
            "source_id": target_source,
            "decision_key": target_key,
            "issue_signature": _clean(job.get("issue_signature", "")),
            "issue": dict(job.get("issue", {})) if isinstance(job.get("issue"), dict) else {},
            "source_only": bool(job.get("source_only", False)),
            "values": dict(job.get("values", {})),
            "completed_at": float(job.get("created_at", time.time())),
            "job_id": job_id,
        })
        # Limit the recovery registry while retaining more than enough history
        # for very large events.
        state["completed_review_decisions"] = records[-5000:]
        try:
            REVIEW_STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
            REVIEW_STATE_FILE.write_text(json.dumps(state, indent=2), encoding="utf-8")
        except Exception:
            pass

    def _remove_journaled_review_save_job(self, job_id: str) -> None:
        jobs = [
            item for item in self._pending_review_jobs_from_state()
            if _clean(item.get("job_id", "")) != _clean(job_id)
        ]
        self._write_pending_review_jobs(jobs)

    def _resume_pending_review_saves(self) -> None:
        """Replay tiny recovery-journal entries left by an interrupted session."""
        for job in self._pending_review_jobs_from_state():
            workbook_value = _clean(job.get("workbook", ""))
            if not workbook_value:
                self._remove_journaled_review_save_job(_clean(job.get("job_id", "")))
                continue
            workbook_path = Path(workbook_value)
            if workbook_path.is_file():
                self._queue_review_save_job(job, record_journal=False)
            else:
                self._remove_journaled_review_save_job(_clean(job.get("job_id", "")))

    def _ensure_review_save_worker(self) -> None:
        if self._review_save_worker_started:
            return
        self._review_save_worker_started = True
        threading.Thread(
            target=self._review_save_worker_loop,
            name="orchid-review-save-worker",
            daemon=True,
        ).start()

    def _queue_review_save_job(self, job: dict, *, record_journal: bool = True) -> None:
        if record_journal:
            self._journal_review_save_job(job)
        self._review_save_pending += 1
        self._ensure_review_save_worker()
        self._review_save_jobs.put(job)
        if not self._review_save_polling:
            self._review_save_polling = True
            self.after(50, self._poll_review_save_results)

    def _review_save_worker_loop(self) -> None:
        """Serialize Excel writes, batching rapid decisions into one workbook save."""
        while True:
            first_job = self._review_save_jobs.get()
            if first_job is None:
                self._review_save_jobs.task_done()
                return
            batch = [first_job]
            # Give the Tk thread time to draw the next decision. The journal is
            # already durable, so delaying the expensive XLSX rewrite is safe.
            time.sleep(0.70)
            while len(batch) < 100:
                try:
                    next_job = self._review_save_jobs.get_nowait()
                except queue.Empty:
                    break
                if next_job is None:
                    self._review_save_jobs.task_done()
                    continue
                batch.append(next_job)
            try:
                self._apply_review_save_jobs(batch)
                results = [
                    {"ok": True, "job_id": job.get("job_id", "")}
                    for job in batch
                ]
            except Exception as error:
                results = [
                    {
                        "ok": False,
                        "job_id": job.get("job_id", ""),
                        "error": str(error) or error.__class__.__name__,
                    }
                    for job in batch
                ]
            for result in results:
                self._review_save_results.put(result)
            for _job in batch:
                self._review_save_jobs.task_done()

    @staticmethod
    def _apply_review_save_job(job: dict) -> None:
        OrchidPurchaseManager._apply_review_save_jobs([job])

    @staticmethod
    def _apply_review_save_jobs(
        jobs: list[dict], *, skip_orphaned_recovery_jobs: bool = False,
    ) -> list[dict]:
        """Apply many completed decisions with one open/save cycle per workbook.

        A completed Purchase Review decision can legitimately become orphaned
        after the corresponding line is removed from a regenerated packet.  An
        orphan must not blank the whole Decoration Color Audit or block final
        report generation.  The optional recovery mode excludes *only* jobs
        whose source, line, decision key, and unique semantic identity are all
        absent from both event sheets.  Anything even possibly current remains
        strict and raises instead of being guessed or silently dropped.
        """
        grouped: dict[str, list[dict]] = {}
        skipped_jobs: list[dict] = []
        for job in jobs:
            workbook_value = _clean(job.get("workbook", ""))
            grouped.setdefault(workbook_value, []).append(job)

        for workbook_value, workbook_jobs in grouped.items():
            if not workbook_value:
                raise ValueError("A saved Purchase Review recovery entry is missing its workbook path.")
            workbook_path = Path(workbook_value)
            if not workbook_path.is_file():
                raise FileNotFoundError(f"Purchase Review workbook was not found: {workbook_path}")
            workbook = load_workbook(workbook_path)
            temp_path = workbook_path.with_name(f".{workbook_path.stem}.orchid-saving{workbook_path.suffix}")

            def header_context(target_sheet):
                if target_sheet is None:
                    return None, {}
                for row_number in range(1, min(target_sheet.max_row, 20) + 1):
                    row_values = [
                        str(target_sheet.cell(row_number, col).value or "").strip()
                        for col in range(1, target_sheet.max_column + 1)
                    ]
                    if "Line ID" in row_values:
                        return row_number, {
                            value: idx + 1 for idx, value in enumerate(row_values) if value
                        }
                return None, {}

            def ensure_header(target_sheet, header_row, headers, field, default_value):
                if target_sheet is None or not header_row or field in headers:
                    return
                column = target_sheet.max_column + 1
                target_sheet.cell(header_row, column).value = field
                for row_number in range(header_row + 1, target_sheet.max_row + 1):
                    target_sheet.cell(row_number, column).value = default_value
                headers[field] = column

            def semantic_keys(values, issue=None):
                values = values or {}
                issue = issue or {}
                quantity = _clean(values.get("Quantity", ""))
                try:
                    number = float(quantity)
                    quantity = str(int(number)) if number.is_integer() else str(number)
                except Exception:
                    pass
                base = "|".join((
                    re.sub(r"\s+", "", _clean(values.get("Product #", ""))).casefold(),
                    _clean(values.get("Garment Color", "")).casefold(),
                    _clean(values.get("Size", "")).casefold(),
                    quantity.casefold(),
                    _clean(values.get("Purchase Vendor", "")).casefold(),
                ))
                order = _clean(issue.get("order", values.get("Order Number", ""))).lstrip("#").casefold()
                employee = _clean(issue.get("employee", values.get("Employee Name", ""))).casefold()
                keys = []
                if order or employee:
                    keys.append(f"{order}|{employee}|{base}")
                keys.append(base)
                return [key for key in dict.fromkeys(keys) if key.strip("|")]

            def audit_identity(values):
                """Return the stable identity used by Decoration Color Audit.

                Audit choices intentionally apply to every line for one
                decoration type, style, and garment color.  Unlike Line ID and
                Purchase Review's per-order semantic key, those three fields
                remain stable when a review workbook is regenerated.
                """
                values = values or {}
                style = re.sub(r"\s+", "", _clean(values.get("Product #", ""))).casefold()
                garment_color = _clean(values.get("Garment Color", "")).casefold()
                decoration_type = _clean(values.get("Decoration Type", "")).casefold()
                if not style or not garment_color or not decoration_type:
                    return ""
                return "|".join((style, garment_color, decoration_type))

            def build_indexes(target_sheet, header_row, headers):
                by_source: dict[str, list[int]] = {}
                by_line: dict[str, list[int]] = {}
                by_key: dict[str, list[int]] = {}
                by_semantic: dict[str, list[int]] = {}
                by_audit_identity: dict[str, list[int]] = {}
                if target_sheet is None or not header_row:
                    return by_source, by_line, by_key, by_semantic, by_audit_identity
                source_col = headers.get("Source ID")
                line_col = headers.get("Line ID")
                key_col = headers.get("Decision Key")
                for row_number in range(header_row + 1, target_sheet.max_row + 1):
                    source_id = _clean(target_sheet.cell(row_number, source_col).value) if source_col else ""
                    line_id = _clean(target_sheet.cell(row_number, line_col).value) if line_col else ""
                    decision_key = _clean(target_sheet.cell(row_number, key_col).value) if key_col else ""
                    if source_id:
                        by_source.setdefault(source_id, []).append(row_number)
                    if line_id:
                        by_line.setdefault(line_id, []).append(row_number)
                    if decision_key:
                        by_key.setdefault(decision_key, []).append(row_number)
                    row_values = {
                        field: target_sheet.cell(row_number, column).value
                        for field, column in headers.items()
                        if field in {
                            "Source ID", "Product #", "Garment Color", "Size", "Quantity", "Purchase Vendor",
                            "Order Number", "Employee Name", "Decoration Type",
                        }
                    }
                    for semantic in semantic_keys(row_values):
                        by_semantic.setdefault(semantic, []).append(row_number)
                    stable_identity = audit_identity(row_values)
                    if stable_identity:
                        by_audit_identity.setdefault(stable_identity, []).append(row_number)
                return by_source, by_line, by_key, by_semantic, by_audit_identity

            def matching_rows(job, by_source, by_line, by_key, by_semantic, by_audit_identity):
                target_source = _clean(job.get("source_id", ""))
                if target_source:
                    # Source ID is authoritative. Never fall back to a generated
                    # Line ID, Decision Key, or semantic product match when it is present.
                    source_rows = by_source.get(target_source, [])
                    return sorted(source_rows) if len(source_rows) == 1 else []

                target_line = _clean(job.get("line_id", ""))
                target_key = _clean(job.get("decision_key", ""))
                # Audit confirmations are line-specific. Purchase Review rows
                # can share a decision key (for example, multiple TLS608
                # garment colors share product:TLS608|), so combining a saved
                # audit line ID with that broad key overwrites every color.
                # Prefer the immutable line ID. If regeneration changes it,
                # fall back only to exact style, garment color, and decoration
                # type—never to the Purchase Review decision key.
                if bool(job.get("audit_override", False)):
                    if target_line:
                        direct_rows = by_line.get(target_line, [])
                        if direct_rows:
                            return sorted(direct_rows)
                    stable_identity = audit_identity(job.get("values", {}))
                    if stable_identity:
                        return sorted(by_audit_identity.get(stable_identity, []))
                    return []
                rows = set()
                if target_line:
                    rows.update(by_line.get(target_line, []))
                if target_key:
                    rows.update(by_key.get(target_key, []))
                if not rows:
                    issue = job.get("issue", {}) if isinstance(job.get("issue"), dict) else {}
                    for semantic in semantic_keys(job.get("values", {}), issue):
                        candidates = by_semantic.get(semantic, [])
                        # Legacy semantic fallback must be unique; ambiguous
                        # matches are rejected rather than guessed.
                        if len(candidates) == 1:
                            rows.update(candidates)
                            break
                return sorted(rows)

            def possible_current_rows(job, by_line, by_key, by_semantic):
                """Return any non-source evidence that a recovery job is current.

                This is deliberately more conservative than a normal match:
                multiple candidates are evidence that the job may still matter,
                not permission to apply it.  They therefore prevent an orphan
                skip and preserve the normal safety error.

                A Source ID is different.  It is the immutable identity of one
                imported event line and is deliberately authoritative in
                ``matching_rows``.  If that exact ID is absent, an older Line
                ID, Decision Key, or semantic product match must *not* keep the
                recovery job alive: those values can legitimately belong to a
                different line after an event is regenerated.  Treating them as
                proof here created a contradiction where a job was declared
                "possibly current" but was forbidden from being applied, which
                blanked Decoration Color Audit with a phantom "1 to Review".
                """
                rows = set()
                if _clean(job.get("source_id", "")):
                    return rows
                target_line = _clean(job.get("line_id", ""))
                target_key = _clean(job.get("decision_key", ""))
                if target_line:
                    rows.update(by_line.get(target_line, []))
                if target_key:
                    rows.update(by_key.get(target_key, []))
                issue = job.get("issue", {}) if isinstance(job.get("issue"), dict) else {}
                for semantic in semantic_keys(job.get("values", {}), issue):
                    rows.update(by_semantic.get(semantic, []))
                return rows

            def is_orphaned_recovery_job(job, visible_rows, source_rows):
                if (
                    not skip_orphaned_recovery_jobs
                    or bool(job.get("audit_override", False))
                    or visible_rows
                    or source_rows
                ):
                    return False
                visible_candidates = possible_current_rows(
                    job, review_by_line, review_by_key, review_by_semantic
                )
                source_candidates = possible_current_rows(
                    job, source_by_line, source_by_key, source_by_semantic
                )
                return not visible_candidates and not source_candidates

            def record_orphaned_job(job):
                skipped_jobs.append({
                    "job_id": _clean(job.get("job_id", "")),
                    "source_id": _clean(job.get("source_id", "")),
                    "line_id": _clean(job.get("line_id", "")),
                    "decision_key": _clean(job.get("decision_key", "")),
                })

            def write_rows(target_sheet, row_numbers, headers, job, *, source=False):
                values = dict(job.get("values", {}))
                is_instruction_decision = bool(job.get("is_instruction_decision", False))
                for row_number in row_numbers:
                    for field, value in values.items():
                        column = headers.get(field)
                        if column:
                            target_sheet.cell(row_number, column).value = value
                    if headers.get("Review Status"):
                        target_sheet.cell(row_number, headers["Review Status"]).value = "Ready"
                    if headers.get("Resolution"):
                        resolution = "Completed in App"
                        if is_instruction_decision:
                            resolution = f"Instruction: {values.get('Decoration Decision', '')}"
                        target_sheet.cell(row_number, headers["Resolution"]).value = resolution
                    if source:
                        permanent_reason = (
                            _clean(target_sheet.cell(row_number, headers.get("Permanent Review Reason")).value)
                            if headers.get("Permanent Review Reason") else ""
                        )
                        if headers.get("Event Review Reason"):
                            target_sheet.cell(row_number, headers["Event Review Reason"]).value = ""
                        if headers.get("Review Reason"):
                            target_sheet.cell(row_number, headers["Review Reason"]).value = permanent_reason
                        if headers.get("Fix In"):
                            target_sheet.cell(row_number, headers["Fix In"]).value = (
                                "Product Master" if permanent_reason else "Purchase Review"
                            )
                        if headers.get("Action Required") and not permanent_reason:
                            target_sheet.cell(row_number, headers["Action Required"]).value = ""

            try:
                if "Review & Edit" not in workbook.sheetnames:
                    raise RuntimeError("Review & Edit sheet is missing from the Purchase Review workbook.")
                review_sheet = workbook["Review & Edit"]
                review_header_row, review_headers = header_context(review_sheet)
                if not review_header_row:
                    raise RuntimeError("Review & Edit headers could not be located.")
                ensure_header(review_sheet, review_header_row, review_headers, "Do Not Outsource", "No")
                ensure_header(review_sheet, review_header_row, review_headers, "Decoration Decision", "")
                review_by_source, review_by_line, review_by_key, review_by_semantic, review_by_audit_identity = build_indexes(
                    review_sheet, review_header_row, review_headers
                )

                source_sheet = workbook["All PO Lines"] if "All PO Lines" in workbook.sheetnames else None
                source_header_row, source_headers = header_context(source_sheet)
                if source_sheet is not None and source_header_row:
                    ensure_header(source_sheet, source_header_row, source_headers, "Do Not Outsource", "No")
                    ensure_header(source_sheet, source_header_row, source_headers, "Decoration Decision", "")
                source_by_source, source_by_line, source_by_key, source_by_semantic, source_by_audit_identity = build_indexes(
                    source_sheet, source_header_row, source_headers
                )

                for job in workbook_jobs:
                    visible_rows = matching_rows(
                        job, review_by_source, review_by_line, review_by_key, review_by_semantic, review_by_audit_identity
                    )
                    allow_missing = bool(job.get("allow_missing", False))
                    source_rows = matching_rows(
                        job, source_by_source, source_by_line, source_by_key, source_by_semantic, source_by_audit_identity
                    )
                    orphaned = is_orphaned_recovery_job(job, visible_rows, source_rows)
                    if orphaned:
                        record_orphaned_job(job)
                        continue
                    if not visible_rows and not bool(job.get("editing_search", False)) and not allow_missing:
                        raise RuntimeError("The matching Purchase Review decision could not be located.")
                    # Direct audit saves may never be recorded as completed if
                    # they could not update the event source.  Previously an
                    # unmatched editing-search job could become a silent no-op,
                    # leaving a stale color underneath the value shown on the
                    # audit screen and trapping the user in this step.
                    if bool(job.get("require_event_match", False)):
                        if source_sheet is not None and not source_rows:
                            raise RuntimeError(
                                "The current-event row for this decoration audit could not be located. "
                                "Nothing was marked complete."
                            )
                        if source_sheet is None and not visible_rows:
                            raise RuntimeError(
                                "The current-event row for this decoration audit could not be located. "
                                "Nothing was marked complete."
                            )
                    if source_sheet is not None and not source_rows and not allow_missing:
                        raise RuntimeError("The matching All PO Lines source record could not be located.")
                    write_rows(review_sheet, visible_rows, review_headers, job)
                    if source_sheet is not None:
                        write_rows(source_sheet, source_rows, source_headers, job, source=True)

                workbook.save(temp_path)
                workbook.close()
                os.replace(temp_path, workbook_path)
            finally:
                try:
                    workbook.close()
                except Exception:
                    pass
                try:
                    if temp_path.exists():
                        temp_path.unlink()
                except Exception:
                    pass
        return skipped_jobs

    def _poll_review_save_results(self) -> None:
        had_result = False
        had_error = False
        while True:
            try:
                result = self._review_save_results.get_nowait()
            except queue.Empty:
                break
            had_result = True
            self._review_save_pending = max(0, self._review_save_pending - 1)
            job_id = _clean(result.get("job_id", ""))
            if result.get("ok"):
                self._remove_journaled_review_save_job(job_id)
            else:
                had_error = True
                self._review_save_last_error = _clean(result.get("error", "Unable to save Purchase Review."))

        if self._review_save_pending:
            self.after(70, self._poll_review_save_results)
            return

        self._review_save_polling = False
        if had_error:
            self._review_exit_after_saves = False
            if hasattr(self, "review_exit_button"):
                self.review_exit_button.configure(state="normal")
            messagebox.showerror(
                "Unable to Finish Purchase Review Save",
                self._review_save_last_error
                + "\n\nThe decision will be restored from Orchid's recovery journal.",
            )
            self._invalidate_fast_review_snapshot()
            if self.active_page == "review":
                self._load_review_page_async(force=True)
            return

        if had_result:
            self._invalidate_review_cache()
            if self._review_exit_after_saves:
                self._review_exit_after_saves = False
                self._invalidate_fast_review_snapshot()
                self.refresh_dashboard()
                self.show_page("import")
            elif self.active_page == "review":
                # The in-memory queue already contains the saved values and the
                # next unresolved item. Avoid rereading Excel and rebuilding the
                # page after every background save.
                if self.last_review_workbook and self.last_review_workbook.exists():
                    self._review_fast_cache_key = _file_stamp(self.last_review_workbook)
                if self._review_fast_snapshot_cache:
                    self._review_fast_snapshot_cache["pending_sync"] = False
                self._warm_dashboard_snapshot_async(self.last_review_workbook)
            else:
                self._invalidate_fast_review_snapshot()
                self._warm_dashboard_snapshot_async(self.last_review_workbook)

    def _optimistic_review_snapshot_after_save(self, issue: dict, values: dict) -> dict:
        """Advance the visible queue immediately while Excel sync runs serially."""
        snapshot = dict(self._review_fast_snapshot_cache or {})
        issues = list(snapshot.get("issues", []))
        target_line = _clean(issue.get("line_id", ""))
        target_key = _clean(issue.get("decision_key", ""))
        target_source = _clean(issue.get("source_id", values.get("Source ID", "")))

        def matches(item: dict) -> bool:
            if _clean(item.get("fix_in", item.get("source", ""))).casefold() == "product master":
                return False
            source_id = _clean(item.get("source_id", ""))
            line_id = _clean(item.get("line_id", ""))
            decision_key = _clean(item.get("decision_key", ""))
            if target_source:
                return source_id == target_source
            return bool(
                (target_line and line_id == target_line)
                or (target_key and decision_key == target_key)
            )

        remaining_issues = [item for item in issues if not matches(item)]
        removed = len(remaining_issues) < len(issues)
        snapshot["issues"] = remaining_issues
        purchase_issues = [
            item for item in remaining_issues
            if _clean(item.get("fix_in", item.get("source", ""))).casefold() != "product master"
        ]
        snapshot["review_count"] = len(purchase_issues)
        snapshot["purchase_review_count"] = len(purchase_issues)
        snapshot["product_master_review_count"] = len(remaining_issues) - len(purchase_issues)
        snapshot["workflow_blocked"] = bool(remaining_issues or snapshot.get("blocked_route_count", 0))
        if removed:
            snapshot["review_completed"] = int(snapshot.get("review_completed", 0)) + 1
            snapshot["review_remaining"] = max(int(snapshot.get("review_remaining", 1)) - 1, 0)

        review_values = dict(snapshot.get("review_values", {}))
        updated_values = dict(self._read_review_line_values(issue))
        updated_values.update(values)
        updated_values["Review Status"] = "Ready"
        if target_key:
            review_values[target_key] = updated_values
        if target_line:
            review_values[target_line] = updated_values
        if target_source:
            review_values[target_source] = updated_values
        snapshot["review_values"] = review_values
        snapshot["pending_sync"] = True
        snapshot["updated"] = time.time()
        self._review_fast_snapshot_cache = snapshot
        self._review_fast_values = review_values
        self._review_fast_cache_key = ("pending", self._review_save_sequence, self._review_save_pending)
        return snapshot

    def save_current_review_decision(self):
        issue = self._active_review_issue
        if not issue or _clean(issue.get("fix_in", issue.get("source", ""))).casefold() == "product master":
            return

        editing_previous = self._review_history_index is not None
        editing_search = getattr(self, "_review_search_issue", None) is not None
        values = {field: _clean(variable.get()) for field, variable in self.review_vars.items()}
        current_values = self._read_review_line_values(issue)
        note_text = current_values.get("Purchase Instructions") or current_values.get("Shopify Order Notes")
        note_key = _clean(note_text).casefold()
        is_outsource_instruction = any(term in note_key for term in OUTSOURCE_INSTRUCTION_TERMS)
        is_waterproof_verification = (
            "waterproof/rainwear item:" in note_key
            or "waterproof style " in note_key
        )
        is_decoration_instruction = decoration_note_requires_review(note_text) or is_waterproof_verification
        is_instruction_decision = bool(
            "customer decision required" in _clean(issue.get("reason", "")).casefold()
        )
        self._review_instruction_kind = (
            "decoration" if is_decoration_instruction else
            "outsource" if is_outsource_instruction else
            "general"
        )
        choice = values.get("Decoration Decision", "")
        if choice == DO_NOT_OUTSOURCE_DECISION:
            values["Include"] = "Yes"
            values["Do Not Outsource"] = "Yes"
            self.review_vars["Include"].set("Yes")
        elif is_instruction_decision:
            values["Do Not Outsource"] = current_values.get("Do Not Outsource", "No") or "No"
        else:
            values["Do Not Outsource"] = current_values.get("Do Not Outsource", "No") or "No"

        include_item = values.get("Include", "Yes").casefold() in {"yes", "y", "true", "1", "include"}
        do_not_outsource = values.get("Do Not Outsource", "No").casefold() in {"yes", "y", "true", "1"}
        if not include_item:
            values["Do Not Outsource"] = "No"
            do_not_outsource = False
        elif do_not_outsource and not self._review_outsourcing_exception_applicable(values.get("Decoration Type", "")):
            messagebox.showwarning(
                "Do Not Outsource Is Not Applicable",
                "This exception is only available when the current decoration route would normally be outsourced.",
            )
            self.review_vars["Do Not Outsource"].set("No")
            self.review_vars["Decoration Decision"].set("Choose an instruction decision")
            return

        if is_instruction_decision:
            choice = values.get("Decoration Decision", "")
            valid_choices = set(self._review_instruction_decision_options(values.get("Decoration Type", ""))[1:])
            if choice not in valid_choices:
                messagebox.showwarning(
                    "Choose an Instruction Decision",
                    "This customer instruction must be acknowledged. Choose one of the available options before saving.",
                )
                self.review_decoration_decision_menu.focus_set()
                return
            if choice == "No Decoration":
                values["Decoration Type"] = BLANK_DECORATION_LABEL
                values["Decoration Color"] = ""
            elif choice in {"Embroidery", "Screen Print"}:
                values["Decoration Type"] = choice
                if include_item and not values.get("Decoration Color", ""):
                    messagebox.showwarning(
                        "Decoration Color Required",
                        f"Enter the thread/ink color for {choice}, then click Save & Next.",
                    )
                    self.review_entries["Decoration Color"].focus_set()
                    return
            elif choice == "Sew On Patch":
                values["Decoration Type"] = SEW_ON_PATCH_LABEL
                values["Decoration Color"] = ""
            elif choice == "Hemming / Alteration":
                values["Decoration Type"] = HEMMING_ALTERATION_LABEL
                values["Decoration Color"] = ""

        if include_item and not is_blank_decoration(values.get("Decoration Type", "")) and not is_in_house_decoration(values.get("Decoration Type", "")):
            if values.get("Decoration Location") == OTHER_CUSTOM and not _clean(self.review_custom_location_var.get()):
                messagebox.showwarning(
                    "Custom Location Required",
                    "Enter the exact custom decoration location before saving this decision.",
                )
                self.review_custom_location_entry.focus_set()
                return
            values["Decoration Location"] = self._resolved_review_location(values.get("Decoration Type", ""))
        else:
            values["Decoration Location"] = ""
            values["Decoration Placement Instructions"] = ""

        try:
            raw_quantity = values.get("Quantity", "")
            numeric_value = float(raw_quantity or 0)
            if include_item and (numeric_value <= 0 or not numeric_value.is_integer()):
                raise ValueError
            values["Quantity"] = int(numeric_value) if numeric_value.is_integer() else raw_quantity
        except (TypeError, ValueError):
            if include_item:
                messagebox.showwarning("Check Quantity", "Quantity must be a whole number greater than zero.")
                quantity_entry = self.review_entries.get("Quantity")
                if quantity_entry is not None:
                    quantity_entry.focus_set()
                return
            values["Quantity"] = 0

        # Validate entirely from the already loaded review record. This avoids
        # opening and scanning the full workbook on the UI thread.
        proposed = dict(current_values)
        proposed.update(values)
        proposed["Review Status"] = "Ready"
        blockers = line_block_reasons(proposed)
        if include_item and blockers:
            field_by_reason = {
                "missing product number and description": "Product #",
                "missing purchase vendor": "Purchase Vendor",
                "missing decoration type": "Decoration Type",
                "missing decoration color": "Decoration Color",
                "missing garment color": "Garment Color",
                "missing size": "Size",
                "quantity must be greater than zero": "Quantity",
                "custom decoration location is incomplete": "Decoration Location",
            }
            first_field = next(
                (field_by_reason.get(reason.casefold()) for reason in blockers if field_by_reason.get(reason.casefold())),
                None,
            )
            messagebox.showwarning(
                "Complete Required Information",
                "This line cannot be marked Ready yet:\n\n" + "\n".join(f"• {reason}" for reason in blockers),
            )
            if first_field and self.review_entries.get(first_field) is not None:
                self.review_entries[first_field].focus_set()
            return

        original_quantity = _clean(current_values.get("Original Quantity", ""))
        try:
            original_quantity_number = int(float(original_quantity))
        except (TypeError, ValueError):
            original_quantity_number = int(values.get("Quantity", 0) or 0)
        values["Quantity Override Confirmed"] = (
            "Yes" if int(values.get("Quantity", 0) or 0) != original_quantity_number else "No"
        )

        self._review_save_sequence += 1
        source_only = bool(issue.get("source_only", False))
        source_id = _clean(current_values.get("Source ID", issue.get("source_id", "")))
        job = {
            "job_id": f"{int(time.time() * 1000)}-{self._review_save_sequence}",
            "workbook": str(self.last_review_workbook),
            "line_id": _clean(issue.get("line_id", "")),
            "source_id": source_id,
            "decision_key": _clean(issue.get("decision_key", "")),
            "issue_signature": self._review_issue_signature(issue),
            "issue": {
                field: (bool(issue.get(field, False)) if field == "source_only" else _clean(issue.get(field, "")))
                for field in ("source_id", "order", "employee", "product", "description", "reason", "instructions", "source_only")
            },
            "values": values,
            "is_instruction_decision": is_instruction_decision,
            # A source-only validation issue was discovered on hidden All PO Lines
            # and has no visible Review & Edit row.  Save it directly to the
            # immutable source row while still requiring an exact Source ID match.
            "editing_search": editing_search or source_only,
            "allow_missing": source_only,
            "require_event_match": bool(source_id),
            "source_only": source_only,
            "created_at": time.time(),
        }
        self._queue_review_save_job(job)
        snapshot = self._optimistic_review_snapshot_after_save(issue, values)
        if (
            int(snapshot.get("purchase_review_count", snapshot.get("review_count", 0)) or 0) == 0
            and int(snapshot.get("product_master_review_count", 0) or 0) == 0
        ):
            self._update_review_state(
                review_complete=True,
                review_completed_at=time.time(),
                review_workbook=str(self.last_review_workbook or ""),
            )
        self._clear_saved_review_draft(issue)

        if editing_search:
            self._review_search_issue = None
        elif editing_previous:
            self._review_history_index = None
        else:
            history_key = _clean(issue.get("decision_key", "")) or _clean(issue.get("line_id", ""))
            last_key = ""
            if self._review_decision_history:
                last = self._review_decision_history[-1]
                last_key = _clean(last.get("decision_key", "")) or _clean(last.get("line_id", ""))
            if history_key and history_key != last_key:
                self._review_decision_history.append(dict(issue))

        # The next decision appears immediately. Workbook serialization continues
        # on one background worker, preserving order even when several decisions
        # are completed faster than Excel can be rewritten.
        self.refresh_review_page(snapshot=snapshot)

    def _archive_release_roots(self) -> list[Path]:
        """Return only archive roots explicitly bound to the active packet.

        Archive folders are historical records.  They must never be discovered
        by event name or modified date because two unrelated runs can share a
        name such as ``Line Boots 2026``.  An unbound archive can be opened
        from Event Archive, but can never become the current event's report.
        """
        preferred = []
        state = self._load_review_state()
        # ``last_purchase_order_dir`` is allowed to change when reports are
        # regenerated from a completed event.  The original completed packet is
        # not.  Retain that immutable archive path separately, then fall back to
        # the legacy archive key for Test 2 sessions created before this field
        # existed.
        for key in (
            "completed_purchase_order_archive_dir",
            "last_purchase_order_archive_dir",
        ):
            saved = _clean(state.get(key, ""))
            if saved:
                preferred.append(Path(saved).expanduser())
        unique = []
        seen = set()
        for candidate in preferred:
            try:
                resolved = candidate.resolve()
            except Exception:
                resolved = candidate
            if resolved in seen or not candidate.is_dir():
                continue
            seen.add(resolved)
            unique.append(candidate)
        return unique

    def _release_belongs_to_active_packet(self, packet_root: Path | None) -> bool:
        """Require a final release to carry this event's exact identity."""
        if self._is_archived_report_reopen():
            return True
        if not packet_root:
            return False
        state = self._load_review_state()
        active_id = _clean(state.get("active_packet_id", ""))
        active_lock_id = _clean(state.get("active_review_packet_lock_id", ""))
        if not active_id or not active_lock_id:
            return False
        try:
            manifest = json.loads((Path(packet_root) / "FINAL_RELEASE_MANIFEST.json").read_text(encoding="utf-8"))
        except (OSError, ValueError, TypeError):
            return False
        return (
            _clean(manifest.get("active_packet_id", "")) == active_id
            and _clean(manifest.get("packet_lock_id", "")) == active_lock_id
        )

    @staticmethod
    def _release_manifest_identity(packet_root: Path | None) -> str:
        """Return the exact manifest identity for one live or archived release."""
        if not packet_root:
            return ""
        manifest_path = Path(packet_root) / "FINAL_RELEASE_MANIFEST.json"
        if not manifest_path.is_file():
            return ""
        try:
            return hashlib.sha256(manifest_path.read_bytes()).hexdigest()
        except OSError:
            return ""

    @staticmethod
    def _manifest_vendor_purchase_order_count(packet_root: Path | None) -> int:
        """Count vendor POs listed in one protected release manifest.

        A report-only reopen can produce a valid manifest with just the three
        supporting reports.  It is a valid report release but not a complete
        purchase-order packet, so Finder must never use it for the Purchase
        Orders action.
        """
        if not packet_root:
            return 0
        manifest_path = Path(packet_root) / "FINAL_RELEASE_MANIFEST.json"
        try:
            manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        except Exception:
            return 0
        entries = manifest.get("files")
        if not isinstance(entries, list):
            return 0
        count = 0
        for entry in entries:
            if is_final_purchase_order_pdf(str(entry.get("path", ""))):
                count += 1
        return count

    def _verified_complete_archive_purchase_orders(self, archive_root: Path | None) -> Path | None:
        """Return an archive's PDF folder only when it includes vendor POs."""
        if not archive_root:
            return None
        archive_root = Path(archive_root)
        valid, _reason = validate_release_packet(archive_root, archived=True)
        archive_pdfs = archive_root / "Purchase Orders"
        if (
            valid
            and archive_pdfs.is_dir()
            and self._release_belongs_to_active_packet(archive_root)
            and self._manifest_vendor_purchase_order_count(archive_root) > 0
        ):
            return archive_pdfs
        return None

    def _purchase_order_release_location(self) -> tuple[Path | None, bool]:
        """Return only the manifest-verified release for the live active event.

        Completed event files are intentionally not a fallback. Once an event is
        finished and cleared, Orchid has no route back to its PDFs or workbook.
        That prevents an older event from ever becoming a current report source.
        """
        if not self._active_packet_is_valid()[0]:
            return None, False
        live = self.last_purchase_order_dir
        if live and live.is_dir():
            valid, _reason = validate_release_packet(live)
            if (
                valid
                and self._release_belongs_to_active_packet(live)
                and self._manifest_vendor_purchase_order_count(live) > 0
            ):
                return live, False
        return None, False

    def _archived_source_purchase_order_location(self) -> Path | None:
        """Return the original completed packet when viewing an archived report copy.

        Regenerating reports from an archived event deliberately produces only
        reports: it must never replace that event's original vendor purchase
        orders in the Finder display.  The archive source workbook identifies
        the immutable original packet, whose manifest validates the complete
        purchase-order set independently of the report-only working copy.
        """
        state = self._load_review_state()
        if not bool(state.get("archived_report_reopen", False)):
            return None
        source_workbook = _clean(state.get("archive_source_workbook", ""))
        if not source_workbook:
            return None
        archive_root = Path(source_workbook).expanduser().parent
        return self._verified_complete_archive_purchase_orders(archive_root)

    def _purchase_order_display_location(self) -> Path | None:
        """Return the complete, manifest-verified folder intended for Finder.

        Test 2 proved that a report-only reopen can become the newest active
        output.  That folder contains only supporting reports, so comparing
        against the mutable active manifest is unsafe.  Prefer the saved
        protected archive, and accept only a manifest that inventories at
        least one vendor purchase order.
        """
        live, _using_archive_fallback = self._purchase_order_release_location()
        if live:
            return live
        return self._archived_source_purchase_order_location()

    def _completed_purchase_packet_display_folder(self) -> Path | None:
        """Create one immutable Finder-friendly copy of the current PO packet.

        The Purchase Orders page may correctly list vendor PDFs while macOS
        Finder is still showing a separate, same-named report-only folder.  Do
        not ask Finder to resolve another live or archive path.  Instead, use
        the exact manifest-verified release that supplied the purchase-order
        list, copy its manifest-listed PDFs into a uniquely named display
        packet, and open that single unambiguous folder.

        This never changes the live release, Event Archive, or any existing
        display copy.  A matching display packet is reused; a mismatched or
        incomplete older display packet is left intact for diagnosis and a new
        uniquely named packet is created.
        """
        pdf_root, _using_archive_fallback = self._purchase_order_release_location()
        if not pdf_root or not pdf_root.is_dir():
            return None

        packet_root = None
        for candidate in (pdf_root, pdf_root.parent):
            if (candidate / "FINAL_RELEASE_MANIFEST.json").is_file():
                packet_root = candidate
                break
        if packet_root is None:
            return None

        valid, _reason = validate_release_packet(packet_root, archived=(packet_root != pdf_root))
        if not valid or self._manifest_vendor_purchase_order_count(packet_root) < 1:
            return None

        try:
            manifest = json.loads((packet_root / "FINAL_RELEASE_MANIFEST.json").read_text(encoding="utf-8"))
            entries = manifest.get("files", [])
        except Exception:
            return None
        if not isinstance(entries, list):
            return None

        source_roots = [pdf_root, packet_root / "Purchase Orders", packet_root]
        source_pdfs: list[tuple[Path, str, int]] = []
        source_names: set[str] = set()
        for entry in entries:
            name = Path(str(entry.get("path", ""))).name
            if not name.casefold().endswith(".pdf"):
                continue
            expected_sha256 = str(entry.get("sha256", "")).strip().casefold()
            try:
                expected_size = int(entry.get("size"))
            except (TypeError, ValueError):
                return None
            if len(expected_sha256) != 64 or expected_size < 1 or name in source_names:
                return None
            source = next((root / name for root in source_roots if (root / name).is_file()), None)
            if source is None:
                return None
            source_names.add(name)
            source_pdfs.append((source, expected_sha256, expected_size))
        if not source_pdfs:
            return None

        manifest_identity = self._release_manifest_identity(packet_root)
        if not manifest_identity:
            return None
        packet_name = safe_filename(self.current_event_name or manifest.get("event_name") or "Purchase Packet")
        display_parent = REPORTS / "Completed Purchase Packets"
        display_parent.mkdir(parents=True, exist_ok=True)
        base_name = f"{packet_name}__Completed_Purchase_Packet__{manifest_identity[:12]}"

        receipt = {
            "source_manifest_sha256": manifest_identity,
            "source_packet": str(packet_root),
            "pdf_count": len(source_pdfs),
            "pdfs": [path.name for path, _sha256, _size in source_pdfs],
            "verified_pdf_count": len(source_pdfs),
            "pdf_checks": [
                {"name": path.name, "sha256": sha256, "size": size}
                for path, sha256, size in source_pdfs
            ],
            "created_at": time.time(),
        }

        def sha256_file(path: Path) -> str:
            digest = hashlib.sha256()
            with path.open("rb") as file_handle:
                for chunk in iter(lambda: file_handle.read(1024 * 1024), b""):
                    digest.update(chunk)
            return digest.hexdigest()

        # Keep the internal verification record hidden in Finder.  It is only
        # used to prove that the visible PDFs are the exact locked release.
        receipt_name = ".orchid_packet_display_receipt.json"

        def complete_display_packet(folder: Path) -> bool:
            try:
                saved = json.loads((folder / receipt_name).read_text(encoding="utf-8"))
                return (
                    saved.get("source_manifest_sha256") == manifest_identity
                    and saved.get("pdfs") == receipt["pdfs"]
                    and saved.get("verified_pdf_count") == len(source_pdfs)
                    and all(
                        (folder / path.name).is_file()
                        and (folder / path.name).stat().st_size == size
                        and sha256_file(folder / path.name) == sha256
                        for path, sha256, size in source_pdfs
                    )
                )
            except Exception:
                return False

        destination = display_parent / base_name
        if destination.is_dir() and complete_display_packet(destination):
            return destination
        if destination.exists():
            suffix = 2
            while (display_parent / f"{base_name}__{suffix}").exists():
                suffix += 1
            destination = display_parent / f"{base_name}__{suffix}"

        staging = display_parent / f".{destination.name}__building_{time.time_ns()}"
        try:
            staging.mkdir(parents=False, exist_ok=False)
            for source, expected_sha256, expected_size in source_pdfs:
                destination_pdf = staging / source.name
                # Use a content-only copy.  macOS metadata copied by copy2 can
                # carry Finder visibility flags from an older report location.
                # A fresh content copy gives this display packet no inherited
                # Finder metadata, then verifies it against the locked manifest.
                with source.open("rb") as source_handle, destination_pdf.open("xb") as destination_handle:
                    shutil.copyfileobj(source_handle, destination_handle, length=1024 * 1024)
                if (
                    destination_pdf.stat().st_size != expected_size
                    or sha256_file(destination_pdf) != expected_sha256
                ):
                    raise OSError(f"Verification failed while creating display copy: {source.name}")
            (staging / receipt_name).write_text(
                json.dumps(receipt, indent=2), encoding="utf-8"
            )
            staging.rename(destination)
            return destination if complete_display_packet(destination) else None
        except Exception as error:
            # Preserve an incomplete staging copy for inspection rather than
            # deleting any evidence or source documents.
            try:
                (staging / "PACKET_DISPLAY_ERROR.txt").write_text(
                    f"Completed Purchase Packet could not be assembled.\n{error}\n",
                    encoding="utf-8",
                )
            except Exception:
                pass
            return None

    def _purchase_order_document_paths(self, root: Path | None = None) -> list[Path]:
        """Return final vendor purchase-order PDFs, excluding decoration-only reports."""
        if root is None:
            root, _archived = self._purchase_order_release_location()
        if not root or not root.exists():
            return []
        documents = [
            candidate for candidate in root.rglob("*.pdf")
            if is_final_purchase_order_pdf(candidate)
        ]
        return sorted(documents, key=lambda item: (item.name.casefold(), str(item)))

    def _open_purchase_document(self, document: Path) -> None:
        candidate = Path(document)
        if candidate.exists():
            subprocess.run(["open", str(candidate)], check=False)
        else:
            messagebox.showinfo("Purchase Order Not Found", "This purchase-order PDF is no longer available at its saved location.")

    def _completed_packet_report_paths(self) -> list[Path]:
        """Return all manifest-verified reports in the exact packet order."""
        display_dir = self._completed_purchase_packet_display_folder()
        if not display_dir:
            return []
        try:
            receipt = json.loads(
                (display_dir / ".orchid_packet_display_receipt.json").read_text(encoding="utf-8")
            )
            names = receipt.get("pdfs", [])
        except Exception:
            return []
        if not isinstance(names, list):
            return []
        documents = [display_dir / str(name) for name in names]
        return documents if documents and all(path.is_file() for path in documents) else []

    def print_all_purchase_reports(self) -> None:
        """Make one print-ready PDF from the verified completed packet.

        The source PDFs are read only.  A temporary combined print file is
        written only after every source page has been read successfully, then
        atomically promoted for Preview to open as one document.
        """
        documents = self._completed_packet_report_paths()
        if not documents:
            messagebox.showwarning(
                "Print Packet Not Ready",
                "Orchid could not assemble the verified completed packet. The existing purchase orders were not changed.",
            )
            return
        try:
            from pypdf import PdfReader, PdfWriter

            print_dir = REPORTS / "Print Packets"
            print_dir.mkdir(parents=True, exist_ok=True)
            content_key = hashlib.sha256(
                "|".join(f"{path.name}:{_file_stamp(path)}" for path in documents).encode("utf-8")
            ).hexdigest()[:12]
            packet_name = safe_filename(self.current_event_name or "Purchase Packet")
            output = print_dir / f"{packet_name}__Print_All_Reports__{content_key}.pdf"
            if not output.is_file():
                staging = print_dir / f".{output.name}.{time.time_ns()}.building"
                writer = PdfWriter()
                for document in documents:
                    reader = PdfReader(str(document))
                    for page in reader.pages:
                        writer.add_page(page)
                with staging.open("wb") as handle:
                    writer.write(handle)
                if not staging.is_file() or staging.stat().st_size < 1:
                    raise OSError("The combined print file was not created.")
                staging.replace(output)
            opened = subprocess.run(["open", "-a", "Preview", str(output)], check=False)
            if opened.returncode != 0:
                raise OSError("Preview could not be opened.")
            messagebox.showinfo(
                "Print All Reports Ready",
                "All reports are combined in one PDF. In Preview, choose File > Print once to print the complete packet.",
            )
        except Exception as error:
            messagebox.showerror("Unable to Prepare Print Packet", str(error))

    def _render_purchase_order_documents(self, documents: list[Path]) -> None:
        if not hasattr(self, "purchase_documents_frame"):
            return
        render_key = tuple((str(path.resolve()), _file_stamp(path)) for path in documents)
        if render_key == getattr(self, "_purchase_documents_render_key", None):
            return
        for child in self.purchase_documents_frame.winfo_children():
            child.destroy()
        self._purchase_documents_render_key = render_key
        count = len(documents)
        self.purchase_document_count.configure(text=str(count))
        if not documents:
            empty = ctk.CTkFrame(
                self.purchase_documents_frame, fg_color="#FBF9FD",
                border_width=1, border_color=BORDER, corner_radius=14,
            )
            empty.pack(fill="x", padx=8, pady=8)
            ctk.CTkLabel(
                empty, text="No purchase-order PDFs have been generated yet.",
                text_color=MUTED, font=ctk.CTkFont(size=13),
            ).pack(padx=18, pady=24)
            self.purchase_document_footer.configure(text="Purchase orders will appear here after Step 5 is completed.")
            return

        from datetime import datetime
        for document in documents:
            row = ctk.CTkFrame(
                self.purchase_documents_frame, fg_color=WHITE,
                border_width=1, border_color=BORDER, corner_radius=12,
            )
            row.pack(fill="x", padx=7, pady=5)
            row.grid_columnconfigure(1, weight=1)
            ctk.CTkLabel(
                row, text="PDF", width=44, height=44, corner_radius=10,
                fg_color="#FFF0F0", text_color="#C52C2C",
                font=ctk.CTkFont(size=11, weight="bold"),
            ).grid(row=0, column=0, rowspan=2, padx=(12, 10), pady=10)
            ctk.CTkLabel(
                row, text=document.stem, text_color=PURPLE_DARK,
                font=ctk.CTkFont(size=14, weight="bold"), anchor="w",
            ).grid(row=0, column=1, padx=(0, 8), pady=(10, 1), sticky="ew")
            try:
                modified = datetime.fromtimestamp(document.stat().st_mtime).strftime("%b %d, %Y  %I:%M %p")
                size_kb = max(1, round(document.stat().st_size / 1024))
                detail = f"{modified}  •  {size_kb:,} KB"
            except OSError:
                detail = str(document.parent)
            ctk.CTkLabel(
                row, text=detail, text_color=MUTED,
                font=ctk.CTkFont(size=11), anchor="w",
            ).grid(row=1, column=1, padx=(0, 8), pady=(1, 10), sticky="ew")
            ctk.CTkButton(
                row, text="Open", width=82, height=34, corner_radius=9,
                fg_color="#F8F4FC", hover_color=PURPLE_LIGHT,
                border_width=1, border_color="#CDBBE5", text_color=PURPLE,
                font=ctk.CTkFont(size=12, weight="bold"),
                command=lambda path=document: self._open_purchase_document(path),
            ).grid(row=0, column=2, rowspan=2, padx=(8, 12), pady=10)
        self.purchase_document_footer.configure(
            text=f"Displaying all {count} purchase-order document{'s' if count != 1 else ''}."
        )

    # ---------- decoration color audit ----------
    def _decoration_audit_cache_key(self) -> tuple:
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            return ()
        return (
            _file_stamp(self.last_review_workbook),
            _file_stamp(live_product_master_path()),
            _file_stamp(REVIEW_STATE_FILE),
        )

    def _build_decoration_audit_cached(self, force: bool = False) -> list[dict]:
        cache_key = self._decoration_audit_cache_key()
        if not cache_key:
            return []
        if (
            not force
            and cache_key == getattr(self, "_decoration_audit_build_cache_key", None)
            and isinstance(getattr(self, "_decoration_audit_build_cache", None), list)
        ):
            return self._decoration_audit_build_cache
        records = self._decoration_audit_records_from_current_event(force=force)
        audit = build_decoration_color_audit(
            records, live_product_master_path(), self._decoration_audit_verified_map()
        )
        self._decoration_audit_build_cache_key = cache_key
        self._decoration_audit_build_cache = audit
        return audit

    def _decoration_audit_item_needs_attention(self, item: dict) -> bool:
        """Use the canonical persisted definition for tracker, metrics, and rows."""
        key = _clean(item.get("key", "")) if isinstance(item, dict) else ""
        pending = getattr(self, "decoration_audit_pending_edits", {}).get(key, {})
        failure = self._decoration_audit_failure_map().get(key, {})
        return audit_item_needs_attention(item, pending, failure)

    def _decoration_audit_status(self, force: bool = False) -> dict:
        """Return cached completion status for the current event color audit."""
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            return {"total": 0, "needs_attention": 0, "verified": 0, "complete": False}
        cache_key = self._decoration_audit_cache_key()
        if (
            not force
            and cache_key == getattr(self, "_decoration_audit_status_cache_key", None)
            and isinstance(getattr(self, "_decoration_audit_status_cache_value", None), dict)
        ):
            return dict(self._decoration_audit_status_cache_value)
        try:
            audit = self._build_decoration_audit_cached(force=force)
            snapshot = build_audit_view_snapshot(
                audit, getattr(self, "decoration_audit_pending_edits", {}),
                self._decoration_audit_failure_map(), "All Decoration", True,
            )
            result = {
                "total": snapshot["total"],
                "needs_attention": snapshot["needs_attention"],
                "verified": snapshot["verified"],
                "complete": snapshot["needs_attention"] == 0,
            }
        except Exception as error:
            result = {
                "total": 0, "needs_attention": 1, "verified": 0,
                "complete": False, "error": str(error),
            }
        self._decoration_audit_status_cache_key = cache_key
        self._decoration_audit_status_cache_value = dict(result)
        return result
    def _require_completed_decoration_audit(self) -> bool:
        status = self._decoration_audit_status(force=True)
        remaining = int(status.get("needs_attention", 0) or 0)
        if status.get("complete"):
            return True
        detail = status.get("error") or (
            f"{remaining} decoration color combination{'s' if remaining != 1 else ''} still need verification."
        )
        messagebox.showwarning(
            "Decoration Color Audit Required",
            "Complete the Decoration Color Audit before generating Purchase Orders.\n\n" + detail,
        )
        self.show_page("audit")
        return False

    def build_decoration_audit_page(self):
        page = self.new_page("audit")
        page.grid_rowconfigure(4, weight=1)
        self._build_workflow_tracker(page, 0, "audit_stage_labels", padx=24)

        header = ctk.CTkFrame(page, fg_color="transparent")
        header.grid(row=1, column=0, sticky="ew", padx=32, pady=(4, 8))
        header.grid_columnconfigure(0, weight=1)
        ctk.CTkLabel(
            header, text="Decoration Color Audit", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=28, weight="bold"), anchor="w",
        ).grid(row=0, column=0, sticky="w")
        ctk.CTkLabel(
            header,
            text=(
                "Review the thread or ink color and decoration location for each style and garment-color "
                "combination in this event. Save once to update Product Master and every matching event line."
            ),
            text_color=MUTED, font=ctk.CTkFont(size=13), anchor="w", justify="left",
        ).grid(row=1, column=0, sticky="w", pady=(3, 0))

        metrics = ctk.CTkFrame(page, fg_color="transparent")
        metrics.grid(row=2, column=0, sticky="ew", padx=32, pady=(4, 6))
        for column in range(4):
            metrics.grid_columnconfigure(column, weight=1)
        self.audit_metric_labels = {}
        for column, (key, title) in enumerate((
            ("groups", "Event combinations"),
            ("issues", "Need attention"),
            ("verified", "Verified"),
            ("lines", "Affected items"),
        )):
            card = ctk.CTkFrame(metrics, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=14)
            card.grid(row=0, column=column, sticky="ew", padx=5, pady=3)
            value = ctk.CTkLabel(card, text="—", text_color=PURPLE, font=ctk.CTkFont(size=22, weight="bold"))
            value.pack(anchor="w", padx=16, pady=(12, 0))
            ctk.CTkLabel(card, text=title, text_color=MUTED, font=ctk.CTkFont(size=11)).pack(
                anchor="w", padx=16, pady=(0, 12)
            )
            self.audit_metric_labels[key] = value

        controls = ctk.CTkFrame(page, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=14)
        controls.grid(row=3, column=0, sticky="ew", padx=32, pady=(4, 8))
        controls.grid_columnconfigure(4, weight=1)
        # Open on the complete unresolved list. Starting on Embroidery could make
        # a Screen Print-only event look empty even while Need Attention was nonzero.
        self.audit_type_var = ctk.StringVar(value="All Decoration")
        ctk.CTkLabel(controls, text="Show", text_color=TEXT, font=ctk.CTkFont(size=12, weight="bold")).grid(
            row=0, column=0, padx=(16, 6), pady=12
        )
        self.audit_type_menu = ctk.CTkOptionMenu(
            controls, variable=self.audit_type_var,
            values=["Embroidery", "Screen Print", "All Decoration"],
            command=lambda _value: self._render_decoration_audit_rows(),
            width=155, fg_color=PURPLE, button_color=PURPLE_DARK,
        )
        self.audit_type_menu.grid(row=0, column=1, padx=(0, 12), pady=12)
        self.audit_issues_only_var = ctk.BooleanVar(value=True)
        ctk.CTkCheckBox(
            controls, text="Show items needing attention", variable=self.audit_issues_only_var,
            command=self._render_decoration_audit_rows, text_color=TEXT,
        ).grid(row=0, column=2, padx=10, pady=12)
        ctk.CTkButton(
            controls, text="Rebuild Audit Display", width=155, height=32,
            fg_color=WHITE, hover_color=PURPLE_LIGHT, text_color=PURPLE,
            border_width=1, border_color="#7C55D9",
            command=self.rebuild_decoration_audit_display,
        ).grid(row=0, column=3, padx=(10, 4), pady=10)
        self.audit_status_label = ctk.CTkLabel(
            controls, text="", text_color=MUTED, anchor="e", font=ctk.CTkFont(size=11),
        )
        self.audit_status_label.grid(row=0, column=4, sticky="e", padx=(8, 16), pady=12)

        table_card = ctk.CTkFrame(page, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=14)
        table_card.grid(row=4, column=0, sticky="nsew", padx=32, pady=(2, 8))
        table_card.grid_rowconfigure(0, weight=1)
        table_card.grid_columnconfigure(0, weight=1)
        self.decoration_audit_scroll = ctk.CTkScrollableFrame(
            table_card, fg_color=WHITE, corner_radius=0,
        )
        self.decoration_audit_scroll.grid(row=0, column=0, sticky="nsew", padx=8, pady=8)

        actions = ctk.CTkFrame(page, fg_color="transparent")
        actions.grid(row=5, column=0, sticky="ew", padx=32, pady=(0, 22))
        actions.grid_columnconfigure(3, weight=1)
        ctk.CTkButton(
            actions, text="Save Audit Changes", height=44, width=210,
            fg_color=PURPLE, hover_color=PURPLE_DARK,
            font=ctk.CTkFont(size=13, weight="bold"),
            command=self.save_decoration_audit_changes,
        ).grid(row=0, column=0, padx=(0, 8), pady=4)
        ctk.CTkButton(
            actions, text="Open Product Master", height=44, width=165,
            fg_color=WHITE, hover_color=PURPLE_LIGHT, text_color=PURPLE,
            border_width=1, border_color="#7C55D9",
            command=self.open_product_master,
        ).grid(row=0, column=1, padx=8, pady=4)
        ctk.CTkButton(
            actions, text="Continue to Regenerate Purchase Orders", height=44, width=245,
            fg_color=WHITE, hover_color=PURPLE_LIGHT, text_color=PURPLE,
            border_width=1, border_color="#7C55D9",
            command=lambda: self.show_page("purchase"),
        ).grid(row=0, column=2, padx=8, pady=4)

    def _decoration_audit_verified_map(self) -> dict:
        state = self._load_review_state()
        value = state.get("decoration_color_audit_verified", {})
        verified = dict(value) if isinstance(value, dict) else {}

        # Audit overrides are written atomically for every successfully saved
        # event line.  Treat those durable, user-approved overrides as
        # confirmation candidates as well.  This both restores an RC13 test
        # state where two confirmations were overwritten under the same rebuilt
        # display key and makes rebuild recovery independent of a visual group
        # merge.  The audit builder deduplicates these candidates against the
        # normal confirmation records by immutable line IDs, color, and
        # location, so a multi-line combination still counts once.
        overrides = state.get("decoration_color_audit_overrides", [])
        # A normal confirmation may represent several source lines.  Only use
        # a durable override to repair a line that has no confirmation at all;
        # otherwise we would double-count a legitimate multi-line group.
        confirmed_line_ids = {
            _clean(line_id)
            for candidate in verified.values()
            if isinstance(candidate, dict)
            for line_id in candidate.get("line_ids", [])
            if _clean(line_id)
        }
        if isinstance(overrides, list):
            for override in overrides:
                if not isinstance(override, dict):
                    continue
                line_id = _clean(override.get("line_id", ""))
                color = _clean(override.get("decoration_color", ""))
                location = _clean(override.get("location", ""))
                if not line_id or line_id in confirmed_line_ids or not color or not location:
                    continue
                stable_identity = "|".join([line_id, color.casefold(), location.casefold()])
                storage_key = "audit-override-confirmation:" + hashlib.sha256(
                    stable_identity.encode("utf-8")
                ).hexdigest()
                verified[storage_key] = {
                    "color": color,
                    "location": location,
                    "line_ids": [line_id],
                    "style": _clean(override.get("style", "")),
                    "garment_color": _clean(override.get("garment_color", "")),
                    "decoration_type": _clean(override.get("decoration_type", "")),
                }
        return verified

    def _decoration_audit_failure_map(self) -> dict:
        state = self._load_review_state()
        value = state.get("decoration_color_audit_failures", {})
        return dict(value) if isinstance(value, dict) else {}

    def rebuild_decoration_audit_display(self):
        """Rebuild the display only; never alter Product Master or event data."""
        self._capture_decoration_audit_edits()
        self._decoration_audit_event_records_cache_key = None
        self._decoration_audit_build_cache_key = None
        self._decoration_audit_status_cache_key = None
        self._decoration_audit_status_cache_value = None
        self.audit_status_label.configure(text="Rebuilding audit display…", text_color=WARNING)
        self.update_idletasks()
        self.refresh_decoration_audit_page()

    def _decoration_audit_records_from_current_event(self, force: bool = False) -> list[dict]:
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            return []
        source = Path(self.last_review_workbook)
        cache_key = (_file_stamp(source), _file_stamp(REVIEW_STATE_FILE))
        if (
            not force
            and cache_key == getattr(self, "_decoration_audit_event_records_cache_key", None)
            and isinstance(getattr(self, "_decoration_audit_event_records_cache", None), list)
        ):
            return self._decoration_audit_event_records_cache
        # Include completed Purchase Review decisions and prior audit overrides.
        if self._completed_review_records_from_state(source) or self._audit_override_jobs_for_workbook(source):
            source = self._build_journal_reconciled_workbook(source)
        records = load_review_lines(source)
        self._decoration_audit_event_records_cache_key = cache_key
        self._decoration_audit_event_records_cache = records
        return records

    def refresh_decoration_audit_page(self):
        if not hasattr(self, "decoration_audit_scroll"):
            return
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            self.decoration_audit_records = []
            self.audit_status_label.configure(text="Create or resume a Purchase Review first.", text_color=WARNING)
            self._render_decoration_audit_rows()
            return
        current_source_stamp = self._decoration_audit_cache_key()
        if current_source_stamp != getattr(self, "_decoration_audit_source_stamp", None):
            self.decoration_audit_pending_edits = {}
            self._decoration_audit_source_stamp = current_source_stamp
        started = time.perf_counter()
        try:
            self.decoration_audit_records = self._build_decoration_audit_cached(force=True)
            elapsed = time.perf_counter() - started
            self.audit_status_label.configure(
                text=f"Loaded {len(self.decoration_audit_records):,} combinations in {elapsed:.1f}s",
                text_color=SUCCESS,
            )
        except Exception as error:
            self.decoration_audit_records = []
            self.audit_status_label.configure(text=f"Unable to load audit: {error}", text_color=DANGER)
        self._render_decoration_audit_rows()

    def _capture_decoration_audit_edits(self) -> None:
        pending = dict(getattr(self, "decoration_audit_pending_edits", {}) or {})
        for key, variables in getattr(self, "decoration_audit_row_vars", {}).items():
            pending[key] = {
                "color": _clean(variables["color"].get()),
                "location": _clean(variables["location"].get()),
                "verified": bool(variables["verified"].get()),
            }
        self.decoration_audit_pending_edits = pending

    def _render_decoration_audit_rows(self):
        try:
            self._render_decoration_audit_rows_impl()
        except Exception as error:
            self._render_decoration_audit_fallback(error)

    def _render_decoration_audit_fallback(self, error: Exception):
        """Never leave a blank table while unresolved rows exist."""
        try:
            for child in self.decoration_audit_scroll.winfo_children():
                child.destroy()
        except Exception:
            return
        records = [item for item in getattr(self, "decoration_audit_records", []) if isinstance(item, dict)]
        snapshot = build_audit_view_snapshot(
            records, getattr(self, "decoration_audit_pending_edits", {}),
            self._decoration_audit_failure_map(), "All Decoration", True,
        )
        self.audit_status_label.configure(
            text="Recovered audit display after a rendering error.", text_color=WARNING
        )
        ctk.CTkLabel(
            self.decoration_audit_scroll,
            text="Audit display recovered. These unresolved combinations are still safe and have not been changed.",
            text_color=WARNING, font=ctk.CTkFont(size=13, weight="bold"),
            anchor="w", justify="left", wraplength=900,
        ).grid(row=0, column=0, sticky="ew", padx=12, pady=(12, 6))
        unresolved = snapshot.get("unresolved", [])
        if not unresolved:
            ctk.CTkLabel(
                self.decoration_audit_scroll, text=f"Unable to render the audit: {error}",
                text_color=DANGER, anchor="w", justify="left", wraplength=900,
            ).grid(row=1, column=0, sticky="w", padx=12, pady=12)
            return
        failures = self._decoration_audit_failure_map()
        for row_number, item in enumerate(unresolved, start=1):
            key = _clean(item.get("key", ""))
            failure = failures.get(key, {}) if isinstance(failures.get(key, {}), dict) else {}
            reason = _clean(failure.get("error", "")) or "; ".join(item.get("issues", [])) or "Save required"
            frame = ctk.CTkFrame(
                self.decoration_audit_scroll, fg_color="#FBF9FE",
                border_width=1, border_color="#EEE8F4", corner_radius=8,
            )
            frame.grid(row=row_number, column=0, sticky="ew", padx=4, pady=3)
            frame.grid_columnconfigure(0, weight=1)
            ctk.CTkLabel(
                frame,
                text=(f"{item.get('decoration_type', '')}  •  "
                      f"{item.get('style') or item.get('product_name') or 'Unknown'}  •  "
                      f"{item.get('garment_color') or 'No garment color'}"),
                text_color=PURPLE_DARK, font=ctk.CTkFont(size=12, weight="bold"), anchor="w",
            ).grid(row=0, column=0, sticky="ew", padx=12, pady=(9, 1))
            ctk.CTkLabel(
                frame, text=reason, text_color=WARNING, anchor="w", justify="left", wraplength=850,
            ).grid(row=1, column=0, sticky="ew", padx=12, pady=(1, 9))

    def _render_decoration_audit_rows_impl(self):
        if not hasattr(self, "decoration_audit_scroll"):
            return
        self._capture_decoration_audit_edits()
        for child in self.decoration_audit_scroll.winfo_children():
            child.destroy()
        self.decoration_audit_row_vars = {}

        records = [item for item in getattr(self, "decoration_audit_records", []) if isinstance(item, dict)]
        selected_type = self.audit_type_var.get() if hasattr(self, "audit_type_var") else "All Decoration"
        issues_only = bool(self.audit_issues_only_var.get()) if hasattr(self, "audit_issues_only_var") else True
        snapshot = build_audit_view_snapshot(
            records, getattr(self, "decoration_audit_pending_edits", {}),
            self._decoration_audit_failure_map(), selected_type, issues_only,
        )
        total_groups = snapshot["total"]
        issue_count = snapshot["needs_attention"]
        verified_count = snapshot["verified"]
        affected_lines = snapshot["affected_lines"]
        visible = list(snapshot["visible"])
        selected_type = snapshot["selected_type"]
        if hasattr(self, "audit_type_var") and self.audit_type_var.get() != selected_type:
            self.audit_type_var.set(selected_type)
        for key, value in (("groups", total_groups), ("issues", issue_count), ("verified", verified_count), ("lines", affected_lines)):
            if hasattr(self, "audit_metric_labels") and key in self.audit_metric_labels:
                self.audit_metric_labels[key].configure(text=f"{value:,}")

        # An attention view must never look complete merely because its decoration-
        # type filter has no unresolved rows. Reveal the canonical unresolved list
        # and synchronize the menu to All Decoration so every row is editable.
        if issues_only and issue_count and not visible:
            visible = list(snapshot["unresolved"])
            selected_type = "All Decoration"
            if hasattr(self, "audit_type_var"):
                self.audit_type_var.set(selected_type)
            self.audit_status_label.configure(
                text=f"Showing all {len(visible)} unresolved row(s) that need attention.",
                text_color=WARNING,
            )

        columns = [
            ("Type", 82), ("Style", 82), ("Garment Color", 112),
            ("Decoration Location", 145), ("Correct Thread / Ink", 145),
            ("Product Master", 130), ("Qty / Orders", 76), ("Attention", 210),
            ("Verified", 66), ("Orders", 60),
        ]
        header = ctk.CTkFrame(self.decoration_audit_scroll, fg_color=PURPLE_LIGHT, corner_radius=9)
        header.grid(row=0, column=0, sticky="ew", pady=(0, 5))
        for column, (label, width) in enumerate(columns):
            header.grid_columnconfigure(column, minsize=width, weight=1 if column == 7 else 0)
            ctk.CTkLabel(
                header, text=label, text_color=PURPLE_DARK,
                font=ctk.CTkFont(size=10, weight="bold"), anchor="w",
            ).grid(row=0, column=column, sticky="ew", padx=5, pady=8)

        if not visible:
            message = "No decoration combinations match this view."
            if records and issues_only:
                selected_remaining = sum(
                    1 for item in records
                    if self._decoration_audit_item_needs_attention(item)
                    and (selected_type == "All Decoration" or item.get("decoration_type") == selected_type)
                )
                if issue_count and not selected_remaining:
                    message = (
                        f"{issue_count} unresolved combination{'s' if issue_count != 1 else ''} remain in another "
                        "decoration type. Choose All Decoration to display them."
                    )
                elif issue_count:
                    message = (
                        f"{issue_count} unresolved combination{'s' if issue_count != 1 else ''} remain. "
                        "Refresh this page to display the current Product Master status."
                    )
                else:
                    message = "No unresolved decoration color issues remain in this view."
            ctk.CTkLabel(
                self.decoration_audit_scroll, text=message, text_color=MUTED,
                font=ctk.CTkFont(size=13), wraplength=850, justify="left",
            ).grid(row=1, column=0, sticky="w", padx=12, pady=24)
            return

        pending_map = self.decoration_audit_pending_edits
        saved_thread_ink_colors = load_thread_ink_colors(
            live_product_master_path(), include_blank=False
        )
        for row_number, item in enumerate(visible, start=1):
            pending = pending_map.get(item["key"], {})
            pending = pending if isinstance(pending, dict) else {}
            row = ctk.CTkFrame(
                self.decoration_audit_scroll,
                fg_color="#FBF9FE" if row_number % 2 else WHITE,
                border_width=1, border_color="#EEE8F4", corner_radius=8,
            )
            row.grid(row=row_number, column=0, sticky="ew", pady=2)
            for column, (_label, width) in enumerate(columns):
                row.grid_columnconfigure(column, minsize=width, weight=1 if column == 7 else 0)

            color_var = ctk.StringVar(value=pending.get("color", item.get("current_color", "")))
            location_var = ctk.StringVar(value=pending.get("location", item.get("location", "")))
            verified_var = ctk.BooleanVar(value=pending.get("verified", item.get("verified", False)))
            self.decoration_audit_row_vars[item["key"]] = {
                "record": item, "color": color_var, "location": location_var, "verified": verified_var,
            }
            ctk.CTkLabel(row, text=item.get("decoration_type", ""), text_color=TEXT, anchor="w", font=ctk.CTkFont(size=10)).grid(row=0, column=0, sticky="ew", padx=4)
            ctk.CTkLabel(row, text=item.get("style", "") or "—", text_color=TEXT, anchor="w", font=ctk.CTkFont(size=10, weight="bold")).grid(row=0, column=1, sticky="ew", padx=4)
            ctk.CTkLabel(row, text=item.get("garment_color", "") or "—", text_color=TEXT, anchor="w", wraplength=112, font=ctk.CTkFont(size=10)).grid(row=0, column=2, sticky="ew", padx=4)
            location_values = list(dict.fromkeys([item.get("location", "")] + DECORATION_LOCATIONS))
            location_values = [value for value in location_values if value] or [LEFT_CHEST]
            location_combo = ctk.CTkComboBox(row, variable=location_var, values=location_values, width=145, height=30)
            location_combo.grid(row=0, column=3, sticky="ew", padx=4, pady=7)
            color_values = list(dict.fromkeys(
                saved_thread_ink_colors
                + list(item.get("colors_found", []))
                + [item.get("master_color", ""), item.get("suggested_color", "")]
            ))
            color_values = [value for value in color_values if value] or [""]
            combo = ctk.CTkComboBox(row, variable=color_var, values=color_values, width=145, height=30)
            combo.grid(row=0, column=4, sticky="ew", padx=4, pady=7)
            master_text = item.get("master_color", "") or "Not saved"
            master_location = _clean(item.get("master_location", ""))
            source = item.get("master_source", "")
            if source and source != "Not saved":
                master_text += f"\n{source}"
            if master_location and master_location.casefold() != _clean(item.get("location", "")).casefold():
                master_text += f"\nLocation: {master_location}"
            ctk.CTkLabel(row, text=master_text, text_color=MUTED, anchor="w", justify="left", wraplength=130, font=ctk.CTkFont(size=9)).grid(row=0, column=5, sticky="ew", padx=4)
            ctk.CTkLabel(
                row, text=f"{item.get('quantity', 0)} / {item.get('order_count', 0)}",
                text_color=TEXT, font=ctk.CTkFont(size=10, weight="bold"),
            ).grid(row=0, column=6, padx=4)
            issue_text = "; ".join(item.get("issues", [])) or "No issue detected"
            failure = self._decoration_audit_failure_map().get(item.get("key", ""), {})
            failure = failure if isinstance(failure, dict) else {}
            if _clean(failure.get("error", "")):
                issue_text = _clean(failure.get("error", "")) + "\n" + issue_text
            issue_color = WARNING if item.get("suspicious") else SUCCESS
            if item.get("suggested_color"):
                issue_text += f"\nSuggestion only: {item['suggested_color']}"
            if pending.get("verified") and self._decoration_audit_item_needs_attention(item):
                issue_text += "\nSave required to complete this row."
            ctk.CTkLabel(
                row, text=issue_text, text_color=issue_color, anchor="w", justify="left",
                wraplength=205, font=ctk.CTkFont(size=9),
            ).grid(row=0, column=7, sticky="ew", padx=4)
            ctk.CTkCheckBox(row, text="", width=24, variable=verified_var).grid(row=0, column=8, padx=10)
            ctk.CTkButton(
                row, text="View", width=52, height=28, fg_color=WHITE,
                hover_color=PURPLE_LIGHT, text_color=PURPLE, border_width=1, border_color=BORDER,
                command=lambda record=item: self.show_decoration_audit_orders(record),
            ).grid(row=0, column=9, padx=4, pady=7)
    def _pending_decoration_audit_changes(self) -> tuple[list[dict], dict[str, dict]]:
        self._capture_decoration_audit_edits()
        record_map = {item["key"]: item for item in self.decoration_audit_records}
        changes: list[dict] = []
        final_values: dict[str, dict] = {}
        for key, pending in self.decoration_audit_pending_edits.items():
            record = record_map.get(key)
            if not record:
                continue
            color = _clean(pending.get("color", record.get("current_color", "")))
            location = _clean(pending.get("location", record.get("location", "")))
            verified = bool(pending.get("verified", record.get("verified", False)))
            event_change_required = bool(
                color.casefold() != _clean(record.get("current_color", "")).casefold()
                or location.casefold() != _clean(record.get("location", "")).casefold()
            )
            master_sync_required = bool(verified and not record.get("master_synchronized"))
            requires_save = event_change_required or master_sync_required
            final_values[key] = {
                "color": color,
                "location": location,
                "verified": verified,
                "record": record,
                "requires_save": requires_save,
                "event_change_required": event_change_required,
            }
            if requires_save:
                changed = dict(record)
                changed["key"] = key
                changed["decoration_color"] = color
                changed["location"] = location
                changed["event_change_required"] = event_change_required
                changes.append(changed)
        return changes, final_values
    def save_decoration_audit_changes(self):
        changes, final_values = self._pending_decoration_audit_changes()
        if self._existing_product_master_pid() and changes:
            messagebox.showwarning(
                "Close Product Master First",
                "Product Master is currently open. Close that window before saving audit changes so neither save can overwrite the other.",
            )
            return

        missing = [
            f"{value['record'].get('style') or value['record'].get('product_name')} / {value['record'].get('garment_color')}"
            for value in final_values.values()
            if (value["verified"] or value["requires_save"]) and not value["color"]
        ]
        if missing:
            messagebox.showwarning(
                "Decoration Color Required",
                "Choose a correct thread or ink color before saving:\n\n" + "\n".join(missing[:12]),
            )
            return
        invalid_locations = [
            f"{value['record'].get('style') or value['record'].get('product_name')} / {value['record'].get('garment_color')}"
            for value in final_values.values()
            if (value["verified"] or value["requires_save"])
            and (not value["location"] or value["location"] == OTHER_CUSTOM)
        ]
        if invalid_locations:
            messagebox.showwarning(
                "Specific Decoration Location Required",
                "Choose a standard location or type the actual custom location for:\n\n"
                + "\n".join(invalid_locations[:12]),
            )
            return

        verification_requests = sum(1 for value in final_values.values() if value["verified"])
        if changes:
            affected_lines = sum(len(item.get("line_records", [])) for item in changes)
            if not messagebox.askyesno(
                "Save Decoration Audit Changes?",
                f"Save {len(changes)} exact Product Master style/color rule(s) and update up to {affected_lines} current-event line(s)?\n\n"
                "The app will create a Product Master backup first. Rows are only completed after both saves succeed.",
            ):
                return
        elif not verification_requests:
            self.audit_status_label.configure(text="No audit changes to save.", text_color=MUTED)
            return

        master_result = {
            "updated": 0, "added": 0, "created_styles": 0,
            "saved_keys": [], "failed": [],
        }
        fully_saved_keys: set[str] = set()
        event_failed: list[dict] = []
        try:
            if changes:
                # This function reloads Product Master from disk, builds fresh
                # indexes, writes all successful rules atomically, and reports
                # individual failures instead of stopping at the first style.
                master_result = apply_product_master_decoration_colors(
                    live_product_master_path(), changes
                )

            master_saved_keys = set(master_result.get("saved_keys", []) or [])
            change_map = {change["key"]: change for change in changes}
            event_changes = [
                change for change in changes
                if change["key"] in master_saved_keys and change.get("event_change_required")
            ]
            master_only_keys = {
                change["key"] for change in changes
                if change["key"] in master_saved_keys and not change.get("event_change_required")
            }
            fully_saved_keys.update(master_only_keys)

            if event_changes:
                jobs = []
                for change in event_changes:
                    for index, record in enumerate(change.get("line_records", []), start=1):
                        values = {
                            "Decoration Color": _clean(change.get("decoration_color", "")),
                            "Product #": _clean(record.get("Product #", change.get("style", ""))),
                            "Description": _clean(record.get("Description", change.get("product_name", ""))),
                            "Garment Color": _clean(record.get("Garment Color", change.get("garment_color", ""))),
                            "Decoration Type": _clean(record.get("Decoration Type", change.get("decoration_type", ""))),
                            "Decoration Location": _clean(change.get("location", record.get("Decoration Location", ""))),
                            "Purchase Vendor": _clean(record.get("Purchase Vendor", "")),
                            "Size": _clean(record.get("Size", "")),
                            "Quantity": _clean(record.get("Quantity", "")),
                            "Order Number": _clean(record.get("Order Number", "")),
                            "Employee Name": _clean(record.get("Employee Name", "")),
                        }
                        jobs.append({
                            "job_id": f"audit-direct-{_clean(record.get('Line ID', '')) or index}-{time.time_ns()}",
                            "workbook": str(self.last_review_workbook),
                            "line_id": _clean(record.get("Line ID", "")),
                            "decision_key": _clean(record.get("Decision Key", "")),
                            "values": values,
                            "issue": {
                                "order": _clean(record.get("Order Number", "")),
                                "employee": _clean(record.get("Employee Name", "")),
                            },
                            "is_instruction_decision": False,
                            "audit_override": True,
                            "require_event_match": True,
                            "editing_search": True,
                            "allow_missing": True,
                        })
                try:
                    if jobs:
                        self._apply_review_save_jobs(jobs)
                    fully_saved_keys.update(change["key"] for change in event_changes)
                except Exception as error:
                    event_failed = [
                        {
                            "key": change["key"],
                            "style": _clean(change.get("style", "")) or _clean(change.get("product_name", "")),
                            "garment_color": _clean(change.get("garment_color", "")),
                            "error": "Current event save failed: " + (str(error) or error.__class__.__name__),
                        }
                        for change in event_changes
                    ]

            saved_changes = [change_map[key] for key in fully_saved_keys if key in change_map]
            if saved_changes:
                self._store_decoration_color_audit_changes(saved_changes)

            state = self._load_review_state()
            verified = state.get("decoration_color_audit_verified", {})
            verified = dict(verified) if isinstance(verified, dict) else {}
            failed_keys = {
                item.get("key", "") for item in (master_result.get("failed", []) or []) + event_failed
            }
            for old_key, value in final_values.items():
                record = value["record"]
                final_key = audit_group_key(
                    record.get("decoration_type", ""), record.get("style", ""),
                    record.get("product_name", ""), record.get("garment_color", ""), value["location"],
                )
                line_ids = list(dict.fromkeys(
                    _clean(line.get("Line ID", ""))
                    for line in record.get("line_records", [])
                    if _clean(line.get("Line ID", ""))
                ))
                # Remove an earlier confirmation for these exact source lines
                # before replacing it.  Do not use the rebuilt display key as
                # storage: different original confirmations can legitimately
                # share that key after reconciliation.
                line_id_set = set(line_ids)
                for verification_key, candidate in list(verified.items()):
                    if not isinstance(candidate, dict):
                        continue
                    candidate_ids = {
                        _clean(line_id) for line_id in candidate.get("line_ids", [])
                        if _clean(line_id)
                    }
                    if candidate_ids and candidate_ids == line_id_set:
                        verified.pop(verification_key, None)
                verification_identity = "|".join([
                    old_key, value["color"].casefold(), value["location"].casefold(), *sorted(line_ids),
                ])
                confirmation_key = "audit-confirmation:" + hashlib.sha256(
                    verification_identity.encode("utf-8")
                ).hexdigest()
                if value["verified"] and (not value["requires_save"] or old_key in fully_saved_keys):
                    verified[confirmation_key] = {
                        "color": value["color"],
                        "location": value["location"],
                        # Display grouping can be recomputed after journal
                        # reconciliation.  Keep immutable line identities so a
                        # successful audit confirmation remains visible even if
                        # its former group key changes or merges.
                        "line_ids": line_ids,
                        "style": _clean(record.get("style", "")),
                        "garment_color": _clean(record.get("garment_color", "")),
                        "decoration_type": _clean(record.get("decoration_type", "")),
                        "verified_at": time.time(),
                    }
                else:
                    # Failed saves can never remain verified. This keeps the
                    # attention count and visible rows in perfect agreement.
                    verified.pop(confirmation_key, None)
            state["decoration_color_audit_verified"] = verified
            failure_map = state.get("decoration_color_audit_failures", {})
            failure_map = dict(failure_map) if isinstance(failure_map, dict) else {}
            for old_key, value in final_values.items():
                record = value["record"]
                final_key = audit_group_key(
                    record.get("decoration_type", ""), record.get("style", ""),
                    record.get("product_name", ""), record.get("garment_color", ""), value["location"],
                )
                if old_key in fully_saved_keys:
                    failure_map.pop(old_key, None)
                    failure_map.pop(final_key, None)
            for item in (master_result.get("failed", []) or []) + event_failed:
                key = _clean(item.get("key", ""))
                if key:
                    failure_map[key] = {
                        "error": _clean(item.get("error", "Unable to save")),
                        "style": _clean(item.get("style", "")),
                        "garment_color": _clean(item.get("garment_color", "")),
                        "failed_at": time.time(),
                    }
            state["decoration_color_audit_failures"] = failure_map
            if fully_saved_keys:
                # The audit writes its exact color/location choice to the
                # current event lines before completing the row. A normal
                # thread/ink confirmation therefore never needs to rerun
                # Purchase Review. The protected routing check still catches
                # a later Product Master vendor/type/location/outsourcing edit.
                state.pop("decoration_color_audit_requires_regeneration", None)
            self._write_review_state_atomic(state)

            if master_saved_keys:
                try:
                    self.product_master_update_marker.write_text(
                        json.dumps({"saved_at": time.time(), "source": "decoration_color_audit"}, indent=2),
                        encoding="utf-8",
                    )
                except Exception:
                    pass
                global _PRODUCT_MASTER_HEALTH_CACHE_KEY, _PRODUCT_MASTER_HEALTH_CACHE_VALUE
                _PRODUCT_MASTER_HEALTH_CACHE_KEY = None
                _PRODUCT_MASTER_HEALTH_CACHE_VALUE = None
                self.last_purchase_order_dir = None

            # Keep failed input values visible for an immediate retry, but force
            # their verification boxes off. Successful rows are cleared.
            failed_pending: dict[str, dict] = {}
            for key in failed_keys:
                value = final_values.get(key)
                if value:
                    failed_pending[key] = {
                        "color": value["color"],
                        "location": value["location"],
                        "verified": False,
                    }
            self.decoration_audit_pending_edits = failed_pending
            self._decoration_audit_event_records_cache_key = None
            self._decoration_audit_build_cache_key = None
            self._decoration_audit_status_cache_key = None
            self._decoration_audit_status_cache_value = None
            self._invalidate_review_cache()
            self._invalidate_fast_review_snapshot()
            self._mission_snapshot_cache_key = None
            # Prevent the normal source-stamp refresh from discarding failed
            # values after a successful Product Master atomic replacement.
            self._decoration_audit_source_stamp = self._decoration_audit_cache_key()
            self.refresh_decoration_audit_page()
            self.refresh_dashboard()

            failures = list(master_result.get("failed", []) or []) + event_failed
            if failures:
                details = "\n".join(
                    f"• {item.get('style', 'Unknown')} / {item.get('garment_color', '')}: {item.get('error', 'Unable to save')}"
                    for item in failures[:12]
                )
                messagebox.showwarning(
                    "Decoration Audit Partially Saved",
                    f"Completed {len(fully_saved_keys)} combination(s).\n"
                    f"{len(failures)} combination(s) still need attention and remain visible.\n\n"
                    f"{details}",
                )
            elif changes:
                affected_lines = sum(
                    len(change_map[key].get("line_records", []))
                    for key in fully_saved_keys if key in change_map
                )
                messagebox.showinfo(
                    "Decoration Audit Saved",
                    f"Updated {len(fully_saved_keys)} combination(s) across {affected_lines} current-event line(s).\n\n"
                    f"Product Master rows updated: {master_result.get('updated', 0)}\n"
                    f"Exact color rows added: {master_result.get('added', 0)}\n"
                    f"New Product Master styles created: {master_result.get('created_styles', 0)}\n\n"
                    "Continue to Purchase Orders and regenerate the documents.",
                )
            else:
                self.audit_status_label.configure(text="Verification progress saved.", text_color=SUCCESS)
        except Exception as error:
            # No pending inputs are discarded on a top-level file or workbook
            # failure, and no verification is recorded.
            for key, value in final_values.items():
                if value.get("requires_save"):
                    self.decoration_audit_pending_edits[key] = {
                        "color": value["color"],
                        "location": value["location"],
                        "verified": False,
                    }
            messagebox.showerror("Unable to Save Decoration Audit", str(error))
    def show_decoration_audit_orders(self, record: dict):
        window = ctk.CTkToplevel(self)
        window.title("Affected Orders")
        window.geometry("900x560")
        window.transient(self)
        window.lift()
        window.focus_force()
        ctk.CTkLabel(
            window,
            text=f"{record.get('style') or record.get('product_name')} — {record.get('garment_color')} — {record.get('location')}",
            text_color=PURPLE_DARK, font=ctk.CTkFont(size=20, weight="bold"), anchor="w",
        ).pack(fill="x", padx=22, pady=(20, 4))
        ctk.CTkLabel(
            window,
            text=(
                f"{record.get('decoration_type')} | {record.get('quantity')} item(s) | "
                f"{record.get('order_count')} order(s)"
            ),
            text_color=MUTED, font=ctk.CTkFont(size=12), anchor="w",
        ).pack(fill="x", padx=22, pady=(0, 12))
        scroll = ctk.CTkScrollableFrame(window, fg_color=WHITE, border_width=1, border_color=BORDER)
        scroll.pack(fill="both", expand=True, padx=22, pady=(0, 14))
        headers = ("Order", "Employee", "Style", "Garment", "Size", "Qty", "Current Color")
        for column, header in enumerate(headers):
            scroll.grid_columnconfigure(column, weight=1)
            ctk.CTkLabel(scroll, text=header, text_color=PURPLE_DARK, font=ctk.CTkFont(size=10, weight="bold")).grid(
                row=0, column=column, sticky="w", padx=6, pady=7
            )
        for row_number, line in enumerate(record.get("line_records", []), start=1):
            values = (
                _clean(line.get("Order Number", "")), _clean(line.get("Employee Name", "")),
                _clean(line.get("Product #", "")), _clean(line.get("Garment Color", "")),
                _clean(line.get("Size", "")), _clean(line.get("Quantity", "")),
                _clean(line.get("Decoration Color", "")),
            )
            for column, value in enumerate(values):
                ctk.CTkLabel(scroll, text=value or "—", text_color=TEXT, font=ctk.CTkFont(size=10), anchor="w").grid(
                    row=row_number, column=column, sticky="ew", padx=6, pady=5
                )
        ctk.CTkButton(window, text="Close", command=window.destroy, width=110).pack(pady=(0, 18))


    def build_purchase_page(self):
        """Build the approved Purchase Orders dashboard layout without cat artwork."""
        page = self.new_page("purchase")
        # new_page reserves row 9 for dashboard-style pages.  Purchase Orders
        # owns row 1, so release that otherwise blank lower area to the PDF list.
        page.grid_rowconfigure(9, weight=0, minsize=0)
        page.grid_rowconfigure(1, weight=1)
        self._build_workflow_tracker(page, 0, "purchase_stage_labels", padx=34)

        workspace = ctk.CTkFrame(page, fg_color="transparent")
        workspace.grid(row=1, column=0, sticky="nsew", padx=28, pady=(2, 22))
        workspace.grid_columnconfigure(0, weight=7, uniform="purchase_workspace")
        workspace.grid_columnconfigure(1, weight=3, uniform="purchase_workspace")
        workspace.grid_rowconfigure(1, weight=1)
        self.purchase_workspace = workspace

        # Completion / readiness card.
        status_shadow = ctk.CTkFrame(workspace, fg_color=SHADOW, corner_radius=18, height=214)
        status_shadow.grid(row=0, column=0, sticky="ew", padx=(0, 12), pady=(6, 8))
        status_shadow.grid_propagate(False)
        status = ctk.CTkFrame(
            workspace, fg_color=WHITE, border_width=1, border_color=BORDER,
            corner_radius=18, height=208,
        )
        status.grid(row=0, column=0, sticky="ew", padx=(0, 12), pady=(1, 13))
        status.grid_propagate(False)
        status.grid_columnconfigure(1, weight=1)
        self.purchase_status_card = status
        self.purchase_status_icon = ctk.CTkLabel(
            status, text="5", width=66, height=66, corner_radius=33,
            fg_color=PURPLE_LIGHT, text_color=PURPLE,
            font=ctk.CTkFont(size=28, weight="bold"),
        )
        self.purchase_status_icon.grid(row=0, column=0, rowspan=2, padx=(24, 18), pady=(24, 10), sticky="n")
        self.purchase_state_title = ctk.CTkLabel(
            status, text="Generate Purchase Orders & Reports", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=22, weight="bold"), anchor="w", justify="left",
        )
        self.purchase_state_title.grid(row=0, column=1, padx=(0, 22), pady=(26, 3), sticky="ew")
        self.purchase_state_body = ctk.CTkLabel(
            status,
            text="Assign one PO number per vendor, create vendor PDFs, and generate the decoration reports.",
            text_color=MUTED, font=ctk.CTkFont(size=13), anchor="w", justify="left", wraplength=570,
        )
        self.purchase_state_body.grid(row=1, column=1, padx=(0, 22), pady=(0, 5), sticky="ew")
        self.purchase_folder_label = ctk.CTkLabel(
            status, text="", text_color=MUTED, font=ctk.CTkFont(size=10),
            anchor="w", justify="left", wraplength=690,
        )
        self.purchase_folder_label.grid(row=2, column=1, padx=(0, 22), pady=(0, 7), sticky="ew")
        divider = ctk.CTkFrame(status, height=1, fg_color="#E9E3EF", corner_radius=0)
        divider.grid(row=3, column=0, columnspan=2, sticky="ew", padx=24, pady=(2, 0))
        metrics = ctk.CTkFrame(status, fg_color="transparent")
        metrics.grid(row=4, column=0, columnspan=2, sticky="ew", padx=22, pady=(9, 14))
        metrics.grid_columnconfigure((0, 1, 2, 3), weight=1, uniform="po_metric")
        self.purchase_metric_labels = {}
        for column, (key, caption) in enumerate((
            ("vendors", "Vendors"),
            ("documents", "Purchase Orders"),
            ("lines", "Line Items"),
            ("status", "Workflow Status"),
        )):
            cell = ctk.CTkFrame(metrics, fg_color="transparent")
            cell.grid(row=0, column=column, sticky="nsew", padx=4)
            if column:
                ctk.CTkFrame(cell, width=1, fg_color="#E4DDEB", corner_radius=0).place(x=0, rely=0.12, relheight=0.76)
            value = ctk.CTkLabel(
                cell, text="—", text_color=PURPLE_DARK,
                font=ctk.CTkFont(size=19, weight="bold"),
            )
            value.pack(pady=(0, 1))
            ctk.CTkLabel(
                cell, text=caption, text_color=MUTED,
                font=ctk.CTkFont(size=10, weight="bold"),
            ).pack()
            self.purchase_metric_labels[key] = value

        # Purchase-order document list.
        list_shadow = ctk.CTkFrame(workspace, fg_color=SHADOW, corner_radius=18)
        list_shadow.grid(row=1, column=0, sticky="nsew", padx=(0, 12), pady=(6, 0))
        list_card = ctk.CTkFrame(
            workspace, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=18,
        )
        list_card.grid(row=1, column=0, sticky="nsew", padx=(0, 12), pady=(1, 5))
        list_card.grid_columnconfigure(0, weight=1)
        list_card.grid_rowconfigure(1, weight=1)
        list_header = ctk.CTkFrame(list_card, fg_color="transparent")
        list_header.grid(row=0, column=0, sticky="ew", padx=20, pady=(16, 7))
        list_header.grid_columnconfigure(1, weight=1)
        ctk.CTkLabel(
            list_header, text="Purchase Orders", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=18, weight="bold"), anchor="w",
        ).grid(row=0, column=0, sticky="w")
        self.purchase_document_count = ctk.CTkLabel(
            list_header, text="0", width=30, height=27, corner_radius=13,
            fg_color=PURPLE_LIGHT, text_color=PURPLE,
            font=ctk.CTkFont(size=12, weight="bold"),
        )
        self.purchase_document_count.grid(row=0, column=1, padx=(9, 8), sticky="w")
        ctk.CTkButton(
            list_header, text="Open Folder", width=118, height=34, corner_radius=9,
            fg_color="#F8F4FC", hover_color=PURPLE_LIGHT,
            border_width=1, border_color="#CDBBE5", text_color=PURPLE,
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.open_latest_purchase_orders,
        ).grid(row=0, column=2, padx=(5, 0), sticky="e")
        self.purchase_print_all_button = ctk.CTkButton(
            list_header, text="Print All Reports", width=132, height=34, corner_radius=9,
            fg_color=PURPLE, hover_color=PURPLE_DARK, text_color=WHITE,
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.print_all_purchase_reports,
        )
        self.purchase_print_all_button.grid(row=0, column=3, padx=(8, 0), sticky="e")
        self.purchase_documents_frame = ctk.CTkScrollableFrame(
            list_card, fg_color="#FCFBFE", border_width=0, corner_radius=13,
            scrollbar_button_color="#CDBBE5", scrollbar_button_hover_color=PURPLE,
        )
        self.purchase_documents_frame.grid(row=1, column=0, sticky="nsew", padx=13, pady=(0, 7))
        self.purchase_document_footer = ctk.CTkLabel(
            list_card, text="", text_color=MUTED,
            font=ctk.CTkFont(size=10), anchor="w",
        )
        self.purchase_document_footer.grid(row=2, column=0, sticky="ew", padx=20, pady=(0, 12))

        # Right-side action rail with a visible orchid watermark and no cats.
        actions = ctk.CTkFrame(workspace, fg_color="transparent")
        actions.grid(row=0, column=1, rowspan=2, sticky="nsew", padx=(10, 0), pady=(1, 5))
        actions.grid_columnconfigure(0, weight=1)
        actions.grid_rowconfigure(6, weight=1)
        self.purchase_actions = actions
        watermark = self._add_orchid_watermark(
            actions, size=(420, 420), relx=0.56, rely=0.78,
            anchor="center", opacity=0.17,
        )
        ctk.CTkLabel(
            actions, text="Purchase Order Actions", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=18, weight="bold"), anchor="w",
        ).grid(row=0, column=0, sticky="ew", padx=4, pady=(16, 10))
        self.purchase_action_primary = ctk.CTkButton(
            actions, text="Open Purchase Orders", command=self.open_latest_purchase_orders,
            height=50, corner_radius=11, fg_color=PURPLE, hover_color="#46109F",
            font=ctk.CTkFont(size=14, weight="bold"),
        )
        self.purchase_action_primary.grid(row=1, column=0, sticky="ew", padx=4, pady=6)
        self.purchase_action_regenerate = ctk.CTkButton(
            actions, text="Regenerate Purchase Orders", command=self.regenerate_current_purchase_orders,
            height=48, corner_radius=11, fg_color=WHITE, hover_color=PURPLE_LIGHT,
            border_width=1, border_color="#7C55D9", text_color=PURPLE,
            font=ctk.CTkFont(size=13, weight="bold"),
        )
        self.purchase_action_regenerate.grid(row=2, column=0, sticky="ew", padx=4, pady=6)
        self.purchase_action_audit = ctk.CTkButton(
            actions, text="Audit Decoration Colors", command=lambda: self.show_page("audit"),
            height=48, corner_radius=11, fg_color=WHITE, hover_color=PURPLE_LIGHT,
            border_width=1, border_color="#7C55D9", text_color=PURPLE,
            font=ctk.CTkFont(size=13, weight="bold"),
        )
        self.purchase_action_audit.grid(row=3, column=0, sticky="ew", padx=4, pady=6)
        self.purchase_action_edit = ctk.CTkButton(
            actions, text="Edit PO Numbers", command=lambda: self.open_po_number_editor(generate_after_save=False),
            height=48, corner_radius=11, fg_color=WHITE, hover_color=PURPLE_LIGHT,
            border_width=1, border_color="#7C55D9", text_color=PURPLE,
            font=ctk.CTkFont(size=13, weight="bold"),
        )
        self.purchase_action_edit.grid(row=4, column=0, sticky="ew", padx=4, pady=6)
        self.purchase_action_new = ctk.CTkButton(
            actions, text="Start New Event", command=self.start_new_purchase_packet,
            height=48, corner_radius=11, fg_color=WHITE, hover_color=PURPLE_LIGHT,
            border_width=1, border_color="#7C55D9", text_color=PURPLE,
            font=ctk.CTkFont(size=13, weight="bold"),
        )
        self.purchase_action_new.grid(row=5, column=0, sticky="ew", padx=4, pady=6)
        self.purchase_action_note = ctk.CTkLabel(
            actions,
            text="Audit thread and ink colors before generating or regenerating Purchase Orders.",
            text_color=MUTED, font=ctk.CTkFont(size=11), justify="left", wraplength=290,
        )
        self.purchase_action_note.grid(row=6, column=0, sticky="ew", padx=8, pady=(10, 0))

        # Compatibility aliases used by older workflow refresh code.
        self.purchase_primary_button = self.purchase_action_primary
        self.purchase_secondary_button = self.purchase_action_regenerate

    def refresh_purchase_page(self):
        if not hasattr(self, "purchase_state_title"):
            return
        review_exists = bool(self.last_review_workbook and self.last_review_workbook.exists())
        routing_status = self._purchase_review_routing_status() if review_exists else {"current": True, "reason": "", "changes": []}
        routing_stale = review_exists and not bool(routing_status.get("current", False))
        review_state = self._load_review_state()
        audit_requires_regeneration = bool(
            review_state.get("decoration_color_audit_requires_regeneration", False)
            and (routing_stale or review_state.get("additional_import_pending", False))
        )
        snapshot = self._fresh_snapshot() if review_exists else {}
        review_count = int(snapshot.get("review_count", 0) or 0)
        blocked_routes = int(snapshot.get("blocked_route_count", 0) or 0)
        routes = [route for route in snapshot.get("routes", []) if _clean(route.get("vendor", ""))]
        line_count = int(snapshot.get("line_count", self.imported_line_count or 0) or 0)
        audit_status = self._decoration_audit_status() if review_exists and not review_count and not blocked_routes else {}
        audit_remaining = int(audit_status.get("needs_attention", 0) or 0)
        release_dir, using_archive_fallback = self._purchase_order_release_location()
        generated = bool(
            release_dir
            and not audit_requires_regeneration
            and not routing_stale
            and audit_status.get("complete")
        )
        documents = self._purchase_order_document_paths(release_dir) if generated else []
        vendor_count = len(routes) or len(documents)

        self.purchase_metric_labels["vendors"].configure(text=f"{vendor_count:,}" if vendor_count else "—")
        self.purchase_metric_labels["documents"].configure(text=f"{len(documents):,}" if documents else "—")
        self.purchase_metric_labels["lines"].configure(text=f"{line_count:,}" if line_count else "—")
        self._render_purchase_order_documents(documents)
        self.purchase_action_note.configure(
            text="Audit thread and ink colors before generating or regenerating Purchase Orders."
        )

        # Reopening a completed archive creates a report-only working copy.
        # This check has to come before the generic completed-review recovery
        # route so the user cannot accidentally prepare purchase orders again.
        if review_exists and self._is_archived_report_reopen():
            self.purchase_status_icon.configure(text="↻", fg_color=PURPLE_LIGHT, text_color=PURPLE)
            self.purchase_state_title.configure(text="Archived Reports Ready")
            self.purchase_state_body.configure(
                text=(
                    "This completed event is open only to update its Outsourced Job / Logo Name, "
                    "upload a logo, and regenerate reports. The original archived event remains unchanged."
                )
            )
            self.purchase_folder_label.configure(text=self.last_review_workbook.name, text_color=MUTED)
            self.purchase_metric_labels["status"].configure(text="Report-only", text_color=PURPLE)
            self.purchase_action_primary.configure(
                text="Open Logo & Regenerate Archived Reports",
                command=self.open_archived_report_regeneration,
                state="normal",
            )
            self.purchase_action_regenerate.configure(
                text="Regenerate Archived Reports",
                command=self.regenerate_current_purchase_orders,
                state="normal",
            )
            self.purchase_action_audit.configure(state="disabled")
            self.purchase_action_edit.configure(text="Edit PO Numbers", state="disabled")
            self.purchase_action_new.configure(state="normal")
            self.purchase_action_note.configure(
                text="This archived event is report-only. Update the job name/logo, then regenerate the reports."
            )
        elif routing_stale:
            changed_count = len(routing_status.get("changes", []) or [])
            self.purchase_status_icon.configure(text="↻", fg_color="#FFF1DC", text_color=WARNING)
            self.purchase_state_title.configure(text="Purchase Review Refresh Needed")
            self.purchase_state_body.configure(
                text=(
                    "Product Master changed an active purchase route. Refresh the review before creating PDFs; "
                    "saved decisions will carry forward and Orchid will return you to this page.\n\n"
                    + _clean(routing_status.get("reason", ""))
                )
            )
            self.purchase_folder_label.configure(text=self.last_review_workbook.name, text_color=WARNING)
            self.purchase_metric_labels["status"].configure(
                text=f"Refresh {changed_count}" if changed_count else "Refresh",
                text_color=WARNING,
            )
            self.purchase_action_primary.configure(
                text="Refresh Routing & Return to POs",
                command=lambda: self.regenerate_current_review(return_to_purchase=True),
                state="normal",
            )
            self.purchase_action_regenerate.configure(
                text="Regenerate Purchase Review",
                command=lambda: self.regenerate_current_review(return_to_purchase=True),
                state="normal",
            )
            self.purchase_action_audit.configure(state="normal")
            self.purchase_action_edit.configure(text="Edit PO Numbers", state="disabled")
            self.purchase_action_new.configure(state="normal")
            self.purchase_action_note.configure(
                text="Only a current-event vendor, decoration type/location, or Never Outsource route requires this refresh."
            )
        elif generated:
            count = len(documents)
            self.purchase_status_icon.configure(text="✓", fg_color="#2ECC71", text_color=WHITE)
            self.purchase_state_title.configure(text="Purchase Orders Generated Successfully!")
            self.purchase_state_body.configure(
                text=(
                    f"{count} purchase-order document{'s have' if count != 1 else ' has'} been created and are ready to review. "
                    + "Open, regenerate, or update PO numbers using the actions on the right."
                )
            )
            self.purchase_folder_label.configure(text=str(release_dir), text_color=MUTED)
            self.purchase_metric_labels["status"].configure(text="Complete", text_color=SUCCESS)
            self.purchase_action_primary.configure(
                text="Open Purchase Orders", command=self.open_latest_purchase_orders, state="normal"
            )
            self.purchase_action_regenerate.configure(
                text="Regenerate Purchase Orders", command=self.regenerate_current_purchase_orders, state="normal"
            )
            self.purchase_action_audit.configure(state="normal")
            self.purchase_action_edit.configure(
                text="Edit PO Numbers", command=lambda: self.open_po_number_editor(generate_after_save=False), state="normal"
            )
            self.purchase_action_new.configure(state="normal")
        elif review_exists and not review_count and not blocked_routes and audit_remaining:
            self.purchase_status_icon.configure(text="!", fg_color="#FFF1DC", text_color=WARNING)
            self.purchase_state_title.configure(text="Decoration Color Audit Required")
            self.purchase_state_body.configure(
                text=(
                    f"{audit_remaining} thread or ink color combination{'s' if audit_remaining != 1 else ''} "
                    "still need verification before Purchase Orders can be generated."
                )
            )
            self.purchase_folder_label.configure(text=self.last_review_workbook.name, text_color=WARNING)
            self.purchase_metric_labels["status"].configure(text="Audit", text_color=WARNING)
            self.purchase_action_primary.configure(
                text="Open Decoration Color Audit", command=lambda: self.show_page("audit"), state="normal"
            )
            self.purchase_action_regenerate.configure(text="Regenerate Purchase Orders", state="disabled")
            self.purchase_action_audit.configure(state="normal")
            self.purchase_action_edit.configure(text="Edit PO Numbers", state="disabled")
            self.purchase_action_new.configure(state="normal")
        elif review_exists and self._review_state_is_authoritatively_complete(self.last_review_workbook):
            self.purchase_status_icon.configure(text="✓", fg_color="#EDE5F7", text_color=PURPLE)
            self.purchase_state_title.configure(text="Purchase Review Complete")
            self.purchase_state_body.configure(
                text="Orchid recorded zero decisions remaining. Prepare Purchase Orders from the recovery journal without repeating completed review decisions."
            )
            self.purchase_folder_label.configure(text=self.last_review_workbook.name, text_color=MUTED)
            self.purchase_metric_labels["status"].configure(text="Ready", text_color=PURPLE)
            self.purchase_action_primary.configure(
                text="Prepare & Assign PO Numbers", command=self._prepare_purchase_orders_from_journal, state="normal"
            )
            self.purchase_action_regenerate.configure(
                text="Reopen Unresolved Review", command=self.reopen_unresolved_purchase_review, state="normal"
            )
            self.purchase_action_audit.configure(state="normal")
            self.purchase_action_edit.configure(text="Edit PO Numbers", state="disabled")
            self.purchase_action_new.configure(state="normal")
        elif review_exists and not review_count and not blocked_routes:
            self.purchase_status_icon.configure(text="5", fg_color=PURPLE_LIGHT, text_color=PURPLE)
            self.purchase_state_title.configure(text="Ready to Generate Purchase Orders")
            self.purchase_state_body.configure(
                text="Purchase Review is complete. Enter one PO number for each vendor, then create the final vendor PDFs and decoration reports."
            )
            self.purchase_folder_label.configure(text=self.last_review_workbook.name, text_color=MUTED)
            self.purchase_metric_labels["status"].configure(text="Ready", text_color=PURPLE)
            self.purchase_action_primary.configure(
                text="Generate Purchase Orders", command=self.generate_latest_purchase_orders, state="normal"
            )
            self.purchase_action_regenerate.configure(
                text="Generate Purchase Orders", command=self.generate_latest_purchase_orders, state="normal"
            )
            self.purchase_action_audit.configure(state="normal")
            self.purchase_action_edit.configure(
                text="Edit PO Numbers", command=lambda: self.open_po_number_editor(generate_after_save=False), state="normal"
            )
            self.purchase_action_new.configure(state="normal")
        elif review_exists:
            self.purchase_status_icon.configure(text="!", fg_color="#FFF1DC", text_color=WARNING)
            self.purchase_state_title.configure(text="Purchase Review Required")
            self.purchase_state_body.configure(
                text="Step 5 is waiting for Product Master, Purchase Review, and Decoration Color Audit completion before final purchase orders can be generated."
            )
            self.purchase_folder_label.configure(
                text=self._blocker_summary(snapshot) or self.last_review_workbook.name,
                text_color=WARNING,
            )
            self.purchase_metric_labels["status"].configure(text="Blocked", text_color=WARNING)
            self.purchase_action_primary.configure(
                text="Return to Purchase Review", command=lambda: self.show_page("review"), state="normal"
            )
            self.purchase_action_regenerate.configure(text="Regenerate Purchase Orders", state="disabled")
            self.purchase_action_audit.configure(state="normal")
            self.purchase_action_edit.configure(text="Edit PO Numbers", state="disabled")
            self.purchase_action_new.configure(state="normal")
        else:
            self.purchase_status_icon.configure(text="5", fg_color=PURPLE_LIGHT, text_color=PURPLE)
            self.purchase_state_title.configure(text="Purchase Review Required")
            self.purchase_state_body.configure(
                text="Import and process the orders, complete Product Master setup, and finish Purchase Review before generating purchase orders."
            )
            self.purchase_folder_label.configure(text="", text_color=MUTED)
            self.purchase_metric_labels["status"].configure(text="Waiting", text_color=MUTED)
            self.purchase_action_primary.configure(
                text="Return to Current Event", command=lambda: self.show_page("import"), state="normal"
            )
            self.purchase_action_regenerate.configure(text="Regenerate Purchase Orders", state="disabled")
            self.purchase_action_audit.configure(state="disabled")
            self.purchase_action_edit.configure(text="Edit PO Numbers", state="disabled")
            self.purchase_action_new.configure(state="normal")

    def build_employee_totals_page(self):
        page = self.new_page("employees")
        # Employee Totals owns row 0; do not leave new_page's dashboard spacer
        # competing for the lower half of the window.
        page.grid_rowconfigure(9, weight=0, minsize=0)
        page.grid_rowconfigure(0, weight=1)

        employee_shadow = ctk.CTkFrame(page, fg_color=SHADOW, corner_radius=18, width=1, height=1)
        employee_shadow.grid(row=0, column=0, sticky="nsew", padx=30, pady=(25, 17))
        employee = ctk.CTkFrame(
            page, fg_color=WHITE, border_width=1,
            border_color=BORDER, corner_radius=18,
        )
        employee.grid(row=0, column=0, sticky="nsew", padx=30, pady=(19, 23))
        self.mc_employee_frame = employee
        employee.grid_columnconfigure(0, weight=1)
        employee.grid_rowconfigure(2, weight=1)

        ctk.CTkLabel(
            employee, text="Employee Totals", text_color=TEXT,
            font=ctk.CTkFont(size=24, weight="bold"), anchor="w",
        ).grid(row=0, column=0, padx=22, pady=(20, 3), sticky="w")
        self.mc_employee_summary = ctk.CTkLabel(
            employee,
            text="Create a Uniform Sizing Event Purchase Review to see employee totals.",
            text_color=MUTED, font=ctk.CTkFont(size=13), anchor="w",
            justify="left", wraplength=760,
        )
        self.mc_employee_summary.grid(row=1, column=0, padx=14, pady=(0, 6), sticky="w")

        employee_buttons = ctk.CTkFrame(employee, fg_color="transparent")
        employee_buttons.grid(row=0, column=1, rowspan=2, padx=18, pady=16, sticky="e")
        self.mc_employee_save_button = ctk.CTkButton(
            employee_buttons, text="Save Totals", width=132, height=40,
            fg_color=PURPLE, hover_color=PURPLE_DARK,
            font=ctk.CTkFont(size=13, weight="bold"),
            command=self.save_dashboard_employee_totals, state="disabled",
        )
        self.mc_employee_save_button.pack(side="left", padx=5)
        ctk.CTkButton(
            employee_buttons, text="View PDF", width=105, height=40,
            fg_color=PURPLE_LIGHT, hover_color="#E4D6F5", text_color=PURPLE_DARK,
            command=self.view_employee_totals_pdf,
        ).pack(side="left", padx=5)
        ctk.CTkButton(
            employee_buttons, text="Share / Email", width=125, height=40,
            fg_color=PURPLE_LIGHT, hover_color="#E4D6F5", text_color=PURPLE_DARK,
            command=self.share_employee_totals_pdf,
        ).pack(side="left", padx=5)
        ctk.CTkButton(
            employee_buttons, text="Open Purchase Review", width=158, height=40,
            fg_color=PURPLE_LIGHT, hover_color="#E4D6F5", text_color=PURPLE_DARK,
            command=self.open_latest_review,
        ).pack(side="left", padx=5)

        self.mc_employee_scroll = ctk.CTkScrollableFrame(
            employee, fg_color="#FCFBFE", corner_radius=10,
        )
        self.mc_employee_scroll.grid(row=2, column=0, columnspan=2, padx=18, pady=(0, 10), sticky="nsew")
        self.mc_employee_scroll.grid_columnconfigure(0, weight=1)
        register_scrollable(self, self.mc_employee_scroll)

        self.mc_employee_grand_total = ctk.CTkLabel(
            employee, text="Grand Total  $0.00", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=17, weight="bold"), anchor="e",
        )
        self.mc_employee_grand_total.grid(row=3, column=0, columnspan=2, padx=22, pady=(2, 18), sticky="e")

    def build_archive_page(self):
        page = self.new_page("archive")
        card = self.section_card(
            page, 0, "Completed Events",
            "Each completed event is stored with its order CSV, final Purchase Review workbook, job logo, and PDFs. Reopen an archived event when you need to add or replace a logo and regenerate reports."
        )
        card.grid_columnconfigure((0, 1), weight=1)
        self.primary_button(card, "Open Event Archive Folder", self.open_event_archive, padx=(22, 8), sticky="ew")
        self.secondary_button(
            card, "Reopen Archived Event for Reports", self.reopen_archived_event_for_reports,
            padx=(8, 22), sticky="ew",
        )

    def build_settings_page(self):
        page = self.new_page("settings")
        scroll = ctk.CTkScrollableFrame(page, fg_color=BG, corner_radius=0)
        self.settings_scroll = scroll
        scroll.grid(row=0, column=0, sticky="nsew")
        scroll.grid_columnconfigure(0, weight=1)
        page.grid_rowconfigure(0, weight=1)
        diagnostics = self.section_card(scroll, 0, "Routing Diagnostics", "Create a workbook explaining how Shopify lines matched Product Master and why information is missing.")
        self.primary_button(diagnostics, "Create Diagnostic Workbook", self.create_routing_diagnostics)

        recovery = self.section_card(
            scroll, 1, "Product Master Backup & Recovery",
            "Create a backup, restore a previous Product Master CSV, or return to the clean catalog included with Professional 4.3.8. The current file is always backed up before a restore."
        )
        recovery.grid_columnconfigure((0, 1, 2), weight=1)
        self.primary_button(recovery, "Back Up Now", self.backup_product_master_now, row=2, column=0, padx=(22, 8), sticky="ew")
        self.secondary_button(recovery, "Restore a Backup", self.restore_product_master_backup, row=2, column=1, padx=8, sticky="ew")
        clean_restore = ctk.CTkButton(
            recovery, text="Restore Clean v10 Master", command=self.restore_clean_product_master,
            height=42, fg_color="#FFF4E8", hover_color="#FBE4CC", border_width=1,
            border_color="#E8B77A", text_color="#8A4B00",
            font=ctk.CTkFont(size=13, weight="bold")
        )
        clean_restore.grid(row=2, column=2, padx=(8, 22), pady=(0, 20), sticky="ew")

        cleanup = self.section_card(scroll, 2, "Catalog Maintenance", "Back up Product Master and remove duplicate records without losing purchasing assignments.")
        self.primary_button(cleanup, "Clean & Backup", self.clean_catalog)

        export = self.section_card(
            scroll, 3, "Optional Excel Export",
            "The normal Purchase Review is completed inside Orchid Purchase Manager. Open the saved Excel workbook only when a detailed offline copy is useful."
        )
        self.primary_button(export, "Open Excel Purchase Review", self.open_latest_review)

        optional = self.section_card(
            scroll, 4, "Safe Product Candidate Preview",
            "Review new product/style numbers found in the selected order export. This separate audit-reconciliation recovery test does not mass-add order lines to Product Master, so a large import cannot create hundreds of incomplete permanent records."
        )
        self.primary_button(optional, "Create Candidate List", self.process_csv)

        legacy = legacy_product_master_path()
        paths = self.section_card(
            scroll, 5, "Storage Locations",
            f"Active 2.x Product Master file: {live_product_master_path()}\nReports: {REPORTS}\nv12 data (left untouched): {v12_data_dir()}\nOlder data (left untouched): {legacy_data_dir()}"
        )
        self.primary_button(paths, "Open 2.x Data Folder", lambda: subprocess.run(["open", str(DATA)], check=False), row=2)
        if legacy.parent.exists():
            self.secondary_button(paths, "Open Older Data Folder", lambda: subprocess.run(["open", str(legacy.parent)], check=False), row=2)

        motion = self.section_card(
            scroll, 6, "Interface Motion",
            "Use the subtle purple radial pulse around the official Orchid flower when navigating or completing actions. Reduced-motion system preferences are always respected."
        )
        self.interaction_animation_var = ctk.BooleanVar(value=bool(self.interaction_animations_enabled))
        ctk.CTkSwitch(
            motion, text="Enable Orchid interaction animation", variable=self.interaction_animation_var,
            command=self._set_interaction_animations_enabled, progress_color=PURPLE,
            font=ctk.CTkFont(size=13, weight="bold"), text_color=TEXT,
        ).grid(row=2, column=0, padx=22, pady=(0, 20), sticky="w")

        about = self.section_card(
            scroll, 7, "About Orchid Purchase Manager",
            "Professional 4.9.104\nAdds one-click Needs Setup navigation in Product Master and renames the Ship to Orchid purchase order."
        )
        ctk.CTkLabel(
            about, text="Version 4.9.104", text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=18, weight="bold"), anchor="w",
        ).grid(row=2, column=0, padx=22, pady=(0, 20), sticky="w")

    def refresh_current_event_page(self, snapshot: dict | None = None):
        if not hasattr(self, "current_event_title"):
            return
        snapshot = snapshot or self.last_snapshot or {}
        self._refresh_current_event_progress(snapshot)
        if self.selected_csv or (self.last_review_workbook and self.last_review_workbook.exists()):
            self._show_current_event_active_layout()
            event = self.current_event_name or snapshot.get("event_name") or "Active Purchase Packet"
            event_length = len(event)
            event_font_size = 33 if event_length <= 24 else 29 if event_length <= 38 else 25
            self.current_event_title.configure(
                text=event, font=ctk.CTkFont(size=event_font_size, weight="bold"),
                wraplength=470,
            )
            self._refresh_current_event_job_logo()
            self.current_event_summary_title.configure(text=event)
            self.dashboard_summary_title.configure(text=event)
            self.header_chip.configure(text=event if len(event) <= 34 else event[:31] + "...")
            self.current_event_body.configure(text="Manage the active order import, add late orders, and continue its Purchase Review from one place.")
            lines = int(snapshot.get("line_count", self.imported_line_count or 0) or 0)
            vendor_count = int(snapshot.get("vendor_count", 0) or 0)
            # Keep the sidebar event-summary area date-only in every workflow state.
            self.shell_event_meta.configure(text=time.strftime("%b %d, %Y"))
            details = [f"CSV: {self.selected_csv.name}" if self.selected_csv else f"Review: {self.last_review_workbook.name}"]
            if lines:
                details.append(f"{lines:,} purchase lines")
            if self.selected_csv and self.last_review_workbook and self.last_review_workbook.exists():
                details.append(f"Review: {self.last_review_workbook.name}")
            self.import_file_label.configure(text="  •  ".join(details[:2]), fg_color="#EAF7EE", text_color="#16803B")
            self.current_event_details.configure(text="\n".join(details[2:]))
            self.current_event_primary.grid_configure(row=0, column=0, columnspan=1, sticky="ew", padx=6, pady=6)
            for button in (
                self.current_event_primary,
                self.current_event_add,
                self.current_event_replace,
                self.current_event_regenerate,
                self.current_event_regenerate_po,
                self.current_event_edit_po,
                self.current_event_master,
            ):
                button.configure(
                    fg_color=PURPLE, hover_color=PURPLE_DARK, text_color=WHITE,
                    border_width=0, corner_radius=10, height=40,
                )
            self.current_event_primary.configure(font=ctk.CTkFont(size=12, weight="bold"))
            if self.last_review_workbook and self.last_review_workbook.exists():
                product_setup_count = sum(
                    1 for issue in snapshot.get("issues", [])
                    if _clean(issue.get("fix_in", issue.get("source", ""))).casefold() == "product master"
                )
                if product_setup_count:
                    self.current_event_primary.configure(
                        text="Open Product Master", command=lambda: self.show_page("master")
                    )
                elif self._review_is_stale():
                    self.current_event_primary.configure(
                        text="Update Purchase Review", command=self.regenerate_current_review
                    )
                else:
                    self.current_event_primary.configure(
                        text="Open Purchase Review", command=lambda: self.show_page("review")
                    )
            else:
                self.current_event_primary.configure(
                    text="Process Orders", command=self.create_review_workbook
                )
            self.current_event_add.grid(row=0, column=1, sticky="ew", padx=6, pady=6)
            self.current_event_replace.grid(row=1, column=0, columnspan=2, sticky="ew", padx=5, pady=5)
            self.current_event_discard.grid(row=4, column=0, columnspan=2, sticky="ew", padx=5, pady=(6, 0))
            if self.last_review_workbook and self.last_review_workbook.exists():
                self.current_event_regenerate.configure(state="normal")
                self.current_event_regenerate.grid(row=2, column=0, columnspan=2, sticky="ew", padx=5, pady=5)
            else:
                self.current_event_regenerate.grid_remove()
            if self.last_purchase_order_dir and self.last_purchase_order_dir.exists():
                self.current_event_regenerate_po.configure(state="normal")
                self.current_event_regenerate_po.grid(row=3, column=0, sticky="ew", padx=5, pady=5)
                self.current_event_edit_po.configure(state="normal")
                self.current_event_edit_po.grid(row=3, column=1, sticky="ew", padx=5, pady=5)
            else:
                self.current_event_regenerate_po.grid_remove()
                self.current_event_edit_po.grid_remove()
        else:
            self._show_current_event_empty_layout()
            self.current_event_job_logo.configure(image=None, text="")
            self.current_event_job_logo.grid_remove()
            self.current_event_title.configure(text="No Active Purchase Packet")
            self.current_event_summary_title.configure(text="No Active Purchase Packet")
            self.dashboard_summary_title.configure(text="No Active Purchase Packet")
            self.header_chip.configure(text="No Active Purchase Packet")
            self.current_event_body.configure(text="Import the Shopify or Report Toaster CSV for the orders you want to process. Product Master and saved settings are preserved.")
            self.shell_event_meta.configure(text=time.strftime("%b %d, %Y"))
            self.import_file_label.configure(text="No order CSV selected", fg_color="#F3F1F6", text_color=TEXT)
            self.current_event_details.configure(text="")
            self.current_event_primary.grid_configure(row=0, column=0, columnspan=2, sticky="", padx=6, pady=6)
            self.current_event_primary.configure(
                text="Import Shopify CSV", command=self.choose_csv,
                font=ctk.CTkFont(size=15, weight="bold"), height=60,
                fg_color=PURPLE, hover_color=PURPLE_DARK, text_color=WHITE,
            )
            self.current_event_add.grid_remove()
            self.current_event_replace.grid_remove()
            self.current_event_regenerate.grid_remove()
            self.current_event_regenerate_po.grid_remove()
            self.current_event_edit_po.grid_remove()
            self.current_event_discard.grid_remove()

    @staticmethod
    def _preview_order_numbers(values, limit: int = 8) -> str:
        values = [str(value).strip() for value in values if str(value).strip()]
        if not values:
            return "None"
        shown = values[:limit]
        suffix = f"\n…and {len(values) - limit} more" if len(values) > limit else ""
        return ", ".join(shown) + suffix

    def _write_review_state_atomic(self, state: dict) -> None:
        REVIEW_STATE_FILE.parent.mkdir(parents=True, exist_ok=True)
        temp = REVIEW_STATE_FILE.with_name(
            f".{REVIEW_STATE_FILE.name}.{os.getpid()}.{int(time.time() * 1000)}.tmp"
        )
        try:
            temp.write_text(json.dumps(state, indent=2), encoding="utf-8")
            with temp.open("rb") as handle:
                os.fsync(handle.fileno())
            os.replace(temp, REVIEW_STATE_FILE)
        finally:
            try:
                temp.unlink(missing_ok=True)
            except Exception:
                pass

    def add_orders_to_current_event(self):
        """Append new orders to the active packet without replacing prior work."""
        if not self.selected_csv or not Path(self.selected_csv).exists():
            messagebox.showinfo(
                "Import Orders First",
                "Start the event with its first Shopify or Report Toaster CSV before adding another order file.",
            )
            return

        downloads = Path.home() / "Downloads"
        path = filedialog.askopenfilename(
            title="Add Orders to Current Event (.CSV)",
            initialdir=str(downloads if downloads.exists() else Path.home()),
            filetypes=[("Order CSV files", "*.csv")],
        )
        if not path:
            return
        incoming_path = Path(path)
        try:
            current_normalized = normalize_order_export(Path(self.selected_csv))
            incoming_normalized = normalize_order_export(incoming_path)
            preview = preview_additional_orders(current_normalized, incoming_normalized)
        except Exception as error:
            messagebox.showerror(
                "Unable to Preview Additional Orders",
                f"{error}\n\nThe current event was not changed.",
            )
            return

        if not preview.new_order_count and not preview.changed_order_count:
            messagebox.showinfo(
                "No New Orders Found",
                f"Every order in {incoming_path.name} is already in the current event.\n\n"
                f"Duplicate orders skipped: {preview.duplicate_order_count}\n\n"
                "Nothing was added and the current event was not changed.",
            )
            return

        preview_text = (
            f"Current event orders: {preview.current_order_count}\n"
            f"Orders in selected CSV: {preview.incoming_order_count}\n\n"
            f"New orders to add: {preview.new_order_count}\n"
            f"Duplicate orders to skip: {preview.duplicate_order_count}\n"
            f"Existing orders that appear changed: {preview.changed_order_count}\n\n"
            f"New order numbers:\n{self._preview_order_numbers(preview.new_orders)}"
        )
        if preview.changed_order_count:
            preview_text += (
                f"\n\nChanged order numbers:\n"
                f"{self._preview_order_numbers(preview.changed_orders)}"
            )
        if not messagebox.askyesno(
            "Preview Additional Orders",
            preview_text + "\n\nContinue with this import?",
            default=messagebox.YES,
        ):
            return

        replace_changed = False
        if preview.changed_order_count:
            changed_choice = messagebox.askyesnocancel(
                "Changed Existing Orders Found",
                "Some Shopify order numbers already exist, but their line items or quantities are different.\n\n"
                "Yes — Replace those existing orders with the newer CSV versions.\n"
                "No — Keep the existing versions and add only brand-new order numbers.\n"
                "Cancel — Leave the current event unchanged.",
                default=messagebox.NO,
            )
            if changed_choice is None:
                return
            replace_changed = bool(changed_choice)

        event_stem = safe_event_stem(self.current_event_name or "Current Event")
        combined_path = DATA / "Current Event Imports" / f"{event_stem}_combined_orders.csv"
        batch_label = time.strftime("%Y-%m-%d %H:%M:%S")
        try:
            merged, preview, details = merge_additional_orders(
                current_normalized,
                incoming_normalized,
                incoming_source_name=incoming_path.name,
                current_source_name=Path(self.selected_csv).name,
                replace_changed_orders=replace_changed,
                batch_label=batch_label,
            )
            if not details.get("lines_added"):
                messagebox.showinfo(
                    "No Orders Added",
                    "No brand-new orders were available to add after applying your changed-order choice. "
                    "The current event was not changed.",
                )
                return
            write_combined_csv_atomic(merged, combined_path)
            parsed = parse_shopify_orders(
                combined_path, live_product_master_path(), normalized_orders=merged
            )
            self._remember_import_frames(combined_path, merged, parsed)
        except Exception as error:
            messagebox.showerror(
                "Unable to Add Orders",
                f"{error}\n\nThe existing event and its original CSV remain unchanged.",
            )
            return

        previous_source = str(self.selected_csv)
        self.selected_csv = combined_path
        self.imported_line_count = len(parsed)
        self.last_purchase_order_dir = None

        state = self._load_review_state()
        batches = state.get("import_batches", [])
        if not isinstance(batches, list):
            batches = []
        if not batches:
            batches.append({
                "source_csv": previous_source,
                "file_name": Path(previous_source).name,
                "kind": "original",
            })
        batches.append({
            "source_csv": str(incoming_path),
            "file_name": incoming_path.name,
            "kind": "additional",
            "added_at": batch_label,
            **details,
        })
        state.update({
            "source_csv": str(combined_path),
            "source_signature": self._source_file_signature(combined_path),
            "import_batches": batches,
            "additional_import_pending": bool(self.last_review_workbook and self.last_review_workbook.exists()),
            "last_additional_import": details,
            "review_complete": False,
            "decoration_color_audit_requires_regeneration": True,
        })
        state.pop("review_completed_at", None)
        try:
            self._write_review_state_atomic(state)
        except Exception as error:
            previous_path = Path(previous_source)
            backup_path = combined_path.with_suffix(combined_path.suffix + ".backup")
            try:
                if previous_path.resolve() == combined_path.resolve() and backup_path.exists():
                    shutil.copy2(backup_path, combined_path)
            except Exception:
                pass
            self.selected_csv = previous_path
            try:
                restored_normalized = normalize_order_export(previous_path)
                restored_parsed = parse_shopify_orders(
                    previous_path, live_product_master_path(), normalized_orders=restored_normalized
                )
                self._remember_import_frames(previous_path, restored_normalized, restored_parsed)
                self.imported_line_count = len(restored_parsed)
            except Exception:
                pass
            messagebox.showerror(
                "Unable to Save Additional Import",
                f"The combined CSV was created, but Orchid could not safely update the active-event state:\n{error}\n\n"
                "Orchid restored the previous active source whenever a combined-file backup was available.",
            )
            return

        if self.last_review_workbook and self.last_review_workbook.exists():
            self._regenerate_after_additional_import(merged, parsed, details, incoming_path.name)
            return

        self.dashboard_status.configure(
            text=(
                f"Added {details['new_orders_added']} new order(s) from {incoming_path.name}; "
                f"skipped {details['duplicate_orders_skipped']} duplicate(s)."
            ),
            text_color=SUCCESS,
        )
        self.refresh_current_event_page({"line_count": self.imported_line_count})
        self.refresh_dashboard()
        messagebox.showinfo(
            "Orders Added to Current Event",
            f"New orders added: {details['new_orders_added']}\n"
            f"Duplicate orders skipped: {details['duplicate_orders_skipped']}\n"
            f"Changed orders replaced: {details['changed_orders_replaced']}\n"
            f"Purchase lines now in event: {details['final_line_count']}\n\n"
            "Process Orders when you are ready to create Purchase Review.",
        )

    def _regenerate_after_additional_import(self, normalized_orders, parsed_orders, details: dict, source_name: str):
        """Rebuild the active review in the background while preserving prior decisions."""
        from modules.review_workbook import generate_review_workbook

        previous_workbook = Path(self.last_review_workbook)
        state = self._load_review_state()
        dialog, status_label, progress = self._show_processing_window(
            "Adding Orders to Current Event",
            "Rebuilding Purchase Review and preserving completed decisions…",
        )
        result_queue = queue.Queue()

        def worker():
            try:
                result_queue.put(("status", "Refreshing known Shopify catalog details without changing Orchid purchasing rules…"))
                catalog_merge = sync_shopify_catalog_enrichment(parsed_orders, live_product_master_path())
                if catalog_merge.get("changed"):
                    try:
                        self.product_master_update_marker.parent.mkdir(parents=True, exist_ok=True)
                        self.product_master_update_marker.write_text(str(time.time()), encoding="utf-8")
                    except Exception:
                        pass
                result_queue.put(("status", "Applying existing Product Master rules to the new orders…"))
                result = generate_review_workbook(
                    Path(self.selected_csv), live_product_master_path(), REPORTS,
                    report_mode=state.get("report_mode") or self.current_mode or GENERAL_SALES_PERIOD,
                    event_name=state.get("event_name") or self.current_event_name or "",
                    active_packet_id=_clean(state.get("active_packet_id", "")),
                    decoration_fulfillment=(
                        state.get("decoration_fulfillment")
                        or self.current_decoration_fulfillment
                        or STANDARD_ORCHID_WORKFLOW
                    ),
                    regenerate_helper=REGENERATE_HELPER,
                    previous_review_path=previous_workbook,
                    normalized_orders=normalized_orders,
                    parsed_orders=parsed_orders,
                )
                first_workbook = Path(result["output_path"])
                result_queue.put(("status", "Checking only the newly introduced product styles…"))
                candidate_sync = sync_review_product_candidates(first_workbook, live_product_master_path())
                if candidate_sync.get("changed"):
                    result_queue.put(("status", "Applying new Product Master setup candidates once…"))
                    result = generate_review_workbook(
                        Path(self.selected_csv), live_product_master_path(), REPORTS,
                        report_mode=state.get("report_mode") or self.current_mode or GENERAL_SALES_PERIOD,
                        event_name=state.get("event_name") or self.current_event_name or "",
                        active_packet_id=_clean(state.get("active_packet_id", "")),
                        decoration_fulfillment=(
                            state.get("decoration_fulfillment")
                            or self.current_decoration_fulfillment
                            or STANDARD_ORCHID_WORKFLOW
                        ),
                        regenerate_helper=REGENERATE_HELPER,
                        previous_review_path=first_workbook,
                        normalized_orders=normalized_orders,
                        parsed_orders=parsed_orders,
                    )
                    final_workbook = Path(result["output_path"])
                    try:
                        if first_workbook.exists() and first_workbook != final_workbook:
                            first_workbook.unlink()
                    except Exception:
                        pass
                result_queue.put(("success", result, candidate_sync, catalog_merge))
            except Exception as error:
                result_queue.put(("error", str(error), traceback.format_exc()))

        def close_dialog():
            try:
                progress.stop()
            except Exception:
                pass
            try:
                dialog.grab_release()
            except Exception:
                pass
            try:
                dialog.destroy()
            except Exception:
                pass

        def finish_success(result, candidate_sync, catalog_merge):
            self.last_review_workbook = Path(result["output_path"])
            self._candidate_sync_workbook = str(self.last_review_workbook.resolve())
            self._reset_review_navigation()
            self.current_event_name = result.get("event_name") or self.current_event_name or "General Sales Period"
            self.current_mode = result.get("report_mode") or self.current_mode
            self.current_decoration_fulfillment = (
                result.get("decoration_fulfillment") or self.current_decoration_fulfillment
            )
            self.imported_line_count = int(result.get("lines", len(parsed_orders)))
            self.review_needed = int(result.get("review_decisions", result.get("review_lines", 0)))
            updated_state = self._load_review_state()
            updated_state["additional_import_pending"] = False
            updated_state["last_additional_import_review"] = str(self.last_review_workbook)
            self._write_review_state_atomic(updated_state)
            try:
                snapshot = load_purchase_review_snapshot(self.last_review_workbook, live_product_master_path())
            except Exception:
                snapshot = {"issues": []}
            snapshot.update({
                "event_name": self.current_event_name,
                "report_mode": self.current_mode,
                "line_count": self.imported_line_count,
                "review_count": self.review_needed or 0,
                "purchase_review_count": int(result.get("purchase_review_decisions", self.review_needed or 0) or 0),
                "product_master_review_count": int(result.get("product_master_decisions", 0) or 0),
            })
            self._cache_fast_review_snapshot(snapshot)
            self.last_snapshot = snapshot
            close_dialog()
            self.dashboard_status.configure(
                text=(
                    f"Added {details['new_orders_added']} new order(s) from {source_name}; "
                    f"{self.review_needed or 0} decision(s) now require review."
                ),
                text_color=WARNING if self.review_needed else SUCCESS,
            )
            self.refresh_current_event_page(snapshot)
            self.refresh_dashboard()
            self._warm_dashboard_snapshot_async(self.last_review_workbook)
            messagebox.showinfo(
                "Orders Added Successfully",
                f"New orders added: {details['new_orders_added']}\n"
                f"Duplicate orders skipped: {details['duplicate_orders_skipped']}\n"
                f"Changed orders replaced: {details['changed_orders_replaced']}\n"
                f"Purchase lines now in event: {details['final_line_count']}\n"
                f"Decisions requiring review: {self.review_needed or 0}\n"
                + (f"Shopify descriptions refreshed: {catalog_merge.get('updated_descriptions', 0)}\n" if catalog_merge.get('updated_descriptions') else "")
                + (f"New Shopify garment colors added for setup: {catalog_merge.get('added_colors', 0)}\n" if catalog_merge.get('added_colors') else "")
                + "\nExisting Product Master setup and completed Purchase Review decisions were preserved. "
                "Previously generated purchase orders must be regenerated.",
                parent=self,
            )
            self.show_page("review" if self.review_needed else "import")

        def finish_error(error, detail):
            close_dialog()
            print(detail, file=sys.stderr)
            self.dashboard_status.configure(
                text="Orders were safely added, but Purchase Review could not be refreshed.",
                text_color=DANGER,
            )
            messagebox.showerror(
                "Orders Added — Review Refresh Needed",
                f"The combined event CSV was saved safely, but Purchase Review could not be regenerated:\n\n{error}\n\n"
                "Use Regenerate Purchase Review to try again. Your previous workbook remains in Reports.",
                parent=self,
            )
            self.refresh_current_event_page(self.last_snapshot or {})
            self.refresh_dashboard()

        def poll_worker():
            try:
                while True:
                    item = result_queue.get_nowait()
                    if item[0] == "status":
                        if status_label.winfo_exists():
                            status_label.configure(text=item[1])
                    elif item[0] == "success":
                        finish_success(item[1], item[2], item[3])
                        return
                    elif item[0] == "error":
                        finish_error(RuntimeError(item[1]), item[2])
                        return
            except queue.Empty:
                pass
            if dialog.winfo_exists():
                self.after(100, poll_worker)

        threading.Thread(target=worker, daemon=True, name="orchid-add-orders-worker").start()
        self.after(100, poll_worker)

    def _clear_active_packet_state(self):
        try:
            ACTIVE_PACKET_RESET_FILE.parent.mkdir(parents=True, exist_ok=True)
            ACTIVE_PACKET_RESET_FILE.write_text(json.dumps({"reset_time": time.time()}, indent=2), encoding="utf-8")
            if REVIEW_STATE_FILE.exists():
                REVIEW_STATE_FILE.unlink()
        except Exception:
            pass
        self.selected_csv = None
        self.current_event_name = ""
        self.current_mode = ""
        self.last_review_workbook = None
        self.last_purchase_order_dir = None
        self.imported_line_count = 0
        self.review_needed = None
        self.last_snapshot = {}

    def discard_current_event(self):
        if not (self.selected_csv or self.last_review_workbook or self.last_purchase_order_dir):
            messagebox.showinfo("No Current Event", "There is no active purchase packet to discard.")
            return
        if not messagebox.askyesno(
            "Discard Current Event",
            "Discard the active purchase packet?\n\nThis removes the imported CSV from Orchid's current workflow and clears its active review and purchase-order references. It does not delete the original CSV, Product Master, settings, or completed archived events.",
        ):
            return
        self._suppress_empty_import_prompt = True
        self._clear_active_packet_state()
        self.refresh_current_event_page({})
        self.refresh_dashboard()
        self.show_page("import")
        self.after_idle(lambda: (self.refresh_current_event_page({}), self.refresh_dashboard()))

    def finish_and_clear_current_event(self):
        """Close the active packet and leave its PDFs outside Orchid's workflow."""
        if not (self.selected_csv or self.last_review_workbook or self.last_purchase_order_dir):
            messagebox.showinfo("No Current Event", "There is no active purchase packet to finish.")
            return
        body = (
            "Before clearing, save any PDFs you want to keep in your normal computer folder.\n\n"
            "This removes the event's CSV, Purchase Review, totals, and purchase-order references from Orchid. "
            "Product Master remains saved. Orchid will not reopen this event or use its reports in a later event."
        )
        if not messagebox.askyesno("Finish & Clear Event", body, default=messagebox.NO):
            return
        self._suppress_empty_import_prompt = True
        self._clear_active_packet_state()
        self.refresh_current_event_page({})
        self.refresh_dashboard()
        self.show_page("import")
        self.after_idle(lambda: (self.refresh_current_event_page({}), self.refresh_dashboard()))

    def start_new_purchase_packet(self):
        if not messagebox.askyesno(
            "Start New Event",
            "Start a new purchase packet?\n\nThe active event will be cleared from Orchid. Product Master stays saved; already-created PDFs remain on your computer but cannot be reopened in Orchid.",
        ):
            return
        self._suppress_empty_import_prompt = True
        self._clear_active_packet_state()
        self.refresh_current_event_page({})
        self.refresh_dashboard()
        self.show_page("import")
        self.after_idle(lambda: (self.refresh_current_event_page({}), self.refresh_dashboard()))

    def open_latest_purchase_orders(self):
        report_dir, _unused_archive_flag = self._purchase_order_release_location()
        if report_dir:
            try:
                opened = subprocess.run(["open", str(report_dir)], check=False)
            except OSError:
                opened = None
            if not opened or opened.returncode != 0:
                messagebox.showwarning(
                    "Unable to Open Purchase Orders",
                    "The current purchase-order folder is ready, but Finder could not be opened automatically:\n"
                    f"{report_dir}",
                )
        else:
            messagebox.showwarning(
                "Purchase Orders Need Attention",
                "Orchid could not verify purchase orders for the active event. The existing PDFs were not changed. Do not regenerate; check the current packet first.",
            )

    # ---------- actions ----------
    def _show_create_review_dashboard(self) -> None:
        """Return to and fully repaint Step 1 after the macOS file picker closes."""
        if self.selected_csv and not self.last_review_workbook:
            self.refresh_dashboard()
            self.show_page("import")
            try:
                self.update_idletasks()
                self.update()
                self.lift()
                self.focus_force()
            except Exception:
                pass

    def choose_csv(self):
        downloads = Path.home() / "Downloads"
        path = filedialog.askopenfilename(
            title="Import Shopify or Report Toaster Orders (.CSV)",
            initialdir=str(downloads if downloads.exists() else Path.home()),
            filetypes=[("Order CSV files", "*.csv")],
        )
        if not path:
            return
        candidate = Path(path)
        try:
            normalized = normalize_order_export(candidate)
            parsed = parse_shopify_orders(
                candidate, live_product_master_path(), normalized_orders=normalized
            )
            self._remember_import_frames(candidate, normalized, parsed)
        except Exception as error:
            messagebox.showerror(
                "Unable to Import Order CSV",
                f"{error}\n\nNo event was created. Update the export columns and try again.",
            )
            return

        candidate_signature = self._source_file_signature(candidate)
        state = self._load_review_state()
        active_review = bool(self.last_review_workbook and self.last_review_workbook.exists())
        saved_source = Path(state.get("source_csv", "")).expanduser() if state.get("source_csv") else None
        saved_signature = _clean(state.get("source_signature", ""))
        if not saved_signature and saved_source and saved_source.exists():
            saved_signature = self._source_file_signature(saved_source)
        same_source = bool(
            active_review and (
                (candidate_signature and saved_signature and candidate_signature == saved_signature)
                or (saved_source and saved_source.exists() and candidate.resolve() == saved_source.resolve())
            )
        )
        if active_review and same_source:
            resume = messagebox.askyesnocancel(
                "Existing Purchase Review Found",
                "This CSV matches the active purchase packet.\n\n"
                "Yes — Resume Existing Purchase Review and keep all saved corrections.\n"
                "No — Replace it and start the Purchase Review over.\n"
                "Cancel — Leave the current packet unchanged.",
                default=messagebox.YES,
            )
            if resume is None:
                return
            if resume:
                self.selected_csv = candidate
                self.imported_line_count = len(parsed)
                self._update_review_state(source_csv=str(candidate), source_signature=candidate_signature)
                self.dashboard_status.configure(
                    text=f"Resumed existing Purchase Review for {candidate.name}. Saved corrections were preserved.",
                    text_color=SUCCESS,
                )
                self.refresh_dashboard()
                self.show_page("review")
                return
            self._suppress_empty_import_prompt = True
            self._clear_active_packet_state()
        elif active_review:
            replace = messagebox.askyesno(
                "Replace Active Purchase Packet?",
                "A different Purchase Review is already active. Replace it with the selected CSV?\n\n"
                "The existing workbook will remain in Reports, but it will no longer be the active packet.",
                default=messagebox.NO,
            )
            if not replace:
                return
            self._suppress_empty_import_prompt = True
            self._clear_active_packet_state()

        self.selected_csv = candidate
        self.current_event_name = ""
        self.current_mode = ""
        self.last_review_workbook = None
        self.last_purchase_order_dir = None
        self.review_needed = None
        self.imported_line_count = len(parsed)
        source_name = str(parsed.attrs.get("import_source", "Order CSV") or "Order CSV")
        removed = int(parsed.attrs.get("removed_rows_excluded", 0) or 0)
        adjusted = int(parsed.attrs.get("quantity_adjustments_applied", 0) or 0)
        status = f"Selected {self.selected_csv.name} • {source_name} • {self.imported_line_count:,} purchase lines"
        if removed:
            status += f" • {removed:,} removed item row(s) excluded"
        if adjusted:
            status += f" • {adjusted:,} edited quantity row(s) adjusted"
        self.dashboard_status.configure(text=status, text_color=SUCCESS)
        self.refresh_dashboard()
        self._show_create_review_dashboard()
        self.after_idle(self._show_create_review_dashboard)
        self.after(100, self._show_create_review_dashboard)
        self.after(260, self._show_create_review_dashboard)

    def process_csv(self):
        if not self.selected_csv:
            self.choose_csv()
        if not self.selected_csv:
            return
        try:
            normalized, parsed = self._cached_import_frames(self.selected_csv)
            if parsed is None:
                normalized = normalize_order_export(self.selected_csv)
                parsed = parse_shopify_orders(
                    self.selected_csv, live_product_master_path(), normalized_orders=normalized
                )
                self._remember_import_frames(self.selected_csv, normalized, parsed)
            result = save_product_candidate_list(live_product_master_path(), parsed, REPORTS)
            self.imported_line_count = len(parsed)
            self.dashboard_status.configure(
                text=f"Candidate preview created with {result['candidate_count']} genuinely new product/style number(s).",
                text_color=SUCCESS,
            )
            messagebox.showinfo(
                "Product Candidate List Created",
                f"Shopify product lines reviewed: {len(parsed)}\n"
                f"New product/style candidates: {result['candidate_count']}\n\n"
                "No products were added to Product Master. This prevents a large order import from creating hundreds of incomplete permanent records.\n\n"
                f"Saved as:\n{result['output_path']}",
            )
            self.refresh_dashboard()
        except Exception as error:
            messagebox.showerror("Unable to Create Candidate List", str(error))

    @staticmethod
    def _pid_is_running(pid: int) -> bool:
        return process_is_running(pid)

    @staticmethod
    def _pid_is_product_master(pid: int) -> bool:
        """Return True only when ``pid`` belongs to Orchid's Product Master.

        A PID file survives a forced quit or an app replacement.  macOS can
        reuse that numeric PID for an unrelated process, so ``os.kill(pid, 0)``
        alone is not enough evidence that Product Master is still open.  The
        editor is always launched with ``--product-master`` in its command
        line, which gives us a safe and specific identity check.
        """
        def command_reader(candidate: int) -> str:
            completed = subprocess.run(
                ["/bin/ps", "-p", str(candidate), "-o", "command="],
                check=False,
                stdout=subprocess.PIPE,
                stderr=subprocess.DEVNULL,
                text=True,
            )
            return completed.stdout if completed.returncode == 0 else ""

        return product_master_process_matches(pid, command_reader)

    @staticmethod
    def _clear_product_master_pid() -> None:
        try:
            PRODUCT_MASTER_PID_FILE.unlink(missing_ok=True)
        except Exception:
            pass

    def _existing_product_master_pid(self) -> int:
        try:
            pid = int(PRODUCT_MASTER_PID_FILE.read_text(encoding="utf-8").strip())
        except Exception:
            return 0
        if self._pid_is_product_master(pid):
            return pid
        self._clear_product_master_pid()
        return 0

    @staticmethod
    def _activate_process(pid: int) -> None:
        if pid <= 0 or sys.platform != "darwin":
            return
        script = (
            'tell application "System Events" to set frontmost of '
            f'(first process whose unix id is {pid}) to true'
        )
        subprocess.run(
            ["/usr/bin/osascript", "-e", script],
            check=False,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )

    def _send_product_master_request(self, prefill_product: str = "", search_product: str = "", setup_product: str = "") -> None:
        action = "show"
        value = ""
        if prefill_product:
            action, value = "new", str(prefill_product)
        elif setup_product:
            action, value = "setup", str(setup_product)
        elif search_product:
            action, value = "search", str(search_product)
        payload = {"action": action, "value": value, "requested_at": time.time()}
        PRODUCT_MASTER_REQUEST_FILE.parent.mkdir(parents=True, exist_ok=True)
        PRODUCT_MASTER_REQUEST_FILE.write_text(json.dumps(payload), encoding="utf-8")

    def _set_product_master_opening_state(self, opening: bool) -> None:
        self.product_master_launching = opening
        widgets = [
            getattr(self, "mc_master_button", None),
            getattr(self, "master_page_open_button", None),
            getattr(self, "master_page_secondary_button", None),
            getattr(self, "mc_banner_button", None),
            getattr(self, "mc_work_open_button", None),
            getattr(self, "mc_action_button", None),
        ]
        for widget in widgets:
            if widget is None:
                continue
            try:
                key = str(widget)
                if opening:
                    originals = getattr(self, "_product_master_button_labels", {})
                    originals.setdefault(key, str(widget.cget("text")))
                    self._product_master_button_labels = originals
                    widget.configure(text="Opening Product Master…", state="disabled")
                else:
                    original = getattr(self, "_product_master_button_labels", {}).get(key)
                    if original is not None:
                        widget.configure(text=original, state="normal")
            except Exception:
                pass

    def _wait_for_product_master_window(self, started_at: float, *, attempts: int = 0) -> None:
        """Finish a Product Master launch only after the editor confirms itself.

        ``open -n`` returns as soon as macOS accepts the application request,
        not when Product Master has finished building its catalog window.  The
        editor therefore writes its own PID file once it is live, which is the
        only reliable success signal for this handoff.
        """
        editor_pid = self._existing_product_master_pid()
        if editor_pid:
            self._product_master_was_open = True
            self._activate_process(editor_pid)
            self.after(300, lambda pid=editor_pid: self._activate_process(pid))
            self.dashboard_status.configure(
                text="Product Master is open. Repeated clicks will bring the same window forward.",
                text_color=PURPLE,
            )
            self._set_product_master_opening_state(False)
            self.refresh_dashboard()
            return

        # A large catalog can take a moment to load, especially immediately
        # after an app update.  Keep the dashboard responsive while allowing
        # the actual Product Master window time to appear.
        if time.monotonic() - started_at < 15.0:
            self.after(250, lambda: self._wait_for_product_master_window(started_at, attempts=attempts + 1))
            return

        self._clear_product_master_pid()
        self.dashboard_status.configure(
            text="Product Master did not open. Please try again after this message closes.",
            text_color=WARNING,
        )
        messagebox.showerror(
            "Product Master Did Not Open",
            "Orchid could not confirm the Product Master window after 15 seconds. "
            "The launch log was saved for support.\n\n"
            "Your current event and Product Master data were not changed.",
        )
        self._set_product_master_opening_state(False)
        self.refresh_dashboard()

    def open_product_master(self, prefill_product: str = "", search_product: str = "", setup_product: str = ""):
        try:
            self._sync_current_product_candidates()
            self._send_product_master_request(prefill_product, search_product, setup_product)
            existing_pid = self._existing_product_master_pid()
            if existing_pid:
                self._product_master_was_open = True
                self._activate_process(existing_pid)
                self.after(180, lambda pid=existing_pid: self._activate_process(pid))
                self.dashboard_status.configure(
                    text="Product Master was already open and has been brought to the front.",
                    text_color=PURPLE,
                )
                return

            if self.product_master_launching:
                self.dashboard_status.configure(
                    text="Product Master is still opening. Please wait…",
                    text_color=PURPLE,
                )
                return

            self._set_product_master_opening_state(True)
            command, is_bundle_handoff = product_master_launch_command(
                Path(__file__),
                sys.executable,
                is_frozen=bool(getattr(sys, "frozen", False)),
                platform=sys.platform,
            )
            if prefill_product:
                command.extend(["--new-product", str(prefill_product)])
            elif setup_product:
                command.extend(["--setup-product", str(setup_product)])
            elif search_product:
                command.extend(["--search-product", str(search_product)])
            self._clear_product_master_pid()
            with open(DATA / "product_master_launch.log", "a", encoding="utf-8") as log_file:
                log_file.write(f"{time.strftime('%Y-%m-%d %H:%M:%S')} Product Master launch: {command!r}\n")
                process = subprocess.Popen(
                    command,
                    cwd=str(PROJECT),
                    stdout=log_file,
                    stderr=log_file,
                    start_new_session=True,
                )
            self.product_master_process = process
            method = "macOS app window" if is_bundle_handoff else "Product Master window"
            self.dashboard_status.configure(text=f"Opening {method}…", text_color=PURPLE)
            self.after(150, lambda: self._wait_for_product_master_window(time.monotonic()))
        except Exception as error:
            self._set_product_master_opening_state(False)
            messagebox.showerror("Unable to Open Product Master", str(error))

    def add_new_product(self):
        value = simpledialog.askstring(
            "Add New Product",
            "Enter the missing product/style number.\n\nExample: 1104",
            parent=self,
        )
        value = " ".join((value or "").split())
        if value:
            self.open_product_master(prefill_product=value)

    def choose_purchase_order_mode(self):
        dialog = ctk.CTkToplevel(self)
        dialog.title("Choose Purchase Order Mode")
        dialog.geometry("620x390")
        dialog.resizable(False, False)
        dialog.configure(fg_color=BG)
        dialog.transient(self)
        dialog.grab_set()
        choice = {"value": None}

        ctk.CTkLabel(dialog, text="Choose Purchase Order Mode", text_color=TEXT,
                     font=ctk.CTkFont(size=24, weight="bold")).pack(anchor="w", padx=28, pady=(26, 6))
        ctk.CTkLabel(dialog, text="This choice controls how final vendor purchase orders are grouped.",
                     text_color=MUTED, font=ctk.CTkFont(size=13)).pack(anchor="w", padx=28, pady=(0, 18))

        def select(value):
            choice["value"] = value
            dialog.destroy()

        uniform = ctk.CTkFrame(dialog, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=12)
        uniform.pack(fill="x", padx=28, pady=7)
        ctk.CTkLabel(uniform, text=UNIFORM_SIZING_EVENT, text_color=PURPLE_DARK,
                     font=ctk.CTkFont(size=16, weight="bold"), anchor="w").pack(fill="x", padx=18, pady=(14, 3))
        ctk.CTkLabel(uniform, text="One report per vendor and decoration type. Decoration colors are sections within each report.",
                     text_color=MUTED, font=ctk.CTkFont(size=12), wraplength=520, justify="left", anchor="w").pack(fill="x", padx=18, pady=(0, 10))
        ctk.CTkButton(uniform, text="Use Uniform Sizing Event", fg_color=PURPLE, hover_color=PURPLE_DARK,
                      command=lambda: select(UNIFORM_SIZING_EVENT)).pack(fill="x", padx=18, pady=(0, 14))

        general = ctk.CTkFrame(dialog, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=12)
        general.pack(fill="x", padx=28, pady=7)
        ctk.CTkLabel(general, text="All Orders", text_color=PURPLE_DARK,
                     font=ctk.CTkFont(size=16, weight="bold"), anchor="w").pack(fill="x", padx=18, pady=(14, 3))
        ctk.CTkLabel(general, text="One combined report per vendor with decoration types and colors separated inside.",
                     text_color=MUTED, font=ctk.CTkFont(size=12), wraplength=520, justify="left", anchor="w").pack(fill="x", padx=18, pady=(0, 10))
        ctk.CTkButton(general, text="Use All Orders", fg_color=PURPLE, hover_color=PURPLE_DARK,
                      command=lambda: select(GENERAL_SALES_PERIOD)).pack(fill="x", padx=18, pady=(0, 14))

        dialog.protocol("WM_DELETE_WINDOW", dialog.destroy)
        dialog.wait_window()
        report_mode = choice["value"]
        if not report_mode:
            return None
        decoration_fulfillment = self.choose_decoration_fulfillment()
        if not decoration_fulfillment:
            return None
        if report_mode == UNIFORM_SIZING_EVENT:
            event_name = simpledialog.askstring(
                "Name Uniform Sizing Event",
                "Enter the event name that should appear on the workbook and final purchase orders.\n\nExample: Water Quality 2026",
                parent=self,
            )
            event_name = " ".join((event_name or "").split())
            if not event_name:
                messagebox.showwarning("Event Name Required", "Please enter a name for the Uniform Sizing Event.", parent=self)
                return None
            return report_mode, event_name, decoration_fulfillment
        if report_mode:
            event_name = simpledialog.askstring(
                "Name All Orders Period",
                "Optional: enter a helpful name for these daily/company orders.\n\nExample: July 8-12 Orders",
                parent=self,
            )
            return report_mode, " ".join((event_name or "").split()), decoration_fulfillment
        return None

    def choose_decoration_fulfillment(self):
        dialog = ctk.CTkToplevel(self)
        dialog.title("Choose Decoration Workflow")
        dialog.geometry("680x430")
        dialog.resizable(False, False)
        dialog.configure(fg_color=BG)
        dialog.transient(self)
        dialog.grab_set()
        choice = {"value": None}

        ctk.CTkLabel(
            dialog, text="How will this order be decorated?", text_color=TEXT,
            font=ctk.CTkFont(size=24, weight="bold"),
        ).pack(anchor="w", padx=28, pady=(26, 6))
        ctk.CTkLabel(
            dialog,
            text="This controls which garments appear on the Receiving & Decoration Report and which appear on the outside-decorator job report.",
            text_color=MUTED, font=ctk.CTkFont(size=13), wraplength=610, justify="left",
        ).pack(anchor="w", padx=28, pady=(0, 18))

        def select(value):
            choice["value"] = value
            dialog.destroy()

        standard = ctk.CTkFrame(dialog, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=12)
        standard.pack(fill="x", padx=28, pady=7)
        ctk.CTkLabel(
            standard, text=STANDARD_ORCHID_WORKFLOW, text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=16, weight="bold"), anchor="w",
        ).pack(fill="x", padx=18, pady=(14, 3))
        ctk.CTkLabel(
            standard,
            text="Normal Orchid workflow: embroidery and blank garments come to Orchid; screen printing is sent to the outside decorator.",
            text_color=MUTED, font=ctk.CTkFont(size=12), wraplength=580, justify="left", anchor="w",
        ).pack(fill="x", padx=18, pady=(0, 10))
        ctk.CTkButton(
            standard, text="Use Standard Orchid Workflow", fg_color=PURPLE, hover_color=PURPLE_DARK,
            command=lambda: select(STANDARD_ORCHID_WORKFLOW),
        ).pack(fill="x", padx=18, pady=(0, 14))

        outsourced = ctk.CTkFrame(dialog, fg_color=WHITE, border_width=1, border_color=BORDER, corner_radius=12)
        outsourced.pack(fill="x", padx=28, pady=7)
        ctk.CTkLabel(
            outsourced, text=ENTIRE_ORDER_OUTSOURCED, text_color=PURPLE_DARK,
            font=ctk.CTkFont(size=16, weight="bold"), anchor="w",
        ).pack(fill="x", padx=18, pady=(14, 3))
        ctk.CTkLabel(
            outsourced,
            text="All decorated garments are sent to the outside decorator. Use Do Not Outsource only for rare lines that must come to Orchid.",
            text_color=MUTED, font=ctk.CTkFont(size=12), wraplength=580, justify="left", anchor="w",
        ).pack(fill="x", padx=18, pady=(0, 10))
        ctk.CTkButton(
            outsourced, text="Use Entire Order Outsourced", fg_color=PURPLE, hover_color=PURPLE_DARK,
            command=lambda: select(ENTIRE_ORDER_OUTSOURCED),
        ).pack(fill="x", padx=18, pady=(0, 14))

        dialog.protocol("WM_DELETE_WINDOW", dialog.destroy)
        dialog.wait_window()
        return choice["value"]

    def _save_review_state(self, report_mode: str, event_name: str, decoration_fulfillment: str):
        existing = self._load_review_state()
        state = new_active_packet_state(Path(self.selected_csv))
        state.update({
            "report_mode": report_mode,
            "event_name": event_name,
            "decoration_fulfillment": normalize_decoration_fulfillment(decoration_fulfillment),
            "app_source": str(Path(__file__).resolve()),
        })
        # Preserve the append-import manifest when orders were combined before
        # the first Purchase Review workbook was created.
        for key in ("import_batches", "last_additional_import"):
            if key in existing:
                state[key] = existing[key]
        self._write_review_state_atomic(state)

    def _show_processing_window(self, title: str, message: str):
        dialog = ctk.CTkToplevel(self)
        dialog.title(title)
        dialog.geometry("560x255")
        dialog.resizable(False, False)
        dialog.configure(fg_color=BG)
        dialog.transient(self)
        dialog.grab_set()
        dialog.protocol("WM_DELETE_WINDOW", lambda: None)
        ctk.CTkLabel(
            dialog, text=title, text_color=TEXT,
            font=ctk.CTkFont(size=23, weight="bold"),
        ).pack(padx=28, pady=(30, 8))
        status = ctk.CTkLabel(
            dialog, text=message, text_color=MUTED,
            font=ctk.CTkFont(size=14), wraplength=490, justify="center",
        )
        status.pack(padx=28, pady=(0, 18))
        progress = ctk.CTkProgressBar(dialog, mode="indeterminate", width=430, height=14)
        progress.pack(padx=28, pady=(0, 16))
        progress.start()
        ctk.CTkLabel(
            dialog,
            text="You can leave Orchid Purchase Manager open while this finishes.",
            text_color=MUTED, font=ctk.CTkFont(size=12),
        ).pack(padx=28)
        try:
            dialog.update_idletasks()
            x = self.winfo_x() + max((self.winfo_width() - dialog.winfo_width()) // 2, 0)
            y = self.winfo_y() + max((self.winfo_height() - dialog.winfo_height()) // 2, 0)
            dialog.geometry(f"+{x}+{y}")
        except Exception:
            pass
        return dialog, status, progress

    def create_review_workbook(self):
        from modules.review_workbook import generate_review_workbook

        if not self.selected_csv or not self.selected_csv.exists():
            messagebox.showinfo(
                "Import Orders First",
                "No Shopify orders have been imported yet.\n\nGo to Import Orders and select the Shopify Orders CSV before processing the import.",
            )
            self.show_page("import")
            return
        selection = self.choose_purchase_order_mode()
        if not selection:
            return
        report_mode, event_name, decoration_fulfillment = selection
        selected_csv = Path(self.selected_csv)
        master_path = live_product_master_path()
        before_signature = product_master_signature(master_path)
        normalized_orders, parsed_orders = self._cached_import_frames(selected_csv)
        self._save_review_state(report_mode, event_name, decoration_fulfillment)
        dialog, status_label, progress = self._show_processing_window(
            "Processing Imported Orders",
            "Building the purchase packet and checking Product Master. Large events can take a few minutes, but the app should remain responsive.",
        )
        if hasattr(self, "review_save_button"):
            self.review_save_button.configure(state="disabled", text="Processing…")

        result_queue = queue.Queue()

        def update_status(message: str):
            result_queue.put(("status", message))

        def worker():
            try:
                update_status("Refreshing known Shopify catalog details without changing Orchid purchasing rules…")
                catalog_merge = sync_shopify_catalog_enrichment(parsed_orders, master_path)
                if catalog_merge.get("changed"):
                    try:
                        self.product_master_update_marker.parent.mkdir(parents=True, exist_ok=True)
                        self.product_master_update_marker.write_text(str(time.time()), encoding="utf-8")
                    except Exception:
                        pass
                update_status("Parsing purchase lines and applying saved product rules…")
                result = generate_review_workbook(
                    selected_csv, master_path, REPORTS,
                    report_mode=report_mode, event_name=event_name,
                    active_packet_id=_clean(self._load_review_state().get("active_packet_id", "")),
                    decoration_fulfillment=decoration_fulfillment,
                    regenerate_helper=REGENERATE_HELPER,
                    normalized_orders=normalized_orders,
                    parsed_orders=parsed_orders,
                )
                first_workbook = Path(result["output_path"])
                update_status("Checking for new or incomplete Product Master styles…")
                candidate_sync = sync_review_product_candidates(first_workbook, live_product_master_path())
                if candidate_sync.get("changed"):
                    try:
                        self.product_master_update_marker.parent.mkdir(parents=True, exist_ok=True)
                        self.product_master_update_marker.write_text(str(time.time()), encoding="utf-8")
                    except Exception:
                        pass
                    update_status("Applying the new Product Master candidates once…")
                    rebuilt = generate_review_workbook(
                        selected_csv, live_product_master_path(), REPORTS,
                        report_mode=report_mode, event_name=event_name,
                        active_packet_id=_clean(self._load_review_state().get("active_packet_id", "")),
                        decoration_fulfillment=decoration_fulfillment,
                        regenerate_helper=REGENERATE_HELPER,
                        previous_review_path=first_workbook,
                        normalized_orders=normalized_orders,
                        parsed_orders=parsed_orders,
                    )
                    result = rebuilt
                    final_workbook = Path(result["output_path"])
                    try:
                        if first_workbook.exists() and first_workbook != final_workbook:
                            first_workbook.unlink()
                    except Exception:
                        pass
                result_queue.put(("success", result, candidate_sync, catalog_merge))
            except Exception as error:
                result_queue.put(("error", str(error), traceback.format_exc()))

        def close_processing_window():
            try:
                progress.stop()
            except Exception:
                pass
            try:
                dialog.grab_release()
            except Exception:
                pass
            try:
                dialog.destroy()
            except Exception:
                pass

        def finish_error(error, detail):
            close_processing_window()
            if hasattr(self, "review_save_button"):
                self.review_save_button.configure(state="normal", text="Process Orders")
            print(detail, file=sys.stderr)
            messagebox.showerror("Unable to Process Imported Orders", str(error), parent=self)

        def finish_success(result, candidate_sync, catalog_merge):
            try:
                self.current_event_name = result["event_name"] or "General Sales Period"
                self.current_mode = result["report_mode"]
                self.current_decoration_fulfillment = result.get("decoration_fulfillment", decoration_fulfillment)
                self.last_review_workbook = Path(result["output_path"])
                self._bind_active_review_workbook(self.last_review_workbook)
                self._candidate_sync_workbook = str(self.last_review_workbook.resolve())
                self._reset_review_navigation()
                self.imported_line_count = int(result["lines"])
                self.review_needed = int(result.get("review_decisions", result["review_lines"]))
                missing_products = result.get("missing_product_numbers", [])
                status_text = f"Created {result['report_mode']} workbook with {result['lines']} purchase lines."
                if missing_products:
                    status_text += f" Added {len(missing_products)} unknown product number(s) to the Product Master setup queue."
                self.dashboard_status.configure(text=status_text, text_color=WARNING if missing_products else SUCCESS)
                close_processing_window()
                # Return control immediately. Current Event only needs the small
                # review queue plus the import totals; the full 1,085-line mission
                # snapshot is warmed on a background thread.
                try:
                    quick_snapshot = load_purchase_review_snapshot(self.last_review_workbook, live_product_master_path())
                except Exception:
                    quick_snapshot = {"issues": []}
                quick_snapshot.update({
                    "event_name": self.current_event_name,
                    "report_mode": self.current_mode,
                    "line_count": self.imported_line_count,
                    "review_count": self.review_needed or 0,
                    "purchase_review_count": int(result.get("purchase_review_decisions", self.review_needed or 0) or 0),
                    "product_master_review_count": int(result.get("product_master_decisions", 0) or 0),
                })
                self._cache_fast_review_snapshot(quick_snapshot)
                self.last_snapshot = quick_snapshot
                self.refresh_current_event_page(quick_snapshot)
                self._warm_dashboard_snapshot_async(self.last_review_workbook)
                summary = (
                    f"Purchase order mode: {result['report_mode']}\n"
                    f"Decoration workflow: {result.get('decoration_fulfillment', STANDARD_ORCHID_WORKFLOW)}\n"
                    + (f"Event name: {result['event_name']}\n" if result['event_name'] else "")
                    + f"Purchase lines: {result['lines']}\n"
                    f"Total quantity: {result['quantity']}\n"
                    f"Product Master products needing setup: {result.get('product_master_decisions', 0)}\n"
                    f"Order-specific Purchase Review decisions: {result.get('purchase_review_decisions', 0)}\n"
                    f"Automatically routed rows: {result.get('inferred_routing_lines', 0)}\n"
                    f"Rows with purchase instructions: {result['notes_lines']}\n"
                    + (f"Protected packet lock: {result.get('packet_lock_id', '')}\n" if result.get('packet_lock_id') else "")
                    + (f"New Product Master candidates: {candidate_sync.get('added_styles', 0)}\n" if candidate_sync.get('added_styles') else "")
                    + (
                        f"Repeat-import duplicate color rows repaired: "
                        f"{int(catalog_merge.get('repaired_duplicate_colors', 0) or 0) + int(candidate_sync.get('repaired_duplicate_colors', 0) or 0)}\n"
                        if int(catalog_merge.get('repaired_duplicate_colors', 0) or 0) + int(candidate_sync.get('repaired_duplicate_colors', 0) or 0)
                        else ""
                    )
                    + (f"Shopify descriptions refreshed: {catalog_merge.get('updated_descriptions', 0)}\n" if catalog_merge.get('updated_descriptions') else "")
                    + (f"New Shopify garment colors added for setup: {catalog_merge.get('added_colors', 0)}\n" if catalog_merge.get('added_colors') else "")
                    + "\nProcessing is complete. Continue inside Orchid Purchase Manager."
                )
                # Keep the success dialog concise; the long file path remains visible on Current Event.
                messagebox.showinfo("Imported Orders Processed", summary, parent=self)
                if self.active_page != "import":
                    self.show_page("import")
                self.after(120, lambda: self.refresh_current_event_page(self.last_snapshot or {}))
                self.after(260, self._play_orchid_success_animation)
            except Exception as error:
                finish_error(error, traceback.format_exc())

        def poll_worker():
            try:
                while True:
                    item = result_queue.get_nowait()
                    kind = item[0]
                    if kind == "status":
                        if status_label.winfo_exists():
                            status_label.configure(text=item[1])
                    elif kind == "success":
                        finish_success(item[1], item[2], item[3])
                        return
                    elif kind == "error":
                        finish_error(RuntimeError(item[1]), item[2])
                        return
            except queue.Empty:
                pass
            if dialog.winfo_exists():
                self.after(100, poll_worker)

        threading.Thread(target=worker, daemon=True, name="orchid-import-worker").start()
        self.after(100, poll_worker)

    def _run_final_purchase_order_generation(
        self,
        workbook_path: Path,
        po_overrides=None,
        active_workbook_path: Path | None = None,
        hold_unresolved_review: bool = False,
        test_mode: bool = False,
    ):
        from modules.final_po_generator import generate_final_purchase_orders

        try:
            if not self._require_active_packet("Purchase Order generation"):
                return
            requested_workbook = Path(workbook_path)
            active_workbook = Path(active_workbook_path) if active_workbook_path else requested_workbook
            # Every report-generation path must use the same reconciled event
            # data.  The initial Generate Purchase Orders flow already supplies
            # a journal-reconciled workbook, but Regenerate Purchase Orders and
            # Regenerate Purchase Review & Reports previously sent the original
            # workbook directly to the PDF generator.  That could resurrect an
            # older Decoration Color (for example JP56 Baby showing as Black)
            # even though the color-specific Product Master/audit decision was
            # saved correctly.
            report_only_archive = False
            effective_workbook = requested_workbook
            if active_workbook_path is None and not report_only_archive:
                effective_workbook = self._build_journal_reconciled_workbook(active_workbook)
            overrides = (
                po_overrides if po_overrides is not None
                else load_po_overrides(DATA, active_workbook)
            )
            job_logo = self._outsourced_job_logo_path()
            outsourced_job_name = self._outsourced_job_name()
            outsourced_cover_artwork = self._outsourced_cover_artwork_entries()
            outsourced_cover_notes = self._outsourced_cover_notes()
            outsourced_to = self._outsourced_job_metadata().get("outsourced_to", {})
            job_logo_metadata = self._outsourced_job_logo_metadata_path()
            result = generate_final_purchase_orders(
                effective_workbook, REPORTS, po_number_overrides=overrides,
                job_logo_path=job_logo,
                outsourced_job_name=outsourced_job_name,
                outsourced_cover_artwork=outsourced_cover_artwork,
                outsourced_cover_notes=outsourced_cover_notes,
                outsourced_to=outsourced_to,
                allow_historical_lock=False,
                hold_unresolved_review=hold_unresolved_review,
                test_mode=test_mode,
            )
            generated_from = effective_workbook
            output_dir = Path(result["output_dir"])
            released_pdfs = [Path(path) for path in result.get("pdf_files", []) if Path(path).is_file()]
            expected_reports = int(result.get("report_count", 0) or 0)
            release_manifest = Path(result.get("release_manifest_path", ""))
            if not output_dir.is_dir():
                raise RuntimeError("Purchase-order generation did not create a final report folder.")
            if expected_reports <= 0:
                raise RuntimeError("Purchase-order generation completed without creating any vendor or internal purchase-order PDFs.")
            if len(released_pdfs) < expected_reports:
                raise RuntimeError(
                    f"Purchase-order generation expected {expected_reports} PDF(s), but only {len(released_pdfs)} verified file(s) were released."
                )
            if not release_manifest.is_file():
                raise RuntimeError("The protected final-release manifest was not created.")
            release_valid, release_reason = validate_release_packet(output_dir)
            if not release_valid:
                raise RuntimeError(f"The protected final-release packet could not be verified: {release_reason}.")

            self.last_review_workbook = active_workbook
            self.last_purchase_order_dir = output_dir
            self.current_event_name = result["event_name"] or self.current_event_name or "General Sales Period"
            self.current_mode = result["report_mode"]
            self.review_needed = int(result["review_lines"])
            state = self._load_review_state()
            state.pop("decoration_color_audit_requires_regeneration", None)
            state["last_purchase_order_dir"] = str(self.last_purchase_order_dir)
            state["active_purchase_order_dir"] = str(self.last_purchase_order_dir)
            state["active_purchase_order_packet_lock_id"] = _clean(
                state.get("active_review_packet_lock_id", "")
            )
            state["last_purchase_order_generated_at"] = time.time()
            state.pop("last_purchase_order_archive_dir", None)
            state.pop("completed_purchase_order_archive_dir", None)
            state.pop("archived_report_reopen", None)
            state.pop("archive_source_workbook", None)
            try:
                self._write_review_state_atomic(state)
            except Exception:
                pass
            held_count = int(result.get("review_lines", 0) or 0)
            if test_mode:
                status_text = (
                    f"Created {result['report_count']} test purchase-order document(s). "
                    f"{held_count} unfinished line{'s' if held_count != 1 else ''} were held out of the orders."
                )
            else:
                status_text = (
                    f"Created {result['report_count']} purchase-order document(s), "
                    f"plus {result.get('decoration_report_count', 0)} decoration report(s)."
                )
            self.dashboard_status.configure(text=status_text, text_color=SUCCESS)
            # Force a fresh workflow/page render after generation.  The
            # dashboard cache must not leave Step 5 showing the pre-generation
            # "Ready" state after protected PDFs have already been released.
            self._dashboard_refresh_key = None
            self.refresh_dashboard()
            self.show_page("purchase")
            self.refresh_purchase_page()
            visible_documents = self._purchase_order_document_paths()
            if not visible_documents:
                raise RuntimeError(
                    "Protected PDFs were released, but the Purchase Orders page could not locate them. "
                    f"Open this folder manually: {self.last_purchase_order_dir}"
                )
            if test_mode:
                hold_text = (
                    f"\n\n{held_count} unfinished line{'s were' if held_count != 1 else ' was'} held from vendor orders in "
                    "TEST_ONLY_Held_for_Review.pdf."
                    if held_count else "\n\nNo unfinished lines were held."
                )
                messagebox.showinfo(
                    "Test Purchase Orders Created",
                    f"Created {len(visible_documents)} test purchase-order PDF(s)."
                    f"{hold_text}\n\n"
                    "The protected release check verified that every included garment appears exactly once "
                    "in a vendor order, the held-items report, or the internal purchase route.\n\n"
                    f"Folder: {self.last_purchase_order_dir}\n\n"
                    "Finder will open this event's verified PDF folder when this message closes.",
                )
            else:
                messagebox.showinfo(
                    "Purchase Orders Created",
                    f"Created {len(visible_documents)} purchase-order PDF(s).\n\n"
                    f"Folder: {self.last_purchase_order_dir}\n\n"
                    "Finder will open this event's verified PDF folder when this message closes.",
                )
            # A successful final release should take the user directly to the
            # one verified folder.  Do this after the confirmation closes so
            # Finder is not hidden behind the modal success message.
            self.after(120, self.open_latest_purchase_orders)
            self.after(220, self._play_orchid_success_animation)
        except Exception as error:
            messagebox.showerror("Unable to Create Final Purchase Orders", str(error))

    def create_final_purchase_orders(self):
        if self._review_save_pending:
            messagebox.showinfo("Finishing Purchase Review Saves", "Wait for the background Purchase Review saves to finish before selecting a final workbook.")
            return
        initial = REPORTS / "Review Workbooks"
        if not initial.exists():
            initial = REPORTS / "review_workbooks"
        path = filedialog.askopenfilename(
            title="Choose Edited Orchid Purchase Review Workbook",
            initialdir=str(initial if initial.exists() else REPORTS),
            filetypes=[("Excel workbooks", "*.xlsx")],
        )
        if not path:
            return
        workbook = Path(path)
        snapshot = load_mission_control_snapshot(workbook, DATA, live_product_master_path())
        review_count = int(snapshot.get("review_count", 0))
        blocked_route_count = int(snapshot.get("blocked_route_count", 0))
        if review_count or blocked_route_count:
            messagebox.showwarning(
                "Purchase Review Incomplete",
                "The selected workbook did not pass the final readiness check.\n\n"
                + (self._blocker_summary(snapshot) or f"{review_count} decision(s) and {blocked_route_count} blocked route(s) remain."),
            )
            return
        self.last_review_workbook = workbook
        self.open_po_number_editor(generate_after_save=True, workbook_path=workbook)

    def open_event_archive(self):
        folder = REPORTS / "Event Archive"
        folder.mkdir(parents=True, exist_ok=True)
        subprocess.run(["open", str(folder)], check=False)

    def _is_archived_report_reopen(self) -> bool:
        """Archived events are no longer available to the live workflow."""
        return False

    def _stage_archived_event_for_reports(self, archived_workbook: Path) -> dict:
        """Create a protected working copy of one archived completed event.

        The archive is never modified. The working copy is placed with normal
        Review Workbooks so it can be reopened again after the app is closed.
        """
        archived_workbook = Path(archived_workbook)
        if not archived_workbook.exists():
            raise FileNotFoundError(f"Archived Purchase Review workbook was not found: {archived_workbook}")

        records = load_review_lines(archived_workbook)
        if not records:
            raise ValueError("The selected archive workbook does not contain any purchase-order lines.")

        event_name = load_event_name(archived_workbook) or archived_workbook.parent.name
        report_mode = load_report_mode(archived_workbook)
        decoration_fulfillment = load_decoration_fulfillment(archived_workbook)
        label = safe_filename(event_name or "Archived Event")
        stamp = time.strftime("%Y%m%d_%H%M%S")
        working_dir = REPORTS / "Review Workbooks"
        working_dir.mkdir(parents=True, exist_ok=True)
        working_workbook = working_dir / f"{label}__Archived_Report_Reopen_{stamp}.xlsx"
        shutil.copy2(archived_workbook, working_workbook)

        archived_csvs = [path for path in archived_workbook.parent.glob("*.csv") if path.is_file()]
        working_csv = None
        if archived_csvs:
            archived_csv = max(archived_csvs, key=lambda path: path.stat().st_mtime_ns)
            working_csv = working_dir / f"{label}__Archived_Source_{stamp}{archived_csv.suffix.lower()}"
            shutil.copy2(archived_csv, working_csv)

        # An archived packet can include a logo from a prior report run. Copy it
        # beside the protected working copy, where the existing upload/replace
        # controls can manage it without touching the archive itself.
        logo_candidates = []
        for suffix in (".png", ".jpg", ".jpeg", ".webp"):
            logo_candidates.append(
                archived_workbook.parent / f"{safe_filename(event_name)}__Outsourced_Job_Logo{suffix}"
            )
        logo_candidates.extend(
            path for path in archived_workbook.parent.iterdir()
            if path.is_file()
            and path.suffix.casefold() in {".png", ".jpg", ".jpeg", ".webp"}
            and "outsourced_job_logo" in path.stem.casefold()
        )
        working_logo = None
        for candidate in logo_candidates:
            if candidate.exists():
                working_logo = working_dir / f"{safe_filename(event_name)}__Outsourced_Job_Logo{candidate.suffix.lower()}"
                shutil.copy2(candidate, working_logo)
                break

        # Restore the sidecar that holds the outside decorator's production
        # name (for example, “Utilities Department”) along with the image.
        metadata_candidates = [
            archived_workbook.parent / f"{safe_filename(event_name)}__Outsourced_Job_Logo_Metadata.json",
        ]
        metadata_candidates.extend(
            path for path in archived_workbook.parent.iterdir()
            if path.is_file()
            and path.suffix.casefold() == ".json"
            and "outsourced_job_logo_metadata" in path.stem.casefold()
        )
        working_logo_metadata = None
        for candidate in metadata_candidates:
            if candidate.exists():
                working_logo_metadata = working_dir / f"{safe_filename(event_name)}__Outsourced_Job_Logo_Metadata.json"
                shutil.copy2(candidate, working_logo_metadata)
                break

        # The report-cover metadata names the exact per-color artwork files.
        # Restore those files beside the staged workbook as well so a later
        # archived report regeneration uses the same artwork reference page.
        if working_logo_metadata:
            try:
                cover_artwork = json.loads(working_logo_metadata.read_text(encoding="utf-8")).get("artwork", {})
            except (OSError, ValueError, TypeError, AttributeError):
                cover_artwork = {}
            if isinstance(cover_artwork, dict):
                for entry in cover_artwork.values():
                    filename = _clean(entry.get("filename", "")) if isinstance(entry, dict) else ""
                    if not filename or Path(filename).name != filename:
                        continue
                    candidate = archived_workbook.parent / filename
                    if candidate.is_file():
                        shutil.copy2(candidate, working_dir / filename)

        return {
            "archive_workbook": archived_workbook,
            "working_workbook": working_workbook,
            "working_csv": working_csv,
            "working_logo": working_logo,
            "working_logo_metadata": working_logo_metadata,
            "event_name": event_name,
            "report_mode": report_mode,
            "decoration_fulfillment": decoration_fulfillment,
            "line_count": len(records),
        }

    def reopen_archived_event_for_reports(self):
        """Reopen one completed archive solely to update its logo and reports."""
        archive_root = REPORTS / "Event Archive"
        archive_root.mkdir(parents=True, exist_ok=True)
        if (
            self.last_review_workbook
            and self.last_review_workbook.exists()
            and not self._is_archived_report_reopen()
            and not messagebox.askyesno(
                "Switch to Archived Event?",
                "This will switch Orchid to a safe working copy of the selected completed event. "
                "Your current event remains saved and can be reopened later. Continue?",
                parent=self,
            )
        ):
            return
        selected = filedialog.askopenfilename(
            parent=self,
            title="Choose Archived Purchase Review Workbook",
            initialdir=str(archive_root),
            filetypes=[("Excel workbooks", "*.xlsx")],
        )
        if not selected:
            return
        try:
            staged = self._stage_archived_event_for_reports(Path(selected))
            self.last_review_workbook = Path(staged["working_workbook"])
            self.selected_csv = Path(staged["working_csv"]) if staged["working_csv"] else None
            self.last_purchase_order_dir = (
                Path(selected).parent / "Purchase Orders"
                if (Path(selected).parent / "Purchase Orders").exists()
                else None
            )
            self.current_event_name = _clean(staged["event_name"])
            self.current_mode = _clean(staged["report_mode"])
            self.current_decoration_fulfillment = normalize_decoration_fulfillment(
                staged["decoration_fulfillment"]
            )
            self.imported_line_count = int(staged["line_count"])
            self.review_needed = 0
            self.last_snapshot = {}
            self._review_save_pending = 0
            self._review_save_last_error = ""
            self._review_decision_history = []
            self._review_history_index = None
            self._review_search_issue = None
            self._invalidate_review_cache()
            self._invalidate_fast_review_snapshot()
            self._mission_snapshot_cache_key = None
            self._mission_snapshot_cache = {}
            self._decoration_audit_source_stamp = None
            self._decoration_audit_event_records_cache_key = None
            self._decoration_audit_event_records_cache = []
            self._decoration_audit_build_cache_key = None
            self._decoration_audit_build_cache = []
            self.decoration_audit_records = []
            self.decoration_audit_pending_edits = {}

            state = {
                "source_csv": str(self.selected_csv or ""),
                "source_signature": self._source_file_signature(self.selected_csv),
                "report_mode": self.current_mode,
                "event_name": self.current_event_name,
                "decoration_fulfillment": self.current_decoration_fulfillment,
                "review_complete": True,
                "review_completed_at": time.time(),
                "archived_report_reopen": True,
                "archive_source_workbook": str(staged["archive_workbook"]),
                "completed_purchase_order_archive_dir": str(Path(selected).parent),
                "reopened_workbook": str(self.last_review_workbook),
                "outsourced_job_name": self._outsourced_job_name(),
                "app_source": str(Path(__file__).resolve()),
            }
            self._write_review_state_atomic(state)
            self.refresh_dashboard()
            self.show_page("job_logo")
            self.after(100, self.refresh_outsourced_job_logo_page)
            messagebox.showinfo(
                "Archived Event Reopened",
                f"{self.current_event_name} is open as a safe report-only working copy.\n\n"
                "Next: confirm the outsourced job name, upload the logo if needed, then click Regenerate Archived Reports. "
                "The original archived event will remain unchanged.",
                parent=self,
            )
        except Exception as error:
            messagebox.showerror("Unable to Reopen Archived Event", str(error), parent=self)

    def backup_product_master_now(self):
        try:
            backup = backup_file(live_product_master_path(), label="manual")
            if not backup:
                messagebox.showinfo("Product Master Backup", "There is no active Product Master file to back up.")
                return
            messagebox.showinfo("Product Master Backed Up", f"Backup saved as:\n{backup}")
        except Exception as error:
            messagebox.showerror("Unable to Back Up Product Master", str(error))

    def restore_product_master_backup(self):
        initial = DATA / "backups"
        initial.mkdir(parents=True, exist_ok=True)
        selected = filedialog.askopenfilename(
            title="Choose Product Master Backup",
            initialdir=str(initial),
            filetypes=[("CSV files", "*.csv")],
        )
        if not selected:
            return
        if not messagebox.askyesno(
            "Restore Product Master?",
            "The current Product Master will be backed up, then replaced with the selected CSV. Continue?",
            parent=self,
        ):
            return
        try:
            result = restore_product_master(live_product_master_path(), Path(selected))
            messagebox.showinfo(
                "Product Master Restored",
                f"Restored {result['records']} records across {result['styles']} products/styles.\n\n"
                f"Previous active file backed up as:\n{result['backup'] or 'No previous file'}",
            )
            self.refresh_dashboard()
        except Exception as error:
            messagebox.showerror("Unable to Restore Product Master", str(error))

    def restore_clean_product_master(self):
        if not messagebox.askyesno(
            "Restore Clean Product Master?",
            "This will replace the active 2.x Product Master with the clean 166-record Product Master included with the original v10.0 package. The current file will be backed up first. Continue?",
            parent=self,
        ):
            return
        try:
            result = restore_seed_product_master(live_product_master_path())
            messagebox.showinfo(
                "Clean Product Master Restored",
                f"Restored {result['records']} records across {result['styles']} products/styles.\n\n"
                f"Previous active file backed up as:\n{result['backup'] or 'No previous file'}",
            )
            self.refresh_dashboard()
        except Exception as error:
            messagebox.showerror("Unable to Restore Clean Product Master", str(error))

    def clean_catalog(self):
        try:
            result = clean_product_master(live_product_master_path(), make_backup=True)
            self.dashboard_status.configure(text=f"Product Master cleaned. {result['after']} unique records remain.", text_color=SUCCESS)
            backup_text = str(result["backup"]) if result.get("backup") else "No backup was needed."
            messagebox.showinfo(
                "Product Master Cleaned",
                f"Rows before: {result['before']}\n"
                f"Rows after: {result['after']}\n"
                f"Duplicates removed: {result['removed']}\n\n"
                f"Backup saved as:\n{backup_text}",
            )
            self.refresh_dashboard()
        except Exception as error:
            messagebox.showerror("Unable to Clean Product Master", str(error))

    def open_latest_review(self):
        if self._review_save_pending:
            messagebox.showinfo("Finishing Purchase Review Saves", "Orchid is still synchronizing completed decisions to Excel. The workbook can be opened when the saves finish.")
            return
        if not self.last_review_workbook or not self.last_review_workbook.exists():
            self.discover_existing_work()
        if self.last_review_workbook and self.last_review_workbook.exists():
            subprocess.run(["open", str(self.last_review_workbook)], check=False)
        else:
            messagebox.showinfo("No Purchase Review", "Create a Purchase Review first.")

    def open_review_folder(self):
        folder = REPORTS / "Review Workbooks"
        if not folder.exists():
            folder = REPORTS / "review_workbooks"
        folder.mkdir(parents=True, exist_ok=True)
        subprocess.run(["open", str(folder)], check=False)

    def open_purchase_orders_folder(self):
        folder = REPORTS / "Purchase Orders"
        folder.mkdir(parents=True, exist_ok=True)
        subprocess.run(["open", str(folder)], check=False)


def _save_and_close_active_excel_workbook() -> None:
    """Best-effort save/close for the workbook that launched regeneration."""
    if sys.platform != "darwin":
        return
    script = (
        'tell application "Microsoft Excel"\n'
        'if it is running then\n'
        'try\n'
        'save active workbook\n'
        'close active workbook saving yes\n'
        'end try\n'
        'end if\n'
        'end tell'
    )
    subprocess.run(["/usr/bin/osascript", "-e", script], check=False,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    time.sleep(0.8)


def _show_regeneration_error(message: str) -> None:
    if sys.platform == "darwin":
        safe = str(message).replace('"', "'")
        subprocess.run([
            "/usr/bin/osascript", "-e",
            f'display alert "Purchase Review could not be regenerated" message "{safe}" as critical',
        ], check=False)
    else:
        print(message, file=sys.stderr)


def regenerate_review_from_saved_state(request_path: Path | None = None):
    """Regenerate from a workbook request file or the app's saved state."""
    from modules.review_workbook import generate_review_workbook

    try:
        if request_path:
            request = Path(request_path).expanduser().resolve()
            if not request.exists():
                raise FileNotFoundError(f"Regeneration request not found: {request}")
            state = json.loads(request.read_text(encoding="utf-8"))
        else:
            if not REVIEW_STATE_FILE.exists():
                raise FileNotFoundError("No saved Purchase Review state was found.")
            state = json.loads(REVIEW_STATE_FILE.read_text(encoding="utf-8"))

        source_csv = Path(state.get("source_csv", "")).expanduser()
        if not source_csv.exists():
            raise FileNotFoundError(f"The saved order CSV no longer exists: {source_csv}")
        previous_workbook_text = state.get("previous_workbook", "")
        previous_workbook = Path(previous_workbook_text).expanduser() if previous_workbook_text else None

        _save_and_close_active_excel_workbook()
        result = generate_review_workbook(
            source_csv, live_product_master_path(), REPORTS,
            report_mode=state.get("report_mode") or GENERAL_SALES_PERIOD,
            event_name=state.get("event_name") or "",
            active_packet_id=_clean(state.get("active_packet_id", "")),
            decoration_fulfillment=state.get("decoration_fulfillment") or STANDARD_ORCHID_WORKFLOW,
            regenerate_helper=REGENERATE_HELPER,
            previous_review_path=previous_workbook,
        )
        subprocess.run(["open", str(result["output_path"])], check=False)
    except Exception as error:
        _show_regeneration_error(str(error))
        raise SystemExit(1) from error


if __name__ == "__main__":
    regeneration_requests = [Path(arg) for arg in sys.argv[1:] if str(arg).casefold().endswith(".orchidregen")]
    if regeneration_requests:
        regenerate_review_from_saved_state(regeneration_requests[0])
    elif "--regenerate-review" in sys.argv:
        regenerate_review_from_saved_state()
    elif "--product-master" in sys.argv:
        from modules.product_master_editor import ProductMasterV2
        prefill = ""
        if "--new-product" in sys.argv:
            try:
                prefill = sys.argv[sys.argv.index("--new-product") + 1]
            except (ValueError, IndexError):
                prefill = ""
        search = ""
        if "--search-product" in sys.argv:
            try:
                search = sys.argv[sys.argv.index("--search-product") + 1]
            except (ValueError, IndexError):
                search = ""
        setup = ""
        if "--setup-product" in sys.argv:
            try:
                setup = sys.argv[sys.argv.index("--setup-product") + 1]
            except (ValueError, IndexError):
                setup = ""
        ProductMasterV2(prefill_product=prefill, search_product=search, setup_product=setup).mainloop()
    else:
        OrchidPurchaseManager().mainloop()
